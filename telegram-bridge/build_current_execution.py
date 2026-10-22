#!/usr/bin/env python3
"""Build the best available cumulative execution workbook for a deal."""
import json
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from build_cumulative_deal_execution import build as build_table_cumulative

BASE = os.path.dirname(os.path.abspath(__file__))
ARCHIVE = os.path.join(BASE, "installer_approved_reports.json")
OUT = os.path.join(BASE, "generated_reports")
NAVY, WHITE, GREEN = "102733", "FFFFFF", "E2F0D9"


def safe(value, limit=120):
    return re.sub(r"[^0-9A-Za-zА-Яа-яЁё№._ -]+", "_", str(value or "")).strip(" ._")[:limit]


def build_generic(deal_id, client, essence, reports, path):
    wb = Workbook()
    ws = wb.active
    ws.title = safe(f"{deal_id} {client} {essence}", 31) or f"Сделка {deal_id}"
    ws["A1"] = f"НАКОПИТЕЛЬНОЕ ВЫПОЛНЕНИЕ — №{deal_id} — {client} — {essence}"
    ws.merge_cells("A1:F1")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=14)
    headers = ["Дата отчёта", "Объект", "Выполненная работа", "Сумма", "Накопительный итог", "Источник"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(3, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
    cumulative = 0.0
    for report in sorted(reports, key=lambda item: datetime.strptime(item["date"], "%d.%m.%Y")):
        for deal in report.get("deals") or []:
            if str(deal.get("deal_id")) != str(deal_id):
                continue
            for row in deal.get("rows") or []:
                amount = float(row.get("amount") or 0)
                cumulative += amount
                ws.append([report.get("date"), row.get("object"), row.get("source"), amount, cumulative, report.get("source_file")])
    ws.append([None, None, "ИТОГО НАКОПИТЕЛЬНО", f"=SUM(D4:D{ws.max_row})", cumulative])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=GREEN)
    for col, width in enumerate([16, 28, 68, 18, 20, 38], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column in (4, 5):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
    ws.freeze_panes = "A4"
    wb.save(path)


def build(spec):
    deal_id = str(spec["deal_id"])
    client = spec.get("client") or "Клиент"
    essence = spec.get("essence") or "Выполнение работ"
    data = json.load(open(ARCHIVE, encoding="utf-8"))
    reports = [report for report in data.get("reports") or [] if any(str(deal.get("deal_id")) == deal_id for deal in report.get("deals") or [])]
    if not reports:
        raise ValueError("по этой сделке пока нет утверждённых отчётов монтажника")
    table_paths = []
    for report in reports:
        matched_deals = [deal for deal in report.get("deals") or [] if str(deal.get("deal_id")) == deal_id]
        for deal in matched_deals:
            table_paths.extend(deal.get("table_sources") or [])
        if len(report.get("deals") or []) == 1:
            table_paths.extend(report.get("table_sources") or [])
    table_path = next((path for path in reversed(table_paths) if os.path.isfile(path)), None)
    os.makedirs(OUT, exist_ok=True)
    filename = safe(f"Накопительное выполнение - №{deal_id} - {client} - {essence}") + ".xlsx"
    path = os.path.join(OUT, filename)
    if table_path:
        try:
            build_table_cumulative(deal_id, table_path, path)
            wb = load_workbook(path)
            ws = wb.active
            ws.title = safe(f"{deal_id} {client} {essence}", 31) or f"Сделка {deal_id}"
            ws["A1"] = f"НАКОПИТЕЛЬНОЕ ВЫПОЛНЕНИЕ — №{deal_id} — {client} — {essence}"
            wb.save(path)
        except ValueError:
            build_generic(deal_id, client, essence, reports, path)
    else:
        build_generic(deal_id, client, essence, reports, path)
    return {"path": path, "filename": filename, "reports_count": len(reports)}


if __name__ == "__main__":
    try:
        print(json.dumps({"ok": True, **build(json.load(sys.stdin))}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        raise SystemExit(1)
