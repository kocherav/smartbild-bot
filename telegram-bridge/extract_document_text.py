"""Read Office invoice cells without empty merged-cell noise; no external writes."""
import json,sys
from pathlib import Path

def extract(path):
    suffix=Path(path).suffix.lower();lines=[]
    if suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            for sheet in book:
                lines.append('Лист: '+sheet.title)
                for row in sheet.iter_rows(values_only=True):
                    cells=[str(x).strip() for x in row if x is not None and str(x).strip()]
                    if cells:lines.append('\t'.join(cells))
        finally:book.close()
    elif suffix=='.xls':
        import xlrd
        book=xlrd.open_workbook(path)
        for sheet in book.sheets():
            lines.append('Лист: '+sheet.name)
            for n in range(sheet.nrows):
                cells=[str(c.value).strip() for c in sheet.row(n) if c.ctype not in (0,6) and str(c.value).strip()]
                if cells:lines.append('\t'.join(cells))
    elif suffix=='.docx':
        from docx import Document
        doc=Document(path);lines.extend(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                seen=set();cells=[]
                for cell in row.cells:
                    if cell._tc in seen:continue
                    seen.add(cell._tc)
                    if cell.text.strip():cells.append(cell.text.strip())
                if cells:lines.append('\t'.join(cells))
    else:raise ValueError('Unsupported Office format')
    return '\n'.join(lines)
if __name__=='__main__':print(extract(json.load(sys.stdin)['path']))
