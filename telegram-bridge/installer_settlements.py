#!/usr/bin/env python3
"""Persistent installer accrual/payment ledger and cumulative Excel export."""
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import math
import durable_json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
LEDGER = os.path.join(BASE, "installer_settlements.json")
INSTALLATIONS_DB=os.path.join(os.path.dirname(BASE),'megaplan-crm','installations.sqlite3')
INSTALLER = "Сергей Деревянко"
NAVY, WHITE, GREEN, YELLOW = "102733", "FFFFFF", "E2F0D9", "FFF2CC"


def load():
    data = durable_json.read(LEDGER,lambda:{"version":1,"operations":[],"allocations":[]})
    data.setdefault("operations", [])
    data.setdefault("allocations", [])
    return data


def save(data):
    durable_json.save(LEDGER,data)


def digest_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def operation_id(kind, identity):
    return hashlib.sha256(f"{kind}|{identity}".encode()).hexdigest()[:20]


@durable_json.transaction(lambda:LEDGER)
def record_accrual(session):
    data = load()
    source_path = session.get("source_path")
    source_hash = digest_file(source_path) if source_path and os.path.isfile(source_path) else operation_id("source", json.dumps(session.get("report_lines") or [], ensure_ascii=False, sort_keys=True))
    oid = operation_id("accrual", source_hash)
    amount = float(session.get("payable_amount") if session.get("payable_amount") is not None else session.get("net_total") or 0)
    if not math.isfinite(amount) or amount<0: raise ValueError('Некорректная сумма начисления')
    existing = next((op for op in data["operations"] if op["id"] == oid), None)
    if existing:
        if round(float(existing['amount']),2)!=round(amount,2):
            raise ValueError('Этот отчёт уже принят с другой суммой; требуется отдельная корректировка, повторная приёмка запрещена')
        return existing, False
    op = {
        "id": oid, "type": "accrual", "installer": INSTALLER,
        "display_name": f"Деревянко Сергей — отчёт от {session.get('report_date') or session.get('report')}",
        "date": session.get("report_date") or session.get("report"),
        "amount": amount, "source_hash": source_hash,
        "source_file": os.path.basename(source_path or ""),
        "report": session.get("report"),
        "deals": sorted({str(item.get("deal_id")) for item in session.get("exceptions") or [] if item.get("deal_id")}),
        "created_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
    }
    data["operations"].append(op)
    allocate_available_advances(data,op)
    save(data)
    return op, True


def allocate_available_advances(data,accrual):
    """Use only unapplied historical credit; never rewrite old allocations/rates."""
    remaining=round(float(accrual['amount']),2)
    for payment in sorted((x for x in data['operations'] if x.get('type')=='payment'),
                          key=lambda x:(date_key(x.get('date')),x.get('created_at') or '')):
        credit=float(payment['settlement_amount'] if payment.get('settlement_amount') is not None else payment['amount'])
        if not math.isfinite(credit) or credit<0:raise ValueError('Некорректная ранее зачтённая выплата')
        consumed=0
        for row in data['allocations']:
            if row['payment_id']!=payment['id']:continue
            amount=float(row['amount'])
            if not math.isfinite(amount) or amount<0:raise ValueError('Некорректное распределение выплаты')
            consumed+=amount
        available=round(credit-consumed,2)
        if available < -0.01:raise ValueError('Распределения превышают ранее зачтённую выплату')
        amount=round(min(max(0,available),remaining),2)
        if amount>0:
            data['allocations'].append({'payment_id':payment['id'],'accrual_id':accrual['id'],'amount':amount})
            remaining=round(remaining-amount,2)
            payment['unallocated_advance']=round(available-amount,2)


def pdf_text(path):
    result = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, timeout=30)
    if result.returncode:
        raise ValueError("не удалось прочитать платёжный документ")
    return result.stdout


def parse_amount(text):
    patterns = [
        r"(?:Сумма|СУММА)\s*(?:платежа)?\s*[: ]\s*([0-9][0-9\s]*(?:[,.][0-9]{2})?)",
        r"([0-9][0-9\s]*(?:[,.][0-9]{2}))\s*(?:руб|RUB)",
    ]
    values = []
    for pattern in patterns:
        for raw in re.findall(pattern, text):
            try:
                values.append(float(raw.replace(" ", "").replace(",", ".")))
            except ValueError:
                pass
    return max(values) if values else None


def parse_payment(path):
    text = pdf_text(path) if path.casefold().endswith(".pdf") else open(path, encoding="utf-8", errors="replace").read()
    date_match = re.search(r"\b([0-3]?\d[./-][01]?\d[./-]20\d{2})\b", text)
    number_match = re.search(r"ПЛАТ[ЕЁ]ЖНОЕ\s+ПОРУЧЕНИЕ\s+([0-9A-Za-zА-Яа-я/-]+)", text, re.IGNORECASE)
    lines = [line.strip() for line in text.splitlines()]
    recipient = None
    purpose = None
    recipient_label = next((i for i, line in enumerate(lines) if line.casefold() == "получатель"), None)
    purpose_label = next((i for i, line in enumerate(lines) if line.casefold() == "назначение платежа"), None)
    if recipient_label is not None:
        candidates = [line for line in lines[max(0, recipient_label - 12):recipient_label] if line]
        recipient = next((line for line in reversed(candidates) if not any(mark in line.casefold() for mark in ("сч. №", "инн", "бик", "банк получателя", "код", "вид оп", "срок плат", "очер", "рез.поле"))), None)
    if recipient_label is not None and purpose_label is not None and purpose_label > recipient_label:
        purpose_lines = [line for line in lines[recipient_label + 1:purpose_label] if line and not line.startswith("НДС")]
        purpose = " ".join(purpose_lines) or None
    amount = parse_amount(text)
    file_hash = digest_file(path)
    result = {
        "path": os.path.realpath(path), "file_hash": file_hash,
        "date": date_match.group(1).replace("/", ".").replace("-", ".") if date_match else None,
        "number": number_match.group(1) if number_match else None,
        "recipient": recipient,
        "purpose": purpose,
        "amount": amount,
    }
    result["missing"] = [field for field in ("date", "number", "recipient", "amount") if not result.get(field)]
    data = load()
    result["duplicate"] = any(op.get("source_hash") == file_hash for op in data["operations"] if op.get("type") == "payment")
    return result


def balances(data):
    allocated = {}
    for row in data["allocations"]:
        allocated[row["accrual_id"]] = allocated.get(row["accrual_id"], 0) + float(row["amount"])
    result = []
    for op in data["operations"]:
        if op["type"] != "accrual":
            continue
        paid = allocated.get(op["id"], 0)
        result.append({**op, "paid": paid, "balance": float(op["amount"]) - paid})
    return result


def date_key(value):
    try:
        return datetime.strptime(value or "", "%d.%m.%Y")
    except ValueError:
        return datetime.max


def reconcile_all(data):
    data["allocations"] = []
    accruals = sorted((op for op in data["operations"] if op["type"] == "accrual"), key=lambda op: (date_key(op.get("date")), op.get("created_at") or ""))
    remaining_debt = {op["id"]: float(op["amount"]) for op in accruals}
    payments = sorted((op for op in data["operations"] if op["type"] == "payment"), key=lambda op: (date_key(op.get("date")), op.get("created_at") or ""))
    for payment in payments:
        method = payment.get("payment_method") or "cashless"
        coefficient = float(payment.get("coefficient") if payment.get("coefficient") is not None else (0.94 if method == "cashless" else 1.0))
        gross = float(payment.get("gross_amount") if payment.get("gross_amount") is not None else payment["amount"])
        payment.update({"gross_amount": gross, "amount": gross, "payment_method": method, "coefficient": coefficient, "settlement_amount": round(gross * coefficient, 2)})
        remaining = payment["settlement_amount"]
        for accrual in accruals:
            debt = remaining_debt[accrual["id"]]
            amount = min(remaining, debt)
            if amount > 0.009:
                data["allocations"].append({"payment_id": payment["id"], "accrual_id": accrual["id"], "amount": amount})
                remaining_debt[accrual["id"]] -= amount
                remaining -= amount
            if remaining <= 0.009:
                break
        payment["unallocated_advance"] = round(max(0, remaining), 2)
    return data


@durable_json.transaction(lambda:LEDGER)
def record_payment(payment):
    data = load()
    for field in ('date','number','recipient','file_hash'):
        if not str(payment.get(field) or '').strip():raise ValueError('Не заполнено поле выплаты: '+field)
    datetime.strptime(str(payment['date']),'%d.%m.%Y')
    if payment.get('payment_method') not in ('cash','cashless'):
        raise ValueError('Выберите наличную или безналичную выплату')
    if payment.get("duplicate") or any(op.get("source_hash") == payment.get("file_hash") for op in data["operations"]):
        raise ValueError("эта платёжка уже учтена")
    if payment.get("missing"):
        raise ValueError("не заполнены поля: " + ", ".join(payment["missing"]))
    identity = f"{payment['number']}|{payment['date']}|{payment['amount']}|{payment['recipient']}"
    oid = operation_id("payment", identity)
    if any(op["id"] == oid for op in data["operations"]):
        raise ValueError("платёж с такими номером, датой, суммой и получателем уже учтён")
    method = payment.get("payment_method") or "cashless"
    coefficient = 0.94 if method == "cashless" else 1.0
    gross_amount = float(payment["amount"])
    if not math.isfinite(gross_amount) or gross_amount<=0: raise ValueError('Некорректная сумма выплаты')
    settlement_amount = round(gross_amount * coefficient, 2)
    op = {
        "id": oid, "type": "payment", "installer": INSTALLER,
        "display_name": (
            f"Наличные от {payment['date']} — {gross_amount:,.2f} ₽".replace(",", " ")
            if method == "cash" else f"КЛИВЭЛ №{payment['number']} от {payment['date']}"
        ),
        "date": payment["date"], "number": payment["number"],
        "recipient": payment["recipient"], "purpose": payment.get("purpose"),
        "amount": gross_amount, "gross_amount": gross_amount,
        "payment_method": method, "coefficient": coefficient,
        "settlement_amount": settlement_amount, "source_hash": payment["file_hash"],
        "source_file": os.path.basename(payment["path"]),
        "created_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
    }
    remaining = op["settlement_amount"]
    allocation_rows = []
    for accrual in sorted((row for row in balances(data) if row["balance"] > 0.009), key=lambda row: (date_key(row.get("date")), row["created_at"])):
        amount = min(remaining, accrual["balance"])
        if amount <= 0:
            break
        allocation = {"payment_id": oid, "accrual_id": accrual["id"], "amount": amount}
        data["allocations"].append(allocation)
        allocation_rows.append(allocation)
        remaining -= amount
    op["unallocated_advance"] = round(remaining, 2)
    data["operations"].append(op)
    save(data)
    return op, allocation_rows


def header(ws, values):
    for col, value in enumerate(values, 1):
        cell = ws.cell(1, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def current_view():
    from settlement_view import combined
    return combined(load(),INSTALLATIONS_DB)


def export_xlsx(data=None):
    if data is None:
        data = current_view()
    accruals = balances(data)
    payments = [op for op in data["operations"] if op["type"] == "payment"]
    wb = Workbook(); ws = wb.active; ws.title = "Сводка"
    rows = [
        ("Монтажник", INSTALLER),
        ("Начислено", sum(float(row["amount"]) for row in accruals)),
        ("Перечислено фактически", sum(float(row.get("gross_amount") or row["amount"]) for row in payments)),
        ("Зачтено во взаиморасчёты", sum(float(row.get("settlement_amount") or row["amount"]) for row in payments)),
        ("Аванс по общему балансу", max(0,sum(float(row.get('settlement_amount',row['amount'])) for row in payments)-sum(float(row['amount']) for row in accruals))),
        ("Текущая задолженность", max(0,sum(float(row['amount']) for row in accruals)-sum(float(row.get('settlement_amount',row['amount'])) for row in payments))),
    ]
    for row in rows: ws.append(row)
    ws.append(('Область учёта','Общие взаиморасчёты монтажника. Исторические выплаты сохранены по отчётам, без искусственного распределения по сделкам (Q-PAY-02).'))
    ws.append(('Сделки в отчёте','Перечень объектов не является распределением суммы выплаты по этим объектам.'))
    ws.append(('Новый кабинет','Новые принятые отчёты и выплаты включены чтением монтажной базы. Исторический реестр и распределения не переписываются. Неразнесённая по отчётам выплата уменьшает общий баланс, но не объявляет конкретный отчёт оплаченным.'))
    ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=22
    for cell in ws[ws.max_row]: cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor=YELLOW)

    accrual_names = {row["id"]: row.get("display_name") or f"Деревянко Сергей — отчёт от {row.get('date')}" for row in accruals}
    payment_names = {row["id"]: row.get("display_name") or f"КЛИВЭЛ №{row.get('number')} от {row.get('date')}" for row in payments}
    ws=wb.create_sheet("Отчёты"); header(ws,["Отчёт","Дата отчёта","Сделки","Начислено","Выплачено","Остаток","Статус","Исходный файл"])
    for row in accruals:
        status="Оплачен" if abs(row["balance"])<0.01 else "Частично оплачен" if row["paid"] else "Не оплачен"
        ws.append([accrual_names[row["id"]],row.get("date"),", ".join(row.get("deals") or []),row["amount"],row["paid"],row["balance"],status,row.get("source_file")])
    ws=wb.create_sheet("Выплаты"); header(ws,["Платёж","Дата","Номер","Получатель","Способ","Перечислено","Коэффициент","Зачтено","Назначение","Аванс","Файл"])
    for row in payments: ws.append([payment_names[row["id"]],row.get("date"),row.get("number"),row.get("recipient"),"Безналичная" if row.get("payment_method") == "cashless" else "Наличная",row.get("gross_amount") or row["amount"],row.get("coefficient"),row.get("settlement_amount") or row["amount"],row.get("purpose"),row.get("unallocated_advance"),row.get("source_file")])
    ws=wb.create_sheet("Распределение выплат"); header(ws,["Платёж","Отчёт","Сумма"])
    for row in data["allocations"]: ws.append([payment_names.get(row["payment_id"],row["payment_id"]),accrual_names.get(row["accrual_id"],row["accrual_id"]),row["amount"]])
    for sheet in wb.worksheets:
        sheet.freeze_panes="A2"
        for col in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(col)].width=max(14,min(45,max(len(str(sheet.cell(r,col).value or "")) for r in range(1,sheet.max_row+1))+2))
        for row in sheet.iter_rows():
            for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
    out=tempfile.mkdtemp(prefix="installer-settlements-")
    path=os.path.join(out,f"Взаиморасчёты - {INSTALLER}.xlsx")
    wb.save(path)
    return path


@durable_json.transaction(lambda:LEDGER)
def main():
    spec=json.load(sys.stdin); action=spec["action"]
    if action=="accrual":
        op,created=record_accrual(spec["session"])
        result={"ok":True,"operation":op,"created":created}
    elif action=="parse_payment": result={"ok":True,"payment":parse_payment(spec["path"])}
    elif action=="record_payment":
        op,allocations=record_payment(spec["payment"]); result={"ok":True,"operation":op,"allocations":allocations,"path":export_xlsx()}
    elif action=="reconcile":
        raise ValueError('Автоматическое перераспределение истории отключено. Нужен отдельный согласованный план миграции; используйте export для сверки без изменений.')
    elif action=="export":
        data=current_view()
        result={"ok":True,"path":export_xlsx(data),"balances":balances(data)}
    else: raise ValueError("неизвестное действие")
    print(json.dumps(result,ensure_ascii=False))


if __name__=="__main__":
    try: main()
    except Exception as exc:
        print(json.dumps({"ok":False,"error":str(exc)},ensure_ascii=False)); raise SystemExit(1)
