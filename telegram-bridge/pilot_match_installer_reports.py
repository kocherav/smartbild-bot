#!/usr/bin/env python3
"""Read-only pilot: map four installer reports to live Megaplan deal positions."""
import html
import io
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from difflib import SequenceMatcher
from pathlib import Path

import urllib.request

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))
sys.path.insert(0, "/home/kocher.av/megaplan-crm")
from build_installer_report_mockups import source_reports
from megaplan_api import MegaplanAPI

OUTPUT = BASE / "installer_report_mockups" / "Пилот — сопоставление отчётов со сделками.xlsx"
NAVY, WHITE, BLUE, YELLOW, RED, GREEN = "102733", "FFFFFF", "D9EAF7", "FFF2CC", "F4CCCC", "E2F0D9"
THIN = Side(style="thin", color="B7B7B7")

DEALS = {
    "Примавера": "862",
    "Дитейлинг — электрика": "861",
    "Дитейлинг — вентиляция": "858",
    "Дитейлинг — отопление": "857",
    "Дитейлинг — кондиционирование": "853",
    "Слава": "864",
}


def plain(value):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html.unescape(str(value or "")))).strip()


def norm(value):
    value = plain(value).casefold().replace("ё", "е")
    replacements = {
        "фреоновая трасса": "фреоновая магистраль",
        "фасонка": "фасонные части",
        "приточка": "вентиляционная установка",
        "опрессовка": "опресовка",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    return re.sub(r"[^a-zа-я0-9]+", " ", value).strip()


def score(a, b):
    na, nb = norm(a), norm(b)
    ta, tb = set(na.split()), set(nb.split())
    token = len(ta & tb) / max(1, len(ta | tb))
    sequence = SequenceMatcher(None, na, nb).ratio()
    return max(token, sequence * 0.78)


def deal_for(report, obj, source):
    text = f"{obj} {source}".casefold().replace("ё", "е")
    if "примавер" in text:
        return "862", "Примавера"
    if "слава" in text:
        return "864", "Слава"
    if "дитейлинг" in text and "электрик" in text:
        return "861", "Дитейлинг — электрика"
    if "дитейлинг" in text and ("вент" in text or "воздуховод" in text):
        return "858", "Дитейлинг — вентиляция"
    if report == "13.07" and "отоплен" in text:
        return "857", "Дитейлинг — отопление"
    if report == "13.07" and "кондиционирован" in text:
        return "853", "Дитейлинг — кондиционирование"
    return None, None


def load_deals():
    def one(item):
        label, deal_id = item
        api = MegaplanAPI()
        deal = api._call("GET", f"/deal/{deal_id}")["data"]
        files = api._call("GET", f"/deal/{deal_id}/allFiles", params={"limit": 100}).get("data") or []
        statements = [f for f in files if "ведомост" in (f.get("name") or "").casefold() and (f.get("name") or "").casefold().endswith((".xlsx", ".xls"))]
        statements.sort(key=lambda f: (((f.get("timeCreated") or {}).get("value") or ""), int(f.get("id") or 0)), reverse=True)
        statement = statements[0] if statements else None
        statement_rows = []
        if statement and statement.get("path"):
            remote = statement["path"]
            url = remote if remote.startswith("http") else f"https://{api.domain}{remote}"
            req = urllib.request.Request(url, headers={"Authorization": f"Bearer {api.token}"})
            import megaplan_transport
            with megaplan_transport.urlopen(req, timeout=60) as response:
                book = load_workbook(io.BytesIO(response.read()), data_only=False)
            sheet = book.active
            for row in range(1, sheet.max_row + 1):
                number, name = sheet.cell(row, 1).value, sheet.cell(row, 2).value
                if not name or not str(number or "").strip().isdigit():
                    continue
                qty_raw, unit, rate_raw = sheet.cell(row, 4).value, sheet.cell(row, 3).value, sheet.cell(row, 24).value
                try:
                    qty = float(str(qty_raw).replace(",", "."))
                except (TypeError, ValueError):
                    qty = None
                cleaned = re.sub(r"[^0-9,.-]", "", str(rate_raw or "").replace("\xa0", "")).replace(",", ".")
                try:
                    rate = float(cleaned)
                except ValueError:
                    rate = None
                statement_rows.append({"number": int(number), "name": plain(name), "qty": qty, "unit": unit, "rate": rate, "total": qty * rate if qty is not None and rate is not None else None})
        return label, {
            "id": deal_id,
            "client": (deal.get("contractor") or {}).get("name") or "",
            "essence": plain(deal.get("description")),
            "positions": api.ledger_from_deal(deal, price_source="megaplan_live"),
            "statement_name": statement.get("name") if statement else None,
            "statement_id": str(statement.get("id")) if statement else None,
            "statement_rows": statement_rows,
        }
    with ThreadPoolExecutor(max_workers=6) as pool:
        return dict(pool.map(one, DEALS.items()))


def suggest_position(line_type, source, positions):
    candidates = sorted(((score(source, p["name"]), p) for p in positions), key=lambda x: x[0], reverse=True)
    best_score, best = candidates[0]
    low = norm(source)
    if "по таблице" in low:
        aggregate = next((p for p in positions if norm(p["name"]) == "монтажные работы"), None)
        return aggregate, 0.45, "Нужна исходная таблица"
    if "работ" in line_type.casefold() and best_score < 0.42:
        aggregate = next((p for p in positions if norm(p["name"]) == "монтажные работы"), None)
        if aggregate:
            return aggregate, 0.40, "Предложен агрегат; подтвердить"
    if best_score >= 0.62:
        return best, best_score, "Высокое совпадение"
    if best_score >= 0.38:
        return best, best_score, "Вероятное совпадение"
    return None, best_score, "Вне плана / выбрать вручную"


def header(ws, row, values):
    for col, value in enumerate(values, 1):
        c = ws.cell(row, col, value)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(color=WHITE, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def body_style(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
            if cell.column in (7, 8, 12, 13):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'


def main():
    deals = load_deals()
    wb = Workbook()
    ws = wb.active
    ws.title = "Привязка объектов"
    cols = ["Объект отчёта", "Сделка", "Клиент", "Суть сделки", "Позиций", "Решение"]
    header(ws, 1, cols)
    for label, deal_id in DEALS.items():
        d = deals[label]
        ws.append([label, int(deal_id), d["client"], d["essence"], len(d["positions"]), "Предлагается подтвердить"])
    for i, width in enumerate([32, 12, 34, 44, 12, 26], 1): ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    body_style(ws)

    ws = wb.create_sheet("Сопоставление строк")
    cols = ["Отчёт", "Объект/блок", "Исходная строка", "Тип", "Кол-во", "Ед.", "Ставка", "Сумма", "Сделка", "Суть сделки", "Предложенная позиция", "План, кол-во", "Цена продажи", "Уверенность", "Решение"]
    header(ws, 1, cols)
    questions = []
    report_records = {report: [] for report in source_reports}
    for report, lines in source_reports.items():
        previous_deal = previous_label = None
        for obj, source, line_type, qty, unit, rate, amount, sign in lines:
            if sign == 0 or line_type in ("Заголовок", "Итог из PDF", "Итог начислений из PDF", "Подытог из PDF"):
                report_records[report].append([report, obj, source, line_type, qty, unit, rate, None, None, None, None, None, None, None, "Контрольная строка PDF — повторно не считается"])
                continue
            if sign < 0 or line_type in ("Аванс", "Удержание", "Премия", "Зачёт оборудования"):
                signed_amount = -amount if sign < 0 and amount is not None else amount
                report_records[report].append([report, obj, source, line_type, qty, unit, rate, signed_amount, None, None, None, None, None, None, "Взаиморасчёты — подтвердить"])
                questions.append([report, obj, source, "Операция взаиморасчётов", "Подтвердить тип, сумму и период; не относить к себестоимости позиции автоматически"])
                continue
            deal_id, label = deal_for(report, obj, source)
            inherited = False
            if not deal_id and line_type == "Материал" and previous_deal:
                deal_id, label, inherited = previous_deal, previous_label, True
            if not deal_id:
                record = [report, obj, source, line_type, qty, unit, rate, amount, None, None, None, None, None, 0, "Нужно выбрать сделку"]
                ws.append(record)
                report_records[report].append(record)
                questions.append([report, obj, source, "Не найдена сделка", "Выбрать сделку из пяти последних сделок исполнителя или найти по адресу"])
                continue
            previous_deal, previous_label = deal_id, label
            d = deals[label]
            position, confidence, decision = suggest_position(line_type, source, d["positions"])
            if inherited:
                confidence, decision = 0.95, "Объект унаследован по порядку отчёта"
            if "по таблице" in norm(source):
                decision = "Ведомость загружена; выбрать закрытые строки" if d["statement_rows"] else "Монтажная ведомость не найдена"
                confidence = 0.90 if d["statement_rows"] else 0.20
            record = [
                report, obj, source, line_type, qty, unit, rate, amount, int(deal_id), d["essence"],
                position.get("name") if position else None,
                position.get("qty") if position else None,
                position.get("sell_price") if position else None,
                confidence, decision,
            ]
            ws.append(record)
            report_records[report].append(record)
            if decision != "Высокое совпадение":
                prompt = f"Сделка №{deal_id}: подтвердить позицию «{position.get('name') if position else 'вне плана'}»"
                if "Ведомость загружена" in decision:
                    prompt = f"Сделка №{deal_id}: отметить в раскрытом составе ведомости фактически выполненные строки и количества; сумма отчёта {amount:,.2f} ₽"
                questions.append([report, obj, source, decision, prompt])
    for i, width in enumerate([11, 24, 54, 20, 10, 10, 14, 16, 11, 34, 48, 14, 16, 14, 30], 1): ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{ws.max_row}"
    for row in range(2, ws.max_row + 1):
        decision = ws.cell(row, 15).value
        fill = GREEN if decision == "Высокое совпадение" else RED if decision == "Нужно выбрать сделку" else YELLOW
        ws.cell(row, 15).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row, 14).number_format = "0%"
    body_style(ws)

    # One sheet per source report. Original report lines stay grouped by object;
    # "по таблице" rows are immediately followed by the live statement contents.
    for report, records in report_records.items():
        rs = wb.create_sheet(f"Отчёт {report}")
        rs.sheet_view.showGridLines = False
        rs.merge_cells("A1:Q1")
        rs["A1"] = f"ОТЧЁТ ОТ {report} — ПИЛОТНАЯ РАЗНОСКА"
        rs["A1"].fill = PatternFill("solid", fgColor=NAVY)
        rs["A1"].font = Font(color=WHITE, bold=True, size=14)
        rs["A1"].alignment = Alignment(horizontal="center")
        row = 3
        last_object = None
        object_start = None
        block_total_rows = []
        for rec in records:
            obj = rec[1]
            if obj != last_object:
                if last_object is not None:
                    rs.cell(row, 2, f"Итого по блоку «{last_object}»")
                    rs.cell(row, 8, f"=SUM(H{object_start}:H{row-1})")
                    rs.cell(row, 2).font = rs.cell(row, 8).font = Font(bold=True)
                    block_total_rows.append(row)
                    row += 2
                rs.merge_cells(start_row=row, start_column=1, end_row=row, end_column=17)
                rs.cell(row, 1, f"ОБЪЕКТ / БЛОК: {obj}")
                rs.cell(row, 1).fill = PatternFill("solid", fgColor="375A6C")
                rs.cell(row, 1).font = Font(color=WHITE, bold=True)
                row += 1
                header(rs, row, ["№", "Объект", "Строка отчёта", "Тип", "Кол-во", "Ед.", "Ставка", "Сумма", "Сделка", "Суть сделки", "Позиция сделки", "Уверенность", "Решение", "Источник", "План по ведомости", "Ед.", "Ставка ведомости"])
                row += 1
                object_start = row
                last_object = obj
            rs.append([None] * 17)
            for col, value in enumerate([row - object_start + 1, obj, rec[2], rec[3], rec[4], rec[5], rec[6], rec[7], rec[8], rec[9], rec[10], rec[13], rec[14], "PDF отчёт", None, None, None], 1):
                rs.cell(row, col, value)
            rs.cell(row, 12).number_format = "0%"
            if rec[8] and "по таблице" in norm(rec[2]):
                label = next((key for key, did in DEALS.items() if str(did) == str(rec[8])), None)
                d = deals.get(label) if label else None
                for statement_line in (d or {}).get("statement_rows", []):
                    row += 1
                    rs.append([None] * 17)
                    rs.cell(row, 3, "↳ " + statement_line["name"])
                    rs.cell(row, 4, "Строка монтажной ведомости")
                    rs.cell(row, 9, rec[8])
                    rs.cell(row, 13, "Кандидат: отметить выполненный объём")
                    rs.cell(row, 14, f"Мегаплан, файл №{d['statement_id']}")
                    rs.cell(row, 15, statement_line["qty"])
                    rs.cell(row, 16, statement_line["unit"])
                    rs.cell(row, 17, statement_line["rate"])
                    for col in range(1, 18):
                        rs.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
            row += 1
        if last_object is not None:
            rs.cell(row, 2, f"Итого по блоку «{last_object}»")
            rs.cell(row, 8, f"=SUM(H{object_start}:H{row-1})")
            rs.cell(row, 2).font = rs.cell(row, 8).font = Font(bold=True)
            block_total_rows.append(row)
            row += 2
        rs.cell(row, 2, "ИТОГО НАЧИСЛЕНИЙ ПО ОТЧЁТУ")
        # Sum block totals so detail rows and their subtotals are never counted twice.
        rs.cell(row, 8, "=" + "+".join(f"H{r}" for r in block_total_rows))
        rs.cell(row, 2).font = rs.cell(row, 8).font = Font(bold=True, size=12)
        stated_totals = [amount for _obj, _source, kind, _qty, _unit, _rate, amount, _sign in source_reports[report] if kind == "Итог из PDF" and amount is not None]
        if stated_totals:
            net_row = row
            row += 1
            rs.cell(row, 2, "КОНТРОЛЬНЫЙ ИТОГ ИЗ PDF")
            rs.cell(row, 8, stated_totals[-1])
            row += 1
            rs.cell(row, 2, "РАСХОЖДЕНИЕ С ПОСТРОЧНЫМ РАСЧЁТОМ")
            rs.cell(row, 8, f"=H{net_row}-H{row-1}")
            rs.cell(row, 2).font = rs.cell(row, 8).font = Font(bold=True, color="C00000")
        for col, width in enumerate([7, 24, 55, 25, 12, 10, 14, 16, 11, 35, 48, 14, 36, 28, 16, 10, 18], 1):
            rs.column_dimensions[get_column_letter(col)].width = width
        rs.freeze_panes = "A5"
        body_style(rs)

    ws = wb.create_sheet("Вопросы")
    header(ws, 1, ["Отчёт", "Объект", "Исходная строка", "Причина", "Что нужно подтвердить"])
    for row in questions: ws.append(row)
    for i, width in enumerate([12, 24, 52, 28, 58], 1): ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{ws.max_row}"
    body_style(ws)

    ws = wb.create_sheet("Позиции сделок")
    header(ws, 1, ["Сделка", "Объект", "Клиент", "Суть сделки", "Позиция", "Кол-во", "Цена продажи"])
    for label, deal_id in DEALS.items():
        d = deals[label]
        for p in d["positions"]:
            ws.append([int(deal_id), label, d["client"], d["essence"], p["name"], p["qty"], p["sell_price"]])
    for i, width in enumerate([12, 30, 34, 40, 65, 14, 18], 1): ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    body_style(ws)

    wb.save(OUTPUT)
    print(json.dumps({"ok": True, "path": str(OUTPUT), "matches": wb["Сопоставление строк"].max_row - 1, "questions": len(questions)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
