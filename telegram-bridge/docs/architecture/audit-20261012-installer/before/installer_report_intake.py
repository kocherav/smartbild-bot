#!/usr/bin/env python3
"""Build a compact, read-only review session for an installer PDF report.

The first production slice intentionally supports the four agreed pilot report
formats.  It produces JSON for the Telegram state machine; it never writes to
Megaplan.  Unknown layouts fail explicitly instead of guessing financial data.
"""
import json
import itertools
import os
import re
import subprocess
import sys
from datetime import datetime

from openpyxl import load_workbook

from build_installer_report_mockups import source_reports

FACT_REGISTRY = os.path.join(os.path.dirname(os.path.abspath(__file__)), "installer_fact_registry.json")


REPORT_MARKERS = {
    "20.08": ("Монтаж кондиционера Илья", "Слава по таблице"),
    "13.07": ("По таблице отопление", "Касета-99.000"),
    "10.08": ("Дитейлинг вент", "Примавера материал"),
    "11.09": ("Такелаж +доставка увлажнителя", "Итого общий:181.040"),
}

OBJECT_DEALS = {
    "Дитейлинг вентиляция": ("858", "Вентиляция 1эт НК", 1.0),
    "Дитейлинг электрика": ("861", "Электрика 1 эт НК", 1.0),
    "Дитейлинг": ("861", "Электрика 1 эт НК", 0.9),
    "Отопление": ("857", "Отопление 1эт НК", 1.0),
    "Кондиционирование": ("853", "Кондиционирование 1эт НК", 1.0),
    "Примавера — материалы": ("862", "ЖК Примавера ОВиК для Александра", 1.0),
    "Примавера — работы": ("862", "ЖК Примавера ОВиК для Александра", 1.0),
    "Примавера": ("862", "ЖК Примавера ОВиК для Александра", 1.0),
    "Слава": ("864", "ЖК Слава", 1.0),
}

CONTROL_TYPES = {"Заголовок", "Итог из PDF", "Итог начислений из PDF", "Подытог из PDF"}
SETTLEMENT_TYPES = {"Аванс", "Удержание", "Премия", "Зачёт оборудования"}

DOMAIN_WORDS = {
    "Отопление": ("котел", "насос", "коллектор", "гидравличес", "радиатор", "теплый пол", "отоплен"),
    "Кондиционирование": ("кондиционер", "daikin", "фреон", "рефнет", "внутренний блок", "наружный блок"),
    "Дитейлинг": ("кабель", "автомат", "щит", "светильник", "электр", "провод", "dali"),
    "Слава": ("воздуховод", "теплоизоляц", "диффузор", "вентиляц", "komfovent"),
}


def pdf_text(path):
    result = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, timeout=30)
    if result.returncode:
        raise ValueError("не удалось прочитать PDF")
    return result.stdout


def identify_report(text):
    normalized = text.replace("ё", "е")
    for report, markers in REPORT_MARKERS.items():
        if all(marker.replace("ё", "е") in normalized for marker in markers):
            return report
    raise ValueError("формат отчёта пока не распознан; нужен ручной разбор")


def mapped_object(obj, last_mapping):
    if obj in OBJECT_DEALS:
        return OBJECT_DEALS[obj], False
    if obj == "Материалы" and last_mapping:
        deal_id, essence, _confidence = last_mapping
        return (deal_id, essence, 0.95), True
    return None, False


def number(value):
    try:
        return float(str(value).replace("\xa0", "").replace(" руб.", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def table_profile(path):
    book = load_workbook(path, data_only=False, read_only=False)
    rows = []
    text_parts = [os.path.basename(path)]
    completion_columns = []
    for sheet in book.worksheets:
        text_parts.append(sheet.title)
        for row in range(1, sheet.max_row + 1):
            values = [sheet.cell(row, col).value for col in range(1, min(sheet.max_column, 24) + 1)]
            text_parts.extend(str(v) for v in values if v not in (None, ""))
            ordinal, name = sheet.cell(row, 1).value, sheet.cell(row, 2).value
            if not name or not str(ordinal or "").strip().isdigit():
                continue
            qty = number(sheet.cell(row, 4).value)
            unit = sheet.cell(row, 3).value
            # The canonical table is A:F: quantity is D, unit rate is E and
            # planned total is F. Columns to the right are reporting-period
            # execution amounts or helper data and must never override the rate.
            rate = number(sheet.cell(row, 5).value)
            if rate is None:
                rate = number(sheet.cell(row, 24).value)
            rows.append({
                "number": int(ordinal), "name": re.sub(r"\s+", " ", str(name)).strip(),
                "qty": qty, "unit": str(unit or ""), "rate": rate,
                "total": qty * rate if qty is not None and rate is not None else None,
                "sheet_row": row,
                "extra_values": {str(col): number(sheet.cell(row, col).value) for col in range(7, sheet.max_column + 1) if number(sheet.cell(row, col).value) is not None},
            })
        total_rows = [row for row in range(1, sheet.max_row + 1) if "итого" in str(sheet.cell(row, 1).value or "").casefold()]
        for col in range(7, sheet.max_column + 1):
            values = [item["extra_values"].get(str(col)) for item in rows if item["extra_values"].get(str(col)) is not None]
            footer = next((number(sheet.cell(row, col).value) for row in total_rows if number(sheet.cell(row, col).value) is not None), None)
            if values or footer is not None:
                completion_columns.append({"column": col, "sum": sum(values), "footer": footer})
    return {
        "path": os.path.realpath(path), "filename": os.path.basename(path),
        "text": " ".join(text_parts).casefold().replace("ё", "е"), "rows": rows,
        "calculated_total": sum(row["total"] or 0 for row in rows),
        "completion_columns": completion_columns,
    }


def table_score(exception, profile):
    obj = exception["object"]
    text = profile["text"]
    score = 0.0
    if obj.casefold().replace("ё", "е") in text:
        score += 5
    deal_id = str(exception.get("deal_id") or "")
    if deal_id and deal_id in profile["filename"]:
        score += 6
    for key, words in DOMAIN_WORDS.items():
        if key.casefold() in obj.casefold():
            score += sum(1 for word in words if word in text)
    return score


def attach_tables(exceptions, table_paths):
    targets = [item for item in exceptions if item.get("type") == "table"]
    profiles = [table_profile(path) for path in table_paths]
    try:
        registry = json.load(open(FACT_REGISTRY, encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        registry = {}
    if not targets or not profiles:
        return
    best = None
    for permutation in itertools.permutations(range(len(profiles)), min(len(targets), len(profiles))):
        total = sum(table_score(targets[i], profiles[p]) for i, p in enumerate(permutation))
        if best is None or total > best[0]:
            best = (total, permutation)
    for index, profile_index in enumerate(best[1] if best else ()):
        target, profile = targets[index], profiles[profile_index]
        raw_score = table_score(target, profile)
        report_amount = float(target.get("amount") or 0)
        candidates = sorted(
            profile.get("completion_columns") or [],
            key=lambda item: min(abs((item.get("footer") or 0) - report_amount), abs((item.get("sum") or 0) - report_amount)),
        )
        completion = candidates[0] if candidates else None
        completion_col = str(completion["column"]) if completion else None
        detailed_rows = []
        for row in profile["rows"]:
            current_amount = row.get("extra_values", {}).get(completion_col) if completion_col else None
            current_qty = current_amount / row["rate"] if current_amount is not None and row.get("rate") not in (None, 0) else 0
            plan_qty = row.get("qty") or 0
            key = re.sub(r"[^a-zа-я0-9]+", " ", str(row.get("name") or "").casefold().replace("ё", "е")).strip()
            previous_qty = float(((registry.get(str(target.get("deal_id"))) or {}).get(key) or {}).get("qty") or 0)
            cumulative_qty = previous_qty + current_qty
            detailed_rows.append({
                "number": row.get("number"), "name": row.get("name"), "unit": row.get("unit"),
                "plan_qty": plan_qty, "rate": row.get("rate"), "plan_amount": row.get("total"),
                "previous_qty": previous_qty, "current_qty": current_qty,
                "current_amount": current_amount or 0, "cumulative_qty": cumulative_qty,
                "remaining_qty": max(0, plan_qty - cumulative_qty),
            })
        target.update({
            "table_file": profile["filename"], "table_path": profile["path"],
            "table_rows": detailed_rows, "table_rows_count": len(detailed_rows),
            "table_calculated_total": profile["calculated_total"],
            "completion_column": completion.get("column") if completion else None,
            "completion_sum": completion.get("sum") if completion else None,
            "completion_footer": completion.get("footer") if completion else None,
            "completion_matches_report": bool(completion and (abs((completion.get("footer") or completion.get("sum") or 0) - report_amount) < 0.01)),
            "table_match_confidence": 0.95 if raw_score >= 4 else 0.8 if raw_score >= 1 else 0.5,
            "reason": "таблица получена; нужно подтвердить выполненные строки или учёт одной суммой",
        })


def date_from_filename(path):
    name = os.path.basename(path)
    match = re.search(r"(?<!\d)([0-3]?\d)[._-]([01]?\d)[._-](20\d{2})(?!\d)", name)
    if match:
        return f"{int(match.group(1)):02d}.{int(match.group(2)):02d}.{match.group(3)}"
    match = re.search(r"(?<!\d)(\d{2})([01]\d)([0-3]\d)(?!\d)", name)
    if match:
        return f"{match.group(3)}.{match.group(2)}.20{match.group(1)}"
    return None


def build_excel_session(path):
    profile = table_profile(path)
    candidates = profile.get("completion_columns") or []
    if not candidates:
        raise ValueError("в таблице не найден отдельный столбец текущего выполнения")
    completion = sorted(
        candidates,
        # In installer workbooks each new reporting period is added as the next
        # execution column to the right. Prefer the rightmost column that has its
        # own footer; do not confuse the larger cumulative/plan columns with the
        # current report merely because their sum is greater.
        key=lambda item: (item.get("footer") is not None, item.get("column") or 0),
        reverse=True,
    )[0]
    completion_col = str(completion["column"])
    detailed_rows = []
    report_lines = []
    for row in profile["rows"]:
        current_amount = row.get("extra_values", {}).get(completion_col)
        current_qty = current_amount / row["rate"] if current_amount is not None and row.get("rate") not in (None, 0) else 0
        plan_qty = row.get("qty") or 0
        detail = {
            "number": row.get("number"), "name": row.get("name"), "unit": row.get("unit"),
            "plan_qty": plan_qty, "rate": row.get("rate"), "plan_amount": row.get("total"),
            "previous_qty": 0, "current_qty": current_qty, "current_amount": current_amount or 0,
            "cumulative_qty": current_qty, "remaining_qty": max(0, plan_qty - current_qty),
        }
        detailed_rows.append(detail)
        if current_amount not in (None, 0):
            report_lines.append({
                "number": len(report_lines) + 1, "object": "Объект из Excel", "source": row.get("name"),
                "kind": "Работа по таблице", "qty": current_qty, "unit": row.get("unit"),
                "rate": row.get("rate"), "amount": current_amount, "sign": 1,
                "deal_id": None, "deal_essence": None, "mapping_confidence": 0.0,
                "inherited_object": False,
            })
    total = completion.get("footer") if completion.get("footer") is not None else sum(row["current_amount"] for row in detailed_rows)
    if not report_lines and not total:
        raise ValueError("в столбце текущего выполнения нет заполненных сумм")
    source = f"Выполнение по таблице — {total:.2f}"
    exception = {
        "id": 1, "object": "Объект из Excel", "source": source, "amount": total,
        "type": "table", "reason": "Excel загружен как основной отчёт; нужно подтвердить выполнение",
        "table_file": profile["filename"], "table_path": profile["path"],
        "table_rows": detailed_rows, "table_rows_count": len(detailed_rows),
        "table_calculated_total": profile["calculated_total"], "completion_column": completion["column"],
        "completion_sum": completion.get("sum"), "completion_footer": completion.get("footer"),
        "completion_matches_report": True, "table_match_confidence": 1.0,
    }
    return {
        "version": 1, "source_kind": "excel", "source_path": os.path.realpath(path),
        "report": date_from_filename(path) or os.path.splitext(os.path.basename(path))[0],
        "report_date": date_from_filename(path),
        "objects": [{"name": "Объект из Excel", "deal_id": None, "essence": None, "confidence": 0.0, "inherited": False, "lines": len(report_lines), "amount": total}],
        "report_lines": report_lines, "automatic_count": len(report_lines), "automatic_sum": total,
        "exceptions": [exception], "required_tables": 0, "attached_tables": 1, "table_requests": [],
        "net_total": total, "declared_total": total, "discrepancy": 0, "payable_amount": total,
        "stage": "objects", "approved": False,
    }


def build_session(path, table_paths=None):
    if os.path.splitext(path)[1].casefold() in (".xlsx", ".xlsm"):
        return build_excel_session(path)
    report = identify_report(pdf_text(path))
    objects = []
    object_index = {}
    exceptions = []
    automatic_count = 0
    automatic_sum = 0.0
    net_total = 0.0
    last_mapping = None
    report_lines = []
    object_totals = {}
    declared_total = None
    subtotal_controls = []

    for number, (obj, source, kind, qty, unit, rate, amount, sign) in enumerate(source_reports[report], 1):
        mapping, inherited = mapped_object(obj, last_mapping)
        if mapping:
            last_mapping = mapping
        signed = (amount or 0) * sign
        if sign:
            net_total += signed
            object_totals[obj] = object_totals.get(obj, 0) + signed
        if kind == "Итог из PDF" and amount is not None:
            declared_total = float(amount)
        if kind == "Подытог из PDF" and amount is not None:
            subtotal_controls.append({"object": obj, "declared": float(amount)})
        if obj not in object_index and obj not in ("Контроль отчёта", "Взаиморасчёты"):
            deal_id, essence, confidence = mapping or (None, None, 0.0)
            object_index[obj] = len(objects)
            objects.append({
                "name": obj, "deal_id": deal_id, "essence": essence,
                "confidence": confidence, "inherited": inherited,
                "lines": 0, "amount": 0.0,
            })
        if obj in object_index and sign > 0:
            objects[object_index[obj]]["lines"] += 1
            objects[object_index[obj]]["amount"] += amount or 0
        report_lines.append({
            "number": number, "object": obj, "source": source, "kind": kind,
            "qty": qty, "unit": unit, "rate": rate, "amount": amount,
            "sign": sign, "deal_id": mapping[0] if mapping else None,
            "deal_essence": mapping[1] if mapping else None,
            "mapping_confidence": mapping[2] if mapping else 0.0,
            "inherited_object": inherited,
        })

        if sign == 0 or kind in CONTROL_TYPES:
            continue
        if kind in SETTLEMENT_TYPES or sign < 0:
            exceptions.append({"id": len(exceptions) + 1, "object": obj, "source": source, "amount": signed, "type": "settlement", "reason": "операция взаиморасчётов требует подтверждения"})
            continue
        if not mapping:
            exceptions.append({"id": len(exceptions) + 1, "object": obj, "source": source, "amount": amount, "type": "deal", "reason": "не выбрана сделка"})
            continue
        if "по таблице" in source.casefold():
            exceptions.append({"id": len(exceptions) + 1, "object": obj, "source": source, "amount": amount, "deal_id": mapping[0], "type": "table", "reason": "нужно выбрать способ учёта ведомости"})
            continue
        if amount is None or "?" in source:
            exceptions.append({"id": len(exceptions) + 1, "object": obj, "source": source, "amount": amount, "deal_id": mapping[0], "type": "amount", "reason": "не указана сумма"})
            continue
        if kind == "Материал" and not inherited and source.casefold().startswith("материал-"):
            exceptions.append({"id": len(exceptions) + 1, "object": obj, "source": source, "amount": amount, "deal_id": mapping[0], "type": "material", "reason": "материал не расшифрован"})
            continue
        automatic_count += 1
        automatic_sum += amount or 0

    attach_tables(exceptions, table_paths or [])
    discrepancy = (declared_total - net_total) if declared_total is not None else 0
    probable = []
    for control in subtotal_controls:
        calculated = object_totals.get(control["object"], 0)
        delta = control["declared"] - calculated
        if abs(delta) >= 0.01:
            probable.append({"object": control["object"], "declared": control["declared"], "calculated": calculated, "difference": delta})
    if declared_total is not None and abs(discrepancy) >= 0.01:
        exceptions.append({
            "id": len(exceptions) + 1, "object": "Контроль отчёта",
            "source": f"Итог монтажника {declared_total:.2f}; построчный итог {net_total:.2f}",
            "amount": discrepancy, "type": "discrepancy",
            "reason": "итог в отчёте не совпадает с суммой всех распознанных строк",
            "declared_total": declared_total, "calculated_total": net_total,
            "probable_causes": probable,
        })
    table_exceptions = [item for item in exceptions if item.get("type") == "table"]
    return {
        "version": 1,
        "source_path": os.path.realpath(path),
        "report": report,
        "objects": objects,
        "report_lines": report_lines,
        "automatic_count": automatic_count,
        "automatic_sum": automatic_sum,
        "exceptions": exceptions,
        "required_tables": len(table_exceptions),
        "attached_tables": len([item for item in table_exceptions if item.get("table_file")]),
        "table_requests": [{"object": item["object"], "source": item["source"], "deal_id": item.get("deal_id")} for item in table_exceptions],
        "net_total": net_total,
        "declared_total": declared_total,
        "discrepancy": discrepancy,
        "payable_amount": net_total if abs(discrepancy) < 0.01 else None,
        "report_date": ({"20.08": "20.08.2026", "13.07": "13.07.2026", "10.08": "10.08.2026", "11.09": "11.09.2026"}.get(report) or report),
        "stage": "objects",
        "approved": False,
    }


def main():
    try:
        path = sys.argv[1]
        print(json.dumps({"ok": True, "session": build_session(path, sys.argv[2:])}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
