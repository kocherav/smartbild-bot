"""Build a readable, self-contained DOCX from the architecture source documents."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
SOURCES = ['README.md', 'PLAN.md', 'SCENARIOS.md', 'SCENARIO_RULES.md', 'DECISIONS.md', 'TEMPLATES.md', 'TEST_MATRIX.md', 'RECOVERY.md', 'HISTORICAL_SETTLEMENT_RECONCILIATION.md', 'CHANGELOG.md']

def plain(text):
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text.replace('**', '').replace('`', '')

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)
doc.add_heading('SmartBild — план стабилизации и документация', 0)
doc.add_paragraph('Единый комплект. Локальные ссылки заменены названиями; для чтения этого комплекта переходить по ним не требуется.')
doc.add_paragraph('Актуальный статус этапов и границы выполненного приведены в PLAN. Результаты проверок и включений — в CHANGELOG; восстановление — в RECOVERY. Промежуточный документ не означает завершения всего плана.')
for filename in SOURCES:
    doc.add_page_break()
    table = None
    for line in (ROOT / filename).read_text(encoding='utf-8').splitlines():
        if line.startswith('|'):
            cells = [plain(x.strip()) for x in line.strip('|').split('|')]
            if all(re.fullmatch(r'[:\- ]+', x) for x in cells):
                continue
            if table is None:
                table = doc.add_table(rows=0, cols=len(cells))
                table.style = 'Table Grid'
            row = table.add_row()
            for cell, text in zip(row.cells, cells):
                cell.text = text
            continue
        table = None
        if not line.strip():
            continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            doc.add_heading(plain(line.lstrip('#').strip()), min(level, 3))
        else:
            doc.add_paragraph(plain(line))
output = ROOT / 'SmartBild — план и документация.docx'
doc.save(output)
check = Document(output)
assert len(check.tables) >= 3
assert any('A1' in p.text for p in check.paragraphs)
print(f'Created and reopened DOCX: {output.name}; tables={len(check.tables)}; bytes={output.stat().st_size}')
