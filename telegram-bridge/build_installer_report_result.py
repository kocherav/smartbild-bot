#!/usr/bin/env python3
"""Create the final approved installer-report workbook from a review session."""
import json
import os
import re
import tempfile
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY, WHITE, BLUE, GREEN, YELLOW, RED = "102733", "FFFFFF", "D9EAF7", "E2F0D9", "FFF2CC", "F4CCCC"


def header(ws, row, values):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, values):
    for col, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def style(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column in (6, 7, 8, 9):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'


def safe_sheet(value, used):
    base = re.sub(r"[\\/*?:\[\]]", "_", value)[:31] or "Таблица"
    name, number = base, 2
    while name in used:
        suffix = f" {number}"
        name = base[:31-len(suffix)] + suffix
        number += 1
    used.add(name)
    return name


def deal_sheet_name(item, used):
    deal_id = str(item.get("deal_id") or "без сделки")
    client = "JetCar" if deal_id in {"853", "857", "858", "861"} else (item.get("client_short") or item.get("object") or "Объект")
    direction = item.get("object") or "Работы"
    return safe_sheet(f"{deal_id} {client} {direction}", used)


def add_deal_sheet(wb, item, name):
    ws = wb.create_sheet(name)
    ws["A1"] = f"Сделка №{item.get('deal_id')} — {item.get('object')} — отчёт от {item.get('report_date')}"
    ws.merge_cells("A1:K1")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True)
    ws["A2"] = f"Файл-основание: {item.get('table_file')} · решение: {item.get('resolution')}"
    ws.merge_cells("A2:K2")
    ws["A3"] = "Цвета: жёлтый — начато; зелёный — план закрыт; красный — превышение плана."
    ws.merge_cells("A3:K3")
    header(ws, 4, ["№", "Позиция монтажной ведомости", "Ед.", "План, кол-во", "Расценка", "Плановая сумма", "Закрыто ранее", "Закрыто в этом отчёте", "Сумма этого отчёта", "Накопительный факт", "Остаток"])
    for row in item.get("table_rows") or []:
        ws.append([row.get("number"), row.get("name"), row.get("unit"), row.get("plan_qty"), row.get("rate"), row.get("plan_amount"), row.get("previous_qty"), row.get("current_qty"), row.get("current_amount"), row.get("cumulative_qty"), row.get("remaining_qty")])
        cumulative, plan = row.get("cumulative_qty") or 0, row.get("plan_qty") or 0
        fill = RED if cumulative > plan + 1e-9 else GREEN if cumulative > 0 and abs(cumulative - plan) <= 1e-9 else YELLOW if cumulative > 0 else None
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
    ws.append([None, "ИТОГО ТЕКУЩЕГО ОТЧЁТА", None, None, None, None, None, None, f"=SUM(I5:I{ws.max_row})"])
    ws.cell(ws.max_row, 2).font = ws.cell(ws.max_row, 9).font = Font(bold=True)
    widths(ws, [8, 68, 10, 15, 16, 18, 16, 20, 20, 20, 16])
    ws.freeze_panes = "A5"
    style(ws)
    return ws


def add_execution_sheet(wb, deal_id, rows, table_items, name, report_date):
    ws = wb.create_sheet(name)
    ws["A1"] = f"ВЫПОЛНЕНИЕ ПО СДЕЛКЕ №{deal_id} — ОТЧЁТ ОТ {report_date}"
    ws.merge_cells("A1:I1")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=14)
    header(ws, 3, ["№", "Объект", "Исходная строка", "Тип", "Кол-во", "Ед.", "Расценка", "Сумма", "Способ привязки"])
    for number, row in enumerate(rows, 1):
        ws.append([
            number, row.get("object"), row.get("source"), row.get("kind"),
            row.get("qty"), row.get("unit"), row.get("rate"), row.get("amount"),
            row.get("mapping_note") or ("автоматически" if row.get("mapping_confidence") else "по решению руководителя"),
        ])
    ws.append([None, None, "ИТОГО ПО СДЕЛКЕ В ЭТОМ ОТЧЁТЕ", None, None, None, None, f"=SUM(H4:H{ws.max_row})"])
    ws.cell(ws.max_row, 3).font = Font(bold=True)
    ws.cell(ws.max_row, 8).font = Font(bold=True)
    ws.cell(ws.max_row, 8).fill = PatternFill("solid", fgColor=GREEN)

    for item in table_items:
        start = ws.max_row + 3
        ws.cell(start, 1, f"Детализация по монтажной ведомости: {item.get('object')}")
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=11)
        ws.cell(start, 1).font = Font(bold=True, color=WHITE)
        ws.cell(start, 1).fill = PatternFill("solid", fgColor=NAVY)
        header(ws, start + 1, ["№", "Позиция монтажной ведомости", "Ед.", "План, кол-во", "Расценка", "Плановая сумма", "Закрыто ранее", "Закрыто сейчас", "Сумма отчёта", "Накопительный факт", "Остаток"])
        detail_start = start + 2
        for row in item.get("table_rows") or []:
            ws.append([row.get("number"), row.get("name"), row.get("unit"), row.get("plan_qty"), row.get("rate"), row.get("plan_amount"), row.get("previous_qty"), row.get("current_qty"), row.get("current_amount"), row.get("cumulative_qty"), row.get("remaining_qty")])
            cumulative, plan = row.get("cumulative_qty") or 0, row.get("plan_qty") or 0
            fill = RED if cumulative > plan + 1e-9 else GREEN if cumulative > 0 and abs(cumulative - plan) <= 1e-9 else YELLOW if cumulative > 0 else None
            if fill:
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill("solid", fgColor=fill)
        ws.append([None, "ИТОГО ПО ТАБЛИЦЕ", None, None, None, None, None, None, f"=SUM(I{detail_start}:I{ws.max_row})"])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        ws.cell(ws.max_row, 9).font = Font(bold=True)
    widths(ws, [8, 30, 62, 24, 12, 10, 16, 18, 24, 18, 16])
    ws.freeze_panes = "A4"
    style(ws)
    return ws


def build(session):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = f"ИТОГОВЫЙ РАЗБОР ОТЧЁТА МОНТАЖНИКА ОТ {session.get('report')}"
    ws.merge_cells("A1:F1")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=14)
    header(ws, 3, ["Объект", "Сделка", "Суть сделки", "Строк", "Сумма отчёта", "Привязка"])
    for obj in session.get("objects") or []:
        ws.append([obj.get("name"), obj.get("deal_id"), obj.get("essence"), obj.get("lines"), obj.get("amount"), "по контексту 95%" if obj.get("inherited") else "подтверждена"])
    ws.append([])
    lines = session.get("report_lines") or []
    settlement_kinds = {"Аванс", "Удержание", "Премия", "Зачёт оборудования"}
    controls = {"Заголовок", "Итог из PDF", "Итог начислений из PDF", "Подытог из PDF"}
    accruals = sum((line.get("amount") or 0) for line in lines if line.get("sign") == 1 and line.get("kind") not in settlement_kinds | controls)
    premiums = sum((line.get("amount") or 0) for line in lines if line.get("kind") == "Премия")
    deductions = sum((line.get("amount") or 0) for line in lines if line.get("sign") == -1)
    ws.append(["Начислено за работы и материалы", None, None, None, accruals])
    ws.append(["Премии", None, None, None, premiums])
    ws.append(["Авансы, удержания и зачёты", None, None, None, -deductions])
    ws.append(["Корректировка расхождения", None, None, None, session.get("adjustment") or 0])
    ws.append(["К оплате по данному отчету", None, None, None, session.get("payable_amount") if session.get("payable_amount") is not None else session.get("net_total")])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill("solid", fgColor=GREEN)
    widths(ws, [28, 13, 45, 12, 18, 24])
    style(ws)

    ws = wb.create_sheet("Строки отчёта")
    header(ws, 1, ["№", "Объект", "Исходная строка", "Тип", "Кол-во", "Ед.", "Ставка", "Сумма", "Знак", "Сделка", "Суть сделки", "Уверенность", "По контексту"])
    for line in session.get("report_lines") or []:
        ws.append([line.get("number"), line.get("object"), line.get("source"), line.get("kind"), line.get("qty"), line.get("unit"), line.get("rate"), line.get("amount"), line.get("sign"), line.get("deal_id"), line.get("deal_essence"), line.get("mapping_confidence"), "да" if line.get("inherited_object") else "нет"])
    widths(ws, [7, 28, 58, 24, 12, 10, 15, 16, 9, 12, 42, 14, 16])
    ws.freeze_panes = "A2"
    style(ws)

    ws = wb.create_sheet("Принятые решения")
    header(ws, 1, ["№", "Объект", "Исходная строка", "Сумма", "Причина проверки", "Решение", "Сделка", "Таблица", "Строк в таблице", "Уверенность"])
    for item in session.get("exceptions") or []:
        ws.append([item.get("id"), item.get("object"), item.get("source"), item.get("amount"), item.get("reason"), item.get("resolution"), item.get("deal_id"), item.get("table_file"), item.get("table_rows_count"), item.get("table_match_confidence")])
        ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor=GREEN if item.get("resolution") and item.get("resolution") != "оставлено вопросом" else YELLOW)
    widths(ws, [7, 26, 55, 16, 45, 28, 12, 35, 16, 14])
    ws.freeze_panes = "A2"
    style(ws)

    miscellaneous = [
        item for item in session.get("exceptions") or []
        if item.get("bucket") == "Всякая всячина" or str(item.get("resolution") or "").startswith("всякая всячина")
    ]
    if miscellaneous:
        ws = wb.create_sheet("Всякая всячина")
        ws["A1"] = f"РАБОТЫ БЕЗ ПРИВЯЗКИ К СДЕЛКАМ — ОТЧЁТ ОТ {session.get('report_date') or session.get('report')}"
        ws.merge_cells("A1:F1")
        ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
        ws["A1"].font = Font(color=WHITE, bold=True, size=14)
        header(ws, 3, ["№", "Объект из отчёта", "Исходная строка", "Сумма", "Решение", "Комментарий"])
        for number, item in enumerate(miscellaneous, 1):
            ws.append([number, item.get("object"), item.get("source"), item.get("amount"), "Начислить без сделки", item.get("reason")])
        ws.append([None, None, "ИТОГО «ВСЯКАЯ ВСЯЧИНА»", f"=SUM(D4:D{ws.max_row})"])
        ws.cell(ws.max_row, 3).font = Font(bold=True)
        ws.cell(ws.max_row, 4).font = Font(bold=True)
        ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=GREEN)
        widths(ws, [8, 28, 58, 18, 26, 44])
        ws.freeze_panes = "A4"
        style(ws)

    used = set(wb.sheetnames)
    deal_files = []
    for item in session.get("exceptions") or []:
        if item.get("type") != "table" or not item.get("table_rows"):
            continue
        item["report_date"] = session.get("report_date") or session.get("report")
        name = deal_sheet_name(item, used)
        add_deal_sheet(wb, item, name)

    out_dir = tempfile.mkdtemp(prefix="installer-report-result-")
    path = os.path.join(out_dir, f"Итоговый разбор отчёта монтажника {session.get('report')}.xlsx")
    wb.save(path)
    rows_by_deal = {}
    for line in session.get("report_lines") or []:
        if line.get("deal_id"):
            rows_by_deal.setdefault(str(line["deal_id"]), []).append(dict(line))
    for item in session.get("exceptions") or []:
        deal_id = str(item.get("deal_id") or "")
        if not deal_id or item.get("type") in {"settlement", "discrepancy"}:
            continue
        rows = rows_by_deal.setdefault(deal_id, [])
        if not any(row.get("source") == item.get("source") and row.get("object") == item.get("object") for row in rows):
            rows.append({
                "object": item.get("object"), "source": item.get("source"),
                "kind": "Работа", "amount": item.get("amount"),
                "mapping_note": "по решению руководителя",
            })
    table_items_by_deal = {}
    for item in session.get("exceptions") or []:
        if item.get("deal_id") and item.get("type") == "table" and item.get("table_rows"):
            table_items_by_deal.setdefault(str(item["deal_id"]), []).append(item)
    object_names = {}
    for obj in session.get("objects") or []:
        if obj.get("deal_id"):
            object_names.setdefault(str(obj["deal_id"]), obj.get("name"))
    for deal_id in sorted(rows_by_deal, key=lambda value: int(value) if value.isdigit() else value):
        deal_book = Workbook()
        deal_book.remove(deal_book.active)
        descriptor = {"deal_id": deal_id, "object": object_names.get(deal_id) or (rows_by_deal[deal_id][0].get("object") if rows_by_deal[deal_id] else "Работы")}
        sheet_name = deal_sheet_name(descriptor, set())
        add_execution_sheet(deal_book, deal_id, rows_by_deal[deal_id], table_items_by_deal.get(deal_id, []), sheet_name, session.get("report_date") or session.get("report"))
        deal_dir = tempfile.mkdtemp(prefix=f"installer-report-deal-{deal_id}-")
        deal_path = os.path.join(deal_dir, f"Выполнение - {session.get('report_date') or session.get('report')}.xlsx")
        deal_book.save(deal_path)
        deal_files.append({"deal_id": deal_id, "path": deal_path, "filename": os.path.basename(deal_path), "sheet": sheet_name})
    return path, deal_files


if __name__ == "__main__":
    try:
        session = json.load(sys.stdin)
        path, deal_files = build(session)
        print(json.dumps({"ok": True, "path": path, "deal_files": deal_files}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        raise SystemExit(1)
