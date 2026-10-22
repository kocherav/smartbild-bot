#!/usr/bin/env python3
"""Build a cumulative plan/fact workbook for one deal from a multi-period table."""
import json
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from installer_report_intake import table_profile

BASE = os.path.dirname(os.path.abspath(__file__))
ARCHIVE = os.path.join(BASE, "installer_approved_reports.json")
FACT_REGISTRY = os.path.join(BASE, "installer_fact_registry.json")
NAVY, WHITE, GREEN, YELLOW, RED = "102733", "FFFFFF", "E2F0D9", "FFF2CC", "F4CCCC"


def normalized(value):
    return re.sub(r"[^a-zа-я0-9]+", " ", str(value or "").casefold().replace("ё", "е")).strip()


def approved_table_totals(deal_id):
    try:
        data = json.load(open(ARCHIVE, encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    result = []
    for report in data.get("reports") or []:
        for deal in report.get("deals") or []:
            if str(deal.get("deal_id")) != str(deal_id):
                continue
            rows = deal.get("rows") or []
            candidates = [row for row in rows if "по таблице" in str(row.get("source") or "").casefold()]
            if not candidates and len(rows) == 1:
                candidates = rows
            for row in candidates:
                amount = row.get("amount")
                if amount is not None:
                    result.append({"date": report.get("date"), "amount": float(amount)})
    unique = {}
    for item in result:
        unique[(item["date"], round(item["amount"], 2))] = item
    return list(unique.values())


def build(deal_id, table_path, output_path):
    profile = table_profile(table_path)
    columns = [item for item in profile.get("completion_columns") or [] if item.get("footer") is not None]
    approved = approved_table_totals(deal_id)
    periods = []
    used_columns = set()
    for item in approved:
        matches = sorted(columns, key=lambda col: abs(float(col.get("footer") or col.get("sum") or 0) - item["amount"]))
        if not matches or abs(float(matches[0].get("footer") or matches[0].get("sum") or 0) - item["amount"]) > 0.01:
            continue
        column = matches[0]
        if column["column"] in used_columns:
            continue
        used_columns.add(column["column"])
        periods.append({"date": item["date"], "amount": item["amount"], "column": str(column["column"])})
    periods.sort(key=lambda item: datetime.strptime(item["date"], "%d.%m.%Y"))
    if not periods:
        raise ValueError("не удалось сопоставить столбцы выполнения с утверждёнными отчётами")

    wb = Workbook()
    ws = wb.active
    ws.title = f"Сделка {deal_id} накопительно"
    last_col = 6 + len(periods) * 2 + 3
    ws["A1"] = f"НАКОПИТЕЛЬНОЕ ВЫПОЛНЕНИЕ ПО СДЕЛКЕ №{deal_id}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=14)
    headers = ["№", "Позиция", "Ед.", "План, кол-во", "Расценка", "Плановая сумма"]
    for period in periods:
        headers += [f"Кол-во {period['date']}", f"Сумма {period['date']}"]
    headers += ["Накопительное кол-во", "Накопительная сумма", "Остаток, кол-во"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(3, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    registry_rows = {}
    for source in profile["rows"]:
        rate = source.get("rate")
        amounts = [float(source.get("extra_values", {}).get(period["column"]) or 0) for period in periods]
        quantities = [(amount / rate if rate not in (None, 0) else 0) for amount in amounts]
        cumulative_qty, cumulative_amount = sum(quantities), sum(amounts)
        plan_qty = float(source.get("qty") or 0)
        values = [source.get("number"), source.get("name"), source.get("unit"), plan_qty, rate, source.get("total")]
        for qty, amount in zip(quantities, amounts):
            values += [qty, amount]
        values += [cumulative_qty, cumulative_amount, max(0, plan_qty - cumulative_qty)]
        ws.append(values)
        fill = RED if cumulative_qty > plan_qty + 1e-9 else GREEN if cumulative_qty > 0 and abs(cumulative_qty - plan_qty) <= 1e-9 else YELLOW if cumulative_qty > 0 else None
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
        registry_rows[normalized(source.get("name"))] = {
            "name": source.get("name"), "qty": cumulative_qty, "updated_report": periods[-1]["date"],
        }

    total_row = ws.max_row + 1
    ws.cell(total_row, 2, "ИТОГО")
    ws.cell(total_row, 2).font = Font(bold=True)
    ws.cell(total_row, 6, f"=SUM(F4:F{total_row-1})")
    for index in range(len(periods)):
        amount_col = 8 + index * 2
        letter = get_column_letter(amount_col)
        ws.cell(total_row, amount_col, f"=SUM({letter}4:{letter}{total_row-1})")
    cumulative_amount_col = 8 + len(periods) * 2
    letter = get_column_letter(cumulative_amount_col)
    ws.cell(total_row, cumulative_amount_col, f"=SUM({letter}4:{letter}{total_row-1})")
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=GREEN)
    widths = [8, 68, 10, 15, 16, 18] + [16, 18] * len(periods) + [22, 22, 18]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column >= 4:
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
    ws.freeze_panes = "A4"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return {"path": output_path, "periods": periods, "registry_rows": registry_rows}


if __name__ == "__main__":
    try:
        result = build(sys.argv[1], sys.argv[2], sys.argv[3])
        print(json.dumps({"ok": True, **result}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        raise SystemExit(1)
