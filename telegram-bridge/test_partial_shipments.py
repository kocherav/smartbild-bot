import copy,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch,Mock
import partial_shipments as p
import closing_runtime as c
from closing_workflow import Registry
from docx import Document

class PartialShipmentIntegrationTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.base=Path(self.tmp.name);self.reg=Registry(self.base/'jobs.json')
  self.patches=[patch.object(p.workspace,'DB',self.base/'shipments.sqlite3'),patch.object(c,'BASE',self.base),patch.object(c,'CRM',self.base/'crm'),patch.object(c,'REGISTRY',self.base/'jobs.json')]
  for x in self.patches:x.start()
  c.CRM.mkdir();(c.BASE/'partial_shipments.enabled').touch()
  self.source={'contract_date_iso':'2026-09-01','total':'100.00','positions':[{'id':'1','kind':'goods','name':'Товар','unit':'шт.','quantity':'2','unit_price':'50.00','total':'100.00','tax_rate':5}]}
  self.saved=p.workspace.save_book('7',self.source)
 def tearDown(self):
  for x in self.patches[::-1]:x.stop()
  self.tmp.cleanup()
 def create(self,n):return p.workspace.confirm({'deal_id':'7','request_id':'request-00000'+str(n),'fingerprint':self.saved['fingerprint'],'transfer_date':'2026-10-13','positions':[{'id':'1','quantity':'1'}]},'owner')['id']
 def test_worker_replay_one_job_and_receipt_projection(self):
  ident=self.create(1)
  p.poll();p.poll();self.assertEqual(len(self.reg.read()['jobs']),1)
  self.assertEqual(self.reg.get(ident)['shipment_id'],ident)
  self.reg.patch(ident,status='awaiting_edo',shipment_finance={'ok':True,'operations':[{'id':'123'}]})
  p.poll();row=p.workspace.shipments('7')[0];self.assertEqual(row['payload']['finance']['operations'][0]['id'],'123')
 def test_first_edo_never_closes_deal_second_needs_finance(self):
  a=self.create(1);p.poll();self.reg.patch(a,edo_confirmation={'actor':'tanya'},shipment_finance={'ok':True})
  self.assertFalse(p.all_edo(self.reg.get(a)))
  b=self.create(2);p.poll();self.reg.patch(b,edo_confirmation={'actor':'tanya'})
  self.assertFalse(p.all_edo(self.reg.get(b)))
  self.reg.patch(b,shipment_finance={'ok':True});self.assertTrue(p.all_edo(self.reg.get(b)))
 def test_partial_stage_watcher_does_not_queue_second_full_act(self):
  a=self.create(1);p.poll();self.reg.observe('7','136')
  (c.CRM/'deal_browser_cache.json').write_text(json.dumps([{'id':'7','program':{'id':'16'},'state':{'id':'138'}}]))
  c.watch_cached();self.assertEqual(list(self.reg.read()['jobs']),[a])
 def test_legacy_full_volume_cannot_be_shipped_again(self):
  j=self.reg.enqueue('7');self.reg.patch(j['id'],status='awaiting_edo',data=self.source)
  p.import_full_jobs();self.assertTrue(p.workspace.complete_volume('7'))
  with self.assertRaises(ValueError):self.create(1)
  self.reg.patch(j['id'],edo_confirmation={'actor':'tanya'},status='awaiting_payment')
  p.poll();row=p.workspace.shipments('7')[0];self.assertEqual(row['state'],'legacy_full');self.assertEqual(row['payload']['edo']['actor'],'tanya');self.assertEqual(row['payload']['lifecycle'],'awaiting_payment')
 def test_finance_is_saved_before_publication_and_failure_does_not_duplicate(self):
  import shipment_finance as f
  j=self.reg.enqueue('7');j=self.reg.patch(j['id'],status='ready',mode='live',data={})
  (c.BASE/'shipment_finance.enabled').touch()
  with patch.object(c,'api',return_value=Mock()),patch.object(c,'validate_current'),patch.object(f,'plans',return_value=[]),patch.object(f,'Backend'),patch.object(f,'ensure',return_value={'ok':True,'operations':[{'id':'x'}]}):
   updated=c.record_transfer(j)
  self.assertEqual(updated['shipment_finance']['operations'][0]['id'],'x')
 def test_single_group_partial_act_no_false_final_payment(self):
  import closing_documents as d
  data={'deal_id':'7','act_number':'7-АПП/01','partial':True,'partial_payment':{},'positions':self.source['positions'],'total':'100.00','act_date':'13.10.2026','number':'1.7','date':'01.09.2026','deadline':'30.10.2026','address':'Тестовый адрес','payer_name':'Тестовый клиент','payer':{'inn':'123','kpp':'456'},'signatory':{'position':'Директор','name':'Иванов Иван Иванович','authority':'Устав'},'payment':{'amount':'100','clause':'6.4.2'},'tax_text':'НДС 5%','transfer_dates':{'goods':'13.10.2026'}}
  def fake_run(*args,**kwargs):
   (self.base/'Акт_7-АПП-01.pdf').write_bytes(b'%PDF-fixture');return Mock(returncode=0)
  with patch.object(d.subprocess,'run',side_effect=fake_run):d.generate_act(data,self.base)
  doc=Document(self.base/'Акт_7-АПП-01.docx');text='\n'.join(x.text for x in doc.paragraphs)
  detail=next(t for t in doc.tables if len(t.columns)==6)
  for row in detail.rows:
   for i in (3,4,5):
    self.assertTrue(all(p.alignment==2 for p in row.cells[i].paragraphs))
  self.assertIsNone(detail.rows[1].cells[1].paragraphs[0].alignment)
  self.assertIn('7-АПП/01',text);self.assertIn('Остальной объём договора настоящим актом не закрывается',text);self.assertNotIn('Окончательный расчёт:',text);self.assertEqual(text.count('ПРИЛОЖЕНИЕ № 1'),1)
if __name__=='__main__':unittest.main()
