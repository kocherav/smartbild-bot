import copy, tempfile, unittest
from pathlib import Path
from unittest.mock import patch, Mock
from decimal import Decimal
from docx import Document
import closing_documents as documents
import partial_shipments as partial
import closing_runtime as runtime
import shipment_workspace as workspace

class SmartBildShipmentTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.out=Path(self.tmp.name)
  # The test supplies its own signature image; no production registries required.
  from PIL import Image
  self.signature=self.out/'signature.png';Image.new('RGBA',(28,10),(0,0,0,255)).save(self.signature)
  self.seller={'id':'kocherzhinsky_ip','name':'ИП Кочержинский А.В.','legal_name':'Индивидуальный предприниматель Кочержинский Артем Вадимович','inn':'782093991468','ogrnip':'324774600759551','signatory':{'name':'Кочержинский Артем Вадимович'},'signature_image':str(self.signature),'accounts':['40802810300007034866'],'bank':{'bik':'044525974','correspondent_account':'30101810145250000974'}}
  self.data={'deal_id':'999','number':'1.999','date':'01.09.2026','act_date':'30.09.2026','act_number':'999-АПП/02','partial':True,'partial_payment':{},'total':'250.00','tax_text':'Без НДС','address':'Адрес объекта','seller':self.seller,'payer_name':'Тестовый заказчик','payer':{'inn':'123','kpp':'456','contractor':{'id':'55'}},'signatory':{'name':'Иванов Иван Иванович','position':'Директор','authority':'Устав'},'positions':[{'id':'1','name':'Товар','kind':'material','unit':'шт.','quantity':'3','unit_price':'50.00','total':'150.00','tax_rate':-1},{'id':'2','name':'Монтажные работы','kind':'work','unit':'этап','quantity':'1','unit_price':'100.00','total':'100.00','tax_rate':-1}]}
 def tearDown(self):self.tmp.cleanup()
 def test_common_generator_routes_ip_uses_approved_brand_and_both_groups(self):
  def convert(*a,**kw):
   Path(a[0][-1]).with_suffix('.pdf').write_bytes(b'%PDF-test');return Mock(returncode=0)
  with patch.object(documents.subprocess,'run',side_effect=convert):
   result=documents.generate_act(self.data,self.out)
  doc=Document(Path(result).with_suffix('.docx'))
  text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
  self.assertIn('999-АПП/02',text);self.assertIn('250,00',text);self.assertIn('30.09.2026',text)
  self.assertIn('ИТОГО ПО ГРУППЕ: 150,00',text);self.assertIn('ИТОГО ПО ГРУППЕ: 100,00',text)
  self.assertNotIn('Кливэл',text);self.assertNotIn('50 000',text);self.assertNotIn('БАРИОН',text)
  self.assertEqual(len(doc.sections[0].header._element.xpath('.//wp:anchor')),1)
  self.assertEqual(len(doc.element.xpath('.//wp:anchor')),1) # signature, no foreign stamp
  for t in doc.tables:
   if len(t.columns)==6:
    for r in t.rows:
     for i in (3,4,5):self.assertEqual(r.cells[i].paragraphs[0].alignment,2)
 def test_long_goods_and_work_act_has_separate_appendices(self):
  import smartbild_act
  d=copy.deepcopy(self.data);d['positions']*=3;d['total']='750.00'
  doc=smartbild_act.document(d,True)
  text='\n'.join(p.text for p in doc.paragraphs)
  self.assertEqual(text.count('ПРИЛОЖЕНИЕ № 1'),2)
  self.assertIn('ТЕСТ',text);self.assertIn('750,00',text)
 def test_bank_validation_uses_selected_company(self):
  text='40802810300007034866 044525974 30101810145250000974'
  runtime.validate_seller_bank(self.seller,text)
  with self.assertRaises(ValueError):runtime.validate_seller_bank(self.seller,'40702810902570001580 044525593 30101810200000000593')
 def test_partial_work_by_amount_preserves_title_and_cannot_overclose(self):
  d=copy.deepcopy(self.data);d['positions'][1].update(quantity='1',unit='шт.',unit_price='100',total='100')
  selected=workspace.select(d,[{'id':'1','quantity':'1'},{'id':'2','amount_mode':True,'stage_amount':'40'}],[])
  self.assertEqual(selected[1]['name'],'Монтажные работы');self.assertEqual(selected[1]['quantity'],'1')
  self.assertEqual(sum(Decimal(r['total']) for r in selected),Decimal('90'))
  with self.assertRaises(ValueError):workspace.select(d,[{'id':'2','amount_mode':True,'stage_amount':'101'}],[])
 def test_selected_date_is_act_date_without_final_invoice(self):
  j={'shipment':{'positions':self.data['positions'],'total':'250','act_number':'999-АПП/02','transfer_date':'30.09.2026'}}
  with patch('document_tax_policy.company_rate',return_value=-1):result=partial.apply(j,{**self.data,'act_date':'19.10.2026'})
  self.assertEqual(result['act_date'],'30.09.2026');self.assertEqual(result['partial_payment'],{})
  self.assertEqual(partial.payment_context(j,'6.4.2 Остаток')['amount'],'0.00')

 def test_ip_transfer_finance_is_zero_cash_and_correct_account(self):
  import shipment_finance as f
  plan=f.plans({'id':'test-job','mode':'live','deal_id':'999','shipment_id':'shipment-unique','data':self.data})[0]
  self.assertEqual(plan['account_id'],'23');self.assertEqual(plan['amount_cents'],25000)
  self.assertEqual(f.form(plan)['Model[Sum][Value]'],'0.00')
  self.assertEqual(f.form(plan)['Model[SecondarySum][Value]'],'250.00')
