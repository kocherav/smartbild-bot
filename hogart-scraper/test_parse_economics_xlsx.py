import os
import tempfile
import unittest

from openpyxl import Workbook

import parse_economics_xlsx as parser


class ProcurementSpecificationTest(unittest.TestCase):
    def test_compact_economics_uses_prices_and_costs_not_stale_profit(self):
        workbook=Workbook();sheet=workbook.active
        sheet.append(['КП для Алексей'])
        sheet.append(['Позиция','Ед.','Кол-во','Цена в КП','Сумма в КП','Себестоимость (ед.)','Себестоимость (сумма)','Прибыль'])
        sheet.append(['Воздуховоды','м2',38,2850,108300,750,28500,79800])
        sheet.append(['Монтажные работы','шт.',1,310330,310330,150000,150000,4000])
        sheet.append(['ИТОГО',None,None,None,418630,None,178500,4000])
        with tempfile.TemporaryDirectory() as directory:
            path=os.path.join(directory,'economics.xlsx');workbook.save(path)
            result=parser.parse(path)
        self.assertEqual(result['total'],418630)
        self.assertEqual(len(result['items']),2)
        self.assertEqual(result['items'][0]['unit_id'],'5')
        self.assertEqual(result['items'][1]['cost_work'],150000)
        self.assertEqual(result['items'][1]['cost_material'],0)

    def split_workbook(self):
        w = Workbook()
        w.active.title = 'Материалы'
        w.active.append(['№', 'Наименование', 'Ед.', 'Кол-во', 'КП / ед.', 'КП итого', 'Закупка / ед.', 'Закупка итого'])
        w.active.append([1, 'Профиль', 'шт', 2, 720, 1440, 640, 1280])
        s = w.create_sheet('Работы')
        s.append(['№', 'Наименование работ', 'Ед.', 'Кол-во', 'Рыночная цена / ед.', 'Рыночная сумма', 'Себестоимость'])
        s.append([1, 'Монтаж профиля', 'м', 200, 375, 75000, 52500])
        w.create_sheet('Сводка').append(['ИТОГО', 76440, 53780])
        w.create_sheet('КП клиенту').append([1, 'Не импортировать повторно'])
        return w

    def test_split_tables_preserve_discounted_sales_and_unit_work_cost(self):
        r = parser.parse_split_calculation(self.split_workbook())
        self.assertEqual(len(r['items']), 2)
        self.assertEqual(r['total'], 76440)
        self.assertEqual(r['cost_total'], 53780)
        self.assertEqual(r['items'][0]['unit_price'], 720)
        self.assertEqual(r['items'][1]['cost_work'], 262.5)
        self.assertEqual(r['items'][1]['cost_material'], 0)

    def test_split_tables_reject_inconsistent_summary(self):
        w = self.split_workbook()
        w['Сводка']['B1'] = 1
        with self.assertRaisesRegex(ValueError, 'сводкой'):
            parser.parse_split_calculation(w)

    def test_split_tables_reject_missing_formula_cache(self):
        w = self.split_workbook()
        w['Работы']['G2'] = None
        with self.assertRaisesRegex(ValueError, 'не является числом'):
            parser.parse_split_calculation(w)

    def test_split_tables_reject_row_total_mismatch(self):
        w = self.split_workbook()
        w['Работы']['F2'] = 100
        with self.assertRaisesRegex(ValueError, 'не совпадает'):
            parser.parse_split_calculation(w)

    def test_procurement_specification_keeps_sale_and_both_costs(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "№", "Наименование проекта / укрупнённой позиции", "Ед.", "Кол-во",
            "Закупка / ед., ₽", "Закупка, сумма ₽", "Работы за ед", "Работы сумма",
            "Цена КП / ед., ₽", "Сумма КП, ₽",
        ])
        sheet.append([1, "Канальный блок", "шт", 2, 1000, 2000, 300, 600, 1800, 3600])
        sheet.append(["Итого"])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as output:
            path = output.name
        try:
            workbook.save(path)
            result = parser.parse(path)
        finally:
            os.unlink(path)

        self.assertTrue(result["ok"])
        self.assertEqual(1, len(result["items"]))
        item = result["items"][0]
        self.assertEqual("Канальный блок", item["name"])
        self.assertEqual(2, item["qty"])
        self.assertEqual(1800, item["unit_price"])
        self.assertEqual(1000, item["cost_material"])
        self.assertEqual(300, item["cost_work"])
        self.assertEqual(3600, result["total"])


if __name__ == "__main__":
    unittest.main()
