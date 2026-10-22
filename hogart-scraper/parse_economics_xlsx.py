#!/usr/bin/env python3
"""Extract every priced row from a Cliwell economics workbook without aggregation.

stdin:  {"path": "/absolute/file.xlsx"}
stdout: {"ok": true, "items": [...], "total": 123.45}

This is intentionally independent from the Telegram/model layer.  A language model
may describe the table, but it must never merge its rows before a CRM write.
"""
import json
import math
import os
import re
import sys

from openpyxl import load_workbook


UNIT_IDS = {
    "м2": "5", "м²": "5", "кв.м": "5",
    "шт": "1", "шт.": "1", "м": "4", "м.": "4", "компл": "8", "компл.": "8",
    "комплект": "8", "к-т": "8", "усл": "9", "усл.": "9", "уп. 100 шт": "8",
}


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split()).strip()


def number(value, field, row_number):
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"строка {row_number}: поле «{field}» не является числом")
    if not math.isfinite(result):
        raise ValueError(f"строка {row_number}: поле «{field}» не является конечным числом")
    return int(result) if result.is_integer() else result


def find_sheet_and_header(workbook):
    required = {"раздел", "позиция", "ед.", "кол-во", "цена кп, ₽", "себестоимость ед., ₽"}
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [clean(value).casefold() for value in row]
            if required.issubset(set(values)):
                return "economics", sheet, row_number, {value: index for index, value in enumerate(values) if value}
            if {"позиция", "ед.", "кол-во", "цена в кп", "себестоимость (ед.)"}.issubset(set(values)):
                return "compact", sheet, row_number, {value: index for index, value in enumerate(values) if value}
            # Procurement specifications are also authoritative source files for a
            # new deal.  They deliberately keep material and work cost in separate
            # columns, unlike the compact economics report above.  Recognise the
            # actual headings, preserving every detail row rather than collapsing
            # the document into a single "equipment" line.
            procurement_required = {
                "наименование проекта / укрупнённой позиции", "ед.", "кол-во",
                "закупка / ед., ₽", "работы за ед", "цена кп / ед., ₽",
            }
            if procurement_required.issubset(set(values)):
                return "procurement", sheet, row_number, {value: index for index, value in enumerate(values) if value}
    raise ValueError("не найдена таблица экономики с обязательными колонками")


def parse(path):
    if not os.path.isfile(path):
        raise ValueError("файл экономики не найден")
    workbook = load_workbook(path, data_only=True, read_only=True)
    if {"Материалы", "Работы", "Сводка"}.issubset(workbook.sheetnames):
        return parse_split_calculation(workbook)
    layout, sheet, header_row, columns = find_sheet_and_header(workbook)
    items = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        position_key = "позиция" if layout in ("economics", "compact") else "наименование проекта / укрупнённой позиции"
        position = clean(row[columns[position_key]] if columns[position_key] < len(row) else None)
        first = clean(row[0] if row else None).casefold()
        if first == "итого":
            break
        if not position:
            continue
        # A numbered blank/separator line in a supplier spreadsheet must not turn
        # into a bogus CRM position.  Conversely, never discard a valid line just
        # because its ordinal number has gaps.
        unit = clean(row[columns["ед."]]) or "шт."
        qty = number(row[columns["кол-во"]], "Количество", row_number)
        if layout == "economics":
            price = number(row[columns["цена кп, ₽"]], "Цена КП", row_number)
            cost = number(row[columns["себестоимость ед., ₽"]], "Себестоимость", row_number)
            section = clean(row[columns["раздел"]]).casefold()
            cost_material = 0 if ("работ" in section or "услуг" in section) else cost
            cost_work = cost if ("работ" in section or "услуг" in section) else 0
            section_value = clean(row[columns["раздел"]])
        elif layout == "compact":
            price = number(row[columns["цена в кп"]], "Цена КП", row_number)
            cost = number(row[columns["себестоимость (ед.)"]], "Себестоимость", row_number)
            is_work = bool(re.search(r'(?i)\bработ\w*\b|^монтаж\b|^пусконалад', position))
            cost_material, cost_work = (0, cost) if is_work else (cost, 0)
            section_value = "Работы" if is_work else "Оборудование и материалы"
            for heading, expected in (("сумма в кп", qty * price), ("себестоимость (сумма)", qty * cost)):
                if heading in columns and abs(number(row[columns[heading]], heading, row_number) - expected) > 0.01:
                    raise ValueError(f"строка {row_number}: {heading} не совпадает с количеством × ценой")
        else:
            price = number(row[columns["цена кп / ед., ₽"]], "Цена КП", row_number)
            cost_material = number(row[columns["закупка / ед., ₽"]], "Закупка", row_number)
            cost_work = number(row[columns["работы за ед"]], "Работы", row_number)
            section_value = "Оборудование и материалы"
        if qty <= 0 or price < 0 or cost_material < 0 or cost_work < 0:
            raise ValueError(f"строка {row_number}: отрицательная цена/себестоимость или нулевое количество")
        items.append({
            "name": position,
            "qty": qty,
            "unit": unit,
            "unit_id": UNIT_IDS.get(unit.casefold(), "1"),
            "unit_price": price,
            "cost_material": cost_material,
            "cost_work": cost_work,
            "section": section_value,
        })
    if not items:
        raise ValueError("в таблице экономики нет позиций")
    return {"ok": True, "items": items, "total": round(sum(float(x["qty"]) * float(x["unit_price"]) for x in items), 2)}


def parse_split_calculation(workbook):
    """Read both detailed tables, not the duplicate client presentation sheet.

    Sales totals already include discounts. Work cost is a ROW total, whereas
    material purchase price is per unit. Keep these semantics explicit.
    """
    items = []
    totals = []
    for sheet_name, name_key, price_key, total_key, cost_key, cost_total_key in (
        ("Материалы", "наименование", "кп / ед.", "кп итого", "закупка / ед.", "закупка итого"),
        ("Работы", "наименование работ", "рыночная цена / ед.", "рыночная сумма", None, "себестоимость"),
    ):
        columns = None
        sale_sum = cost_sum = 0
        for rn, row in enumerate(workbook[sheet_name].values, 1):
            values = [clean(x).casefold() for x in row]
            required = {name_key, price_key, total_key, cost_total_key, "кол-во", "ед."}
            if required.issubset(values):
                columns = {v: i for i, v in enumerate(values) if v}
                continue
            if columns is None or not isinstance(row[0], (int, float)):
                continue
            def val(key):
                return row[columns[key]]
            qty = number(val("кол-во"), "Количество", rn)
            price = number(val(price_key), "Цена за единицу", rn)
            total = number(val(total_key), "Сумма строки", rn)
            cost_total = number(val(cost_total_key), "Себестоимость строки", rn)
            name = clean(val(name_key))
            if not name or qty <= 0 or min(price, total, cost_total) < 0:
                raise ValueError(f"{sheet_name}, строка {rn}: некорректные данные")
            if abs(round(qty * price, 2) - total) > 0.011:
                raise ValueError(f"{sheet_name}, строка {rn}: цена × количество не совпадает с суммой")
            if cost_key and abs(round(qty * number(val(cost_key), cost_key, rn), 2) - cost_total) > 0.011:
                raise ValueError(f"{sheet_name}, строка {rn}: закупочная цена не совпадает с суммой")
            unit = clean(val("ед."))
            items.append({"name": name, "qty": qty, "unit": unit,
                          "unit_id": UNIT_IDS.get(unit.casefold(), "1"),
                          "unit_price": total / qty,
                          "cost_material": cost_total / qty if cost_key else 0,
                          "cost_work": 0 if cost_key else cost_total / qty,
                          "section": sheet_name, "source_row": rn,
                          "line_total": total, "cost_total": cost_total})
            sale_sum += total
            cost_sum += cost_total
        if columns is None or sale_sum <= 0:
            raise ValueError(f"Не найдена заполненная таблица «{sheet_name}»")
        totals.append((sale_sum, cost_sum))
    summary = None
    for row in workbook["Сводка"].values:
        if clean(row[0]).casefold() == "итого":
            summary = (number(row[1], "Итог КП", 0), number(row[2], "Итог себестоимости", 0))
            break
    sale = round(sum(t[0] for t in totals), 2)
    cost = round(sum(t[1] for t in totals), 2)
    if summary is None or abs(sale - summary[0]) > 0.011 or abs(cost - summary[1]) > 0.011:
        raise ValueError("Суммы подробных строк не совпадают со сводкой калькуляции")
    return {"ok": True, "items": items, "total": sale, "cost_total": cost}


if __name__ == "__main__":
    try:
        payload = json.load(sys.stdin)
        print(json.dumps(parse(payload["path"]), ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        raise SystemExit(1)
