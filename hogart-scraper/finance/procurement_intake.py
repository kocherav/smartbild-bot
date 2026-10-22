"""Canonical supplier document intake. Parsing never posts money or changes sales."""
import copy,hashlib,json,re,sys,subprocess,tempfile,time,datetime as dt
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import allocation_workspace as aw
ROOT=Path(__file__).resolve().parent
CRM=ROOT.parent.parent/'megaplan-crm'
sys.path.insert(0,str(CRM))
from registry_transaction import locked
import review_registry as queue

def read(path,default):
 try:return json.loads(Path(path).read_text())
 except FileNotFoundError:return default

def digest(value):return hashlib.sha256(str(value).encode()).hexdigest()[:20]
def money(value):return aw.cents(str(value).replace('\xa0','').replace(' ','').replace(',','.'))
def extract(path):
 p=Path(path)
 if p.suffix.lower()=='.pdf':
  text=subprocess.run(['pdftotext','-layout',str(p),'-'],capture_output=True,check=True,timeout=30).stdout.decode()
  if len(text.strip())>=80:return text
  info=subprocess.run(['pdfinfo',str(p)],capture_output=True,check=True,timeout=15).stdout.decode()
  pages=re.search(r'^Pages:\s*(\d+)',info,re.M)
  if not pages or int(pages[1])>20:raise ValueError('Скан требует ручного разбора: не более 20 страниц для OCR')
  with tempfile.TemporaryDirectory(prefix='supplier-ocr-') as folder:
   prefix=str(Path(folder)/'page')
   subprocess.run(['pdftoppm','-r','180','-png',str(p),prefix],capture_output=True,check=True,timeout=90)
   texts=[]
   for image in sorted(Path(folder).glob('page-*.png')):
    texts.append(subprocess.run(['tesseract',str(image),'stdout','-l','rus+eng','--psm','6'],capture_output=True,check=True,timeout=45).stdout.decode())
   text='\n'.join(texts)
  if len(text.strip())<80:raise ValueError('Скан не читается. Загрузите более чёткий файл или заполните строки вручную')
  return text
 if p.suffix.lower()=='.xlsx':
  import openpyxl
  w=openpyxl.load_workbook(p,data_only=True,read_only=True)
  try:return '\n'.join(' | '.join(str(c or '') for c in row) for sh in w for row in sh.values)
  finally:w.close()
 if p.suffix.lower()=='.xls':
  import xlrd
  w=xlrd.open_workbook(str(p));return '\n'.join(' | '.join(map(str,sh.row_values(i))) for sh in w.sheets() for i in range(sh.nrows))
 if p.suffix.lower()=='.docx':
  import docx
  w=docx.Document(p);return '\n'.join(x.text for x in w.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in w.tables for row in t.rows)
 raise ValueError('Нужен PDF, XLS, XLSX или DOCX')

def allocate_goods_discount(result,discount):
 """Keep service charges intact; allocate whole-invoice discount to goods in cents."""
 result=copy.deepcopy(result);lines=result['lines']
 goods=[x for x in lines if str(x.get('unit','')).casefold() not in ('услуга','усл','усл.') and not re.search(r'доставк|транспортн.*услуг',x['name'],re.I)]
 base=sum(x['amount_cents'] for x in goods)
 if not isinstance(discount,int) or discount<=0 or discount>base:raise ValueError('Скидка превышает стоимость товаров или не подтверждена')
 if sum(x['amount_cents'] for x in lines)-discount!=result['total_cents']:raise ValueError('Скидка не объясняет итог счёта')
 shares={x['line_no']:discount*x['amount_cents']//base for x in goods}
 ordered=sorted(goods,key=lambda x:(-(discount*x['amount_cents']%base),x['line_no']))
 for x in ordered[:discount-sum(shares.values())]:shares[x['line_no']]+=1
 for x in lines:
  x['gross_amount_cents']=x['amount_cents'];x['discount_cents']=shares.get(x['line_no'],0);x['amount_cents']-=x['discount_cents']
 result['invoice_discount']={'policy':'goods_only','amount_cents':discount}
 return result

def parse(text,meta):
 """Strict common PDF layouts; unsupported tables proceed to structured extraction."""
 lines=[];current=None
 if 'Код' in text and 'товара ЭТМ' in text:
  for raw in text.splitlines():
   m=re.match(r'\s*(\d+)\s+(ETM\S+)\s+(.+?)\s+([SNZСН])\s*(шт|м|кг|уп|компл)\s+(\d+(?:[.,]\d*)?)\s+(.+)',raw)
   if m:
    nums=re.findall(r'\d+\.\d{2}',m[7])
    if len(nums)<3:raise ValueError('Не распознана сумма строки ЭТМ')
    current={'line_no':int(m[1]),'sku':m[2],'name':m[3].strip(),'unit':m[5],'qty':m[6].rstrip('.').replace(',','.'),'amount_cents':money(nums[-1])};lines.append(current)
   elif current and re.match(r'^\s{14,25}\S',raw) and not re.search(r'Итого|Всего|Лист|Счет|Счёт|ИНН|руб\.|НДС|N\s+Код',raw,re.I):current['name']+=' '+raw.strip()
 elif 'Сантехкомплект' in text:
  for raw in text.splitlines():
   m=re.match(r'\s*(\d+)\s+(\d{3}-\d+)\s+(.+?)\s{2,}(\d+(?:[.,]\d+)?)\s+(шт|м|кг|уп|компл)\s+(.+)',raw)
   if m:
    nums=re.findall(r'\d[\d ]*,\d{2}',m[6])
    if len(nums)<3:raise ValueError('Не распознана сумма строки Сантехкомплект')
    current={'line_no':int(m[1]),'sku':m[2],'name':m[3].strip(),'unit':m[5],'qty':m[4].replace(',','.'),'amount_cents':money(nums[2])};lines.append(current)
   elif current and re.match(r'^\s{20,30}\S',raw) and not re.search(r'Итого|Всего|Страница|Счет|Счёт|НДС',raw,re.I):current['name']+=' '+raw.strip()
 elif re.search(r'Скидка:\s*[\d ]+[,.]\d{2}',text) and 'Товары (работы, услуги)' in text:
  for raw in text.splitlines():
   m=re.match(r'\s*(\d+)\s+(\S+)\s{2,}(.+?)\s{2,}(?:(\d+(?:[.,]\d+)?)\s+)?(шт|м|услуга|Услуга)\s+([\d ]+[,.]\d{2})\s{2,}([\d ]+[,.]\d{2})\s*$',raw)
   if m:
    current={'line_no':int(m[1]),'sku':m[2],'name':m[3].strip(),'qty':(m[4] or '1').replace(',','.'),'unit':m[5].lower(),'amount_cents':money(m[7])};lines.append(current)
   elif current and re.match(r'^\s{20,30}\S',raw) and not re.search(r'Скидка|Итого|НДС|Всего|Оплатить|Руководитель',raw,re.I):current['name']+=' '+raw.strip()
 elif 'Сумма с НДС' in text and 'евро' in text:
  previous=''
  for raw in text.splitlines():
   m=re.match(r'\s*(\d+)\s+(\S+)\s{2,}(шт|м|кг|компл)\s+(\d+[,.]\d{2})\s+(.+)',raw)
   if m:
    values=re.findall(r'\d+(?: \d{3})*[,.]\d{2}',m[5])
    if len(values)!=5:raise ValueError('Не распознаны колонки валютной строки')
    lines.append({'line_no':int(m[1]),'sku':m[2],'name':previous+' '+m[2],'unit':m[3],'qty':m[4].replace(',','.'),'amount_cents':money(values[-1])})
   previous=raw.strip()
 total=meta.get('amount_cents')
 if not total and 'товара ЭТМ' in text:
  totals=re.findall(r'^\s*Итого\s*:[^\n]+',text,re.M|re.I)
  if totals:
   values=re.findall(r'\d+\.\d{2}',totals[-1]);total=money(values[-1]) if values else None
 if not total:
  candidates=re.findall(r'(?:Всего к оплате|Всего с НДС|Итого к оплате|ИТОГО:|Всего к перечислению)[^\n]*?([\d][\d ]*[.,]\d{2})\s*(?:руб|$)',text,re.I|re.M)
  if candidates:total=money(candidates[-1])
 date=re.search(r'от\s*(\d{2}[./]\d{2}[./]\d{4})',text,re.I)
 result={'lines':lines,'total_cents':total,'date':date[1].replace('/','.') if date else ''}
 if lines and total and sum(x['amount_cents'] for x in lines)!=total:
  discount=re.search(r'Скидка:\s*([\d ]+[,.]\d{2})',text)
  if discount and meta.get('discount_policy')=='goods_only':result=allocate_goods_discount(result,money(discount[1]))
 validate(result,meta,text);return result

def validate(result,meta,text=None):
 lines=result.get('lines');total=result.get('total_cents')
 if not isinstance(lines,list) or not lines or len(lines)>1000:raise ValueError('Таблица счёта требует разбора')
 if not isinstance(total,int) or total<=0:raise ValueError('Не подтверждена итоговая сумма счёта')
 if meta.get('amount_cents') and abs(total-meta['amount_cents'])>1:raise ValueError('Итог не совпадает с подтверждённым документом')
 difference=sum(x['amount_cents'] for x in lines)-total
 if difference:
  conversion=result.get('currency_conversion') or {}
  rate=aw.number(conversion.get('rate',0))
  foreign=conversion.get('foreign_total_cents')
  def converted(cents):return int((Decimal(cents)*rate).quantize(Decimal(1),rounding=ROUND_HALF_UP))
  if (rate<=0 or not isinstance(foreign,int) or any(not isinstance(x.get('foreign_amount_cents'),int) for x in lines)
      or sum(x['foreign_amount_cents'] for x in lines)!=foreign or converted(foreign)!=total
      or any(converted(x['foreign_amount_cents'])!=x['amount_cents'] for x in lines)
      or difference!=conversion.get('rounding_cents') or abs(difference)>len(lines)):
   raise ValueError('Сумма строк отличается от итога счёта')
 if len({x['line_no'] for x in lines})!=len(lines):raise ValueError('Повтор номера строки')
 if result.get('invoice_discount'):
  discount=result['invoice_discount']
  if discount.get('policy')!='goods_only':raise ValueError('Не подтверждена политика скидки')
  gross=copy.deepcopy(result)
  for x in gross['lines']:x['amount_cents']=x.get('gross_amount_cents',-1)
  expected=allocate_goods_discount(gross,discount.get('amount_cents'))
  if any(x['amount_cents']!=y['amount_cents'] or x.get('discount_cents')!=y['discount_cents'] for x,y in zip(lines,expected['lines'])):raise ValueError('Неверное распределение скидки')
 for x in lines:
  if isinstance(x.get('qty'),str) and re.fullmatch(r'\d+(?:[ \xa0]\d{3})*(?:[.,]\d+)?',x['qty'].strip()):x['qty']=re.sub(r'[ \xa0]','',x['qty'].strip()).replace(',','.')
  if not isinstance(x['line_no'],int) or x['line_no']<=0 or not str(x.get('name','')).strip() or not str(x.get('unit','')).strip() or aw.number(x['qty'])<=0 or not isinstance(x['amount_cents'],int) or x['amount_cents']<0:raise ValueError('Некорректная строка счёта')
  x['key']=str(x['line_no'])+':'+str(x.get('sku') or digest(x['name'])).casefold()
 if text:
  compact=lambda value:''.join(c for c in str(value).casefold() if c.isalnum())
  if meta.get('number') and compact(meta['number']) not in compact(text):raise ValueError('Номер счёта не найден в исходном документе')
  for key in ('supplier_inn','buyer_inn'):
   if meta.get(key) and meta[key] not in re.sub(r'\s','',text):raise ValueError('ИНН документа не совпадает с карточкой счёта')
  amounts={money(m) for m in re.findall(r'(?<![\w.,])\d+(?:[ \u00a0]\d{3})*(?:[.,]\d{1,2})?(?![\w.,])',text)}
  for x in lines:
   amount=format(Decimal(x['amount_cents'])/100,'.2f')
   source_amount=x.get('gross_amount_cents') if result.get('invoice_discount') else x['amount_cents']
   if source_amount not in amounts:raise ValueError('Сумма строки не найдена в исходном документе')
  if result.get('invoice_discount'):
   found=re.search(r'Скидка:\s*([\d ]+[,.]\d{2})',text)
   if not found or money(found[1])!=result['invoice_discount']['amount_cents']:raise ValueError('Скидка не найдена в оригинале')
 return result

def source_meta(inv):
 currency=inv.get('currency','RUB')
 amount=inv.get('foreign_total') if currency!='RUB' else inv.get('total')
 if currency!='RUB' and (not amount or not inv.get('plan_rate')):raise ValueError('Для валютного счёта укажите сумму оригинала и согласованный курс в карточке счёта')
 return {'number':inv['number'],'supplier_inn':inv.get('supplier_inn'),'buyer_inn':inv.get('buyer_inn'),'amount_cents':aw.cents(amount) if amount else None,'discount_policy':inv.get('discount_policy')}

def convert_source(parsed,inv):
 """Keep foreign original and the already agreed RUB rate; never use today's FX."""
 if inv.get('currency','RUB')=='RUB':return parsed
 validate(parsed,source_meta(inv));result=copy.deepcopy(parsed);rate=aw.number(inv['plan_rate'])
 if rate<=0:raise ValueError('Не подтверждён курс валютного счёта')
 convert=lambda cents:int((Decimal(cents)*rate).quantize(Decimal(1),rounding=ROUND_HALF_UP))
 foreign=parsed['total_cents'];result['total_cents']=convert(foreign)
 if abs(result['total_cents']-aw.cents(inv['total']))>1:raise ValueError('Согласованный курс не объясняет рублёвую сумму; требуется сверка курса')
 for line in result['lines']:
  line['foreign_amount_cents']=line['amount_cents'];line['amount_cents']=convert(line['amount_cents'])
 result['currency_conversion']={'currency':inv['currency'],'rate':str(rate),'foreign_total_cents':foreign,'rounding_cents':sum(l['amount_cents'] for l in result['lines'])-result['total_cents']}
 validate(result,{});return result

def archive_source(entry,target):
 """Fetch an already archived original once, with a mandatory content hash."""
 target=Path(target);target.parent.mkdir(parents=True,exist_ok=True)
 if target.is_file() and hashlib.sha256(target.read_bytes()).hexdigest()==entry['sha256']:return target
 if not entry.get('onedrive_remote'):raise ValueError('В архиве отсутствует адрес оригинала')
 with tempfile.TemporaryDirectory(prefix='supplier-archive-') as folder:
  temporary=Path(folder)/target.name
  subprocess.run(['rclone','copyto',entry['onedrive_remote'],str(temporary),'--retries','1','--low-level-retries','1'],capture_output=True,check=True,timeout=60)
  if hashlib.sha256(temporary.read_bytes()).hexdigest()!=entry['sha256']:raise ValueError('Контрольная сумма архивного оригинала не совпадает')
  target.write_bytes(temporary.read_bytes());target.chmod(0o660)
 return target

def discover(root=ROOT,crm=CRM,downloader=archive_source):
 """Merge evidence indexes by SHA; archived copies never become new purchases."""
 catalog=read(root/'evidence/posting/archive_catalog.json',[])
 # Earlier mailbox reviews retained originals outside the posting index. Recover
 # missing sources against existing invoice identity and the complete table,
 # never against a mere mention of an invoice number in correspondence.
 legacy=read(crm/'procurement_registry.json',{}).get('invoices',[])
 evidence=read(root/'evidence/deep_review/additional_documents.json',[])
 for inv in legacy:
  if inv.get('source_path') or not inv.get('total'):continue
  hits={}
  for item in evidence:
   if inv['number'] not in item.get('text',''):continue
   p=Path(item.get('path',''));p=p if p.is_absolute() else root.parent/p
   if not p.is_file():continue
   try:
    text=extract(p);parsed=parse(text,source_meta(inv))
   except (ValueError,OSError,subprocess.SubprocessError):continue
   sha=hashlib.sha256(p.read_bytes()).hexdigest()
   hits[sha]={'source':str(p),'sha256':sha,'number':inv['number'],'supplier_inn':inv['supplier_inn'],'buyer_inn':inv.get('buyer_inn',''),'supplier':inv['supplier'],'amount_cents':parsed['total_cents'],'deal_ids':sorted({str(a['deal_id']) for a in inv.get('allocations',[]) if a.get('deal_id')})}
  if len(hits)==1:catalog.append(next(iter(hits.values())))
 rules=read(root/'user_rules.json',{}).get('confirmed_payment_allocations',{})
 reviewed=read(root/'evidence/deep_review/reviewed_allocations.json',{})
 for r in [*reviewed.values(),*rules.values()]:
  if r.get('confirmed') and r.get('invoice_source'):catalog.append({'source':r['invoice_source'],'number':r['number'],'supplier_inn':r['inn'],'sha256':r.get('sha256'),'deal_ids':[str(a['deal_id']) for a in r['allocations'] if a.get('deal_id')],'note':r.get('source','')})
 archive=read(crm/'financial_archive.json',{}).get('documents',[]);docs={}
 bank_rows=read(root/'reports/ledger.json',[])
 normalize=lambda v:re.sub(r'[\s№/\-]','',str(v)).casefold()
 downloads=0
 for a in archive:
  if a.get('category')!='supplier_invoice' or a.get('status')!='uploaded' or not a.get('document_number') or not re.fullmatch(r'[a-f0-9]{64}',str(a.get('sha256',''))):continue
  if any(c.get('sha256')==a['sha256'] for c in catalog):continue
  suffix=Path(a.get('filename','')).suffix.lower();source=root/'evidence/kav'/(a['sha256']+suffix)
  if suffix not in ('.pdf','.xls','.xlsx','.docx'):continue
  matches=[r for r in bank_rows if r.get('direction')=='out' and (normalize(r.get('document_numbers',''))==normalize(a['document_number']) or any(normalize(p.get('invoice_number',''))==normalize(a['document_number']) for p in r.get('confirmed_payment_allocations',[])))]
  inns={r.get('inn') for r in matches if r.get('inn')}
  if len(inns)!=1:continue
  if not source.is_file() or hashlib.sha256(source.read_bytes()).hexdigest()!=a['sha256']:
   p=root/'data/review_registry.json';p.parent.mkdir(parents=True,exist_ok=True);key='archive:'+a['sha256']
   previous=read(p,{}).get('tasks',{}).get(key,{})
   if previous.get('attempts',0)>=3 or time.time()-previous.get('last_attempt',0)<300:continue
   if downloads>=2:continue # Bound first-time archive work per background cycle.
   downloads+=1
   try:downloader(a,source)
   except (ValueError,OSError,subprocess.SubprocessError) as exc:
    # Preserve a concrete retryable task in the canonical review registry.
    with locked(p):
     state=read(p,{});attempts=previous.get('attempts',0)+1;state.setdefault('tasks',{})[key]={'id':key,'kind':'archive_source','state':'open','manual':True,'owner':'assistant','bank_ids':[],'supplier':matches[0].get('counterparty',''),'invoices':a['document_number'],'action':('Не удалось получить оригинал из архива. Автоматический повтор через 5 минут, максимум 3 попытки.' if attempts<3 else 'Три попытки получения оригинала из архива не удались. Загрузите файл через «Добавить новый счёт» в пульте; ассистент должен проверить доступ к архиву.'),'attempts':attempts,'last_attempt':time.time(),'error':type(exc).__name__};queue.atomic(p,state)
    continue
   if previous:
    with locked(p):
     state=read(p,{});state['tasks'][key].update(state='resolved',resolution='Оригинал получен, SHA256 совпал');queue.atomic(p,state)
  catalog.append({'source':str(source),'sha256':a['sha256'],'number':a['document_number'],'supplier_inn':next(iter(inns)),'supplier':matches[0].get('counterparty') or next(iter(inns)),'deal_ids':[str(a['deal_id'])]})
 for source in catalog:
  if source.get('category','supplier_invoice')!='supplier_invoice':continue
  p=Path(source.get('source',''))
  if not p.is_file():continue
  sha=hashlib.sha256(p.read_bytes()).hexdigest()
  if source.get('sha256') and sha!=source['sha256']:continue
  d=docs.setdefault(sha,{**source,'sha256':sha,'source':str(p),'deal_ids':[]})
  d['deal_ids']=sorted(set(d['deal_ids'])|set(source.get('deal_ids',[]))|{str(a['deal_id']) for a in archive if a.get('sha256')==sha and a.get('status')=='uploaded'})
  if source.get('amount_cents'):d['amount_cents']=source['amount_cents']
 for row in bank_rows:
  if row.get('direction')!='out':continue
  for d in docs.values():
   if row.get('inn')!=d.get('supplier_inn'):continue
   if not d.get('supplier') and row.get('counterparty'):d['supplier']=row['counterparty']
   refs=str(row.get('document_numbers',''));normalize=lambda v:re.sub(r'[\s№/\-]','',str(v)).casefold()
   parts=row.get('confirmed_payment_allocations',[])
   explicit=[p for p in parts if p.get('invoice_number')]
   if explicit:
    matched=[p for p in explicit if normalize(p['invoice_number'])==normalize(d['number'])]
    if not matched:continue
    amount=sum(p['amount_cents'] for p in matched);scope=matched
   else:
    if normalize(refs)!=normalize(d['number']) or ';' in refs or ',' in refs:continue
    amount=row['amount_cents'];scope=parts
   if d.get('amount_cents') and amount>d['amount_cents']:
    d['payment_issue']='Платёж превышает сумму счёта; нужна документальная разбивка';continue
   if str(row.get('verified_deal_id','')).isdigit():d['deal_ids']=sorted(set(d['deal_ids'])|{str(row['verified_deal_id'])})
   for part in scope:
    if part.get('deal_id'):d['deal_ids']=sorted(set(d['deal_ids'])|{str(part['deal_id'])})
   d.setdefault('payments',[]).append({'transaction_id':row['id'],'amount':amount/100,'date':row['date']})

 for r in rules.values():
  if r.get('confirmed') and r.get('sha256') in docs:
   docs[r['sha256']]['deal_ids']=sorted({str(a['deal_id']) for a in r['allocations'] if a.get('deal_id')})
 return list(docs.values())

def register(meta,path=queue.PROCUREMENT):
 """Create an intake card, not a financial allocation."""
 with locked(path):
  data=read(path,{});invoices=data.setdefault('invoices',[])
  existing=[i for i in invoices if i.get('supplier_inn')==meta.get('supplier_inn') and i.get('number')==meta['number']]
  same=next((i for i in existing if i.get('source_sha256')==meta['sha256']),None)
  if not same and len(existing)==1 and not existing[0].get('source_sha256'):
   old=existing[0]
   if not meta.get('amount_cents') or aw.cents(old.get('total'))==meta['amount_cents']:same=old
  if same:
   changed=False
   if meta.get('supplier') and str(same.get('supplier','')).isdigit():same['supplier']=meta['supplier'];changed=True
   imported=all(re.fullmatch(r'[a-f0-9]{64}',str(p.get('transaction_id',''))) for p in same.get('payments',[]))
   if meta.get('payments') and (not same.get('payments') or imported) and same.get('payments')!=meta['payments']:
    same['payments']=meta['payments'];same['paid_total']=sum(p['amount'] for p in meta['payments']);changed=True
   for k,v in [('source_path',meta['source']),('source_sha256',meta['sha256']),('candidate_deal_ids',sorted(set(same.get('candidate_deal_ids',[]))|set(meta['deal_ids'])))]:
    if same.get(k)!=v:same[k]=v;changed=True
   if 'intake' not in same:same['intake']={'state':'queued','attempts':0};changed=True
   if changed:queue.atomic(path,data)
   return same['document_key']
  key=digest(meta['supplier_inn']+'|'+meta['number']+'|'+meta['sha256'])
  inv={'id':key,'document_key':key,'number':meta['number'],'supplier':meta.get('supplier',meta['supplier_inn']),'supplier_inn':meta['supplier_inn'],'buyer_inn':meta.get('buyer_inn',''),'buyer':'','total':meta.get('amount_cents',0)/100,'source_path':meta['source'],'source_sha256':meta['sha256'],'candidate_deal_ids':meta['deal_ids'],'source_note':meta.get('note',''),'allocations':[],'payments':meta.get('payments',[]),'paid_total':sum(p['amount'] for p in meta.get('payments',[])),'intake':{'state':'version_review' if existing else 'queued','attempts':0,'error':'Найдена другая версия счёта; выберите действующую' if existing else ''}}
  invoices.append(inv);queue.atomic(path,data);return key

def install(key,parsed,path=queue.PROCUREMENT,actor='parser'):
 """Reconcile parsed source rows against existing usage, preserving every allocation."""
 validate(parsed,{})
 with locked(path):
  data=read(path,{});inv=next(i for i in data['invoices'] if i['document_key']==key)
  if inv.get('intake',{}).get('state') in ('version_review','superseded'):raise ValueError('Сначала выберите действующую версию счёта')
  if inv.get('total') and abs(aw.cents(inv['total'])-parsed['total_cents'])>1:raise ValueError('Версия счёта отличается по сумме; существующие распределения сохранены')
  enrich_usage(inv,parsed['lines'])
  stocks=data.setdefault('unallocated_stock',[]);new=[];unknown=[];used=set();recognized_stocks=set()
  for line in parsed['lines']:
   existing=next((s for s in stocks if s.get('document_key')==key and (str(s.get('source_line_key','')).casefold()==line['key'] or (line.get('sku') and str(s.get('sku','')).casefold()==str(line['sku']).casefold() and sum(str(l.get('sku','')).casefold()==str(line['sku']).casefold() for l in parsed['lines'])==1))),None)
   if existing:existing['source_line_key']=line['key'];recognized_stocks.add(existing['id'])
   shares=[l for a in inv.get('allocations',[]) for l in (a.get('source_lines') or a.get('updates') or []) if str(l.get('key') or l.get('source_line_key') or '').casefold()==line['key']]
   qty=sum((aw.number(l.get('source_qty_used',l.get('qty_used',0))) for l in shares),Decimal(0));amount=sum(aw.cents(l.get('amount_used',0)) for l in shares)
   if shares:used.add(line['key'])
   if existing:
    qty+=sum((aw.number(a['qty']) for a in existing.get('allocations',[])),Decimal(0));amount+=sum(aw.cents(a['amount']) for a in existing.get('allocations',[]))
   remaining=aw.number(line['qty'])-qty;left=line['amount_cents']-amount
   if remaining<0 or left<0:unknown.append('Превышено распределение строки '+str(line['line_no']));continue
   if existing:
    if aw.number(existing['available_qty'])!=remaining or aw.cents(existing['available_amount'])!=left:unknown.append('Остаток строки '+str(line['line_no'])+' расходится с сохранённым')
   elif remaining or left:new.append({'id':digest(key+'|'+line['key']),'document_key':key,'source_line_key':line['key'],'name':line['name'],'sku':line.get('sku',''),'unit':line['unit'],'initial_qty':float(remaining),'available_qty':float(remaining),'initial_amount':left/100,'available_amount':left/100,'allocations':[]})
  if any(s['document_key']==key and s['id'] not in recognized_stocks for s in stocks):unknown.append('Прежний складской остаток не сопоставлен со строкой оригинала')
  direct=sum(aw.cents(a.get('amount')) for a in inv.get('allocations',[]))
  known=sum(aw.cents(l.get('amount_used',0)) for a in inv.get('allocations',[]) for l in (a.get('source_lines') or a.get('updates') or []))
  if abs(direct-known)>1:unknown.append('Ранее записана сумма без расшифровки строк; нужна сверка состава')
  if direct and len(used)!=len({str(l.get('key') or l.get('source_line_key') or '').casefold() for a in inv.get('allocations',[]) for l in (a.get('source_lines') or a.get('updates') or [])}):unknown.append('Старые строки не совпали с номерами/артикулами оригинала')
  inv.update(lines=parsed['lines'],date=parsed.get('date',''),total=parsed['total_cents']/100)
  if parsed.get('currency_conversion'):inv['currency_conversion']=parsed['currency_conversion']
  if parsed.get('invoice_discount'):inv['invoice_discount']=parsed['invoice_discount']
  inv['intake']={'state':'review' if unknown else 'ready','error':'; '.join(unknown),'attempts':inv.get('intake',{}).get('attempts',0),'actor':actor,'updated_at':dt.datetime.now(dt.timezone.utc).isoformat()}
  if not unknown:stocks.extend(new)
  data.setdefault('intake_audit',[]).append({'document_key':key,'actor':actor,'state':inv['intake']['state'],'source_sha256':inv.get('source_sha256')})
  queue.atomic(path,data);return inv['intake']


def enrich_usage(inv,lines):
 """Recover keys/row ids from existing receipts; never create new cost allocations."""
 for allocation in inv.get('allocations',[]):
  source=allocation.get('source_lines') or allocation.get('updates') or []
  normalized=[]
  for old in source:
   item=copy.deepcopy(old);sku=str(item.get('source_sku') or '').casefold()
   hits=[l for l in lines if sku and str(l.get('sku','')).casefold()==sku]
   if not hits and item.get('source_name'):
    from product_matching import canonical
    hits=[l for l in lines if canonical(l['name'])==canonical(item['source_name'])]
   if not hits and not item.get('source_name'):
    amount=aw.cents(item.get('amount_used',item.get('material_cost_used',0)));qty=aw.number(item.get('qty_used',0))
    hits=[l for l in lines if l['amount_cents']==amount and aw.number(l['qty'])==qty]
   if len(hits)==1:
    line=hits[0];item.update(key=line['key'],source_sku=line.get('sku',''),source_name=line['name'],source_qty=line['qty'],source_unit=line['unit'],amount_used=item.get('amount_used',item.get('material_cost_used',0)))
   if not item.get('offer_row_id'):
    key=str(old.get('key') or old.get('source_line_key') or '').casefold()
    matches=[u for u in allocation.get('updates',[]) if key and key==str(u.get('source_line_key') or '').casefold()]
    if len(matches)==1:item.update(offer_row_id=matches[0].get('offer_row_id'),offer_id=matches[0].get('offer_id'))
   normalized.append(item)
  if normalized:allocation['source_lines']=normalized

def repair_usage(key,mappings,path=queue.PROCUREMENT,actor='user'):
 """Attach historical monetary receipts to original source rows, without reposting."""
 with locked(path):
  data=read(path,{});inv=next(i for i in data['invoices'] if i['document_key']==key)
  parsed=inv.get('lines',[]);by_no={str(l['line_no']):l for l in parsed};qty_by={};cost_by={}
  if not parsed:raise ValueError('Сначала разберите или заполните строки оригинала')
  for stock in data.get('unallocated_stock',[]):
   if stock['document_key']!=key:continue
   no=mappings.get('stock:'+stock['id'])
   if no is not None:
    line=by_no.get(str(no))
    if not line:raise ValueError('Выберите строку оригинала для прежнего складского остатка')
    stock.update(source_line_key=line['key'],sku=line.get('sku',''),name=line['name'],unit=line['unit'])
  for ai,alloc in enumerate(inv.get('allocations',[])):
   original=alloc.get('source_lines') or alloc.get('updates') or []
   if not original:
    supplied=mappings.get('allocation:'+str(ai))
    if not isinstance(supplied,list) or not supplied:raise ValueError('Укажите строки, количества и суммы прежнего распределения')
    original=[]
    for li,entry in enumerate(supplied):
     if aw.number(entry.get('qty'))<=0 or not isinstance(entry.get('amount_cents'),int) or entry['amount_cents']<=0:raise ValueError('Некорректное количество или сумма прежней строки')
     original.append({'qty_used':entry['qty'],'amount_used':entry['amount_cents']/100})
     mappings[str(ai)+':'+str(li)]=entry.get('line_no')
   fixed=[]
   for li,old in enumerate(original):
    split=mappings.get('split:'+str(ai)+':'+str(li))
    if split is not None:
     if not isinstance(split,list) or len(split)<2:raise ValueError('Укажите не менее двух компонентов комплекта')
     parts=[]
     for entry in split:
      no=str(entry.get('line_no'));line=by_no.get(no)
      if not line:raise ValueError('Выберите строку оригинала для каждого компонента')
      qty=aw.number(entry.get('source_qty'));covered=aw.number(entry.get('covered_qty'));amount=entry.get('amount_cents');reason=str(entry.get('reason','')).strip()
      if qty<=0 or covered<0 or not isinstance(amount,int) or amount<=0 or not reason:raise ValueError('Укажите количество, сумму и основание состава комплекта')
      qty_by[no]=qty_by.get(no,Decimal(0))+qty;cost_by[no]=cost_by.get(no,0)+amount
      parts.append({**old,'key':line['key'],'source_name':line['name'],'source_sku':line.get('sku',''),'source_qty':line['qty'],'source_unit':line['unit'],'source_qty_used':float(qty),'qty_used':float(covered),'amount_used':amount/100,'identity_confirmed':True,'identity_reason':reason})
     if sum(aw.cents(l['amount_used']) for l in parts)!=aw.cents(old.get('amount_used',old.get('material_cost_used',0))):raise ValueError('Суммы компонентов не совпадают с прежней стоимостью комплекта')
     if sum(aw.number(l['qty_used']) for l in parts)!=aw.number(old.get('qty_used',0)):raise ValueError('Количество комплектов не должно изменяться при расшифровке')
     fixed.extend(parts);continue
    no=str(mappings.get(str(ai)+':'+str(li),''));line=by_no.get(no)
    if not line:raise ValueError('Выберите строку оригинала для каждого прежнего распределения')
    amount=aw.cents(old.get('amount_used',old.get('material_cost_used',0)));qty=aw.number(old.get('source_qty_used',old.get('qty_used',0)))
    qty_by[no]=qty_by.get(no,Decimal(0))+qty;cost_by[no]=cost_by.get(no,0)+amount
    fixed.append({**old,'key':line['key'],'source_name':line['name'],'source_sku':line.get('sku',''),'source_qty':line['qty'],'source_unit':line['unit'],'amount_used':amount/100})
   if sum(aw.cents(l['amount_used']) for l in fixed)!=aw.cents(alloc['amount']):raise ValueError('Строки прежней записи не объясняют её сумму')
   alloc['source_lines']=fixed
  for no,qty in qty_by.items():
   if qty>aw.number(by_no[no]['qty']) or cost_by[no]>by_no[no]['amount_cents']:raise ValueError('Прежнее распределение превышает выбранную строку; проверьте версию счёта')
  data.setdefault('intake_audit',[]).append({'document_key':key,'action':'repair_usage','actor':str(actor),'mappings':mappings})
  queue.atomic(path,data)
 return install(key,{'lines':parsed,'total_cents':aw.cents(inv['total']),'date':inv.get('date','')},path,actor)

def restore_bundle_receipts(path=queue.PROCUREMENT,loader=None):
 """Recover a whole kit only from the exact supplier SKUs already stored locally.

 No heuristic names, cost changes or newly inferred customer quantities. Each
 component consumes its supplier quantity, while the kit coverage is counted once.
 """
 import deal_ledger
 loader=loader or deal_ledger.load;changed=[]
 for inv in read(path,{}).get('invoices',[]):
  lines=inv.get('lines',[]);allocs=inv.get('allocations',[])
  if inv.get('intake',{}).get('state')!='review' or len(lines)<2 or len(allocs)!=1:continue
  alloc=allocs[0];old=alloc.get('source_lines') or alloc.get('updates') or []
  if len(old)!=1 or not old[0].get('offer_row_id'):continue
  if aw.cents(alloc['amount'])!=sum(l['amount_cents'] for l in lines):continue
  rows=loader(str(alloc['deal_id']));hits=[r for r in rows if str(r.get('offer_row_id'))==str(old[0]['offer_row_id'])]
  if len(hits)!=1:continue
  row=hits[0];skus=[str(l.get('sku','')).strip().casefold() for l in lines]
  stored=str(row.get('supplier_sku','')).strip().casefold().split('/')
  if not all(skus) or sorted(skus)!=sorted(stored) or len(set(skus))!=len(skus):continue
  quantity=aw.number(old[0].get('qty_used',0))
  if quantity<=0 or any(aw.number(l['qty'])!=quantity for l in lines):continue
  parts=[{'line_no':l['line_no'],'source_qty':l['qty'],'covered_qty':str(quantity if n==0 else 0),'amount_cents':l['amount_cents'],'reason':'Состав комплекта по сохранённому артикулу поставщика: '+row['supplier_sku']} for n,l in enumerate(lines)]
  repair_usage(inv['document_key'],{'split:0:0':parts},path,actor='restore_exact_bundle');changed.append(inv['document_key'])
 return changed


def restore_added_receipts(path=queue.PROCUREMENT,loader=None):
 """Recover legacy added-item receipts from exact source, quantity and money proof."""
 from product_matching import canonical
 import deal_ledger
 loader=loader or deal_ledger.load;changed=[]
 with locked(path):
  data=read(path,{})
  for inv in data.get('invoices',[]):
   if inv.get('intake',{}).get('state')!='review' or not inv.get('lines'):continue
   restored=False
   for alloc in inv.get('allocations',[]):
    if alloc.get('source_lines') or alloc.get('updates') or not alloc.get('additions'):continue
    ledger=loader(str(alloc['deal_id']));result=[]
    for addition in alloc['additions']:
     qty=aw.number(addition.get('qty',0));amount=aw.cents(aw.number(addition.get('cost_material',0))*qty)
     lines=[l for l in inv['lines'] if canonical(l['name'])==canonical(addition.get('source_name') or addition.get('name')) and aw.number(l['qty'])==qty and l['amount_cents']==amount]
     targets=[e for e in ledger if canonical(e['name'])==canonical(addition['name']) and aw.number(e['qty'])==qty and e.get('offer_row_id') and e.get('offer_id')]
     if len(lines)!=1 or len(targets)!=1:break
     l=lines[0];e=targets[0]
     result.append({'key':l['key'],'source_name':l['name'],'source_sku':l.get('sku',''),'source_qty':l['qty'],'source_unit':l['unit'],'qty_used':float(qty),'amount_used':amount/100,'offer_row_id':str(e['offer_row_id']),'offer_id':str(e['offer_id'])})
    else:
     if sum(aw.cents(x['amount_used']) for x in result)==aw.cents(alloc['amount']):alloc['source_lines']=result;restored=True
   if restored:changed.append(inv['document_key'])
  if changed:
   data.setdefault('intake_audit',[]).append({'action':'restore_added_receipts','documents':changed,'money_changed':False});queue.atomic(path,data)
 for key in changed:
  inv=next(i for i in data['invoices'] if i['document_key']==key)
  install(key,{'lines':inv['lines'],'date':inv.get('date',''),'total_cents':aw.cents(inv['total'])},path,actor='confirmed_addition_receipt')
 return changed
