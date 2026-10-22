"""Write plain XLSX worksheets with XlsxWriter, without structured-table metadata."""
import datetime as dt
import xlsxwriter
from openpyxl.utils.cell import column_index_from_string

def save_workbook(source, path):
    with xlsxwriter.Workbook(str(path), {'strings_to_formulas': False, 'strings_to_urls': False}) as out:
        formats={}
        for original in source.worksheets:
            ws=out.add_worksheet(original.title)
            ws.hide_gridlines(2)
            ws.freeze_panes(1,min(2,original.max_column-1))
            for letter,dim in original.column_dimensions.items():
                col=column_index_from_string(letter)-1
                ws.set_column(col,col,dim.width or 20,None,{'hidden':bool(dim.hidden)})
            ws.set_row(0,34)
            if original.max_row>1:ws.autofilter(0,0,original.max_row-1,original.max_column-1)
            for row in original:
                for cell in row:
                    v=cell.value
                    if v is None:continue
                    idx=cell.row-1;col=cell.column-1
                    key=(idx==0,idx%2==1,cell.number_format)
                    if key not in formats:
                        props={'valign':'top','text_wrap':True}
                        if idx==0:props.update(bold=True,font_color='#FFFFFF',bg_color='#153E4D')
                        elif idx%2==1:props['bg_color']='#EFF5F7'
                        if cell.number_format!='General':props['num_format']=cell.number_format
                        formats[key]=out.add_format(props)
                    fmt=formats[key]
                    if isinstance(v,str):ws.write_string(idx,col,v,fmt)
                    elif isinstance(v,(dt.date,dt.datetime)):ws.write_datetime(idx,col,v,fmt)
                    elif isinstance(v,bool):ws.write_boolean(idx,col,v,fmt)
                    elif isinstance(v,(int,float)):ws.write_number(idx,col,v,fmt)
                    else:raise TypeError(type(v))
