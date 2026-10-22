import json,subprocess,zipfile,re
from pathlib import Path
import openpyxl,xlrd
root=Path(__file__).resolve().parent/'evidence/audit45';out=[]
eligible=set()
for msgpath in root.glob('*_*.json'):
 msg=json.loads(msgpath.read_text())
 if not isinstance(msg,dict):continue
 for f in msg.get('files',[]):
  if re.search(r'сч|invoice|пкб|акт|сверк|упд|наклад|заказ|кп|^\d',f['name'],re.I):eligible.add(f['path'])
for filename in sorted(eligible):
 p=Path(filename)
 if p.suffix not in ('.pdf','.xlsx','.xls','.docx','.txt'):continue
 dest=p.with_suffix(p.suffix+'.extracted')
 if not dest.exists():
  try:
   if p.suffix=='.pdf':
    x=subprocess.run(['pdftotext','-layout',str(p),'-'],capture_output=True,timeout=20);text=x.stdout.decode(errors='replace')
   elif p.suffix=='.xlsx':
    w=openpyxl.load_workbook(p,data_only=True,read_only=True);text='\n'.join(' | '.join(str(c) for c in row if c is not None) for sh in w for row in sh.values);w.close()
   elif p.suffix=='.xls':
    w=xlrd.open_workbook(p);text='\n'.join(' | '.join(str(c) for c in sh.row_values(i) if c!='') for sh in w.sheets() for i in range(sh.nrows))
   elif p.suffix=='.docx':
    import docx
    w=docx.Document(p);text='\n'.join(x.text for x in w.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in w.tables for row in t.rows)
   else:text=p.read_text(errors='replace')
   dest.write_text(text)
  except Exception as e:print(p.name,type(e).__name__);continue
 out.append({'path':str(p),'text':dest.read_text()})
(root/'documents.json').write_text(json.dumps(out,ensure_ascii=False));print('documents',len(out))
