import sys,json,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook,load_workbook
from report import sheet
from xlsx_compat import save_workbook
R=Path(__file__).resolve().parents[1];D=R/'evidence/final111';a=json.loads((D/'audit.json').read_text());ar=json.loads((D/'supplier_archive_audit.json').read_text());clients=json.loads((D/'client_archive.json').read_text());sync=json.loads((D/'registry_sync.json').read_text());assert not a['errors'] and len(a['operations'])==111 and not ar['missing'] and len(clients)==14
rows={x['id']:x for x in json.loads((R/'reports/ledger.json').read_text())};w=Workbook();w.remove(w.active)
sheet(w,'Итог',['Показатель','Результат'],[['Проверено финансовых операций',111],['Банковских платежей',109],['Оплаты по сделкам (товарное движение0)',49],['Банковские комиссии (услуга уже списана)',51],['Проценты по депозитам (нетоварный доход)',11],['Клиентских счетов: проверены обе стороны связи',14],['Оплат поставщикам: счета в архиве подтверждены',34],['Исправлена связь счёта со сделкой','05487 (ID798) → сделка851; сумма и оплата сохранены'],['Ошибки контрольной проверки',0],['Область завершения','111 уже созданных финансовых операций; не вся выписка из337 строк'],['Приёмка и отгрузка','Для49 торговых оплат товарное движение0. Проводятся отдельно по документам.'],['Даичи DAER-023701','Фактическая оплата2561559,07 ₽ по указанному банком курсу84,5863; план2548903,64 ₽ сохранён; разница12655,43 ₽ не задолженность.']])
sheet(w,'Все 111 операций',['Операция CRM','Дата','Компания','Платёж №','Контрагент','Приход, ₽','Расход, ₽','Сделка','Счёт','Тип учёта','Контроль'],[[p['operation_id'],p['date'],rows[p['bank_operation_id']]['company'],p['document_number'],rows[p['bank_operation_id']]['counterparty'],p['amount_cents']/100 if p['kind']=='income' else 0,p['amount_cents']/100 if p['kind']=='expense' else 0,p.get('deal_id',''),p['invoice_number'],{'cash_only':'Оплата; товарное движение0','bank_service_settled':'Комиссия; банковская услуга оказана','nontrade_income':'Проценты; нетоварный доход'}[p['payment_semantics']],'Сумма, дата, счёт, контрагент, связи и отсутствие дубля проверены'] for p in a['operations']],(6,7))
sheet(w,'Клиентские счета',['ID','Номер','Сумма, ₽','Оплачено, ₽','Статус'],[[i['id'],i['number'],i['total_cents']/100,i['paid_cents']/100,'Оплачен' if i['status']=='paid' else 'Частично оплачен'] for i in a['invoices']],(3,4))
sheet(w,'Синхронизация закупок',['Счёт','Оплачено, ₽','Остаток, ₽','Статус','Добавленные связи CRM','Разница с планом, ₽'],[[x['invoice'],x['paid_total'],x['payment_balance'],x['status'],', '.join(x['added_operation_ids']),x['plan_difference']] for x in sync],(2,3,6))
sheet(w,'Архив поставщиков',['Операция','Счёт','Сделка','OneDrive'],[[x['operation_id'],x['invoice'],x['deal_id'],d['remote'].removeprefix('onedrive:')] for x in ar['operations'] for d in x['documents']])
sheet(w,'Архив клиентских счетов',['Счёт','Сделка','OneDrive'],[[x['number'],x['deal_id'],x['remote'].removeprefix('onedrive:')] for x in clients])
p=R/'reports/Проверка_111_операций.xlsx';save_workbook(w,p);t=load_workbook(p,read_only=True);assert t['Все 111 операций'].max_row==112;t.close()
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for n in z.namelist():
  if n.endswith('.xml'):ET.fromstring(z.read(n))
print(p)
