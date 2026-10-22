import sys,json,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook,load_workbook
from report import sheet
from xlsx_compat import save_workbook
R=Path(__file__).resolve().parents[1];d=json.loads((R/'evidence/payment_correction/verified.json').read_text());w=Workbook();w.remove(w.active)
sheet(w,'Что исправлено',['Показатель','Результат'],[['Проверено финансовых операций',len(d['operations'])],['Новых финансовых операций при исправлении',len(d['new_operation_ids'])],['Счета клиентов со связанными оплатами',len(d['invoices'])],['Полностью оплаченных счетов',sum(x['status']=='paid' for x in d['invoices'])],['Частично оплаченных счетов',sum(x['status']=='drawn' for x in d['invoices'])],['Движение товара по банковским операциям','0: отгрузка, приёмка и выполнение работ подтверждаются отдельно'],['Новые карточки поставщиков по документам',d['supplier_cards_created']],['Контрольная дата','06.10.2026 — внутренняя дата проверки незакрытых товарных обязательств; не договорный срок поставки'],['Время проверки',d['verified_at']]])
sheet(w,'Счета клиентов',['Счёт','ID счёта','Сумма счёта, ₽','Оплачено, ₽','Осталось, ₽','Статус','Финансовые операции','Ссылка'],[[x['number'],x['id'],x['total_cents']/100,x['paid_cents']/100,(x['total_cents']-x['paid_cents'])/100,'Оплачен' if x['status']=='paid' else 'Частичная оплата; не закрыт',', '.join(x['operation_ids']),'https://cclimat.megaplan.ru/invoices/'+x['id']+'/card/'] for x in d['invoices']],(3,4,5))
sheet(w,'Исправленные операции',['ID операции','Сделка','Направление','Дата оплаты','Платёж №','Деньги, ₽','Товарное движение, ₽','Контрагент ID','Счёт','ID счёта CRM'],[[x['operation_id'],x['deal_id'],x['kind'],x['date'],x['bank_number'],x['cash_cents']/100,x['goods_cents']/100,x['contractor_id'],x['invoice_number'],x['crm_invoice_id'] or 'Счёт поставщика в архиве/реестре закупок'] for x in d['operations']],(6,7))
p=R/'reports/Проверка_оплат_без_отгрузок_06.10.2026.xlsx';save_workbook(w,p);v=load_workbook(p,read_only=True);assert v['Исправленные операции'].max_row==38;assert v['Счета клиентов'].max_row==15;v.close()
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.endswith('.xml'):ET.fromstring(z.read(name))
print(p)
