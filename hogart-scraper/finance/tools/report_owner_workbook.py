from pathlib import Path as _QueuePath
if (_QueuePath(__file__).resolve().parents[1]/'data/review_registry.json').exists():
    raise SystemExit('Архивный обработчик отключён: рабочий реестр — finance/data/review_registry.json; обновление — finance/review_registry.py')
import sys,json,zipfile,xml.etree.ElementTree as ET,datetime,collections
from pathlib import Path
sys.path[:0]=['/home/kocher.av/.report-venv/lib/python3.14/site-packages','/home/kocher.av/hogart-scraper/finance']
from openpyxl import Workbook,load_workbook
from report import sheet,HEADERS,op_values
from xlsx_compat import save_workbook
R=Path(__file__).resolve().parents[1];D=R/'evidence/owner_workbook_20261007'
v=json.loads((D/'verified.json').read_text());assert not v['errors'] and len(v['operations'])==17
answers=json.loads((D/'answers.json').read_text());answered={a['number']:a for a in answers if a['answer']};rows=json.loads((R/'reports/ledger.json').read_text());byid={r['id']:r for r in rows};parts=list(json.loads((R/'data/posting_journal.json').read_text())['operations'].values());assert all(p['state']=='linked' for p in parts)
qpath=R/'evidence/finish_all/owner_questions.json';q=json.loads(qpath.read_text());before=json.loads((D/'questions_before.json').read_text());resolved={g['number']:g for g in q['resolved_groups']}
for g in q['groups']:
 if g['number'] in answered:
  ans=answered[g['number']];ops=[p['operation_id'] for p in parts if p['bank_operation_id'] in g['bank_ids']]
  g.update(owner_answer=ans['answer'],answer=ans['answer'],owner_question_closed=True,resolved=True,resolution='Ответ принят из Excel07.10.2026. '+('Платёж проведён/проверен в CRM.' if ops else 'Назначение вне проектов зафиксировано; техническое оформление отдельно, повторный вопрос об объекте не требуется.'),crm_operation_ids=ops)
  resolved[g['number']]=g
for num,ans in answered.items():
 if num in resolved:resolved[num].update(owner_answer=ans['answer'],answer=ans['answer'])
q['groups']=[g for g in q['groups'] if g['number'] not in answered];q['resolved_groups']=list(resolved.values());assert len(q['groups'])==19
posted={p['bank_operation_id'] for p in parts};install={r['id'] for r in rows if r.get('installer_payment_id')};q.update(posted_bank_rows=len(posted),unposted_rows=len(rows)-len(posted)-len(install));qpath.write_text(json.dumps(q,ensure_ascii=False,indent=2))
openids={oid for g in q['groups'] for oid in g['bank_ids']};technical=[r for r in rows if r['id'] not in posted|install|openids]
def reason(r):
 if r['category']=='Личные расходы':return 'Личный характер установлен ранее. Отдельное оформление расчётов с владельцем; к проектам не относить.'
 if r['inn']=='7728168971':return 'Назначение вне проектов принято. Для корректного проведения нужна проверка реквизитов фактического продавца по чеку: банковский ИНН Альфа-Банка нельзя подставлять вместо продавца.'
 if r['category']=='Перевод между своими счетами' or r['group']!='Внешняя':return 'Внутреннее перемещение; сверка собственного счёта и встречной стороны без повторного расхода.'
 if not r.get('owner_answer'):return 'Новая операция, отсутствовавшая в присланной версии: ассистенту найти счёт и установить объект. Ваш ответ по этой операции не получен и не приписывается.'
 return 'Ответ о собственных нуждах принят; уточнение объекта не требуется. Ассистенту завершить идентификацию контрагента и оформление без проектной сделки.'
changes=json.loads((D/'actions.json').read_text());newdeal=json.loads((D/'new_deal_state.json').read_text())['deal_id'];archives=json.loads((D/'archive_results.json').read_text())
w=Workbook();w.remove(w.active)
sheet(w,'Итог',['Показатель','Результат'],[
 ['Дата обработки','07.10.2026'],['Ответов в присланном файле',len(answered)],['Вопросов до обработки с учётом прежней работы',len(before['groups'])],['Закрыто вопросов из текущего остатка',len(before['groups'])-len(q['groups'])],['Осталось вопросов к владельцу',len(q['groups'])],['Банковских строк в оставшихся вопросах',len(openids)],['Новых финансовых операций по этим ответам',len(parts)-244],['Исправлено прежних привязок платежей',2],['Новая сделка',f'№{newdeal} — Квартира 44 м², электрика. Создана по ответу №4; продажная цена не выдумана.'],['Всего банковских строк в актуальном реестре',len(rows)],['Отличие от присланной версии','В исходном файле337 строк; актуальный реестр включает ещё1 поступившую позднее строку.'],['Строк банка проведено в CRM',len(posted)],['Финансовых операций CRM с учётом разбивок',len(parts)],['Монтажники уже учтены отдельно',len(install)],['Без записи CRM: вопросы владельцу',len(openids)],['Без записи CRM: оформление отдельно',len(technical)],['Контроль полноты',f'{len(posted)} CRM + {len(install)} монтажники + {len(openids)} строки вопросов + {len(technical)} оформление = {len(rows)}'],['Собственные нужды','Ответы учтены вне проектных сделок. Это не автоматическое признание всех платежей личными или налоговыми расходами.'],['Оплата и отгрузка','Новые/исправленные платежи имеют нулевое товарное движение; отгрузка и приёмка не подтверждались.'],['Как продолжать','Заполняйте только лист «Только ваши вопросы». Уже принятые ответы сохранены на отдельном листе. Технический остаток не является повторной анкетой.']])
ws=w['Итог'];ws.column_dimensions['A'].width=56;ws.column_dimensions['B'].width=110
heads=['№ вопроса','Поставщик / получатель','Счёт / заказ','Платежи','Сумма платежей, ₽','Что осталось уточнить','Уже проверено','Ваш ответ']
vals=[[g['number'],g['supplier'],g['invoices'],g['payments'],g['amount_cents']/100,g['question'],g['evidence'],''] for g in q['groups']]
ws=sheet(w,'Только ваши вопросы',heads,vals,(5,))
for col,width in zip('ABCDEFGH',[12,33,30,45,19,70,85,45]):ws.column_dimensions[col].width=width
accepted=[]
for ans in answers:
 if not ans['answer']:continue
 rr=[byid[k] for k in ans['bank_ids']];pp=[p for p in parts if p['bank_operation_id'] in ans['bank_ids']];scope='; '.join(sorted({p['deal_id'] for p in pp if p.get('deal_id')}))
 if any(a['question']==ans['number'] and a.get('old_deal') for a in changes):status='Прежняя операция исправлена; деньги повторно не создавались'
 elif pp:status='Проведено / ранее проведено, ответ учтён'
 else:status='Назначение вне проектов принято; оформление отдельно'
 detail=''
 if ans['number']==3:detail='890:144356,49 ₽; 894:345908,44 ₽. Часть99390,96 ₽ перенесена с861 на894; расчёт по вводам44:76 сохранён.'
 if ans['number']==7:detail='Вся оплата107230,90 ₽ на862. Ответ не выбирает итоговую версию счёта; долг и приёмка не определялись.'
 if ans['number']==9:detail='20827,55 ₽: исправлено883 →862.'
 if ans['number']==34:detail='Сохранена установленная ранее разбивка: счёт1280 18000 ₽ →856, счёт1299 59760 ₽ →854; оба Зоовет.'
 if ans['number']==70:detail='Оплата уже на858. Расхождение версий счёта и доплаты остаётся в технической сверке документов; объект больше не спрашивается.'
 accepted.append([ans['number'],ans['supplier'],ans['invoice'],sum(r['amount_cents'] for r in rr)/100,ans['answer'],ans['note'],status,scope or 'Без проектной сделки',', '.join(p['operation_id'] for p in pp),detail])
ws=sheet(w,'Обработанные ответы',['№ вопроса','Поставщик / получатель','Счёт / заказ','Сумма, ₽','Ваш ответ — сохранён','Ваше примечание','Результат обработки','Сделки','Операции CRM','Пояснение'],accepted,(4,))
for col,width in zip('ABCDEFGHIJ',[12,32,28,18,55,32,60,24,26,85]):ws.column_dimensions[col].width=width
techvalues=[[r['company'],r['date'],r['number'],r['counterparty'],r['amount_cents']/100,r.get('owner_answer',''),reason(r),r['id']] for r in technical]
ws=sheet(w,'Технический остаток',['Компания','Дата','№ банка','Получатель','Сумма, ₽','Принятый ответ','Что осталось оформить — не вопрос об объекте','ID'],techvalues,(5,));ws.column_dimensions['D'].width=35;ws.column_dimensions['F'].width=35;ws.column_dimensions['G'].width=110;ws.column_dimensions['H'].hidden=True
sheet(w,'Сверка документов',['Счёт','Сделка','Что уже завершено','Осталось проверить'],[['ПКБ-00002943','862','Весь платёж107230,90 ₽ проведён на Примаверу по ответу владельца.','Сохранены версии164460,55 и114370,90 ₽. Ответ «ДА целиком» подтверждает объект платежа, но не итоговую стоимость заказа. Сверить версии перед расчётом задолженности.'],['208098','858','Обе оплаты87642,35 ₽ уже на вентиляции1этажаJetCar.','Найденный счёт74835,70 ₽ меньше оплат на12806,65 ₽. Проверить окончательную версию/изменение заказа; не создавать повторный платёж.']])
sheet(w,'Операции',HEADERS,[op_values(r) for r in rows],(6,7))
sheet(w,'Проведено в CRM',['Операция CRM','Компания','Дата','№ банка','Получатель','Сумма, ₽','Сделка','Счёт','Статья','Товарное движение, ₽'],[[p['operation_id'],byid[p['bank_operation_id']]['company'],p['date'],p['document_number'],byid[p['bank_operation_id']]['counterparty'],p['amount_cents']/100,p.get('deal_id') or '',p['invoice_number'],p.get('category','Оплата по сделке'),p.get('secondary_cents',0)/100] for p in parts],(6,10))
sheet(w,'Монтажники без дублей',['Дата','№ банка','Получатель','Сумма, ₽','Запись монтажника'],[[r['date'],r['number'],r['counterparty'],r['amount_cents']/100,r['installer_payment_id']] for r in rows if r.get('installer_payment_id')],(4,))
source=json.loads((D/'source.json').read_text())['source'];original=load_workbook(source,read_only=True)
for name in ['Клиентские счета','Архив документов']:
 originalrows=list(original[name].values);sheet(w,name,list(originalrows[0]),[list(v) for v in originalrows[1:]])
archive=w['Архив документов']
for x in archives:archive.append([x['number'],x['deal_id'],'',x['remote'].removeprefix('onedrive:'),'Назначение по ответам владельца07.10.2026. Проверено наличие файла.'])
# Old invoice copies are evidence history, not new allocations. Explicitly flag superseded locations.
for row in archive.iter_rows(min_row=2):
 if str(row[0].value)=='212/099912362' and str(row[1].value)=='861':row[4]='Историческая копия: после ответа07.10 назначение исправлено на894/890; на861 не учитывать.'
 if str(row[0].value)=='ПКБ-00003931' and str(row[1].value)=='883':row[4]='Историческая копия: после ответа07.10 платёж относится к862 Примавера; на883 не учитывать.'
original.close()
p=R/'reports/Выписка_после_ваших_ответов_07.10.2026.xlsx';save_workbook(w,p)
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.endswith('.xml'):ET.fromstring(z.read(name))
check=load_workbook(p,read_only=True);assert check['Только ваши вопросы'].max_row==20 and check['Обработанные ответы'].max_row==47 and check['Операции'].max_row==len(rows)+1
assert len(posted)+len(install)+len(openids)+len(technical)==len(rows)
assert [row[0] for row in list(check['Только ваши вопросы'].values)[1:]]==[g['number'] for g in q['groups']]
check.close()
summary={'workbook':str(p),'answers_processed':46,'questions_before':47,'questions_closed':28,'questions_remaining':19,'new_crm_operations':len(parts)-244,'corrected_existing':2,'deal_created':newdeal,'bank_rows':len(rows),'posted_bank_rows':len(posted),'crm_parts':len(parts),'technical_rows_without_crm':len(technical),'questions_bank_rows':len(openids),'xlsx_zip_xml_openpyxl_validated':True}
(D/'completion.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2));print(json.dumps(summary,ensure_ascii=False))
