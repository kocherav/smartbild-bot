from pathlib import Path as _QueuePath
if (_QueuePath(__file__).resolve().parents[1]/'data/review_registry.json').exists():
    raise SystemExit('Архивный обработчик отключён: рабочий реестр — finance/data/review_registry.json; обновление — finance/review_registry.py')
import sys,json,zipfile,xml.etree.ElementTree as ET,collections
from pathlib import Path
sys.path[:0]=['/home/kocher.av/.report-venv/lib/python3.14/site-packages','/home/kocher.av/hogart-scraper/finance']
from openpyxl import Workbook,load_workbook
from report import sheet,HEADERS,op_values
from xlsx_compat import save_workbook
from post_confirmed import atomic
R=Path(__file__).resolve().parents[1];D=R/'evidence/final_answers_20261007';v=json.loads((D/'verified.json').read_text());assert not v['errors'] and len(v['operations'])==21
answers=json.loads((D/'answers.json').read_text());answered={a['number']:a for a in answers if a['answer']};rows=json.loads((R/'reports/ledger.json').read_text());byid={r['id']:r for r in rows};parts=list(json.loads((R/'data/posting_journal.json').read_text())['operations'].values());assert len(parts)==278 and all(p['state']=='linked' for p in parts)
qpath=R/'evidence/finish_all/owner_questions.json';q=json.loads((D/'questions_before.json').read_text());resolved={g['number']:g for g in q['resolved_groups']};remain=[]
for g in q['groups']:
 num=g['number'];ans=answered.get(num)
 if ans and num!=37:
  ops=[p['operation_id'] for p in parts if p['bank_operation_id'] in g['bank_ids']];assert ops
  g.update(owner_answer=ans['answer'],answer=ans['answer'],owner_question_closed=True,resolved=True,resolution='Ответ принят из обновлённого Excel07.10.2026; денежные операции проведены и проверены в CRM. Товарное движение0; дополнительных удержаний с монтажника не создавалось.',crm_operation_ids=ops);resolved[num]=g
 else:
  if num==37:
   g['owner_answer']=ans['answer'];g['question']='Для Автошума.нет это кондиционер Илье (сделка869) или другой объект? Если другой — достаточно названия. В счетах AS20HQJ, а в869 указанHSU-09; по всей истории869 совпадения AS20 нет. Зоовет — кабинет директора и одну штуку на собственные нужды уже сохранил.';g['evidence']+=' Новое указание: по1шт Зоовет/Автошума/собственные нужды. Проверены все сделки Автошумы, живые855/869 и полная история24/33события. Не менял клиентские модели; не распределял разные закупочные цены без подтверждённой связи.'
  remain.append(g)
etm=next(r for r in rows if r['date']=='2026-10-06' and r['number']=='145');remain.append({'number':71,'supplier':etm['counterparty'],'invoices':'212/099912439','payments':'ИНСИС 06.10.2026 №145','amount_cents':etm['amount_cents'],'bank_ids':[etm['id']],'question':'Новая выписка06.10: это электрика6этажаJetCar, перенос компрессора или другой объект? В счёте30наконечников35мм², гофра50мм15м, мелкие наконечники и доставка. Если несколько объектов — достаточно указать, что куда.','evidence':'Счёт11690,38 найден в info@innsis.ru UID13739; сумма, ИНН продавца7842224734 и покупателя7743926877 сверены. Письмо содержит запрос счета и снимок заказа, но не название объекта. Рядом шли два проекта с кабелем35мм²; номер сделки не выдуман.'})
posted={p['bank_operation_id'] for p in parts};install={r['id'] for r in rows if r.get('installer_payment_id')};openids={oid for g in remain for oid in g['bank_ids']};technical=[r for r in rows if r['id'] not in posted|install|openids];assert len(posted)+len(install)+len(openids)+len(technical)==len(rows)
for oid in posted:assert sum(p['amount_cents'] for p in parts if p['bank_operation_id']==oid)==byid[oid]['amount_cents']
q.update(groups=remain,resolved_groups=list(resolved.values()),posted_bank_rows=len(posted),unposted_rows=len(rows)-len(posted)-len(install));atomic(qpath,q)
def reason(r):
 if r['category']=='Перевод между своими счетами':return 'Нужны встречные выписки Т-Банка/Озон Банка для подтверждения зачислений. Назначение установлено: внутренний перевод, не расход по объекту.'
 if r['inn']=='7728168971':return 'Назначение уже принято: '+('личное' if r['category']=='Личные расходы' else 'собственные нужды')+'. По выписке виден карточный терминал, юридический продавец/документ не подтверждён. В доступных электронных чеках точного соответствия не найдено. Не подставлен Альфа-Банк как продавец; нужен чек/отчёт по корпоративной карте.'
 return 'Назначение «собственные нужды» принято. Получатель и ИНН есть в банке, но для карточки CRM не найден обязательный адрес/первичный документ. Повторно объект не спрашивается.'
w=Workbook();w.remove(w.active);newdeal=json.loads((D/'new_deal_state.json').read_text())['deal_id'];arc=json.loads((D/'archive_results.json').read_text());outside=json.loads((D/'outside_archive_results.json').read_text());assert len(arc)==4 and len(outside)==7
summary={'date':'07.10.2026','answers_explicit':14,'answers_including_linked_43':15,'old_questions_before':19,'old_questions_closed':14,'old_questions_remaining':5,'new_statement_questions':1,'questions_remaining':6,'new_crm_operations':21,'new_bank_rows_posted':20,'crm_parts':len(parts),'posted_bank_rows':len(posted),'bank_rows':len(rows),'installer_rows':len(install),'question_bank_rows':len(openids),'technical_rows':len(technical),'new_deal':newdeal,'archived_documents':len(arc)+len(outside)}
sheet(w,'Итог',['Показатель','Результат'],[['Дата обработки','07.10.2026'],['Ответов из нового файла','14 прямых ответов + пункт43 по прямому указанию в42'],['Закрыто прежних вопросов','14 из19'],['Осталось прежних вопросов',5],['Добавлен вопрос по новой выписке06.10',1],['Новые финансовые операции CRM',21],['Новых строк банка проведено',20],['Новая сделка',f'№{newdeal} — JetCar, закупки Элевел061898/065961, по вашему указанию'],['Всего строк банка',len(rows)],['Строк банка проведено в CRM',len(posted)],['Операций CRM с разбивками',len(parts)],['Банковские выплаты монтажникам уже учтены отдельно',len(install)],['Строк в вопросах владельцу',len(openids)],['Строк с установленным назначением без записи CRM',len(technical)],['Контроль полноты',f'{len(posted)}+{len(install)}+{len(openids)}+{len(technical)}={len(rows)}'],['Полное завершение','Пока нет: остались6 вопросов и оформление подтверждённых внепроектных/внутренних движений по первичным документам. Они не скрыты в числе завершённых.'],['Оплата ≠ отгрузка','У всех21новых финансовых операций SecondarySum=0; деньги проведены, товарная отгрузка/приёмка не выдумана.'],['Деревянко','5закупок у Трио проведены отдельно. Удержания126000 и31000 уже есть в утверждённых отчётах; повторного уменьшения оплаты монтажнику не создавалось.'],['Элевел054832/054835','98894,70 на сделку851;19832,18 на собственные нужды. Доставка900 из первого документа не покрыта этим платежом.'],['Архив OneDrive',f'{len(arc)}документа по сделкам + {len(outside)}внепроектных/Деревянко; наличие и размер проверены.'],['Что нужно от вас','Ответить только на6вопросов. Для полного документального закрытия: встречные выписки по переводам и доступные чеки/авансовый отчёт по картам; назначение уже отвеченных покупок повторно не спрашивается.']])
w['Итог'].column_dimensions['A'].width=57;w['Итог'].column_dimensions['B'].width=115
heads=['№ вопроса','Поставщик / получатель','Счёт / заказ','Платежи','Сумма платежей, ₽','Что осталось уточнить','Уже проверено','Ваш ответ']
ws=sheet(w,'Только ваши вопросы',heads,[[g['number'],g['supplier'],g['invoices'],g['payments'],g['amount_cents']/100,g['question'],g['evidence'],''] for g in remain],(5,))
for c,wide in zip('ABCDEFGH',[12,33,28,45,19,85,100,45]):ws.column_dimensions[c].width=wide
accepted=[]
for a in answers:
 if not a['answer']:continue
 ps=[p for p in parts if p['bank_operation_id'] in a['bank_ids']];target='; '.join(sorted({p['deal_id'] or p.get('category','Без сделки') for p in ps}));status='Проведено и проверено в CRM' if ps else 'Ответ сохранён; осталось уточнение конкретной сделки Автошумы'
 accepted.append([a['number'],a['supplier'],a['invoices'],a['amount_cents']/100,a['answer'],status,target,', '.join(p['operation_id'] for p in ps)])
sheet(w,'Новые ответы',['№','Поставщик','Счета','Сумма, ₽','Ваш ответ — сохранён','Результат','Сделка / назначение','Операции CRM'],accepted,(4,))
sheet(w,'Остаток оформления',['Компания','Дата','№ банка','Получатель / терминал','Сумма, ₽','Установленное назначение','Что требуется для завершения','ID'],[[r['company'],r['date'],r['number'],r['purpose'] if r['inn']=='7728168971' else r['counterparty'],r['amount_cents']/100,r.get('owner_answer') or r['category'],reason(r),r['id']] for r in technical],(5,));w['Остаток оформления'].column_dimensions['H'].hidden=True;w['Остаток оформления'].column_dimensions['G'].width=110
sheet(w,'Деревянко — сверка',['Вопрос','Счета','Оплата поставщику, ₽','Удержание по отчёту, ₽','Результат'],[[15,'52475/52550',63838.56,None,'По ответу владельца: закупка для Деревянко. Дополнительное удержание не создано; факт закрытия взаиморасчёта не подтверждён.'],[16,'58454/58733',124781.88,126000,'Кассета99000+сплит27000 в утверждённом отчёте13.07.2026. Закупочная стоимость отличается на1218,12. Повторное удержание запрещено.'],[17,'60665',88268.40,None,'Закупка для Деревянко. Дополнительное удержание не создано; факт закрытия взаиморасчёта не подтверждён.'],[19,'74173',30950.40,31000,'Кондеи мои31000 в утверждённом отчёте11.09.2026. Разница закупки49,60. Повторного удержания нет.'],[38,'44472',34207.49,None,'Закупка для Деревянко. Дополнительное удержание не создано; факт закрытия взаиморасчёта не подтверждён.']],(3,4))
sheet(w,'Операции',HEADERS,[op_values(r) for r in rows],(6,7))
sheet(w,'Проведено в CRM',['Операция CRM','Компания','Дата','№ банка','Получатель','Сумма, ₽','Сделка','Счёт','Статья','Товарное движение, ₽'],[[p['operation_id'],byid[p['bank_operation_id']]['company'],p['date'],p['document_number'],byid[p['bank_operation_id']]['counterparty'],p['amount_cents']/100,p.get('deal_id') or '',p['invoice_number'],p.get('category','Оплата по сделке'),p.get('secondary_cents',0)/100] for p in parts],(6,10))
sheet(w,'Монтажники без дублей',['Дата','№ банка','Получатель','Сумма, ₽','Запись монтажника'],[[r['date'],r['number'],r['counterparty'],r['amount_cents']/100,r['installer_payment_id']] for r in rows if r.get('installer_payment_id')],(4,))
original=load_workbook(json.loads((D/'source.json').read_text())['source'],read_only=True)
for name in ['Обработанные ответы','Клиентские счета','Архив документов','Сверка документов']:
 vals=list(original[name].values);sheet(w,name,list(vals[0]),[list(v) for v in vals[1:]])
original.close()
for x in arc:w['Архив документов'].append([x['number'],x['deal_id'],'',x['remote'].removeprefix('onedrive:'),'Проверено после нового ответа07.10'])
for x in outside:w['Архив документов'].append([x['number'],x['scope'],'',x['remote'].removeprefix('onedrive:'),'Вне проектных сделок; проверено наличие'])
w['Сверка документов'].append(['ЭК00-054832','851','98894,70 за товары проведено; счёт в архиве.','В документе дополнительно900доставка; этим платежом не покрыта, подтверждение отмены/отдельной оплаты не найдено. Не объявлено полной оплатой119626,88 по обоим счетам.'])
p=R/'reports/Выписка_обработано_07.10.2026_остаток.xlsx';save_workbook(w,p)
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for n in z.namelist():
  if n.endswith('.xml'):ET.fromstring(z.read(n))
c=load_workbook(p,read_only=True);assert c['Только ваши вопросы'].max_row==7 and c['Новые ответы'].max_row==16 and c['Операции'].max_row==339;c.close();summary['workbook']=str(p);summary['xlsx_validated']=True;atomic(D/'completion.json',summary);print(json.dumps(summary,ensure_ascii=False))
