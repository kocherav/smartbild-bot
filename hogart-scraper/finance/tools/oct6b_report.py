import sys,json,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from report import sheet
from xlsx_compat import save_workbook
R=Path(__file__).resolve().parents[1];D=R/'evidence/continuation_oct6b';v=json.loads((D/'verified.json').read_text());w=load_workbook(R/'reports/Разнесение_выписок.xlsx');rows={r['id']:r for r in json.loads((R/'reports/ledger.json').read_text())}
sheet(w,'Новые оплаты 06.10',['Операция Мегаплан','Дата','Поставщик','Счёт','Оплачено, ₽','Сделка','Основание','Товарное движение'],[[p['operation_id'],p['date'],rows[p['bank_operation_id']]['counterparty'],p['invoice_number'],p['amount_cents']/100,p['deal_id'],p['evidence'],0] for p in v['operations']],(5,8))
sheet(w,'Уточнённые статьи',['Дата','Поставщик','Оплачено, ₽','Статья','Основание','Осталось проверить'],[[x['date'],rows[x['id']]['counterparty'],x['amount_cents']/100,x['category'],x['evidence'],x['remaining_flag']] for x in json.loads((D/'classifications.json').read_text())],(3,))
sheet(w,'Новые файлы в архиве',['Счёт','Сделка','OneDrive'],[[x['number'],x['deal_id'],x['remote'].removeprefix('onedrive:')] for x in json.loads((D/'archive_results.json').read_text())])
p=R/'reports/Разнесение_выписок_06.10.2026_обновлено.xlsx';save_workbook(w,p)
check=load_workbook(p,read_only=True);assert check['Операции'].max_row==338;assert check['Новые оплаты 06.10'].max_row==4;check.close()
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for n in z.namelist():
  if n.endswith('.xml'):ET.fromstring(z.read(n))
print(p)
