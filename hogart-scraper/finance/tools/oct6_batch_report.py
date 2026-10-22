import sys,json,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook,load_workbook
from report import sheet
from xlsx_compat import save_workbook
R=Path(__file__).resolve().parents[1];D=R/'evidence/continuation_oct6';v=json.loads((D/'verified.json').read_text());w=Workbook();w.remove(w.active);rows={r['id']:r for r in json.loads((R/'reports/ledger.json').read_text())};names={'858':'Подвойского,1-й этаж НК — вентиляция','861':'Подвойского,1-й этаж НК — электроснабжение','862':'ЖК Примавера'}
sheet(w,'Результат',['Показатель','Результат'],[['Дополнительно проведено платежей',len(v['operations'])],['Дополнительно разнесено, ₽',v['new_cash_cents']/100],['Всего финансовых операций по выписке',v['total_crm_parts']],['Банковских платежей, охваченных проведением',v['bank_payments']],['Операций в исходной выписке',len(rows)],['Осталось с открытыми замечаниями',sum(r['status']=='Требует уточнения' for r in rows.values())],['Товарное движение','0: оплата отдельно от отгрузки и приёмки'],['Себестоимость','Распределение строк закупок не менялось'],['Новые вопросы владельцу','Нет. Ранее заданные вопросы по ПКБ3931/3937 и2943 сохраняются.']])
sheet(w,'Новые оплаты',['ID Мегаплан','Дата','Платёж №','Компания','Поставщик','Счёт','Сумма, ₽','Объект','Сделка'],[[p['operation_id'],p['date'],p['document_number'],rows[p['bank_operation_id']]['company'],rows[p['bank_operation_id']]['counterparty'],p['invoice_number'],p['amount_cents']/100,names[p['deal_id']],p['deal_id']] for p in v['operations']],(7,))
x=json.loads((D/'elevel_after_payment_sync.json').read_text());sheet(w,'Элевел',['Показатель','Результат'],[['Счёт',x['number']],['Сумма счёта, ₽',x['total']],['Подтверждено оплат, ₽',x['paid_total']],['Разница счёт минус оплаты, ₽',x['payment_balance']],['Уже распределено себестоимости, ₽',x['allocated_total']],['Ещё не распределено себестоимости, ₽',x['unallocated_total']],['Примечание','Остаток оплаты и нераспределённая себестоимость — разные величины. Приёмка товара банковской оплатой не подтверждена.']])
a=json.loads((D/'archive_results.json').read_text());sheet(w,'Архив',['Счёт','Сделка','Путь OneDrive','Проверено'],[[x['number'],x['deal_id'],x['remote'].removeprefix('onedrive:'),'Да'] for x in a])
q=json.loads((D/'research_queue.json').read_text())['held_candidates'];sheet(w,'Продолжаю проверять',['Поставщик','Счёт','Сумма, ₽','Причина'],[[x['supplier'],x['invoice'],x['amount'],x['reason']] for x in q],(3,))
p=R/'reports/Дополнительное_разнесение_06.10.2026.xlsx';save_workbook(w,p);z=load_workbook(p,read_only=True);assert z['Новые оплаты'].max_row==7;z.close()
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.endswith('.xml'):ET.fromstring(z.read(name))
print(p)
