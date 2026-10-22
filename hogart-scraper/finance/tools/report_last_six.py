from pathlib import Path as _QueuePath
if (_QueuePath(__file__).resolve().parents[1]/'data/review_registry.json').exists():
    raise SystemExit('Архивный обработчик отключён: рабочий реестр — finance/data/review_registry.json; обновление — finance/review_registry.py')
import sys,json,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
sys.path[:0]=['/home/kocher.av/.report-venv/lib/python3.14/site-packages','/home/kocher.av/hogart-scraper/finance']
from openpyxl import Workbook,load_workbook
from post_confirmed import atomic,ROOT,JOURNAL
from report import sheet,HEADERS,op_values
from xlsx_compat import save_workbook
D=ROOT/'evidence/last_six_20261007';v=json.loads((D/'verified.json').read_text());assert not v['errors'] and len(v['operations'])==9
rows=json.loads((ROOT/'reports/ledger.json').read_text());byid={r['id']:r for r in rows};parts=list(json.loads(JOURNAL.read_text())['operations'].values());assert len(parts)==287 and all(p['state']=='linked' for p in parts);posted={p['bank_operation_id'] for p in parts};inst={r['id'] for r in rows if r.get('installer_payment_id')};left=[r for r in rows if r['id'] not in posted|inst];assert len(posted)==276 and len(left)==57 and len(inst)==5
for oid in posted:assert sum(p['amount_cents'] for p in parts if p['bank_operation_id']==oid)==byid[oid]['amount_cents']
q=json.loads((D/'questions_before.json').read_text());answers=json.loads((D/'answers.json').read_text());old=json.loads((D/'questions_before.json').read_text());resolved={g['number']:g for g in q['resolved_groups']}
for g in q['groups']:
 ps=[p for p in parts if p['bank_operation_id'] in g['bank_ids']];answer=answers[str(g['number'])];g.update(owner_answer=answer,answer=answer,resolved=True,owner_question_closed=True,crm_operation_ids=[p['operation_id'] for p in ps],resolution='Назначение подтверждено владельцем. '+('Денежные операции проведены и проверены, без товарного движения.' if ps else 'Собственные нужды; до получения чека/реквизитов продавца запись CRM не создавалась.'));resolved[g['number']]=g
q.update(groups=[],resolved_groups=list(resolved.values()),posted_bank_rows=len(posted),unposted_rows=len(left));atomic(ROOT/'evidence/finish_all/owner_questions.json',q)
w=Workbook();w.remove(w.active)
sheet(w,'Итог',['Показатель','Результат'],[['Дата','07.10.2026'],['Принято новых ответов',6],['Осталось вопросов о назначении платежей',0],['Новые финансовые операции CRM',9],['Новых строк банка проведено',7],['Всего строк банка',len(rows)],['Строк банка проведено в CRM',len(posted)],['Операций CRM с учётом разбивок',len(parts)],['Выплаты монтажникам учтены отдельно',len(inst)],['Назначение известно, оформление по документам не завершено',len(left)],['Контроль полноты','276 CRM +5 монтажники +57 оформление =338 строк'],['Полное завершение','Все ответы о назначении обработаны. Полностью оформленными все338строк не объявляются: остаётся57строк без записи CRM.'],['Остаток документов','50карточных списаний: нужны чек/реквизиты продавца;6переводов между своими счетами: встречные выписки;1платёж Шишаевой: счёт34023676 с реквизитами. Повторное распределение по объектам не требуется.'],['Новогорск','98000 ₽ →790;73000 ₽ →собственные нужды. Первый платёж101000:98000+3000; второй70000 —собственные нужды.'],['Трио42965/44245','Последовательное распределение одинаковых единиц:26215,56 →855Зоовет;26215,56 →869Илья;27488,16 →собственные нужды. Это распределение стоимости, не серийных номеров.'],['Оплата и отгрузка','Все9новых операций проверены в CRM; товарное движение0. Клиентские спецификации и модели не менялись.'],['Архив','Счёт42965 сохранён в папках855/869; ЭТМ212/099912439 —861;44245 —в собственных нуждах.']])
w['Итог'].column_dimensions['A'].width=57;w['Итог'].column_dimensions['B'].width=112
sheet(w,'Только ваши вопросы',['№ вопроса','Поставщик','Счёт','Платежи','Сумма, ₽','Вопрос','Проверено','Ваш ответ'],[])
notes={37:'По порядку закупок:855=26215,56;869=26215,56;собственные нужды27488,16. Клиентские модели не менялись.',44:'861 —электрика1этажа',51:'847 —инженерные проекты1этажаJetCar',52:'790Новогорск98000;собственные нужды73000. По хронологии оплат.',54:'Собственные нужды. Остался чек для определения юридического продавца; не подставлен банк.',71:'861 —электрика1этажа, по прямому ответу владельца.'}
sheet(w,'Последние 6 ответов',['№','Поставщик','Счёт','Сумма, ₽','Ваш ответ','Результат','Операции CRM'],[[g['number'],g['supplier'],g['invoices'],g['amount_cents']/100,answers[str(g['number'])],notes[g['number']],', '.join(p['operation_id'] for p in parts if p['bank_operation_id'] in g['bank_ids'])] for g in old['groups']],(4,))
def reason(r):
 if r['category']=='Перевод между своими счетами':return 'Встречная выписка Т-Банка/Озон Банка. Внутренний перевод, не проектный расход.'
 if r['inn']=='7728168971':return 'Назначение сохранено. Нужен чек/отчёт по карте для реквизитов продавца; банковский ИНН не подставлялся вместо продавца.'
 return 'Нужен счёт34023676 с реквизитами получателя; назначение «собственные нужды» принято.'
sheet(w,'Остаток оформления',['Компания','Дата','№ банка','Получатель / терминал','Сумма, ₽','Установленное назначение','Что требуется для завершения','ID'],[[r['company'],r['date'],r['number'],r['purpose'] if r['inn']=='7728168971' else r['counterparty'],r['amount_cents']/100,r.get('owner_answer') or r['category'],reason(r),r['id']] for r in left],(5,));w['Остаток оформления'].column_dimensions['H'].hidden=True
sheet(w,'Операции',HEADERS,[op_values(r) for r in rows],(6,7))
sheet(w,'Проведено в CRM',['Операция CRM','Компания','Дата','№ банка','Получатель','Сумма, ₽','Сделка','Счёт','Статья','Товарное движение, ₽'],[[p['operation_id'],byid[p['bank_operation_id']]['company'],p['date'],p['document_number'],byid[p['bank_operation_id']]['counterparty'],p['amount_cents']/100,p.get('deal_id') or '',p['invoice_number'],p.get('category','Оплата по сделке'),p.get('secondary_cents',0)/100] for p in parts],(6,10))
src=load_workbook(json.loads((D/'source.json').read_text())['source'],read_only=True)
for name in src.sheetnames:
 if name in w.sheetnames:continue
 vals=list(src[name].values);sheet(w,name,list(vals[0]),[list(x) for x in vals[1:]])
src.close()
for x in json.loads((D/'archive_results.json').read_text())+json.loads((D/'outside_archive_results.json').read_text()):w['Архив документов'].append([x['number'],x.get('deal_id') or x['scope'],'',x['remote'].removeprefix('onedrive:'),'Проверено после последних6ответов07.10'])
p=ROOT/'reports/Выписка_07.10.2026_все_ответы_учтены.xlsx';save_workbook(w,p)
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for n in z.namelist():
  if n.endswith('.xml'):ET.fromstring(z.read(n))
c=load_workbook(p,read_only=True);assert c['Последние 6 ответов'].max_row==7 and c['Остаток оформления'].max_row==58 and c['Операции'].max_row==339 and c['Только ваши вопросы'].max_row==1;c.close()
summary={'workbook':str(p),'new_answers':6,'new_crm_parts':9,'new_bank_rows':7,'remaining_allocation_questions':0,'crm_parts':287,'posted_bank_rows':276,'installer_rows':5,'documentary_remaining':57,'bank_rows':338,'xlsx_validated':True};atomic(D/'completion.json',summary);print(json.dumps(summary,ensure_ascii=False))
