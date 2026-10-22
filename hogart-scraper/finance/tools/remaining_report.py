import sys,json,zipfile,collections,xml.etree.ElementTree as ET
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from report import sheet
from xlsx_compat import save_workbook
R=Path(__file__).resolve().parents[1];D=R/'evidence/remaining_oct6';v=json.loads((D/'company_verified.json').read_text());rows=json.loads((R/'reports/ledger.json').read_text());byid={x['id']:x for x in rows};w=load_workbook(R/'reports/Разнесение_выписок.xlsx');fees=[p for p in v['operations'] if p['kind']=='expense'];interest=[p for p in v['operations'] if p['kind']=='income']
sheet(w,'Новые операции компании',['Мегаплан ID','Компания','Дата','Банковский документ','Статья','Приход, ₽','Расход, ₽','Учёт'],[[p['operation_id'],byid[p['bank_operation_id']]['company'],p['date'],p['document_number'],p['category'],p['amount_cents']/100 if p['kind']=='income' else 0,p['amount_cents']/100 if p['kind']=='expense' else 0,'Завершённая банковская услуга' if p['kind']=='expense' else 'Нетоварный доход; не продажа'] for p in v['operations']],(6,7))
sheet(w,'Итог продолжения',['Показатель','Результат'],[['Новых комиссий',len(fees)],['Сумма комиссий, ₽',sum(p['amount_cents'] for p in fees)/100],['Поступлений процентов банка',len(interest)],['Проценты банка, ₽',sum(p['amount_cents'] for p in interest)/100],['Всего банковских платежей в журнале проведения CRM',v['bank_payments']],['Всего финансовых операций в журнале CRM',v['total_parts']],['Строк исходной выписки',len(rows)],['Строк с открытыми замечаниями',sum(bool(r['flags']) for r in rows)],['Из открытых: межфирменные движения',34],['Проверка межфирменных движений','17 пар;14727000 ₽ прихода=14727000 ₽ расхода. Исключены из внешнегоДДС группы. Объекты не назначены.'],['Связь','4 замечания по ставке НДС сняты по первичным документам. Банк:20%, документы:22%; оригиналы не изменены.'],['Ограничение','Это не полное завершение всех337 строк. Лист Рабочая очередь — задачи сверки, лист Вам — конкретные вопросы владельцу.']])
sheet(w,'Документы по связи',['Файл','OneDrive','Проверено'],[[Path(x['source']).name,x['remote'].removeprefix('onedrive:'),'Да'] for x in json.loads((D/'telecom_archive.json').read_text())])
requests=json.loads((R/'reports/document_requests.json').read_text())
sheet(w,'Остаток по счетам',['Поставщик','ИНН','Компания','Счёт','Платежи','Количество','Оплачено, ₽','Документ','Проверено'],[[*x[:7],'Найден; требуется распределение',x[8]] for x in requests['found']]+[[*x[:7],'Продолжается проверка источников',x[8]] for x in requests['pending_assistant_review']],(7,))
p=R/'reports/Разнесение_выписок_06.10.2026_продолжение.xlsx';save_workbook(w,p)
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for n in z.namelist():
  if n.endswith('.xml'):ET.fromstring(z.read(n))
c=load_workbook(p,read_only=True);assert c['Операции'].max_row==338;assert c['Новые операции компании'].max_row==63;c.close();print(p)
