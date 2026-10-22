import json,re,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import xlsxwriter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1];D=R/'evidence/deep_review';q=json.loads((R/'evidence/finish_all/owner_questions.json').read_text());rows={r['id']:r for r in json.loads((R/'reports/ledger.json').read_text())};groups=sorted(q['groups'],key=lambda g:(g['supplier'].lower(),g['number']))
def clean(t):
 t=re.sub(r'([а-яё])([А-ЯЁA-Z])',r'\1 \2',t)
 t=re.sub(r'([А-Яа-яЁё])(?=\d)',r'\1 ',t)
 t=re.sub(r'(?<=\d)(?=[А-Яа-яЁё])',' ',t)
 t=re.sub(r',(?=[^\s\d])',', ',t)
 t=re.sub(r'(?<=[;:])(?=\S)',' ',t)
 t=re.sub(r'\.(?=[А-ЯЁ])','. ',t)
 return t
p=R/'reports/Только_неясные_позиции_после_проверки_06.10.2026.xlsx'
with xlsxwriter.Workbook(p,{'strings_to_formulas':False,'strings_to_urls':False}) as w:
 head=w.add_format({'bold':True,'font_color':'white','bg_color':'#173F4F','text_wrap':True,'valign':'vcenter'})
 txt=w.add_format({'text_wrap':True,'valign':'top','font_size':11});answer=w.add_format({'text_wrap':True,'valign':'top','bg_color':'#FFF4CC'});money=w.add_format({'num_format':'#,##0.00','valign':'top'});note=w.add_format({'text_wrap':True,'valign':'top','font_color':'#475569','font_size':10})
 ws=w.add_worksheet('Только вопросы');ws.hide_gridlines(2);ws.freeze_panes(3,2)
 ws.merge_range(0,0,0,7,'51 группа / 77 банковских строк / 32 получателя. Это реальный остаток вопросов; до 15–20 пока не сокращён.',head)
 ws.merge_range(1,0,1,7,'Номер — из прежнего списка. Если несколько покупок у поставщика на один объект, достаточно одного ответа и списка исключений. Счета, которые уже найдены, повторно не нужны.',txt);ws.set_row(0,34);ws.set_row(1,35)
 headers=['Старый №','Поставщик / продавец','Счёт или заказ','Сумма оплат, ₽','Что нужно уточнить','Ваш ответ','Что установлено при повторной проверке','Оплаты / даты']
 for c,h in enumerate(headers):ws.write(2,c,h,head)
 ws.set_row(2,33)
 for c,width in enumerate([9,29,25,17,65,35,77,43]):ws.set_column(c,c,width)
 for i,g in enumerate(groups,3):
  vals=[g['number'],g['supplier'],g['invoices'],g['amount_cents']/100,clean(g['question']),'',clean(g['evidence']),g['payments']]
  for c,v in enumerate(vals):ws.write(i,c,v,answer if c==5 else money if c==3 else note if c==6 else txt)
  ws.set_row(i,max(85,min(160,20*(len(vals[6])//75+2))))
 ws.autofilter(2,0,len(groups)+2,7)
 ws=w.add_worksheet('Даты оплат');ws.hide_gridlines(2);ws.freeze_panes(1,2)
 headers=['Старый № вопроса','Поставщик / продавец','Компания','Дата банка','№ банка','Сумма, ₽','Счёт / заказ','Назначение банка','Ваше исключение / объект']
 for c,h in enumerate(headers):ws.write(0,c,h,head)
 for c,width in enumerate([12,29,15,14,15,17,25,90,40]):ws.set_column(c,c,width)
 k=1
 for g in groups:
  for oid in g['bank_ids']:
   r=rows[oid];values=[g['number'],g['supplier'],r['company'],r['date'],r['number'],r['amount_cents']/100,g['invoices'],r['purpose'],'']
   for c,v in enumerate(values):ws.write(k,c,v,money if c==5 else answer if c==8 else txt)
   ws.set_row(k,62);k+=1
 ws.autofilter(0,0,k-1,8)
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for n in z.namelist():
  if n.endswith('.xml'):ET.fromstring(z.read(n))
t=load_workbook(p,read_only=True);assert t['Только вопросы'].max_row==len(groups)+3;assert t['Даты оплат'].max_row==78;t.close()
(D/'questions_workbook_validation.json').write_text(json.dumps({'file':str(p),'question_groups':len(groups),'bank_rows':77,'excel_zip_xml_openpyxl_validated':True},ensure_ascii=False,indent=2));print(p)
