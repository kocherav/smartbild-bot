import json,sys,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook,load_workbook
from report import sheet
from xlsx_compat import save_workbook
R=Path(__file__).resolve().parents[1]
a=json.loads((R/'evidence/posting/archive_results.json').read_text());j=list(json.loads((R/'data/posting_journal.json').read_text())['operations'].values())
w=Workbook();w.remove(w.active)
sheet(w,'Итог',['Показатель','Результат'],[['Банковских операций в исходной выписке',337],['Банковских платежей проведено в этом пакете',len({r['bank_operation_id'] for r in j})],['Новых финансовых операций',35],['Использовано ранее существовавших',2],['Привязано к сделкам',len(j)],['Документов в OneDrive с учётом версий',len({r['sha256'] for r in a})],['Копий в папках сделок',sum(bool(r['deal_id']) for r in a)],['Файлов без подтверждённого объекта',sum(not r['deal_id'] for r in a)],['Важно','Версии счетов не складывать. Суммы счетов не равны оплатам. Этот отчёт не является расчётом прибыли.']])
sheet(w,'Проведено в Мегаплан',['ID операции','Сделка','Компания','Дата','Банк №','Направление','Сумма оплаты, ₽','Счёт','Состояние','Основание'],[[r['operation_id'],r['deal_id'],'ИНСИС' if r['account_id']=='20' else 'КЛИВЭЛ',r['date'],r['document_number'],r['kind'],r['amount_cents']/100,r['invoice_number'],r['state'],r['evidence']] for r in j],(7,))
sheet(w,'Архив OneDrive',['Поставщик','ИНН','Покупатель ИНН','Счёт','Сумма документа, ₽','Сделка','Статус','Путь OneDrive','Примечание','SHA256'],[[r['supplier'],r['supplier_inn'],r.get('buyer_inn',''),r['number'],r.get('amount_cents',0)/100,r.get('deal_id') or '', 'Объект подтверждён' if r.get('deal_id') else 'Объект требует сверки',r['remote'].removeprefix('onedrive:'),r.get('note',''),r['sha256']] for r in a],(5,))
sheet(w,'Новые 7 счетов',['Счёт','Сумма счёта, ₽','Объект','Результат'],[['Бюлер557',1252012.87,'ЖК Реномэ,865','Оплата проведена'],['АЕР ГРУПП307',351400,'ЖК Реномэ,865','Оплата проведена'],['ПКБ-00003053',15656.75,'ЖК Примавера,862','Оплата проведена'],['П008-07-26',171080,'Квартира44 м²,867','Проведён аванс130000 ₽'],['ПКБ-00002943',164460.55,'ЖК Примавера — по новому файлу','Платёж107230,90 ₽ пока не привязан: конфликт версий'],['ПКБ-00003931',20827.55,'Не установлен','Сумма платежа совпадает; нужен объект'],['ПКБ-00003937',3355.80,'Не установлен','Платёж в исходной выписке не найден']],(2,))
sheet(w,'Только ваши вопросы',['Вопрос','Контекст'],[['Для каких объектов ПКБ-00003931 и ПКБ-00003937?','20827,55 ₽ и3355,80 ₽'],['ПКБ-00002943:107230,90 ₽ — частичная оплата целиком для Примаверы?','Новый счёт164460,55 ₽; другая версия114370,90 ₽.']])
p=R/'reports/Итоги_разнесения_и_архив_05.10.2026.xlsx';save_workbook(w,p)
v=load_workbook(p,read_only=True);assert v['Проведено в Мегаплан'].max_row==38;assert v['Архив OneDrive'].max_row==65;v.close()
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.endswith('.xml'):ET.fromstring(z.read(name))
print(p)
