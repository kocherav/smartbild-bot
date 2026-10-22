import unittest
from docx import Document

from generate_contract import TEMPLATE, compact_signature_tables, compact_specification_table, contract_data, contract_filename_stem, format_appendix_heading, is_full_person_name, replace_in_runs, safe_filename_stem, subject_for_body, subject_for_title, truncate_filename_utf8


class ContractFilenameTests(unittest.TestCase):
    def test_discounted_contract_uses_net_unit_price(self):
        deal={'id':'1','payer':{},'price':{'value':1500},'positions':[{'name':'Example','quantity':2,'price':{'value':1000},'sum':{'value':1500},'discount':{'value':25}}]}
        data=contract_data(deal,{})
        self.assertEqual(750,data['positions'][0]['unit_price'])
        deal['positions'][0]['sum']['value']=0; deal['price']['value']=0
        self.assertEqual(0,contract_data(deal,{})['positions'][0]['unit_price'])

    def test_reference_only_units_reuse_the_deal_dictionary(self):
        rows=[{'id':str(i),'name':'Row '+str(i),'quantity':1,'sum':{'value':1},'unit':u} for i,u in enumerate([{'id':'8'},{'id':'8','name':'к-т'},{'id':'3','name':'м'},{'id':'3'}])]
        self.assertEqual(['к-т','к-т','м','м'],[r['unit'] for r in contract_data({'id':'1','positions':rows,'price':{'value':4}}, {})['positions']])

    def test_contract_rejects_unfixed_foreign_price(self):
        deal={'id':'1','positions':[{'name':'Example','quantity':1,'price':{'value':10,'currency':'EUR'},'sum':{'value':10,'currency':'EUR'}}]}
        with self.assertRaises(ValueError): contract_data(deal,{})

    def test_truncates_cyrillic_by_bytes_without_breaking_utf8(self):
        value = "Договор — " + "очень длинное название позиции " * 20
        result = truncate_filename_utf8(value, 170)
        self.assertLessEqual(len(result.encode("utf-8")), 170)
        self.assertTrue(result.startswith("Договор"))

    def test_keeps_short_name(self):
        self.assertEqual(truncate_filename_utf8("Договор №888", 170), "Договор №888")

    def test_product_model_slash_cannot_create_a_subdirectory(self):
        result = safe_filename_stem(
            "Договор №1.892-1909/2026 — ООО Зоовет — HSU-07HQJ103/R3-W"
        )
        self.assertNotIn("/", result)
        self.assertNotIn("\\", result)
        self.assertEqual(
            result,
            "Договор №1.892-1909-2026 — ООО Зоовет — HSU-07HQJ103-R3-W",
        )

    def test_all_forbidden_filename_characters_are_removed(self):
        result = safe_filename_stem('A:B*C?D"E<F>G|H/I\\J\nK')
        self.assertFalse(any(char in result for char in '\\/:*?"<>|\n'))

    def test_contract_filename_uses_essence_not_equipment_or_sku(self):
        result = contract_filename_stem(
            "1.892-1909/2026", "ООО Зоовет",
            "Поставка и монтаж кондиционера в 103-й кабинет",
        )
        self.assertEqual(
            result,
            "Договор №1.892-1909-2026 — КЛИВЭЛ — ООО Зоовет — Поставка и монтаж кондиционера в 103-й кабинет",
        )
        self.assertNotIn("HSU", result)
        self.assertNotIn("R3", result)

    def test_contract_filename_uses_payer_with_printable_quotes(self):
        result = contract_filename_stem("1.894-2109/2026", 'ООО "МК"', "Поставка оборудования")
        self.assertIn("ООО «МК»", result)

    def test_director_is_reused_from_payer_custom_field(self):
        deal = {
            "id": "892", "payer": {"id": "55", "Category715CustomFieldFioRukovoditelya": "Бутылкина Марина Владимировна"},
            "positions": [], "price": {"value": 0},
        }
        data = contract_data(deal, {"customer_signatory": {"position": "Генеральный директор", "authority": "Устав"}})
        self.assertEqual(data["signatory"]["name"], "Бутылкина Марина Владимировна")

    def test_signatory_from_another_payer_is_never_reused(self):
        deal = {
            "id": "894", "payer": {"id": "1003023", "Category715CustomFieldFioRukovoditelya": "Ильинов С.А."},
            "positions": [], "price": {"value": 0},
        }
        data = contract_data(deal, {
            "customer_signatory_payer_id": "1002999",
            "customer_signatory": {"name": "Бутылкина Марина Владимировна", "position": "Генеральный директор", "authority": "Устав"},
        })
        self.assertEqual(data["signatory"]["name"], "")
        self.assertIn("ФИО подписанта", data["missing"])

    def test_full_signatory_collected_for_current_payer_is_accepted(self):
        deal = {
            "id": "894", "payer": {"id": "1003023", "Category715CustomFieldFioRukovoditelya": "Ильинов С.А."},
            "positions": [], "price": {"value": 0},
        }
        data = contract_data(deal, {
            "customer_signatory_payer_id": "1003023",
            "customer_signatory": {"name": "Ильинов Сергей Александрович", "position": "Генеральный директор", "authority": "Устав"},
        })
        self.assertEqual(data["signatory"]["name"], "Ильинов Сергей Александрович")

    def test_initials_are_not_a_full_contract_signatory_name(self):
        self.assertFalse(is_full_person_name("Ильинов С.А."))
        self.assertTrue(is_full_person_name("Ильинов Сергей Александрович"))

    def test_contract_subject_grammar_is_stable(self):
        source = "Проектирование, поставка, монтаж и пусконаладка системы кондиционирования"
        self.assertEqual(
            subject_for_title(source),
            "проектирование, поставку, монтаж и пусконаладку системы кондиционирования",
        )
        self.assertEqual(
            subject_for_body(source),
            "проектирование, поставка, монтаж и пусконаладка системы кондиционирования",
        )

    def test_simple_supply_subject(self):
        self.assertEqual(subject_for_title("Поставка оборудования"), "поставку оборудования")
        self.assertEqual(subject_for_body("Поставка оборудования"), "поставка оборудования")

    def test_only_customer_position_is_replaced_in_preamble(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Поставщик в лице Генерального директора Кочержинской С.В., ")
        paragraph.add_run("действующей на основании Устава, и Заказчик в лице ")
        paragraph.add_run("Генерального директора Медведко А.С., действующего на основании Устава")
        replace_in_runs(paragraph, "Генерального директора", "Директора", occurrence=2)
        replace_in_runs(paragraph, "на основании Устава", "на основании доверенности", occurrence=2)
        self.assertIn("Поставщик в лице Генерального директора", paragraph.text)
        self.assertIn("Заказчик в лице Директора", paragraph.text)
        self.assertIn("действующей на основании Устава", paragraph.text)
        self.assertIn("действующего на основании доверенности", paragraph.text)

    def test_signature_rows_cannot_split_and_images_are_untouched(self):
        document = Document()
        for _ in range(3):
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Поставщик"
            table.cell(0, 1).text = "Заказчик"
        drawings_before = len(document.element.xpath(".//w:drawing | .//w:pict"))
        compact_signature_tables(document)
        drawings_after = len(document.element.xpath(".//w:drawing | .//w:pict"))
        for index in (0, 2):
            self.assertEqual(len(document.tables[index].rows[0]._tr.xpath("./w:trPr/w:cantSplit")), 1)
        self.assertEqual(drawings_before, drawings_after)

    def test_specification_rows_have_no_artificial_height_or_spacing(self):
        document = Document()
        table = document.add_table(rows=19, cols=2)
        table.rows[1].height = 355600
        table.cell(1, 0).text = "Позиция"
        table.cell(1, 1).text = "1"
        compact_specification_table(table)
        self.assertFalse(table.rows[1]._tr.xpath("./w:trPr/w:trHeight"))
        for cell in table.rows[1].cells:
            paragraph = cell.paragraphs[0]
            self.assertEqual(paragraph.paragraph_format.space_before.pt, 0)
            self.assertEqual(paragraph.paragraph_format.space_after.pt, 0)
            self.assertEqual(paragraph.paragraph_format.line_spacing, 1.0)
        self.assertTrue(table.rows[3].cells[0].paragraphs[0].paragraph_format.page_break_before)

    def test_appendix_always_starts_on_new_page(self):
        document = Document()
        paragraph = document.add_paragraph("Приложение № 1")
        format_appendix_heading(paragraph)
        self.assertTrue(paragraph.paragraph_format.page_break_before)

    def test_real_template_has_no_blank_lines_before_signatories(self):
        document = Document(str(TEMPLATE))
        drawing = document.element.xpath(".//w:drawing")[0]
        extent_before = dict(drawing.xpath(".//wp:extent")[0].attrib)
        offset_before = int(drawing.xpath(".//wp:positionV/wp:posOffset")[0].text)

        compact_signature_tables(document)

        supplier = document.tables[0].rows[0].cells[0]
        supplier_text = [p.text.strip() for p in supplier.paragraphs]
        supplier_position = supplier_text.index("Генеральный директор")
        self.assertTrue(supplier_text[supplier_position - 1])

        customer = document.tables[0].rows[0].cells[1].tables[0].cell(0, 0)
        customer_text = [p.text.strip() for p in customer.paragraphs]
        customer_position = next(
            index for index, text in enumerate(customer_text) if "директор" in text.casefold()
        )
        self.assertTrue(customer_text[customer_position - 1])

        self.assertEqual(dict(drawing.xpath(".//wp:extent")[0].attrib), extent_before)
        offset_after = int(drawing.xpath(".//wp:positionV/wp:posOffset")[0].text)
        self.assertEqual(offset_before - offset_after, 1860000)


if __name__ == "__main__":
    unittest.main()
