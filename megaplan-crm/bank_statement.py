#!/usr/bin/env python3
"""Import Alfa Bank XLSX statements and keep an auditable deal-assignment registry.

The importer is intentionally independent from Megaplan writes.  It parses and
deduplicates transactions; bridge.py performs an explicit confirmed CRM comment and
only then calls ``assign`` to persist the association locally.
"""
import hashlib
import json
import os
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from registry_transaction import transaction

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REGISTRY_PATH = os.path.join(BASE_DIR, "bank_statement_registry.json")


def text(value):
    return "" if value is None else str(value).strip()


def amount(value):
    if value in (None, ""):
        return None
    try:
        parsed=Decimal(str(value).replace(" ", "").replace(",", "."))
        if not parsed.is_finite() or parsed<0: raise ValueError('Некорректная сумма выписки')
        return format(parsed.quantize(Decimal("0.01")), "f")
    except InvalidOperation as exc:
        raise ValueError(f"некорректная сумма: {value}") from exc


def date_text(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    value = text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).strftime("%d.%m.%Y")
        except ValueError:
            pass
    raise ValueError(f"некорректная дата операции: {value}")


def load_registry():
    try:
        with open(REGISTRY_PATH, encoding="utf-8") as source:
            data = json.load(source)
    except FileNotFoundError:
        data = {"version": 1, "transactions": [], "imports": []}
    data.setdefault("transactions", [])
    data.setdefault("imports", [])
    return data


def save_registry(data):
    tmp = REGISTRY_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as target:
        json.dump(data, target, ensure_ascii=False, indent=2, allow_nan=False)
    os.replace(tmp, REGISTRY_PATH)


def fingerprint(item):
    parts = [item.get(key, "") for key in (
        "account", "date", "debit", "credit",
        "counterparty_inn", "counterparty_account", "purpose",
    )]
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()[:20]


def parse_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    account = text(ws["B1"].value).replace(" ", "")
    period = text(ws["B2"].value)
    owner = text(ws["B3"].value)
    owner_inn = text(ws["B4"].value)
    if text(ws["A1"].value).casefold() != "выписка по счёту" or not account:
        raise ValueError("не распознан формат банковской выписки")
    header_row = None
    # Some bank exports declare a bogus worksheet dimension of A1 even though the
    # actual XML contains the full table, so read a bounded explicit range.
    for row in range(1, 51):
        if text(ws.cell(row, 1).value) == "Дата" and text(ws.cell(row, 2).value) == "Номер документа":
            header_row = row
            break
    if not header_row:
        raise ValueError("не найдена таблица банковских операций")
    operations = []
    for row in ws.iter_rows(min_row=header_row + 2, max_row=10000, max_col=13, values_only=True):
        if not any(value not in (None, "") for value in row):
            continue
        if not row[0]:
            continue
        debit, credit = amount(row[2]), amount(row[3])
        if debit is None and credit is None:
            continue
        item = {
            "account": account, "date": date_text(row[0]),
            "document_number": text(row[1]), "debit": debit or "",
            "credit": credit or "", "direction": "expense" if debit is not None else "income",
            "amount": debit or credit, "counterparty": text(row[4]),
            "counterparty_inn": text(row[5]), "counterparty_kpp": text(row[6]),
            "counterparty_account": text(row[7]), "counterparty_bik": text(row[8]),
            "counterparty_bank": text(row[9]), "purpose": text(row[10]),
            "document_type": text(row[12]), "status": "unassigned",
        }
        item["id"] = fingerprint(item)
        operations.append(item)
    debit_total = sum(Decimal(item["debit"] or "0") for item in operations)
    credit_total = sum(Decimal(item["credit"] or "0") for item in operations)
    expected_debit = Decimal(amount(ws["E7"].value) or "0")
    expected_credit = Decimal(amount(ws["E8"].value) or "0")
    if debit_total != expected_debit or credit_total != expected_credit:
        raise ValueError(
            f"контрольные суммы не совпали: дебет {debit_total}/{expected_debit}, "
            f"кредит {credit_total}/{expected_credit}"
        )
    return {"account": account, "period": period, "owner": owner, "owner_inn": owner_inn,
            "debit_total": str(debit_total), "credit_total": str(credit_total),
            "operations": operations}


def read_1c_text(path):
    with open(path, "rb") as source:
        raw = source.read()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            decoded = raw.decode(encoding)
            if decoded.startswith("1CClientBankExchange"):
                return decoded
        except UnicodeDecodeError:
            pass
    raise ValueError("не удалось определить кодировку файла 1С")


def parse_pairs(lines):
    result = {}
    for line in lines:
        if "=" in line:
            key, value = line.split("=", 1)
            result[key.strip()] = value.strip()
    return result


def parse_1c(path):
    lines = [line.strip() for line in read_1c_text(path).splitlines()]
    header_end = next((i for i, line in enumerate(lines) if line == "СекцияРасчСчет"), None)
    if header_end is None:
        raise ValueError("в файле 1С отсутствует секция расчётного счёта")
    header = parse_pairs(lines[1:header_end])
    account = header.get("РасчСчет", "").replace(" ", "")
    account_end = next((i for i in range(header_end + 1, len(lines)) if lines[i] == "КонецРасчСчет"), None)
    if not account or account_end is None:
        raise ValueError("в файле 1С не определён расчётный счёт")
    totals = parse_pairs(lines[header_end + 1:account_end])
    documents = []
    index = account_end + 1
    while index < len(lines):
        if not lines[index].startswith("СекцияДокумент="):
            index += 1
            continue
        document_type = lines[index].split("=", 1)[1]
        end = next((i for i in range(index + 1, len(lines)) if lines[i] == "КонецДокумента"), None)
        if end is None:
            raise ValueError("повреждена секция документа в файле 1С")
        fields = parse_pairs(lines[index + 1:end])
        fields["_type"] = document_type
        documents.append(fields)
        index = end + 1
    operations, owner, owner_inn = [], "", ""
    for fields in documents:
        payer_account = fields.get("ПлательщикСчет", "").replace(" ", "")
        receiver_account = fields.get("ПолучательСчет", "").replace(" ", "")
        if payer_account == account:
            direction, debit, credit, prefix = "expense", amount(fields.get("Сумма")), "", "Получатель"
            owner = owner or fields.get("Плательщик", "")
            owner_inn = owner_inn or fields.get("ПлательщикИНН", "")
        elif receiver_account == account:
            direction, debit, credit, prefix = "income", "", amount(fields.get("Сумма")), "Плательщик"
            owner = owner or fields.get("Получатель", "")
            owner_inn = owner_inn or fields.get("ПолучательИНН", "")
        else:
            raise ValueError(f"документ №{fields.get('Номер', '?')} не относится к счёту {account}")
        operation_date = (fields.get("ДатаСписано") if direction == "expense" else fields.get("ДатаПоступило")) or fields.get("Дата")
        item = {
            "account": account, "date": date_text(operation_date),
            "document_number": fields.get("Номер", ""), "debit": debit or "",
            "credit": credit or "", "direction": direction, "amount": debit or credit,
            "counterparty": fields.get(prefix, ""), "counterparty_inn": fields.get(prefix + "ИНН", ""),
            "counterparty_kpp": fields.get(prefix + "КПП", ""), "counterparty_account": fields.get(prefix + "Счет", ""),
            "counterparty_bik": fields.get(prefix + "БИК", ""),
            "counterparty_bank": fields.get(prefix + "Банк1", "") or fields.get(prefix + "Банк", ""),
            "purpose": fields.get("НазначениеПлатежа", ""), "document_type": fields.get("_type", ""),
            "status": "unassigned",
        }
        item["id"] = fingerprint(item)
        operations.append(item)
    debit_total = sum(Decimal(item["debit"] or "0") for item in operations)
    credit_total = sum(Decimal(item["credit"] or "0") for item in operations)
    expected_debit = Decimal(amount(totals.get("ВсегоСписано")) or "0")
    expected_credit = Decimal(amount(totals.get("ВсегоПоступило")) or "0")
    if debit_total != expected_debit or credit_total != expected_credit:
        raise ValueError(f"контрольные суммы не совпали: дебет {debit_total}/{expected_debit}, кредит {credit_total}/{expected_credit}")
    return {"account": account, "period": f"c {header.get('ДатаНачала', '')} по {header.get('ДатаКонца', '')}",
            "owner": owner, "owner_inn": owner_inn, "debit_total": str(debit_total),
            "credit_total": str(credit_total), "operations": operations}


def parse_statement(path):
    extension = os.path.splitext(path)[1].casefold()
    if extension == ".xlsx":
        return parse_xlsx(path)
    if extension == ".txt":
        return parse_1c(path)
    raise ValueError("поддерживаются выписки Excel (.xlsx) и 1С (.txt)")


def commission_group(item):
    purpose = (item.get("purpose") or "").casefold()
    if not re.search(r"комисси|ком-я|ком\.", purpose):
        return None
    return tuple(item.get(key, "") for key in (
        "account", "date", "direction", "counterparty_inn", "counterparty_account",
    ))


def remove_cross_format_duplicates(incoming, existing):
    """Handle exact rows plus XLSX bank commissions aggregated from several 1C rows."""
    known_ids = {item["id"] for item in existing}
    fresh = [item for item in incoming if item["id"] not in known_ids]
    existing_commissions = {}
    incoming_commissions = {}
    for item in existing:
        key = commission_group(item)
        if key:
            existing_commissions[key] = existing_commissions.get(key, Decimal("0")) + Decimal(item["amount"])
    for item in incoming:
        key = commission_group(item)
        if key:
            incoming_commissions[key] = incoming_commissions.get(key, Decimal("0")) + Decimal(item["amount"])
    covered_groups = {key for key, total in incoming_commissions.items()
                      if key in existing_commissions and existing_commissions[key] == total}
    return [item for item in fresh if commission_group(item) not in covered_groups]


@transaction(lambda:REGISTRY_PATH)
def import_statement(path):
    parsed = parse_statement(path)
    data = load_registry()
    new = remove_cross_format_duplicates(parsed["operations"], data["transactions"])
    data["transactions"].extend(new)
    with open(path, "rb") as source:
        file_hash = hashlib.sha256(source.read()).hexdigest()
    if file_hash not in {item.get("sha256") for item in data["imports"]}:
        data["imports"].append({"sha256": file_hash, "name": os.path.basename(path),
                                "account": parsed["account"], "period": parsed["period"],
                                "imported_at": datetime.now().isoformat(timespec="seconds")})
    save_registry(data)
    return {"ok": True, **{k: parsed[k] for k in (
        "account", "period", "owner", "debit_total", "credit_total")},
        "read_count": len(parsed["operations"]), "new_count": len(new),
        "duplicate_count": len(parsed["operations"]) - len(new)}


def list_transactions(status=None, limit=30):
    items = load_registry()["transactions"]
    if status:
        items = [item for item in items if item.get("status") == status]
    return {"ok": True, "transactions": list(reversed(items[-int(limit):]))}


def get_transaction(transaction_id):
    item = next((item for item in load_registry()["transactions"] if item["id"] == transaction_id), None)
    return {"ok": bool(item), "transaction": item, "error": None if item else "операция не найдена"}


@transaction(lambda:REGISTRY_PATH)
def assign(transaction_id, deal_id, deal_name=""):
    data = load_registry()
    item = next((item for item in data["transactions"] if item["id"] == transaction_id), None)
    if not item:
        return {"ok": False, "error": "операция не найдена"}
    item.update({"status": "assigned", "deal_id": str(deal_id), "deal_name": deal_name,
                 "assigned_at": datetime.now().isoformat(timespec="seconds")})
    save_registry(data)
    return {"ok": True, "transaction": item}


def main():
    spec = json.load(sys.stdin)
    action = spec.get("action")
    if action == "import":
        result = import_statement(spec["path"])
    elif action == "list":
        result = list_transactions(spec.get("status"), spec.get("limit", 30))
    elif action == "get":
        result = get_transaction(spec["transaction_id"])
    elif action == "assign":
        result = assign(spec["transaction_id"], spec["deal_id"], spec.get("deal_name", ""))
    else:
        result = {"ok": False, "error": "неизвестное действие"}
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        raise SystemExit(1)
