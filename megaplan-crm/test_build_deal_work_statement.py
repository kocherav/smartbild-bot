import os
import tempfile
import unittest

from openpyxl import load_workbook

import build_deal_work_statement as report


ENTRIES = [
    {"name": "Оборудование без работ", "unit": "шт.", "qty": 2, "sell_price": 100000, "cost_material": 70000, "cost_work": 0},
    {"name": "Монтаж", "unit": "шт.", "qty": 3, "sell_price": 15000, "cost_material": 1000, "cost_work": 2500},
]


class WorkStatementTest(unittest.TestCase):
    def test_rows_are_filtered_only_after_signing(self):
        self.assertEqual(len(report.statement_rows(ENTRIES, "Подписание договора")), 2)
        self.assertEqual(len(report.statement_rows(ENTRIES, "Выставление счета")), 1)
        self.assertEqual(len(report.statement_rows(ENTRIES, "Отказ")), 2)

    def test_xlsx_contains_only_work_money_columns(self):
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            result = report.build_xlsx({"path": path, "deal_id": "888", "client": "Тест",
                                        "status": "Выставление счета", "entries": ENTRIES})
            self.assertEqual(result["row_count"], 1)
            ws = load_workbook(path, data_only=False).active
            text = " ".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None)
            self.assertNotIn("Цена продажи", text)
            self.assertNotIn("Себестоимость материалов", text)
            self.assertIn("Себестоимость работ/ед.", text)
            self.assertNotIn("Оборудование без работ", text)
            self.assertIn("Монтаж", text)
            total_row = ws.max_row
            self.assertEqual(ws.cell(total_row, 5).value, f"=SUM(E4:E{total_row - 1})")
            self.assertEqual(ws.cell(4, 1).value, "Монтаж")
        finally:
            os.remove(path)


if __name__ == "__main__":
    unittest.main()
