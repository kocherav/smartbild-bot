import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import json
from openpyxl import load_workbook
import installation_materials as m


class MaterialsTests(unittest.TestCase):
    def setUp(self):
        self.items = [dict(offer_row_id='1', offer_id='10', name='Кабель ВВГнг 3×2,5', qty=100, unit='м')]
        self.update = dict(offer_row_id='1', source_name='Другое название', qty_used=40, source_qty=200,
                           source_unit='м', source_total=99999)
        self.a = dict(id='a', deal_id='894', updates=[self.update])
        self.data = dict(invoices=[dict(document_key='i', number='42', allocations=[self.a])])

    def test_confirmed_not_whole_invoice_and_readonly(self):
        before = copy.deepcopy(self.data)
        rows, _ = m.build_view(894, self.items, self.data)
        self.assertEqual(rows[0]['confirmed'], 40)
        self.assertNotIn('99999', json.dumps(rows))
        self.assertEqual(self.data, before)

    def test_duplicates_and_other_deal(self):
        self.data['invoices'][0]['allocations'] += [self.a, dict(id='b', deal_id='861', updates=[self.update])]
        self.assertEqual(m.build_view(894, self.items, self.data)[0][0]['confirmed'], 40)

    def test_fuzzy_does_not_count_and_conflict_is_visible(self):
        self.a['updates'] = []
        self.a['additional_expenses'] = [dict(source_name='Кабель ВВГнг 3×1,5', qty=20, unit='м')]
        rows, extra = m.build_view(894, self.items, self.data)
        self.assertEqual(rows[0]['confirmed'], 0)
        self.assertIn('Характеристики различаются', rows[0]['candidates'][0]['warning'])
        self.assertEqual(len(extra), 1)

    def test_missing_units_warn(self):
        self.update.pop('source_unit')
        rows, _ = m.build_view(894, self.items, self.data)
        self.assertIn('не сохранена', rows[0]['links'][0]['warning'])

    def test_stock_scoped_and_remaining_not_confirmed(self):
        self.data['unallocated_stock'] = [dict(id='s', document_key='i', name=self.items[0]['name'], unit='м',
            available_qty=20, allocations=[dict(id='s1', deal_id='894', offer_row_id='1', qty=5),
                                          dict(id='s2', deal_id='861', offer_row_id='1', qty=60)])]
        rows, extra = m.build_view(894, self.items, self.data)
        self.assertEqual(rows[0]['confirmed'], 45)
        self.assertEqual(extra[0]['qty'], 20)

    def test_ambiguous_offer_id_not_counted(self):
        self.items.append({**self.items[0], 'offer_row_id': '2'})
        self.update.pop('offer_row_id')
        self.update['offer_id'] = '10'
        rows, _ = m.build_view(894, self.items, self.data)
        self.assertTrue(all(r['confirmed'] == 0 for r in rows))

    def test_workbook_no_prices_or_formulas(self):
        self.update['source_name'] = '=HYPERLINK("bad")'
        with tempfile.TemporaryDirectory() as tmp:
            registry = Path(tmp)/'registry.json'
            registry.write_text(json.dumps(self.data))
            out = Path(tmp)/'materials.xlsx'
            with patch.object(m.procurement, 'REGISTRY', str(registry)), patch.object(m.procurement, 'procurement_items', return_value=self.items):
                m.export(894, out)
            wb = load_workbook(out)
            values = [c.value for s in wb for row in s for c in row]
            self.assertNotIn(99999, values)
            self.assertFalse(any(c.data_type == 'f' for s in wb for row in s for c in row))
            self.assertEqual(wb['Сопоставления']['E6'].value, 40)
            mobile = wb.worksheets[0]
            self.assertEqual(mobile.title, 'Обзор для телефона')
            self.assertEqual(mobile.max_column, 4)
            self.assertIn('A1:D1', str(mobile.merged_cells))
            self.assertFalse(mobile.sheet_view.showGridLines)
            self.assertTrue(all(mobile.row_dimensions[n].height for n in range(1, mobile.max_row+1)))
            self.assertIn('A1:J1', str(wb['Сопоставления'].merged_cells))
            self.assertEqual(mobile['A5'].value, 'Частично')
            self.assertEqual(mobile['D5'].value, '40 м')
            self.assertEqual(mobile['B5'].hyperlink.target, "#'Сопоставления'!B6")
            self.assertNotEqual(mobile['A5'].fill.fgColor.rgb, mobile['B5'].fill.fgColor.rgb)
            self.assertEqual(mobile.max_row, 6)  # title/notes/header + one item + footer
            with patch.object(m.procurement, 'REGISTRY', str(registry)), patch.object(m.procurement, 'procurement_items', return_value=self.items):
                m.export(894, out, mobile_only=True)
            compact = load_workbook(out)
            self.assertEqual(len(compact.worksheets), 1)
            self.assertEqual(compact.active.max_column, 4)
            self.assertEqual(compact.active['D5'].value, '40 м')
            self.assertTrue(compact.active.column_dimensions['E'].hidden)
            self.assertEqual(compact.active.column_dimensions['E'].max, 16384)
            self.assertFalse(any(c.hyperlink for row in compact.active for c in row))

    def test_unmatched_exact_names_visible_in_summary(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = 'Сопоставления'
        names = ['S201 C10 6кА Автоматический выключатель 1-полюсный', '=EXTERNAL()']
        extras = [dict(name=name, qty=5, unit='', invoice='2507', suggested=False,
                       kind='Учтено в прочей себестоимости' if i == 0 else 'Нераспределённый остаток — НЕ закреплён за объектом') for i, name in enumerate(names)]
        extras.append(dict(name='Похожая строка', qty=1, unit='м', invoice='2', suggested=True, kind='Учтено в прочей себестоимости'))
        m.mobile_sheet(wb, 894, [], extras)
        sheet = wb.worksheets[0]
        cells = [c for row in sheet for c in row]
        for name in names:
            cell = next(c for c in cells if c.value == name)
            self.assertEqual(cell.data_type, 's')
            self.assertIsNotNone(cell.hyperlink)
        self.assertNotIn('Похожая строка', [c.value for c in cells])
        self.assertIn('Прочая себестоимость', [c.value for c in cells])
        self.assertIn('Свободный остаток — не объект', [c.value for c in cells])


if __name__ == '__main__':
    unittest.main()
