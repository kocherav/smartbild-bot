import copy,tempfile,unittest
from pathlib import Path
from unittest.mock import Mock,patch
import document_order as flow
from workflow_runtime import Store

class DocumentOrderTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
  self.deal={'id':'77','number':'1.77','contractor':{'contentType':'ContractorCompany','id':'4'},'payer':{'id':'wrong'},'price':{'value':200},'positions':[{'id':'r','name':'Светильник','quantity':2,'sum':{'value':200},'tax':{'rate':5}}]}
  self.payer={'id':'8','name':'ООО Клиент','inn':'123','legalAddress':'Адрес','currentAccount':'123','bank':'Банк','bik':'1','correspondentAccount':'2'}
  self.details={'date':'20.10.2026','address':'Москва','payer':'ООО Клиент','name':'Иванов Иван Иванович','position':'Директор','authority':'Устав'}
  def call(method,path,**kwargs):
   if method=='POST':self.deal.update(kwargs['json_body']);return {'data':self.deal}
   if path=='/deal/77':return {'data':copy.deepcopy(self.deal)}
   if path=='/contractorCompany/4':return {'data':{'payers':[self.payer]}}
   if path=='/payer/8':return {'data':self.payer}
   raise AssertionError(path)
  self.api=Mock();self.api._call.side_effect=call
  for name,value in [('get_signatory',{}),('require_document_tax',[5])]:
   p=patch.object(flow,name,return_value=value);p.start();self.addCleanup(p.stop)
 def test_no_automatic_payer_or_mutation(self):
  result=flow.plan({'deal_id':'77'},self.api)
  self.assertFalse(result['ready']);self.assertEqual(result['missing'],['Дата','Адрес','Плательщик'])
  self.assertTrue(all(c.args[0]=='GET' for c in self.api._call.call_args_list))
 def test_explicit_details_and_real_contract_parser(self):
  result=flow.plan({'deal_id':'77','details':self.details},self.api)
  self.assertTrue(result['ready'],result['missing']);self.assertEqual(result['data']['payer_id'],'8');self.assertEqual(result['data']['total'],200)
 def test_wrong_payer_stays_question(self):
  r=flow.plan({'deal_id':'77','details':{**self.details,'payer':'Чужой'}},self.api);self.assertFalse(r['ready'])
 def test_end_to_end_steps_and_delivery_retry_do_not_recreate_invoice(self):
  doc=self.root/'contract.docx';pdf=self.root/'contract.pdf';doc.write_bytes(b'doc');pdf.write_bytes(b'pdf')
  store=Store(self.root/'runtime.sqlite3');send=Mock(side_effect=[{'ok':True},{'ok':True},RuntimeError('network'),{'ok':True}])
  upload=Mock(return_value={'ok':True,'files':['a','b']});fetch=Mock(return_value={'ok':True,'path':str(pdf),'filename':'Счёт №10500 от 20.10.2026.pdf'})
  spec={'deal_id':'77','details':self.details,'chat_id':'owner'}
  with patch.object(flow,'validate_confirmed_source'),patch.object(flow,'selected_seller',return_value={}),patch.object(flow,'generate',return_value=(doc,pdf)) as render,patch('contract_invoice.preview_full',return_value={'fingerprint':'fp'}),patch('contract_invoice.create_full',return_value={'ok':True,'invoice_id':'9'}) as invoice:
   with self.assertRaises(RuntimeError):flow.fulfill(spec,self.api,store,send,fetch,upload)
   result=flow.fulfill(spec,self.api,store,send,fetch,upload)
   self.assertTrue(result['complete']);invoice.assert_called_once();upload.assert_called_once();render.assert_called_once();self.assertEqual(send.call_count,4)
   writes=sum(c.args[0]=='POST' for c in self.api._call.call_args_list)
   with self.assertRaisesRegex(ValueError,'уже выпущен'):flow.fulfill({**spec,'details':{**self.details,'address':'Другой адрес'}},self.api,store,send,fetch,upload)
   self.assertEqual(writes,sum(c.args[0]=='POST' for c in self.api._call.call_args_list))
 def test_full_payment_contract_for_work_uses_total_not_zero_advance(self):
  from types import SimpleNamespace
  from docx import Document
  import generate_contract as gc
  self.deal['positions'][0]['name']='Монтажные работы'
  data=flow.plan({'deal_id':'77','details':self.details},self.api)['data']
  self.assertEqual(data['works'],200)
  data.update(payment_advance=200,payment_final=0,_seller={'id':'cliwell'})
  def convert(args,**kwargs):
   Path(args[-1]).with_suffix('.pdf').write_bytes(b'%PDF-1.4 test transport')
   return SimpleNamespace(returncode=0,stdout='',stderr='')
  with patch('document_renderer.office_command',return_value='office'),patch.object(gc.subprocess,'run',side_effect=convert):
   docx,pdf=gc.generate(data)
  self.addCleanup(__import__('shutil').rmtree,str(Path(docx).parent),True)
  document=Document(docx)
  paragraphs=[p.text for p in gc.all_paragraphs(document)]
  advance=next(p for p in paragraphs if p.startswith('6.4.1'))
  self.assertFalse(any(p.startswith('6.4.2') for p in paragraphs))
  self.assertIn('100% от цены договора',advance)
  self.assertIn('200 рублей 00 копеек',advance)
  # Both implicit zero (goods only) and nonzero final instalments use the same renderer.
  for works,expected in [(0,False),(200,True)]:
   case={**data,'works':works};case.pop('payment_final');case.pop('payment_advance')
   with patch('document_renderer.office_command',return_value='office'),patch.object(gc.subprocess,'run',side_effect=convert):
    other,_=gc.generate(case)
   self.addCleanup(__import__('shutil').rmtree,str(Path(other).parent),True)
   finals=[p.text for p in gc.all_paragraphs(Document(other)) if p.text.startswith('6.4.2')]
   self.assertEqual(bool(finals),expected)
   if expected:self.assertIn('200 рублей 00 копеек',finals[0])
 def test_fingerprint_ignores_payer_runtime_metadata_but_not_legal_terms(self):
  data=flow.plan({'deal_id':'77','details':self.details},self.api)['data']
  with patch.object(flow,'selected_seller',return_value={'id':'cliwell'}):
   before=flow.document_fingerprint(data)
   data['payer']['timeUpdated']={'value':'new'};data['payer']['balance']=999
   self.assertEqual(before,flow.document_fingerprint(data))
   data['payer']['currentAccount']='changed'
   self.assertNotEqual(before,flow.document_fingerprint(data))
 def test_upload_no_effect_retry_and_uncertain_reconciliation(self):
  doc=self.root/'contract.docx';pdf=self.root/'contract.pdf';doc.write_bytes(b'doc');pdf.write_bytes(b'pdf')
  spec={'deal_id':'77','details':self.details,'chat_id':'owner'}
  store=Store(self.root/'recovery.sqlite3')
  upload=Mock(side_effect=[{'ok':False,'no_effect':True,'error':'missing dependency'}, {'ok':False,'error':'timeout'}, {'ok':True}])
  send=Mock(return_value={'ok':True});fetch=Mock(return_value={'ok':True,'path':str(pdf),'filename':'Счёт №10500 от 20.10.2026.pdf'})
  with patch.object(flow,'validate_confirmed_source'),patch.object(flow,'selected_seller',return_value={}),patch.object(flow,'generate',return_value=(doc,pdf)) as render,patch('contract_invoice.preview_full',return_value={'fingerprint':'fp'}),patch('contract_invoice.create_full',return_value={'ok':True,'invoice_id':'9'}) as invoice:
   with self.assertRaises(flow.OrderStepError):flow.fulfill(spec,self.api,store,send,fetch,upload)
   self.payer['timeUpdated']='updated after request'
   with self.assertRaises(flow.OrderStepError):flow.fulfill(spec,self.api,store,send,fetch,upload)
   with self.assertRaises(flow.OrderStepError) as blocked:flow.fulfill(spec,self.api,store,send,fetch,upload)
   self.assertTrue(blocked.exception.result['requires_reconciliation']);self.assertEqual(upload.call_count,2);invoice.assert_not_called()
   op=store.unresolved('owner')[0]
   store.reconcile(op['id'],op['revision'],'operator','failed','Test transport confirms no remote write')
   self.assertTrue(flow.fulfill(spec,self.api,store,send,fetch,upload)['complete'])
   self.assertTrue(flow.fulfill(spec,self.api,store,send,fetch,upload)['complete'])
   self.assertEqual(upload.call_count,3);invoice.assert_called_once();render.assert_called_once();self.assertEqual(send.call_count,3)
 def test_upload_import_failure_is_no_effect_before_network(self):
  import builtins,deal_browser
  original=builtins.__import__
  def fail(name,*args,**kwargs):
   if name=='generate_contract':raise ModuleNotFoundError("No module named 'pymorphy3'")
   return original(name,*args,**kwargs)
  with patch('builtins.__import__',side_effect=fail):
   result=deal_browser.contract_upload_preflight()
  self.assertTrue(result['no_effect']);self.assertFalse(result['ok'])

 def test_invoice_delivery_preserves_filename_and_bytes(self):
  src=self.root/'random.pdf';src.write_bytes(b'%PDF exact original')
  p=Path(flow.named_invoice_file({'path':str(src),'filename':'Счёт №10500 от 20.10.2026.pdf'}))
  self.addCleanup(__import__('shutil').rmtree,str(p.parent),True)
  self.assertEqual(p.name,'Счёт №10500 от 20.10.2026.pdf');self.assertEqual(p.read_bytes(),src.read_bytes())
  with self.assertRaises(ValueError):flow.named_invoice_file({'path':str(src),'filename':'../unsafe.pdf'})
 def test_delivery_term_only_changes_specification_deadline(self):
  import generate_contract as gc
  doc=gc.Document(gc.TEMPLATE)
  before=[p.text for p in gc.all_paragraphs(doc)]
  gc.apply_delivery_term(doc,{'delivery_workdays':20,'works':0})
  after=[p.text for p in gc.all_paragraphs(doc)]
  changed=[(a,b) for a,b in zip(before,after) if a!=b]
  self.assertEqual(len(changed),1)
  self.assertIn('срок поставки продукции: 20 (двадцать) рабочих дней',changed[0][1])
  self.assertIn('момента получения авансового платежа',changed[0][1])
