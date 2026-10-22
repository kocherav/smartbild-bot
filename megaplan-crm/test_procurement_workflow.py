import json
import os
import tempfile
import unittest
from unittest import mock
from openpyxl import load_workbook

import procurement_workflow as pw


class ProcurementPaymentTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.registry = os.path.join(self.temp.name, "procurement.json")
        self.bank = os.path.join(self.temp.name, "bank.json")
        self.patch = mock.patch.multiple(pw, REGISTRY=self.registry, BANK_REGISTRY=self.bank)
        self.patch.start()

    def tearDown(self):
        self.patch.stop(); self.temp.cleanup()

    def write(self, path, value):
        with open(path, "w", encoding="utf-8") as target:
            json.dump(value, target, ensure_ascii=False)

    def test_stock_consumption_updates_invoice_summary_once(self):
        self.write(self.registry, {'invoices':[{'document_key':'doc','total':100,
            'allocations':[{'amount':60}], 'allocated_total':60}],
            'unallocated_stock':[{'id':'stock','document_key':'doc','available_qty':4,
            'available_amount':40,'allocations':[]}]})
        pw.consume_unallocated_stock('stock','861',3,offer_row_id='row')
        pw.consume_unallocated_stock('stock','861',3,offer_row_id='row')
        data=pw.load_json(self.registry,{})
        self.assertEqual(90,data['invoices'][0]['allocated_total'])
        self.assertEqual(10,data['invoices'][0]['unallocated_total'])
        self.assertEqual(1,len(data['invoices'][0]['allocations']))
        self.assertEqual(1,len(data['unallocated_stock'][0]['allocations']))
        pw.save_registry(data)
        self.assertEqual(90,pw.load_json(self.registry,{})['invoices'][0]['allocated_total'])

    def test_repeat_invoice_sees_consumed_internal_stock(self):
        doc = {"supplier_inn": "123", "buyer_inn": "456", "number": "A-1",
               "date": "01.01.2026", "total": 1000}
        key = pw.invoice_document_key(doc)
        self.write(self.registry, {"deals": {}, "invoices": [{
            "document_key": key, "allocated_total": 0, "allocations": []
        }], "unallocated_stock": [{
            "document_key": key, "source_line_key": "1:SKU", "initial_qty": 10,
            "initial_amount": 1000, "allocations": [{"deal_id": "1", "qty": 10}]
        }]})
        spec = {"invoice": doc, "updates": [{
            "source_line_key": "1:SKU", "source_qty": 10, "source_total": 1000,
            "qty_used": 1, "material_cost_used": 100
        }]}
        status = pw.invoice_allocation_status(spec)
        self.assertEqual(10, status["source_usage"]["1:sku"]["qty"])
        self.assertFalse(pw.prepare_invoice_allocation(spec)["ok"])

    @mock.patch.object(pw, "snapshot")
    def test_exact_expense_matches_supplier_invoice(self, snapshot):
        self.write(self.registry, {"deals": {}, "invoices": [{
            "id": "inv1", "deal_id": "890", "supplier_inn": "123", "number": "A-77",
            "total": 1000, "status": "привязан",
            "purchasing_company": {"id": "cliwell", "name": "ООО «КЛИВЭЛ»"},
        }]})
        self.write(self.bank, {"imports": [], "transactions": [{
            "id": "tx1", "account": "40702810902570001580", "direction": "expense", "counterparty_inn": "123",
            "purpose": "Оплата по счету № А-77", "amount": "1000.00", "date": "22.09.2026",
        }]})
        result = pw.reconcile_bank()
        self.assertEqual(1, len(result["matched"]))
        with open(self.registry, encoding="utf-8") as source:
            saved = json.load(source)["invoices"][0]
        self.assertEqual("оплачен", saved["status"])
        self.assertEqual(1000, saved["paid_total"])
        with open(self.bank, encoding="utf-8") as source:
            self.assertEqual("assigned", json.load(source)["transactions"][0]["status"])

    @mock.patch.object(pw, "snapshot")
    def test_number_without_matching_inn_is_not_applied(self, snapshot):
        self.write(self.registry, {"deals": {}, "invoices": [{
            "id": "inv1", "deal_id": "890", "supplier_inn": "123", "number": "77", "total": 1000,
        }]})
        self.write(self.bank, {"imports": [], "transactions": [{
            "id": "tx1", "direction": "expense", "counterparty_inn": "999",
            "purpose": "Оплата по счету №77", "amount": "1000", "date": "22.09.2026",
        }]})
        self.assertFalse(pw.reconcile_bank()["matched"])

    def test_partial_and_full_receipt(self):
        self.write(self.registry, {"deals": {"890": {"deal_id": "890", "items": [{
            "offer_row_id": "10", "name": "Кабель", "qty": 5, "unit": "м",
            "received_qty": 0, "remaining_qty": 5, "status": "запросить счёт",
        }]}}, "invoices": []})
        first = pw.record_receipt("890", "10", 2)
        self.assertEqual("частично получено", first["status"])
        self.assertEqual(3, first["remaining_qty"])
        second = pw.record_receipt("890", "10", 3)
        self.assertEqual("получено", second["status"])
        self.assertEqual(0, second["remaining_qty"])

    def test_receipt_cannot_exceed_ordered_quantity(self):
        self.write(self.registry, {"deals": {"890": {"deal_id": "890", "items": [{
            "offer_row_id": "10", "name": "Кабель", "qty": 5, "received_qty": 4,
        }]}}, "invoices": []})
        with self.assertRaisesRegex(ValueError, "превышает остаток"):
            pw.record_receipt("890", "10", 2)

    def test_closure_requires_invoice_payment_and_receipt(self):
        self.write(self.registry, {"deals": {"890": {"deal_id": "890", "items": [{
            "offer_row_id": "10", "name": "Кабель", "qty": 5, "actual_qty": 5,
            "received_qty": 5,
        }]}}, "invoices": [{
            "id": "inv", "deal_id": "890", "number": "7", "total": 1000, "paid_total": 500,
        }]})
        check = pw.closure_check("890")
        self.assertFalse(check["ready"])
        self.assertEqual("payment", check["problems"][0]["type"])

    def test_ready_procurement_can_be_closed(self):
        self.write(self.registry, {"deals": {"890": {"deal_id": "890", "items": [{
            "offer_row_id": "10", "name": "Кабель", "qty": 5, "actual_qty": 5,
            "received_qty": 5,
        }]}}, "invoices": [{
            "id": "inv", "deal_id": "890", "number": "7", "total": 1000, "paid_total": 1000,
            "closing_documents": [{"id": "upd-1", "filename": "УПД.pdf"}],
        }]})
        result = pw.close_procurement("890")
        self.assertTrue(result["closed"])
        with open(self.registry, encoding="utf-8") as source:
            self.assertEqual("закрыта", json.load(source)["deals"]["890"]["status"])

    def test_procurement_status_shows_invoice_payment_and_receipt(self):
        self.write(self.registry, {"deals": {"890": {"items": [{
            "qty": 5, "actual_qty": 5, "received_qty": 2
        }]}}, "invoices": [{
            "deal_id": "890", "number": "77", "date": "22.09.2026", "supplier": "Поставщик",
            "total": 1000, "paid_total": 400
        }]})
        status = pw.procurement_status("890")
        self.assertEqual("частично оплачен", status["invoices"][0]["payment_status"])
        self.assertEqual(5, status["invoiced_qty"])
        self.assertEqual(2, status["received_qty"])

    def test_closure_requires_supplier_closing_document(self):
        self.write(self.registry, {"deals": {"890": {"deal_id": "890", "items": [{
            "offer_row_id": "10", "name": "Кабель", "qty": 5, "actual_qty": 5, "received_qty": 5,
        }]}}, "invoices": [{
            "id": "inv", "deal_id": "890", "number": "7", "total": 1000, "paid_total": 1000,
        }]})
        check = pw.closure_check("890")
        self.assertFalse(check["ready"])
        self.assertEqual("closing", check["problems"][0]["type"])

    def test_closing_document_is_idempotent(self):
        self.write(self.registry, {"deals": {"890": {"deal_id": "890", "items": []}}, "invoices": [{
            "id": "inv", "deal_id": "890", "number": "7", "total": 1000,
        }]})
        first = pw.record_closing_document("890", "inv", "УПД.pdf", "same-sha")
        second = pw.record_closing_document("890", "inv", "УПД.pdf", "same-sha")
        self.assertFalse(first["duplicate"])
        self.assertTrue(second["duplicate"])

    @mock.patch.object(pw, "snapshot")
    def test_invoice_records_intercompany_purchase(self, snapshot):
        self.write(self.registry, {"deals": {"887": {"selling_company": {
            "id": "cliwell", "name": "ООО «КЛИВЭЛ»", "inn": "7743151323"
        }}}, "invoices": []})
        row = pw.record_invoice({"deal_id": "887", "invoice": {
            "supplier": "Поставщик", "supplier_inn": "123", "buyer": "ООО Инсис",
            "buyer_inn": "7743926877", "number": "77", "date": "22.09.2026",
            "total": 1200, "vat": 200,
        }, "updates": [{"offer_id": "1", "offer_row_id": "2", "qty_used": 1,
                          "material_cost_used": 1200, "source_line_key": "sku-1"}],
           "additions": []}, {})
        self.assertTrue(row["is_intercompany"])
        self.assertEqual("insis", row["purchasing_company"]["id"])
        self.assertEqual("не компенсировано", row["intercompany_status"])
        self.assertEqual(1200, pw.accounting_hints("887")["balances"][0]["balance"])

    @mock.patch.object(pw, "snapshot")
    def test_bank_payment_must_use_invoice_buyer_company_account(self, snapshot):
        self.write(self.registry, {"deals": {}, "invoices": [{
            "id": "inv1", "deal_id": "887", "supplier_inn": "123", "number": "77",
            "total": 1000, "purchasing_company": {"id": "insis", "name": "ООО «Инсис»"},
        }]})
        self.write(self.bank, {"imports": [], "transactions": [{
            "id": "tx1", "account": "40702810902570001580", "direction": "expense",
            "counterparty_inn": "123", "purpose": "Оплата по счету №77", "amount": "1000",
            "date": "22.09.2026",
        }]})
        result = pw.reconcile_bank()
        self.assertFalse(result["matched"])
        self.assertIn("счёт выставлен", result["review"][0]["reason"])

    @mock.patch.object(pw, "snapshot")
    def test_one_invoice_can_be_allocated_to_two_deals_without_duplication(self, snapshot):
        self.write(self.registry, {"deals": {}, "invoices": []})
        doc = {"supplier": "Поставщик", "supplier_inn": "123", "buyer": "ООО Инсис",
               "buyer_inn": "7743926877", "number": "99", "date": "23.09.2026", "total": 3000}
        first = {"deal_id": "887", "invoice": doc, "updates": [{"offer_id": "1", "offer_row_id": "11",
                 "qty_used": 1, "material_cost_used": 1800, "source_line_key": "sku-a"}], "additions": []}
        second = {"deal_id": "890", "invoice": doc, "updates": [{"offer_id": "2", "offer_row_id": "22",
                  "qty_used": 1, "material_cost_used": 1200, "source_line_key": "sku-b",
                  "source_qty": 1, "source_total": 1200}], "additions": []}
        first["updates"][0].update({"source_qty": 1, "source_total": 1800})
        pw.record_invoice(first, {}); pw.record_invoice(second, {})
        with open(self.registry, encoding="utf-8") as source:
            saved = json.load(source)["invoices"]
        self.assertEqual(1, len(saved))
        self.assertEqual(2, len(saved[0]["allocations"]))
        self.assertEqual(3000, saved[0]["allocated_total"])
        self.assertEqual(0, saved[0]["unallocated_total"])
        duplicate = pw.prepare_invoice_allocation(first)
        self.assertFalse(duplicate["ok"])
        self.assertIn("больше количества", duplicate["error"])

    def test_distribution_prefills_nearby_deals_and_keeps_remainder(self):
        self.write(self.registry, {"deals": {
            "894": {"items": []},
            "901": {"items": [{"offer_id": "o1", "offer_row_id": "r1", "name": "Кабель ВВГнг 3x2,5",
                                  "sku": "VVG-325", "qty": 60, "actual_qty": 0}]},
            "905": {"items": [{"offer_id": "o2", "offer_row_id": "r2", "name": "Кабель ВВГнг 3x2,5",
                                  "sku": "VVG-325", "qty": 25, "actual_qty": 0}]},
        }, "invoices": [], "unallocated_stock": []})
        spec = {"deal_id": "894", "invoice": {"number": "1", "total": 36000},
                "updates": [], "unallocated": [{"source_line_key": "cable", "name": "Кабель ВВГнг 3x2,5",
                                                   "sku": "VVG-325", "qty": 100, "amount": 18000}]}
        plan = pw.suggest_distribution(spec)
        self.assertEqual(["905", "901"], [x["deal_id"] for x in plan["suggestions"]])
        self.assertAlmostEqual(15, plan["inventory_remainder"][0]["qty"])
        self.assertAlmostEqual(2700, plan["inventory_remainder"][0]["amount"])

    def test_user_can_keep_all_remainder_in_internal_stock(self):
        self.write(self.registry, {"deals": {}, "invoices": []})
        spec = {"deal_id": "894", "invoice": {"supplier": "Поставщик", "supplier_inn": "123",
                "buyer": "ООО Инсис", "buyer_inn": "7743926877", "number": "77",
                "date": "23.09.2026", "total": 18000}}
        rows = [{"source_line_key": "cable", "name": "Кабель", "sku": "VVG", "qty": 100,
                 "amount": 18000, "unit": "м"}]
        created = pw.put_unallocated_stock(spec, rows)
        self.assertEqual(100, created[0]["available_qty"])
        # Replaying the same confirmation is idempotent.
        self.assertEqual(created[0]["id"], pw.put_unallocated_stock(spec, rows)[0]["id"])
        self.assertEqual(1, len(pw.list_unallocated_stock("insis")))

    def test_invoice_2507_partial_repeat_moves_exact_four_rows_to_stock(self):
        doc = {"supplier": "ИП Панькин Олег Юрьевич", "supplier_inn": "775114176346",
               "buyer": "ООО «КЛИВЭЛ»", "buyer_inn": "7743151323", "number": "2507",
               "date": "23.09.2026", "total": 59971}
        key = pw.invoice_document_key(doc)
        self.write(self.registry, {"deals": {}, "invoices": [{
            "id": "inv2507", "document_key": key, **doc, "allocated_total": 35595,
            "allocations": [{"deal_id": "894", "amount": 35595, "source_lines": [
                {"key": "matched-a", "qty_used": 5, "amount_used": 15175},
                {"key": "matched-b", "qty_used": 1, "amount_used": 6644},
                {"key": "matched-c", "qty_used": 1, "amount_used": 13776},
            ]}]
        }], "unallocated_stock": [], "allocation_audit": [], "mapping_memory": []})
        rows = [
            {"source_line_key": "1SDA066807R1", "name": "XT1B", "qty": 2, "amount": 20316},
            {"source_line_key": "2CDS251001R0104", "name": "S201", "qty": 5, "amount": 2120},
            {"source_line_key": "2CDL220001R1012", "name": "PS2/12", "qty": 2, "amount": 1258},
            {"source_line_key": "2CDL210001R1012", "name": "PS1/12", "qty": 2, "amount": 682},
        ]
        status = pw.invoice_allocation_status({"invoice": doc})
        self.assertEqual(35595, status["allocated_total"])
        self.assertEqual(24376, status["invoice_total"] - status["allocated_total"])
        created = pw.put_unallocated_stock({"invoice": doc}, rows, actor="test")
        self.assertEqual(4, len(created))
        self.assertEqual(24376, sum(x["available_amount"] for x in created))
        self.assertEqual(11, sum(x["available_qty"] for x in created))
        replayed = pw.put_unallocated_stock({"invoice": doc}, rows, actor="test")
        self.assertEqual([x["id"] for x in created], [x["id"] for x in replayed])
        self.assertEqual(4, len(pw.list_unallocated_stock("cliwell")))

    def test_export_fills_supplier_and_highlights_fully_invoiced_rows(self):
        self.write(self.registry, {"deals": {"894": {"deal_id": "894", "items": [
            {"offer_id": "o1", "offer_row_id": "r1", "name": "Кабель", "qty": 10,
             "actual_qty": 10, "actual_total": 1000, "remaining_qty": 10},
            {"offer_id": "o2", "offer_row_id": "r2", "name": "Крепёж", "qty": 5,
             "actual_qty": 2, "actual_total": 100, "remaining_qty": 5},
        ]}}, "invoices": [{"supplier": "ООО Поставщик", "allocations": [{"deal_id": "894",
            "updates": [{"offer_id": "o1", "offer_row_id": "r1", "qty_used": 10},
                        {"offer_id": "o2", "offer_row_id": "r2", "qty_used": 2}]}]}]})
        path = pw.export("894")
        wb = load_workbook(path); ws = wb["Закупка"]
        self.assertEqual("ООО Поставщик", ws["A2"].value)
        self.assertEqual("полностью привязано к счетам", ws["L2"].value)
        self.assertEqual("00C6EFCE", ws["A2"].fill.fgColor.rgb)
        self.assertEqual("ООО Поставщик", ws["A3"].value)
        self.assertNotEqual("00C6EFCE", ws["A3"].fill.fgColor.rgb)

    def test_internal_stock_can_be_partially_used_by_later_deal(self):
        self.write(self.registry, {"deals": {}, "invoices": [], "unallocated_stock": [{
            "id": "stock1", "name": "Кабель", "unit": "м", "available_qty": 100,
            "available_amount": 18000, "allocations": [], "purchasing_company": {"id": "insis"},
        }]})
        used = pw.consume_unallocated_stock("stock1", "910", 15, "o", "r")
        self.assertEqual(2700, used["amount"])
        remaining = pw.list_unallocated_stock()[0]
        self.assertEqual(85, remaining["available_qty"])
        self.assertEqual(15300, remaining["available_amount"])


if __name__ == "__main__":
    unittest.main()
