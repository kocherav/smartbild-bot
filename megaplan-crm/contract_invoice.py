#!/usr/bin/env python3
"""Create a one-line Megaplan invoice from a payment clause in the latest deal contract."""
import json
import os
import re
import base64
import email.utils
import hashlib
import hmac
import subprocess
import sys
import tempfile
import urllib.parse
import urllib.request
import megaplan_transport
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from docx import Document

from megaplan_api import MegaplanAPI

RECIPIENT_ID = "1000560"       # ООО «КЛИВЭЛ»
INVOICE_TEMPLATE_ID = "3"      # Кливэл
NO_VAT_TAX_ID = "7"
SERVICE_UNIT_ID = "9"
CONTRACT_SOURCES = Path(__file__).resolve().parent / "contracts"


def money(value):
    return f"{value:,.2f}".replace(",", " ").replace(".", ",")


def megaplan_date(value):
    if not isinstance(value, dict) or not value.get("year"):
        return ""
    return f"{int(value.get('day', 0)):02d}.{int(value.get('month', 0)) + 1:02d}.{int(value['year']):04d}"


def contract_basis(number, contract_date):
    number = str(number or "").strip()
    contract_date = str(contract_date or "").strip()
    if not number or not contract_date:
        raise ValueError("в сделке не заполнены номер или дата договора для поля «Основание» счёта")
    return f"Договор № {number} от {contract_date}"


def basis_from_deal(deal):
    number = deal.get("Category1000068CustomFieldNomerDogovora") or deal.get("number")
    contract_date = megaplan_date(deal.get("Category1000068CustomFieldDataDogovora"))
    return contract_basis(number, contract_date)


def document_text(path):
    suffix = Path(path).suffix.casefold()
    if suffix == ".docx":
        document = Document(path)
        # Requisites may be inside nested tables, hyperlinks or content controls.
        # Top-level document.tables/cell.paragraphs silently omit those parties.
        # Preserve soft line breaks and tabs, including those in nested content.
        # Dropping w:br can merge adjacent payment clauses and even select the
        # amount from a different clause.
        chunks = [''.join((node.text or '') if node.tag.endswith('}t')
                          else '\t' if node.tag.endswith('}tab') else '\n'
                          for node in paragraph.xpath('.//w:t | .//w:br | .//w:cr | .//w:tab'))
                  for paragraph in document.element.xpath('.//w:p')]
        return "\n".join(chunks)
    if suffix == ".pdf":
        result = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, timeout=30)
        if result.returncode:
            raise ValueError("не удалось прочитать PDF договора")
        return result.stdout
    raise ValueError("последний договор должен быть в DOCX или PDF")


def parse_contract(text, clause):
    normalized = text.replace("\xa0", " ")
    number_match = re.search(r"ДОГОВОР(?:\s+ПОДРЯДА)?\s*№\s*([^\n]+)", normalized, re.I)
    party_start=re.search(r'(?mi)^\s*(?:Общество\s+с\s+ограниченной|Индивидуальн\w*\s+предприниматель|ООО\s|ИП\s)',normalized)
    header=normalized[:party_start.start()] if party_start else normalized[:(number_match.end()+600 if number_match else 600)]
    date_match = re.search(r"\b(\d{2}\.\d{2}\.\d{4})\b", header)
    clause_match = re.search(
        rf"(?ms)^\s*{re.escape(clause)}\s+(.+?)(?=^\s*6\.4\.\d+\s+|^\s*6\.5(?:\.|\s)|\Z)",
        normalized,
    )
    if not number_match or not date_match:
        raise ValueError("не удалось определить номер или дату последнего договора")
    day,month,year=map(int,date_match.group(1).split('.'))
    date(year,month,day) # reject malformed header dates before any financial use
    if not clause_match:
        raise ValueError(f"в последнем договоре не найден пункт {clause}")
    clause_text = re.sub(r"\s+", " ", clause_match.group(1)).strip()
    amount_match = (
        re.search(r"(?:что\s+)?составляет\s+([\d ]+(?:[,.]\d{1,2})?)\s*(?:руб|\()", clause_text, re.I)
        or re.search(r"(?:размере|сумме)\s+([\d ]+(?:[,.]\d{1,2})?)\s*(?:руб|\()", clause_text, re.I)
    )
    if not amount_match:
        raise ValueError(f"в пункте {clause} не найдена сумма платежа")
    amount = float(amount_match.group(1).replace(" ", "").replace(",", "."))
    number = re.split(r'\s+от\s+\d{2}\.\d{2}\.\d{4}',number_match.group(1),maxsplit=1,flags=re.I)[0].strip()
    if re.search(r"окончательн", clause_text, re.I):
        label = "Окончательный платеж"
    elif re.search(r"предоплат|авансов", clause_text, re.I):
        label = "Авансовый платеж"
    else:
        label = "Платеж"
    line = f"{label}, согласно п.{clause} Договора № {number} от {date_match.group(1)}, сумма {money(amount)} руб."
    contract_date = date_match.group(1)
    return {"clause": clause, "contract_number": number, "contract_date": contract_date,
            "basis": contract_basis(number, contract_date),
            "amount": amount, "line": line, "tax_id": NO_VAT_TAX_ID}


def specification_text(text, deal_id):
    """Apply an explicit owner clarification only to the exact reviewed document.

    The original PDF remains unchanged. The included line is not billed twice.
    """
    path=CONTRACT_SOURCES / (str(deal_id)+'-source.json')
    source=json.loads(path.read_text()) if path.exists() else {}
    approved=source.get('included_specification_lines')
    if not approved:return text
    if hashlib.sha256(text.encode()).hexdigest()!=approved['text_sha256']:
        raise ValueError('Изменилась редакция договора с включёнными расходными материалами')
    start=text.upper().rfind('СПЕЦИФИКАЦИЯ')
    if start<0:raise ValueError('Не найдена спецификация')
    lines=text[start:].splitlines(keepends=True)
    for approved_line in approved['lines']:
        matches=[i for i,line in enumerate(lines) if line.rstrip('\r\n')==approved_line]
        if len(matches)!=1:raise ValueError('Согласованная включённая строка не найдена однозначно')
        del lines[matches[0]]
    return text[:start]+''.join(lines)


def contract_scope(positions, deal_id):
    path=CONTRACT_SOURCES / (str(deal_id)+'-source.json')
    source=json.loads(path.read_text()) if path.exists() else {}
    extra=source.get('supplemental_rows') or {}
    byid={str(r['id']):r for r in positions}
    for ident,approved in extra.items():
        row=byid.get(ident)
        if not row or row['name']!=approved['name'] or any(Decimal(str(row[k]))!=Decimal(str(approved[k])) for k in ('quantity','total')):
            raise ValueError('Изменился согласованный объём дополнительных работ: '+ident)
    return [r for r in positions if str(r['id']) not in extra]


def latest_contract(api, deal_id, context=None):
    source_path=CONTRACT_SOURCES / (str(deal_id)+'-source.json')
    source=json.loads(source_path.read_text()) if source_path.exists() else {}
    rejected=set(map(str,source.get('rejected_file_ids',[])))
    files=[]; params={'limit':100}; seen=set()
    while True:
        response=api._call('GET',f'/deal/{deal_id}/allFiles',params=params)
        page=response.get('data') or []; files.extend(page)
        if not response.get('meta',{}).get('pagination',{}).get('hasMoreNext'):break
        if not page or page[-1]['id'] in seen:raise ValueError('Не получен полный список договоров сделки')
        seen.add(page[-1]['id']);params={'limit':100,'pageAfter':{'contentType':'File','id':page[-1]['id']}}
    candidates = [f for f in files if str(f.get("id")) not in rejected and "договор" in (f.get("name") or "").casefold()
                  and not re.search(r'сч[её]т[\s_–—-]*договор',f.get('name') or '',re.I)
                  and Path(f.get("name") or "").suffix.casefold() in (".docx", ".pdf")]
    if source.get('signed_file_id'):
        candidates=[f for f in candidates if str(f.get('id'))==str(source['signed_file_id'])]
        if not candidates:raise ValueError('Закреплённый договор отсутствует во вложениях сделки; подмена черновиком запрещена')
    if not candidates:
        if source.get('confirmed_total'):
            raise ValueError('Сумма '+money(Decimal(source['confirmed_total']))+' ₽ подтверждена владельцем. Доступная старая редакция исключена. Подписанный договор ещё не найден среди вложений сделки; требуется восстановить привязку ранее загруженного файла.')
        raise ValueError("в сделке нет загруженного договора DOCX/PDF")
    candidates.sort(key=lambda f: (str(f.get("id"))==str(source.get("signed_file_id")), (f.get("timeCreated") or {}).get("value") or "",
                                   Path(f.get("name") or "").suffix.casefold() == ".docx",
                                   int(f.get("id") or 0)), reverse=True)
    for selected in candidates:
        suffix = Path(selected.get("name") or "").suffix.casefold()
        fd, path = tempfile.mkstemp(prefix=f"deal-{deal_id}-invoice-contract-", suffix=suffix)
        os.close(fd)
        try:
            url = selected["path"] if selected["path"].startswith("http") else f"https://{api.domain}{selected['path']}"
            request = urllib.request.Request(url, headers={"Authorization": f"Bearer {api.token}"})
            with megaplan_transport.urlopen(request, timeout=60) as response, open(path, "wb") as output:
                output.write(response.read())
            if str(selected.get('id'))==str(source.get('signed_file_id')) and source.get('sha256'):
                if hashlib.sha256(Path(path).read_bytes()).hexdigest()!=source['sha256']:
                    raise ValueError('Содержимое закреплённого договора изменилось')
            if context:
                text=document_text(path)
                header=re.search(r'ДОГОВОР(?:\s+ПОДРЯДА)?\s*№\s*([^\n]+)',text,re.I)
                normalize=lambda v:re.sub(r'\s+','',str(v).replace('№',''))
                number=re.split(r'\s+от\s+\d{2}\.\d{2}\.\d{4}',header[1],flags=re.I)[0] if header else ''
                if normalize(number)!=normalize(context['number']) or any(not re.search(r'(?<!\d)'+re.escape(str(context[k]))+r'(?!\d)',text) for k in ('seller_inn','payer_inn')):
                    os.unlink(path);continue
            return selected,path
        except BaseException:
            Path(path).unlink(missing_ok=True);raise
    raise ValueError('Не найден договор с заказчиком с номером и сторонами этой сделки; договоры поставщиков исключены')


def preview_fingerprint(data):
    return hashlib.sha256(json.dumps({k:v for k,v in data.items() if k not in ('ok','fingerprint')},sort_keys=True,ensure_ascii=False,allow_nan=False).encode()).hexdigest()


def validate_contract_context(text,parsed,deal,seller,payer):
    """A recent filename is not proof that a contract belongs to current parties."""
    for label,party in (('нашей организации',seller),('плательщика',payer)):
        inn=str(party.get('inn') or '').strip()
        if not re.fullmatch(r'\d{10}|\d{12}',inn):
            raise ValueError('Не подтверждён ИНН '+label+' для проверки договора')
        if not re.search(r'(?<!\d)'+re.escape(inn)+r'(?!\d)',text):
            raise ValueError('Последний договор не содержит текущий ИНН '+label+'. Сначала сформируйте и прикрепите актуальный договор.')
    expected_number=str(deal.get('Category1000068CustomFieldNomerDogovora') or deal.get('number') or '').strip()
    expected_date=megaplan_date(deal.get('Category1000068CustomFieldDataDogovora'))
    if not expected_number or not expected_date:
        raise ValueError('Для проверки договора заполните его номер и дату в сделке')
    normalized=lambda value:re.sub(r'\s+','',str(value).replace('№',''))
    if normalized(parsed['contract_number'])!=normalized(expected_number) or parsed['contract_date']!=expected_date:
        raise ValueError('Номер или дата последнего договора не совпадают с текущей сделкой. Требуется актуальная версия договора.')


def preview(spec, api):
    deal_id = str(spec["deal_id"])
    clause = str(spec.get("clause") or "6.4.1").strip()
    if not re.fullmatch(r"6\.4\.\d+", clause):
        raise ValueError("пункт оплаты должен иметь вид 6.4.1")
    deal = api._call("GET", f"/deal/{deal_id}", params={"fields": ["payer", "invoices", "positions", "number",
        "Category1000068CustomFieldNomerDogovora", "Category1000068CustomFieldDataDogovora"]})["data"]
    from document_tax_policy import require_document_tax, payment_tax
    require_document_tax(deal,api)
    payer = deal.get("payer") or {}
    if not payer.get("id"):
        raise ValueError("к сделке не привязан плательщик")
    from seller_documents import selected_seller
    payer_full=api._call('GET',f"/payer/{payer['id']}")['data']
    seller=selected_seller(deal_id)
    selected, path = latest_contract(api, deal_id,context={'number':deal.get('Category1000068CustomFieldNomerDogovora') or deal.get('number'),'seller_inn':seller['inn'],'payer_inn':payer_full['inn']})
    try:
        text=document_text(path)
        parsed = parse_contract(text, clause)
        parsed = payment_tax(parsed,deal,api)
        from seller_documents import selected_seller
        validate_contract_context(text,parsed,deal,selected_seller(deal_id),payer_full)
    finally:
        os.remove(path)
    from seller_documents import invoice_settings
    recipient_id,template_id=invoice_settings(deal_id)
    data={"ok": True, "deal_id": deal_id, "payer_id": str(payer["id"]),
            "payer_name": payer.get("name") or "", "contract_file_id": str(selected.get("id")),
            "contract_filename": selected.get("name") or "", 'recipient_id':recipient_id,'template_id':template_id, **parsed}
    return {**data,'fingerprint':preview_fingerprint(data)}


def payment_offer(api, line, tax_id=NO_VAT_TAX_ID):
    article = "CONTRACT-PAYMENT-" + hashlib.sha1(line.encode()).hexdigest()[:12].upper()
    return api._call("POST", "/offer", json_body={
        "contentType": "Offer", "name": line, "article": article,
        "price": {"contentType": "Money", "currency": "RUB", "value": 0},
        "unit": {"contentType": "Unit", "id": SERVICE_UNIT_ID},
        "tax": {"contentType": "Tax", "id": str(tax_id)}, "quantity": 0,
    })["data"]


def v1_call(api, path, form, auth=None):
    if auth is None:
        auth_data = v1_call(api, "/BumsCommonApiV01/User/authorize.api", {
            "Login": api.username, "Password": hashlib.md5(api.password.encode()).hexdigest(),
        }, auth=False)["data"]
        auth = (auth_data.get("AccessId") or auth_data["access_id"],
                auth_data.get("SecretKey") or auth_data["secret_key"])
    body = urllib.parse.urlencode(form).encode()
    headers = {"Accept": "application/json", "Content-Type": "application/x-www-form-urlencoded"}
    if auth is not False:
        request_date = email.utils.formatdate(usegmt=True)
        canonical = f"POST\n\napplication/x-www-form-urlencoded\n{request_date}\n{api.domain}{path}"
        digest = hmac.new(auth[1].encode(), canonical.encode(), hashlib.sha1).hexdigest()
        signature = base64.b64encode(digest.encode()).decode()
        headers.update({"Date": request_date, "X-Authorization": f"{auth[0]}:{signature}"})
    request = urllib.request.Request(f"https://{api.domain}{path}", data=body, headers=headers)
    with megaplan_transport.urlopen(request, timeout=30) as response:
        result = json.loads(response.read())
    if (result.get("status") or {}).get("code") != "ok":
        raise ValueError((result.get("status") or {}).get("message") or "ошибка API Мегаплана")
    return result


def build_invoice_v1_forms(data, offer_id, invoice_date=None):
    day = (invoice_date or date.today()).isoformat()
    common = {"Model[Recipient]": RECIPIENT_ID, "Model[Payer]": data["payer_id"],
              "Model[Consignee]": data["payer_id"], "Model[Date]": day, "Deal": data["deal_id"]}
    common["Model[InvoiceTemplate]"] = INVOICE_TEMPLATE_ID
    common["Model[PaymentOrderNumber]"] = data["basis"]
    row = {"Model[Unit]": SERVICE_UNIT_ID, "Model[Quantity]": 1, "Model[OrderPos]": 1,
           "Model[Name]": data["line"], "Model[Offer]": str(offer_id),
           "Model[Tax]": data["tax_id"], "Model[Price]": data["amount"]}
    return {**common, "Model[Status]": 1}, row, {**common, "Model[Status]": 2}


def v1_auth(api):
    auth_data = v1_call(api, "/BumsCommonApiV01/User/authorize.api", {
        "Login": api.username, "Password": hashlib.md5(api.password.encode()).hexdigest(),
    }, auth=False)["data"]
    return (auth_data.get("AccessId") or auth_data["access_id"],
            auth_data.get("SecretKey") or auth_data["secret_key"])


def normalized_invoice_text(value):
    return ' '.join(str(value or '').split())


def verify_created_invoice(api, invoice_id, data, rows, auth=None):
    """Read back an issued invoice; safe both after creation and during recovery."""
    from seller_documents import invoice_settings
    recipient_id, template_id = invoice_settings(data['deal_id'])
    invoice = api._call('GET', f'/invoice/{invoice_id}', params={'fields': ['rows', 'deal', 'invoiceTemplate']})['data']
    saved = v1_call(api, '/BumsInvoiceApiV01/Invoice/card.api', {
        'Id': str(invoice_id), 'RequestedFields[0]': 'PaymentOrderNumber',
    }, auth).get('data', {}).get('invoice', {})
    if (str((invoice.get('deal') or {}).get('id')) != str(data['deal_id'])
            or str((invoice.get('payer') or {}).get('id')) != str(data['payer_id'])
            or str((invoice.get('recipient') or {}).get('id')) != str(recipient_id)
            or str((invoice.get('invoiceTemplate') or {}).get('id')) != str(template_id)
            or invoice.get('status') not in ('drawn', 'paid', 'partially_paid')
            or len(invoice.get('rows') or []) != len(rows)
            or normalized_invoice_text(saved.get('PaymentOrderNumber')) != normalized_invoice_text(data['basis'])):
        raise ValueError('Мегаплан создал счёт, но не сохранил требуемые параметры')
    for actual, expected in zip(invoice['rows'], rows):
        if (normalized_invoice_text(actual.get('name')) != normalized_invoice_text(expected['name'])
                or Decimal(str(actual.get('quantity'))) != Decimal(str(expected['quantity']))
                or Decimal(str((actual.get('price') or {}).get('value'))) != Decimal(str(expected['price']))
                or str((actual.get('tax') or {}).get('id')) != str(expected.get('tax_id') or NO_VAT_TAX_ID)):
            raise ValueError('Строки созданного счёта не совпадают с подтверждёнными')
    expected_sum = sum((Decimal(str(r['quantity'])) * Decimal(str(r['price']))).quantize(Decimal('.01'), rounding=ROUND_HALF_UP) for r in rows)
    saved_sum = invoice.get('sum') or {}
    actual_sum = Decimal(str(saved_sum.get('value')))
    if not actual_sum.is_finite() or actual_sum != expected_sum or saved_sum.get('currency') not in (None, 'RUB'):
        raise ValueError('Итог созданного счёта не совпадает с подтверждёнными рублёвыми строками')
    return invoice


def create_v1_invoice(api, data, rows):
    from seller_documents import invoice_settings
    recipient_id, template_id = invoice_settings(data['deal_id'])
    if (data.get('recipient_id') and data['recipient_id']!=recipient_id) or (data.get('template_id') and data['template_id']!=template_id):
        raise ValueError('Наша организация или шаблон счёта изменились после проверки')
    recipient = api._call('GET', f'/payer/{recipient_id}')['data']
    if not recipient.get('currentAccount'):
        raise ValueError('У получателя счёта отсутствует расчётный счёт')
    common = {"Model[Recipient]": recipient_id, "Model[Payer]": data["payer_id"],
              "Model[Consignee]": data["payer_id"], "Model[Date]": date.today().isoformat(),
              "Model[InvoiceTemplate]": template_id,
              "Model[PaymentOrderNumber]": data["basis"], "Deal": data["deal_id"]}
    auth = v1_auth(api)
    invoice_id = None
    try:
        created = v1_call(api, "/BumsInvoiceApiV01/Invoice/save.api",
                          {**common, "Model[Status]": 1}, auth)
        invoice_data = (created.get("data") or {}).get("Invoice") or {}
        invoice_id = str(invoice_data["Id"])
        for index, row in enumerate(rows, 1):
            form = {"Model[Invoice]": invoice_id, "Model[Unit]": str(row.get("unit_id") or SERVICE_UNIT_ID),
                    "Model[Quantity]": row.get("quantity") or 1, "Model[OrderPos]": index,
                    "Model[Name]": row["name"], "Model[Offer]": str(row["offer_id"]),
                    "Model[Tax]": str(row.get("tax_id") or NO_VAT_TAX_ID), "Model[Price]": row["price"]}
            v1_call(api, "/BumsInvoiceApiV01/InvoiceRow/save.api", form, auth)
        v1_call(api, "/BumsInvoiceApiV01/Invoice/save.api",
                {**common, "Model[Status]": 2, "Id": invoice_id}, auth)
        return verify_created_invoice(api, invoice_id, data, rows, auth)
    except Exception as exc:
        if invoice_id:
            # A timeout can happen after issue succeeded. Deleting a potentially
            # issued invoice is not a rollback. Preserve the exact ID for review.
            raise ValueError(f'Счёт ID {invoice_id} создан частично либо требует сверки; повторное создание запрещено до проверки ({type(exc).__name__}).') from exc
        raise


def create(spec, api):
    data = preview(spec, api)
    if spec.get('expected_fingerprint') and data['fingerprint']!=spec['expected_fingerprint']:
        return {'ok':False,'no_effect':True,'error':'Договор, плательщик или наша организация изменены. Подготовьте счёт заново.'}
    from seller_documents import invoice_settings
    recipient_id, _ = invoice_settings(data['deal_id'])
    # Idempotency: never create the same contractual payment twice silently.
    deal = api._call("GET", f"/deal/{data['deal_id']}", params={"fields": ["invoices"]})["data"]
    for ref in deal.get("invoices") or []:
        invoice = api._call("GET", f"/invoice/{ref['id']}", params={"fields": ["rows"]})["data"]
        if (str((invoice.get('recipient') or {}).get('id')) == recipient_id and
                any((row.get("name") or "") == data["line"] for row in invoice.get("rows") or [])):
            raise ValueError(f"счёт по пункту {data['clause']} этого договора уже существует: №{invoice.get('number')}")
    offer = payment_offer(api, data["line"], data["tax_id"])
    invoice = create_v1_invoice(api, data, [{"name": data["line"], "offer_id": offer["id"],
                                             "quantity": 1, "price": data["amount"], "tax_id": data["tax_id"]}])
    if (invoice.get("rows") or [{}])[0].get("name") != data["line"]:
        raise ValueError("Мегаплан изменил текст строки счёта")
    return {**data, "invoice_id": str(invoice["id"]), "invoice_number": invoice.get("number"),
            "invoice_status": invoice.get("status"), "public_download_link": invoice.get("publicDownloadLink")}


def full_invoice_rows(deal):
    """Use net RUB row totals (after discounts), never catalogue/list prices."""
    rows, total = [], Decimal('0')
    for position in deal.get('positions') or []:
        offer = position.get('offer') or {}
        amount = position.get('sum') or {}
        try:
            qty = Decimal(str(position.get('quantity')))
            net = Decimal(str(amount.get('value')))
        except (InvalidOperation, ValueError):
            raise ValueError('В позиции счёта отсутствует количество или итоговая сумма') from None
        if not qty.is_finite() or qty <= 0 or not net.is_finite() or net < 0:
            raise ValueError('Неверное количество или сумма позиции счёта')
        if amount.get('currency') not in ('RUB', None):
            raise ValueError('Для счёта нужна зафиксированная рублёвая спецификация')
        total += net
        price = net / qty
        rows.append({'name': position.get('name') or offer.get('name') or 'Позиция',
                     'offer_id': offer['id'], 'quantity': str(qty), 'price': str(price),
                     'unit_id': (position.get('unit') or offer.get('unit') or {}).get('id') or SERVICE_UNIT_ID,
                     'tax_id': (position.get('tax') or offer.get('tax') or {}).get('id') or NO_VAT_TAX_ID})
    expected = deal.get('price') or {}
    try:
        expected_value = Decimal(str(expected.get('value')))
    except (InvalidOperation, ValueError):
        raise ValueError('В сделке отсутствует итоговая сумма') from None
    if not rows or not expected_value.is_finite() or expected.get('currency') not in ('RUB', None) or total != expected_value:
        raise ValueError('Рублёвые суммы строк не совпадают с итогом сделки; счёт не создан')
    return rows


def preview_full(spec, api):
    deal_id = str(spec["deal_id"])
    deal = api._call("GET", f"/deal/{deal_id}", params={"fields": ["positions", "payer"]})["data"]
    payer = deal.get("payer") or {}
    if not payer.get("id"):
        raise ValueError("к сделке не привязан плательщик")
    positions = deal.get("positions") or []
    if not positions:
        raise ValueError("в сделке нет позиций для полного счёта")
    from document_tax_policy import settings,require_document_tax
    if settings().get("companies"):require_document_tax(deal,api)
    rows = full_invoice_rows(deal)
    basis = basis_from_deal(deal)
    from seller_documents import invoice_settings
    recipient_id,template_id=invoice_settings(deal_id)
    data={'ok':True,'deal_id':deal_id,'payer_id':str(payer['id']),'payer_name':payer.get('name') or '',
          'basis':basis,'rows':rows,'total':str((deal.get('price') or {}).get('value')),
          'recipient_id':recipient_id,'template_id':template_id}
    return {**data,'fingerprint':preview_fingerprint(data)}


def create_full(spec, api):
    data=preview_full(spec,api)
    if spec.get('expected_fingerprint') and data['fingerprint']!=spec['expected_fingerprint']:
        return {'ok':False,'no_effect':True,'error':'Спецификация, плательщик или наша организация изменены. Подготовьте счёт заново.'}
    deal_id=data['deal_id']; rows=data['rows']; basis=data['basis']
    invoice = create_v1_invoice(api, data, rows)
    return {"ok": True, "deal_id": deal_id, "invoice_id": str(invoice["id"]),
            "invoice_number": invoice.get("number"), "invoice_status": invoice.get("status"),
            "rows_count": len(rows), "basis": basis,
            "public_download_link": invoice.get("publicDownloadLink")}


def list_invoices(spec, api):
    deal_id = str(spec["deal_id"])
    refs = api._call("GET", f"/deal/{deal_id}/invoices", params={"limit": 100}).get("data") or []
    result = []
    for ref in refs:
        invoice = api._call("GET", f"/invoice/{ref['id']}", params={"fields": ["rows", "deal", "invoiceTemplate"]})["data"]
        rows = invoice.get("rows") or []
        invoice_date = invoice.get("invoiceDate") or {}
        date_text = (f"{invoice_date.get('day', 0):02d}.{invoice_date.get('month', 0) + 1:02d}.{invoice_date.get('year')}"
                     if invoice_date else "")
        result.append({"id": str(invoice["id"]), "number": invoice.get("number"), "date": date_text,
                       "status": invoice.get("status"), "sum": (invoice.get("sum") or {}).get("value"),
                       "rows_count": len(rows), "line": rows[0].get("name") if len(rows) == 1 else None,
                       "template": (invoice.get("invoiceTemplate") or {}).get("name")})
    result.sort(key=lambda item: int(item["id"]), reverse=True)
    return {"ok": True, "deal_id": deal_id, "invoices": result}


def delete_invoice(spec, api):
    deal_id, invoice_id = str(spec["deal_id"]), str(spec["invoice_id"])
    invoice = api._call("GET", f"/invoice/{invoice_id}")["data"]
    if (invoice.get("deal") or {}).get("id") != deal_id:
        raise ValueError("счёт не относится к выбранной сделке")
    api._call("DELETE", f"/invoice/{invoice_id}")
    return {"ok": True, "deal_id": deal_id, "invoice_id": invoice_id}


def main():
    try:
        spec = json.load(sys.stdin)
        api = MegaplanAPI()
        action = spec.get("action")
        if action == "preview":
            result = preview(spec, api)
        elif action == 'preview_full':
            result = preview_full(spec,api)
        elif action == "create":
            result = create(spec, api)
        elif action == "create_full":
            result = create_full(spec, api)
        elif action == "list":
            result = list_invoices(spec, api)
        elif action == 'validate_bundle':
            from document_bundle import validate
            result=validate(spec,api)
        elif action == "delete":
            result = delete_invoice(spec, api)
        else:
            raise ValueError("неизвестная операция со счётом")
        print(json.dumps(result, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
