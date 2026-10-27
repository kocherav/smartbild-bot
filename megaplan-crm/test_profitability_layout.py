import tempfile
import unittest
from pathlib import Path

import openpyxl

import build_profitability_report as report


def row(name, qty, price, material, work, profit=None):
    revenue = round(qty * price / 1.05, 2)
    cost = material + work
    return {"name": name, "qty": qty, "sell_price": price, "revenue": revenue, "plan_material": material,
            "plan_work": work, "plan_cost": cost, "fact_qty": 0, "fact_material": 0,
            "remaining_plan_material": material, "fact_work_qty": 0, "fact_work": 0, "remaining_plan_work": work,
            "forecast_cost": cost, "deviation": 0, "profit": round(revenue - cost, 2) if profit is None else profit}


class ProfitabilityLayoutTest(unittest.TestCase):
    """Owner 27.10 (deal 903): totals under every money column and a reason for every loss."""

    def build(self, rows):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        path = Path(folder.name) / 'e.xlsx'
        report.build_xlsx(rows, '903', str(path), [], None)
        ws = openpyxl.load_workbook(path).worksheets[0]
        return {r[0].value: [c.value for c in r] for r in ws.iter_rows() if r[0].value}, ws

    def test_totals_row_sums_every_money_column(self):
        rows = [row('Шкаф', 1, 10000, 7000, 1000), row('Кабель', 100, 100, 6000, 2000)]
        table, _ = self.build(rows)
        totals = table['Итого по позициям']
        self.assertEqual(round(10000 / 1.05 + 10000 / 1.05, 2), totals[3])
        self.assertEqual(13000, totals[4])
        self.assertEqual(3000, totals[5])
        self.assertEqual(16000, totals[6])
        self.assertEqual(round(sum(r['profit'] for r in rows), 2), totals[15])

    def test_every_loss_has_a_reason(self):
        rows = [row('Без цены', 2, 0, 0, 1000), row('Дороже', 10, 100, 1200, 0), row('Без наценки', 1, 500, 500, 0),
                row('Нормальная', 1, 1000, 500, 100)]
        table, ws = self.build(rows)
        self.assertEqual('нет цены продажи: затраты без выручки', table['Без цены'][16])
        self.assertEqual('закупка дороже цены продажи', table['Дороже'][16])
        self.assertEqual('без наценки: закупка равна цене продажи', table['Без наценки'][16])
        self.assertIn(table['Нормальная'][16], ('', None))
        self.assertIn('выручка без НДС − прогноз себестоимости', ws.cell(row=3, column=16).value)
        summary = {x['remark']: x for x in report.remarks_summary(rows)}
        self.assertEqual(1000, summary['нет цены продажи: затраты без выручки']['amount'])
        self.assertEqual(200, summary['закупка дороже цены продажи']['amount'])


if __name__ == '__main__':
    unittest.main()
