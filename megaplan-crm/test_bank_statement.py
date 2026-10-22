import os
import tempfile
import unittest
from openpyxl import Workbook
import bank_statement


class BankStatementTest(unittest.TestCase):
    def make_book(self):
        wb = Workbook(); ws = wb.active
        ws["A1"] = "Выписка по счёту"; ws["B1"] = "40702 810 9 00000 000001"
        ws["B2"] = "c 01.09.2026 по 02.09.2026"; ws["B3"] = "ООО Тест"; ws["B4"] = "123"
        ws["A7"] = "Остаток входящий"; ws["D7"] = "Обороты дебет"; ws["E7"] = 100
        ws["A8"] = "Остаток исходящий"; ws["D8"] = "Обороты кредит"; ws["E8"] = 250
        headers = ["Дата", "Номер документа", "Дебет", "Кредит", "Контрагент", "ИНН", "КПП", "Счёт", "БИК", "Наименование банка", "Назначение платежа", "Код дебитора", "Тип документа"]
        for col, value in enumerate(headers, 1): ws.cell(11, col, value)
        ws.cell(13, 1, "01.09.2026"); ws.cell(13, 2, "1"); ws.cell(13, 3, 100); ws.cell(13, 5, "Поставщик"); ws.cell(13, 11, "Договор №1")
        ws.cell(14, 1, "02.09.2026"); ws.cell(14, 2, "2"); ws.cell(14, 4, 250); ws.cell(14, 5, "Клиент"); ws.cell(14, 11, "Счет 5")
        fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd); wb.save(path); return path

    def test_parse_both_directions_and_totals(self):
        path = self.make_book()
        try:
            result = bank_statement.parse_xlsx(path)
            self.assertEqual(2, len(result["operations"]))
            self.assertEqual("expense", result["operations"][0]["direction"])
            self.assertEqual("income", result["operations"][1]["direction"])
            self.assertEqual("100.00", result["debit_total"])
            self.assertEqual("250.00", result["credit_total"])
            self.assertNotEqual(result["operations"][0]["id"], result["operations"][1]["id"])
        finally: os.remove(path)

    def test_parse_1c_cp1251_both_directions(self):
        body = """1CClientBankExchange
ВерсияФормата=1.03
Кодировка=Windows
ДатаНачала=01.09.2026
ДатаКонца=02.09.2026
РасчСчет=40700000000000000001
СекцияРасчСчет
РасчСчет=40700000000000000001
ВсегоПоступило=250.00
ВсегоСписано=100.00
КонецРасчСчет
СекцияДокумент=Платежное поручение
Номер=1
Дата=01.09.2026
Сумма=100.00
ПлательщикСчет=40700000000000000001
Плательщик=ООО Мы
ПолучательСчет=40700000000000000002
Получатель=ООО Поставщик
ПолучательИНН=2
НазначениеПлатежа=По договору 1
КонецДокумента
СекцияДокумент=Платежное поручение
Номер=2
Дата=02.09.2026
Сумма=250.00
ПлательщикСчет=40700000000000000003
Плательщик=ООО Клиент
ПлательщикИНН=3
ПолучательСчет=40700000000000000001
Получатель=ООО Мы
НазначениеПлатежа=Оплата счета 2
КонецДокумента
КонецФайла
"""
        fd, path = tempfile.mkstemp(suffix=".txt"); os.close(fd)
        try:
            with open(path, "wb") as target: target.write(body.encode("cp1251"))
            result = bank_statement.parse_statement(path)
            self.assertEqual(2, len(result["operations"]))
            self.assertEqual("expense", result["operations"][0]["direction"])
            self.assertEqual("ООО Поставщик", result["operations"][0]["counterparty"])
            self.assertEqual("income", result["operations"][1]["direction"])
            self.assertEqual("ООО Клиент", result["operations"][1]["counterparty"])
        finally: os.remove(path)

    def test_aggregated_excel_commission_matches_split_1c_rows(self):
        base = {"account": "1", "date": "08.09.2026", "direction": "expense",
                "counterparty_inn": "2", "counterparty_account": "3"}
        existing = [{**base, "id": "excel", "amount": "78.00", "purpose": "Комиссия банка"}]
        incoming = [
            {**base, "id": "1c-a", "amount": "39.00", "purpose": "Комиссия за перевод №1"},
            {**base, "id": "1c-b", "amount": "39.00", "purpose": "Комиссия за перевод №2"},
        ]
        self.assertEqual([], bank_statement.remove_cross_format_duplicates(incoming, existing))


if __name__ == "__main__": unittest.main()
