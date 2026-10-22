#!/usr/bin/env python3
"""Export every live Megaplan deal for one contractor to a readable XLSX table."""
import json
import os
import sys
from collections import Counter
from concurrent.futures import ThreadPoolExecutor

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from megaplan_api import MegaplanAPI
import deal_browser


def run(spec):
    contractor_id = str(spec["contractor_id"])
    output = spec["output"]
    shown_ids = {str(value) for value in spec.get("shown_ids") or []}
    api = MegaplanAPI()
    compact = [
        deal for deal in deal_browser.fetch_all(api)
        if str((deal.get("contractor") or {}).get("id")) == contractor_id
    ]

    def fetch(deal):
        return MegaplanAPI()._call("GET", f"/deal/{deal['id']}")["data"]

    with ThreadPoolExecutor(max_workers=min(12, len(compact) or 1)) as pool:
        live = list(pool.map(fetch, compact))
    rows = [deal_browser.summary(deal) for deal in live]
    rows.sort(key=lambda row: int(row["id"]), reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Сделки JetCar"
    headers = [
        "ID", "Номер / название", "Дата создания", "Программа", "Статус", "Адрес",
        "Исполнители", "Сумма, ₽", "Позиций", "Ссылка", "Была в выдаче бота",
    ]
    ws.append(headers)
    for row in rows:
        created = (row.get("created_at") or "")[:10]
        ws.append([
            int(row["id"]), row.get("name"), created, row.get("program"), row.get("status"),
            row.get("address"), ", ".join(row.get("executors") or []), row.get("price"),
            row.get("positions_count"), f"https://cclimat.megaplan.ru/deals/{row['id']}/card/",
            "Да" if row["id"] in shown_ids else "Нет",
        ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 8).number_format = '#,##0.00'
        ws.cell(row, 10).hyperlink = ws.cell(row, 10).value
        ws.cell(row, 10).style = "Hyperlink"
        if ws.cell(row, 11).value == "Нет":
            ws.cell(row, 11).fill = PatternFill("solid", fgColor="FCE4D6")
    widths = [9, 24, 14, 24, 24, 34, 28, 16, 10, 52, 20]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    summary_ws = wb.create_sheet("Сравнение")
    statuses = Counter(row.get("status") or "не указан" for row in rows)
    summary_ws.append(["Показатель", "Значение"])
    summary_ws.append(["Всего сделок клиента в Мегаплане", len(rows)])
    summary_ws.append(["Показано ботом на скриншоте", len(shown_ids & {r['id'] for r in rows})])
    summary_ws.append(["Пропущено ботом", len(rows) - len(shown_ids & {r['id'] for r in rows})])
    summary_ws.append([])
    summary_ws.append(["Статус", "Количество"])
    for status, count in statuses.most_common():
        summary_ws.append([status, count])
    for cell in summary_ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    summary_ws.column_dimensions["A"].width = 42
    summary_ws.column_dimensions["B"].width = 16
    os.makedirs(os.path.dirname(output), exist_ok=True)
    wb.save(output)
    return {"ok": True, "path": output, "count": len(rows), "shown": len(shown_ids),
            "missing": len(rows) - len(shown_ids), "statuses": statuses}


if __name__ == "__main__":
    try:
        result = run(json.load(sys.stdin))
        result["statuses"] = dict(result["statuses"])
        print(json.dumps(result, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        raise SystemExit(1)
