import os
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

import installation_workflow as iw


class InstallationWorkflowTests(unittest.TestCase):
    def test_missing_installation_has_distinct_error(self):
        with self.assertRaises(iw.InstallationNotPrepared):iw.get_installation('absent')

    def test_sync_executor_does_not_create_missing_installation(self):
        with patch.object(iw, 'deal_source', return_value={'installer_name':'Новый Монтажник','installer_ref':'77'}), \
             patch.object(iw, 'backup_db') as backup:
            result=iw.sync_executor('999')
        self.assertFalse(result['installation_exists'])
        backup.assert_not_called()

    def test_sync_executor_updates_existing_installation_and_reads_it_back(self):
        with patch.object(iw, 'deal_source', return_value={'installer_name':'Новый Монтажник','installer_ref':'77'}), \
             patch.object(iw, 'backup_db'):
            result=iw.sync_executor('894','owner')
        self.assertTrue(result['installation_exists'])
        self.assertEqual('Новый Монтажник',result['result']['installer_name'])
        self.assertEqual('77',result['result']['installer_ref'])

    def test_human_folder_essence(self):
        self.assertEqual('электроснабжение 6-й эт НК',iw.brief_essence(
            'поставка и монтаж системы электроснабжения',
            'г.Москва, ул.Подвойского, д.5/19, 6-й этаж НК'))

    def test_project_classifier_rejects_kp_and_accepts_real_project(self):
        self.assertFalse(iw.is_project_document('КП Кондиционирование 1эт НК v2.pdf'))
        self.assertFalse(iw.is_project_document('непонятный файл.pdf'))
        self.assertTrue(iw.is_project_document('Рабочая документация ОВ.pdf'))
        self.assertTrue(iw.is_project_document('Система вентиляции.dwg'))

    def setUp(self):
        reachable=patch.object(iw.deal_browser,'megaplan_reachable',return_value=True)
        reachable.start(); self.addCleanup(reachable.stop)
        self.tmp=tempfile.TemporaryDirectory()
        self.old_db,self.old_dir=iw.DB,iw.WORKBOOK_DIR
        iw.DB=os.path.join(self.tmp.name,'installations.sqlite3')
        iw.WORKBOOK_DIR=os.path.join(self.tmp.name,'files')
        db=iw.connect(); stamp=iw.now()
        db.execute('''INSERT INTO installations(deal_id,client,essence,address,contact_name,contact_phone,installer_name,installer_ref,status,project_required,folder_name,folder_remote,created_at,updated_at)
                      VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',('894','Автосалон JetCar','электроснабжение 6-й эт НК','Москва','Иван','79990000000','Деревянко Сергей','1000142','transferred',0,'folder','remote',stamp,stamp))
        db.execute('INSERT INTO work_items(deal_id,source_key,name,unit,planned_qty,rate,sort_order) VALUES(?,?,?,?,?,?,?)',('894','a','Монтаж кабельной линии','м',100,500,1))
        db.execute('INSERT INTO work_items(deal_id,source_key,name,unit,planned_qty,rate,sort_order) VALUES(?,?,?,?,?,?,?)',('894','b','Монтаж электрического щита','шт.',1,10000,2))
        db.commit(); db.close()

    def tearDown(self):
        iw.DB,self.old_db=self.old_db,iw.DB
        iw.WORKBOOK_DIR=self.old_dir
        self.tmp.cleanup()

    def test_publication_retry_does_not_repeat_financial_acceptance(self):
        db=iw.connect(); iw.enqueue_publication(db,'894'); db.commit(); db.close()
        with patch.object(iw,'build_workbook'), patch.object(iw,'sync_installation'), patch.object(iw,'publish_workbook'), patch.object(iw,'backup_db'), patch.object(iw,'refresh_economics',side_effect=TimeoutError()):
            self.assertIn('economics',iw.publish_pending('894'))
        self.assertEqual({},iw.retry_publications())
        db=iw.connect(); db.execute("UPDATE publication_jobs SET updated_at='2020-01-01T00:00:00+00:00'"); db.commit(); db.close()
        with patch.object(iw,'build_workbook'), patch.object(iw,'sync_installation'), patch.object(iw,'publish_workbook'), patch.object(iw,'backup_db'), patch.object(iw,'refresh_economics'), patch.object(iw,'approve_report') as accept:
            self.assertEqual({'894':{}},iw.retry_publications())
            self.assertEqual({},iw.retry_publications()); accept.assert_not_called()

    def test_submit_does_not_wait_for_archive_and_preserves_item_ids(self):
        report=iw.create_report('894','77','Монтаж кабельной линии 10 м')
        with patch.object(iw,'sync_report') as sync,patch.object(iw,'backup_db') as backup:
            submitted=iw.submit_report(report['id'])
            repeated=iw.submit_report(report['id'])
            sync.assert_not_called(); backup.assert_not_called()
        self.assertEqual('submitted',submitted['report']['status'])
        self.assertEqual([x['id'] for x in submitted['items']],[x['id'] for x in repeated['items']])
        db=iw.connect(); self.assertEqual(1,db.execute('SELECT COUNT(*) FROM report_publications').fetchone()[0]); db.close()

    def test_media_survives_temporary_source_removal_and_is_frozen_on_submit(self):
        report=iw.create_report('894','77','Монтаж кабельной линии 10 м')
        source=os.path.join(self.tmp.name,'photo.jpg')
        with open(source,'wb') as file:file.write(b'fixture')
        iw.add_media(report['id'],source,name='photo.jpg'); os.unlink(source)
        db=iw.connect(); retained=db.execute('SELECT local_path FROM media').fetchone()[0]; db.close()
        self.assertTrue(os.path.isfile(retained))
        iw.submit_report(report['id'])
        with self.assertRaises(ValueError):iw.add_media(report['id'],retained,name='extra.jpg')

    def test_report_publication_cannot_clear_newer_revision(self):
        report=iw.create_report('894','77','Монтаж кабельной линии 10 м')
        iw.submit_report(report['id'])
        def concurrent_return(report_id):iw.return_report(report_id,'Уточнить объём','manager')
        with patch.object(iw,'sync_report',side_effect=concurrent_return),patch.object(iw,'backup_db'):
            iw.retry_report_publications()
        db=iw.connect(); job=db.execute('SELECT * FROM report_publications').fetchone(); db.close()
        self.assertEqual('pending',job['status']); self.assertEqual(2,job['revision'])

    def test_return_report_commits_without_cloud_wait(self):
        report=iw.create_report('894','77','Монтаж кабельной линии 10 м')
        iw.submit_report(report['id'])
        with patch.object(iw,'sync_report') as sync,patch.object(iw,'backup_db') as backup:
            returned=iw.return_report(report['id'],'Уточнить объём','manager')
            sync.assert_not_called(); backup.assert_not_called()
        self.assertEqual('returned',returned['status'])
        self.assertEqual([],iw.pending_reports())

    def test_publication_update_timeout_never_creates_another_comment(self):
        db=iw.connect(); db.execute("UPDATE installations SET megaplan_comment_id='900' WHERE deal_id='894'"); db.commit(); db.close()
        with patch.object(iw,'MegaplanAPI') as factory:
            api=factory.return_value; api.upload_file.return_value={'id':'901'}
            api._call.side_effect=TimeoutError()
            with self.assertRaises(TimeoutError): iw.publish_workbook('894')
            self.assertEqual(1,api._call.call_count)
            self.assertEqual('/comment/900',api._call.call_args.args[1])
        db=iw.connect(); self.assertEqual('900',db.execute('SELECT megaplan_comment_id FROM installations').fetchone()[0]); db.close()

    def test_human_client_and_description_not_contract_number(self):
        from unittest.mock import Mock
        api=Mock();api._call.side_effect=[{'data':{'contractor':{'id':'1','contentType':'ContractorHuman','lastName':'Example','firstName':'Person'},'name':'1.123-2026','description':'<p>Монтаж вентиляции</p>'}},{'data':[]}]
        result=iw.deal_source('1',api)
        self.assertEqual('Example Person',result['client']);self.assertEqual('Монтаж вентиляции',result['essence'])

    def test_reference_only_client_hydrated_by_id(self):
        from unittest.mock import Mock
        api=Mock();api._call.side_effect=[{'data':{'contractor':{'id':'1','contentType':'ContractorCompany'},'name':'Deal'}},{'data':{'name':'Company'}},{'data':[]}]
        result=iw.deal_source('1',api)
        self.assertEqual('Company',result['client']);self.assertEqual('/contractorCompany/1',api._call.call_args_list[1].args[1])

    def test_first_publication_timeout_requires_reconciliation(self):
        with patch.object(iw,'MegaplanAPI') as factory:
            api=factory.return_value; api.upload_file.return_value={'id':'901'}
            api._call.side_effect=TimeoutError()
            with self.assertRaises(TimeoutError): iw.publish_workbook('894')
            with self.assertRaisesRegex(ValueError,'сверка'): iw.publish_workbook('894')
            self.assertEqual(1,api._call.call_count)
            self.assertEqual(1,api.upload_file.call_count)

    def test_readonly_old_comment_creates_one_version_and_uncertain_stops_retry(self):
        db=iw.connect();db.execute("UPDATE installations SET megaplan_comment_id='900' WHERE deal_id='894'");db.commit();db.close()
        with patch.object(iw,'MegaplanAPI') as factory:
            api=factory.return_value;api.upload_file.return_value={'id':'901'}
            api._call.side_effect=[iw.MegaplanAPIError(403,{}),{'data':{'subject':{'id':'894'},'possibleActions':['act_read']}},TimeoutError()]
            with self.assertRaises(TimeoutError):iw.publish_workbook('894')
            with self.assertRaisesRegex(ValueError,'сверка'):iw.publish_workbook('894')
            self.assertEqual(3,api._call.call_count)
            self.assertEqual('/deal/894/comments',api._call.call_args.args[1])
            db=iw.connect();row=db.execute('SELECT megaplan_comment_id,publication_create_state FROM installations').fetchone();db.close()
            self.assertEqual((None,'uncertain'),tuple(row))

    def test_forbidden_comment_for_other_deal_never_creates_version(self):
        db=iw.connect();db.execute("UPDATE installations SET megaplan_comment_id='900' WHERE deal_id='894'");db.commit();db.close()
        with patch.object(iw,'MegaplanAPI') as factory:
            api=factory.return_value;api.upload_file.return_value={'id':'901'}
            api._call.side_effect=[iw.MegaplanAPIError(403,{}),{'data':{'subject':{'id':'999'},'possibleActions':['act_read']}}]
            with self.assertRaises(iw.MegaplanAPIError):iw.publish_workbook('894')
            self.assertEqual(2,api._call.call_count)

    def test_new_publication_revision_is_not_lost_during_sync(self):
        db=iw.connect(); iw.enqueue_publication(db,'894'); db.commit(); db.close()
        def concurrent_change(_deal):
            db=iw.connect(); iw.enqueue_publication(db,'894'); db.commit(); db.close()
        with patch.object(iw,'build_workbook'), patch.object(iw,'sync_installation',side_effect=concurrent_change), patch.object(iw,'publish_workbook'), patch.object(iw,'backup_db'), patch.object(iw,'refresh_economics'):
            iw.publish_pending('894')
        db=iw.connect(); job=db.execute('SELECT * FROM publication_jobs').fetchone(); db.close()
        self.assertEqual('pending',job['status']); self.assertEqual(2,job['revision'])

    def test_invalid_workbook_rolls_back_earlier_rows(self):
        path=iw.build_workbook('894'); wb=load_workbook(path); ws=wb['Монтажная ведомость']
        ws.cell(3,5).value=999; ws.cell(4,4).value=-1; wb.save(path)
        with self.assertRaises(ValueError): iw.import_workbook('894',path)
        db=iw.connect(); rate=db.execute('SELECT rate FROM work_items ORDER BY id LIMIT 1').fetchone()[0]; db.close()
        self.assertEqual(500,rate)

    @patch.object(iw,'backup_db',lambda:None)
    @patch.object(iw,'sync_report',lambda _id:'remote')
    def test_text_report_and_cumulative_columns(self):
        report=iw.create_report('894','77')
        iw.append_report_text(report['id'],'Монтаж кабельной линии 40 м')
        parsed=iw.parse_report(report['id'])
        self.assertEqual(1,len(parsed['items']))
        self.assertEqual('Монтаж кабельной линии',parsed['items'][0]['name'])
        with patch.object(iw,'sync_report',lambda _id:'remote'):
            iw.submit_report(report['id'])
        with patch.object(iw,'sync_installation',lambda _id:'remote'), patch.object(iw,'publish_workbook',lambda _id:{}), patch.object(iw,'refresh_economics',lambda _id:{}):
            result=iw.approve_report(report['id'],'manager')
        self.assertEqual(20000,result['accepted_total'])
        self.assertTrue(os.path.isfile(result['workbook_path']))
        wb=load_workbook(result['workbook_path'],data_only=False); ws=wb['Монтажная ведомость']
        self.assertEqual('ИТОГО ПО ПЛАНУ',ws.cell(ws.max_row,5).value)
        self.assertEqual('Сумма принятого, ₽',wb['Журнал отчётов'].cell(1,6).value)
        self.assertEqual(20000,wb['Журнал отчётов'].cell(2,6).value)

    @patch.object(iw,'backup_db',lambda:None)
    def test_natural_text_can_contain_multiple_clauses_and_word_number(self):
        report=iw.create_report('894','77')
        iw.append_report_text(report['id'],'Смонтировано 18 м кабельной линии, установлен один электрический щит.')
        parsed=iw.parse_report(report['id'])
        self.assertEqual(2,len(parsed['items']))
        self.assertEqual([18,1],[x['declared_qty'] for x in parsed['items']])

    @patch.object(iw,'backup_db',lambda:None)
    def test_overfulfilment_is_blocked(self):
        report=iw.create_report('894','77','Монтаж кабельной линии 120 м')
        iw.parse_report(report['id'])
        with self.assertRaisesRegex(ValueError,'превышает остаток'):
            iw.approve_report(report['id'],'manager')

    @patch.object(iw,'backup_db',lambda:None)
    @patch.object(iw,'sync_installation',lambda _id:'remote')
    @patch.object(iw,'publish_workbook',lambda _id:{})
    @patch.object(iw,'refresh_economics',lambda _id:{})
    def test_workbook_can_remove_and_edit_rows(self):
        path=iw.build_workbook('894')
        from openpyxl import load_workbook
        wb=load_workbook(path); ws=wb['Монтажная ведомость']
        ws.cell(3,5,value=600); ws.delete_rows(4); edited=os.path.join(self.tmp.name,'edited.xlsx'); wb.save(edited)
        result=iw.import_workbook('894',edited,'manager')
        self.assertEqual(1,len(result['works']))
        self.assertEqual(600,result['works'][0]['rate'])
        self.assertTrue(any(x.get('removed') for x in result['changes']))

    @patch.object(iw,'backup_db',lambda:None)
    def test_payment_cannot_exceed_accepted_debt(self):
        with self.assertRaisesRegex(ValueError,'превышает задолженность'):
            iw.record_payment('894','bank',1,'26.09.2026','bank-1')

    def test_acceptance_preserves_cost_when_rate_changes_and_cannot_repeat(self):
        with patch.object(iw,'backup_db'), patch.object(iw,'sync_installation'), patch.object(iw,'publish_workbook'), patch.object(iw,'refresh_economics'):
            report=iw.create_report('894','77','Монтаж кабельной линии 10 м')
            iw.parse_report(report['id']); iw.approve_report(report['id'],'manager')
            db=iw.connect(); db.execute('UPDATE work_items SET rate=9999 WHERE source_key=?',('a',)); db.commit(); db.close()
            self.assertEqual(5000,iw.get_installation('894')['accepted_total'])
            with self.assertRaisesRegex(ValueError,'уже принят'): iw.approve_report(report['id'],'manager')
            with self.assertRaisesRegex(ValueError,'нельзя вернуть'): iw.return_report(report['id'],'undo','manager')

    @patch.object(iw,'sync_report',side_effect=TimeoutError('simulated archive outage'))
    def test_publication_failure_does_not_undo_financial_acceptance(self, sync_report):
        with patch.object(iw,'backup_db'), patch.object(iw,'sync_installation',side_effect=TimeoutError), patch.object(iw,'publish_workbook'), patch.object(iw,'refresh_economics'):
            report=iw.create_report('894','77','Монтаж кабельной линии 10 м')
            iw.parse_report(report['id']); result=iw.approve_report(report['id'],'manager')
            self.assertEqual(5000,result['accepted_total'])
            self.assertTrue(result['publication_pending'])
            self.assertIn('onedrive',iw.retry_publications()['894'])
            self.assertEqual(5000,iw.get_installation('894')['accepted_total'])

    def test_duplicate_work_lines_cannot_overrun_remaining_quantity(self):
        with patch.object(iw,'backup_db'):
            report=iw.create_report('894','77','Монтаж кабельной линии 60 м')
            iw.parse_report(report['id'])
            db=iw.connect(); item=db.execute('SELECT * FROM report_items WHERE report_id=?',(report['id'],)).fetchone()
            db.execute('INSERT INTO report_items(report_id,work_item_id,declared_qty) VALUES(?,?,?)',(report['id'],item['work_item_id'],60)); db.commit(); db.close()
            with self.assertRaisesRegex(ValueError,'превышает остаток'): iw.approve_report(report['id'],'manager')
            self.assertEqual(0,iw.get_installation('894')['accepted_total'])

    @patch.object(iw,'backup_db',lambda:None)
    def test_recorded_payment_refreshes_economics(self):
        with patch.object(iw,'get_installation',side_effect=[{'debt':100}, {'debt':0}]), \
             patch.object(iw,'refresh_economics',return_value={'ok':True}) as refreshed:
            result=iw.record_payment('894','bank',100,'26.09.2026','bank-refresh')
        refreshed.assert_called_once_with('894')
        self.assertTrue(result['economics_refreshed'])

    def test_new_bank_payment_credits_94_percent_without_rewriting_history(self):
        db=iw.connect()
        db.execute("INSERT INTO payments(deal_id,source,amount,payment_date,external_id,created_at) VALUES('894','bank',100,'01.01.2026','old',?)",(iw.now(),)); db.commit(); db.close()
        self.assertEqual(100,iw.get_installation('894')['paid_total'])
        with patch.object(iw,'refresh_economics'):
            result=iw.record_payment('894','bank',100,'29.09.2026','new',allow_overpay=True)
            repeated=iw.record_payment('894','bank',100,'29.09.2026','new',allow_overpay=True)
        self.assertEqual(200,result['paid_gross_total'])
        self.assertEqual(194,result['paid_total'])
        self.assertEqual(94,result['payment_settlement_amount'])
        self.assertEqual(194,repeated['paid_total'])
        self.assertTrue(repeated['duplicate'])
        db=iw.connect(); self.assertIsNone(db.execute("SELECT coefficient FROM payments WHERE external_id='old'").fetchone()[0]); db.close()

    @patch.object(iw,'backup_db',lambda:None)
    def test_returned_report_never_changes_cumulative_total(self):
        report=iw.create_report('894','77','Монтаж кабельной линии 10 м')
        iw.parse_report(report['id'])
        returned=iw.return_report(report['id'],'Исправить объём','manager')
        self.assertEqual('Исправить объём',returned['manager_comment'])
        self.assertEqual(0,iw.get_installation('894')['accepted_total'])

    @patch.object(iw,'backup_db',lambda:None)
    def test_legacy_import_is_exact_and_idempotent(self):
        # Use a clean DB because the legacy importer bootstraps the installation itself.
        os.remove(iw.DB)
        path=os.path.join(self.tmp.name,'legacy.xlsx'); wb=Workbook(); ws=wb.active
        ws.append(['Номер','Наименование работ и затрат','Единица измерения','Кол-во единиц','Цена на ед. изм. руб.','ВСЕГО , руб.','Факт'])
        ws.append(['1','Монтаж трассы','м',10,1000,10000,6000]); ws.append(['ИТОГО ПО АКТУ','','','','',10000,6000]); wb.save(path)
        source={'deal_id':'853','client':'Клиент','essence':'кондиционирование','address':'Москва','contact_name':'Иван','contact_phone':'7999',
                'installer_name':'Деревянко Сергей','installer_ref':'1000142','status':'Монтаж','project_files':[]}
        with patch.object(iw,'deal_source',lambda *args,**kwargs:source):
            result=iw.import_legacy_workbook('853',path,'13.07.2026','legacy-853','manager')
            self.assertEqual(6000,result['historical_accepted_total'])
            self.assertEqual(6,result['works'][0]['planned_qty']-4)  # source plan stays 10
            with self.assertRaisesRegex(ValueError,'уже импортирован'):
                iw.import_legacy_workbook('853',path,'13.07.2026','legacy-853','manager')

    def test_list_mine_matches_cyrillic_installer_name(self):
        source={'deal_id':'894','installer_name':'Деревянко Сергей','status':'Монтаж'}
        with patch.object(iw,'deal_source',lambda *args,**kwargs:source):
            self.assertEqual('894',iw.list_mine('Деревянко')[0]['deal_id'])

    def test_list_mine_keeps_transferred_object_when_crm_state_changes(self):
        # The hand-off record, not a later/stale Megaplan state, controls the
        # installer's access to an already transferred object.
        source={'deal_id':'894','installer_name':'Другой исполнитель','status':'Выставление счета'}
        with patch.object(iw,'deal_source',lambda *args,**kwargs:source):
            self.assertEqual('894',iw.list_mine('Деревянко')[0]['deal_id'])

    def test_approval_requires_explicit_project_decision(self):
        db=iw.connect(); db.execute('UPDATE installations SET project_required=NULL WHERE deal_id=?',('894',)); db.commit(); db.close()
        with self.assertRaisesRegex(ValueError,'решение по проектной документации'):
            iw.approve('894','manager')

    @patch.object(iw,'backup_db',lambda:None)
    def test_without_project_is_an_explicit_decision(self):
        result=iw.set_without_project('894','manager')
        self.assertEqual(0,result['project_required'])

    @patch.object(iw,'backup_db',lambda:None)
    def test_installation_contact_is_selected_by_megaplan_id(self):
        class API:
            def _call(self,method,path,params=None,json_body=None):
                if path=='/deal/894': return {'data':{'contractor':{'contentType':'ContractorCompany','id':'55'}}}
                if path=='/contractorCompany/55/contacts': return {'data':[{'id':'77','lastName':'Иванов','firstName':'Иван','middleName':'Иванович','contactInfo':[{'type':'phone','value':'+7 900 000-00-00'}]}]}
                if path.startswith('/deal/894/comments'): return {'data':{'id':'900'}}
                raise AssertionError(path)
        result=iw.select_installation_contact('894','77','manager',API())
        self.assertEqual('77',result['contact_ref'])
        self.assertEqual('Иванов Иван Иванович',result['contact_name'])
        self.assertEqual('+7 900 000-00-00',result['contact_phone'])

    def test_contact_without_phone_cannot_be_assigned(self):
        class API:
            def _call(self,method,path,params=None,json_body=None):
                if path=='/deal/894': return {'data':{'contractor':{'contentType':'ContractorCompany','id':'55'}}}
                if path=='/contractorCompany/55/contacts': return {'data':[{'id':'77','lastName':'Иванов','firstName':'Иван','contactInfo':[]}]}
                raise AssertionError(path)
        with self.assertRaisesRegex(ValueError,'не указан телефон'):
            iw.select_installation_contact('894','77','manager',API())


if __name__=='__main__': unittest.main()
