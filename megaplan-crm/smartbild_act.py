"""Approved SmartBild IP act layout, populated from any confirmed shipment."""
from decimal import Decimal
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent

def document(data, test=False):
    from closing_documents import floating_picture, money, act_number, group_tax
    from generate_contract import surname_initials
    seller = data['seller']
    if seller['id'] != 'kocherzhinsky_ip':
        raise ValueError('Бланк SmartBild предназначен для ИП Кочержинского')
    rows = data['positions']
    if not rows or sum(Decimal(str(r['total'])) for r in rows) != Decimal(str(data['total'])):
        raise ValueError('Состав передачи не совпадает с итогом акта')
    groups = [('Товары и оборудование', [r for r in rows if r['kind'] != 'work']),
              ('Выполненные работы и услуги', [r for r in rows if r['kind'] == 'work'])]
    groups = [(title, items) for title, items in groups if items]
    d = Document(); s = d.sections[0]
    s.page_height = Mm(297); s.page_width = Mm(210)
    s.top_margin = Mm(36); s.bottom_margin = Mm(16)
    s.left_margin = s.right_margin = Mm(18)
    normal = d.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(6)
    def p(text='', size=9, bold=False):
        a = d.add_paragraph(); r = a.add_run(text)
        r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string('0B0F10')
        return a
    def signatures():
        t = d.add_table(rows=4, cols=2); t.autofit = False
        values = [('ПОСТАВЩИК / ИСПОЛНИТЕЛЬ', 'ЗАКАЗЧИК'),
                  (seller['name'], data['payer_name']),
                  (surname_initials(seller['signatory']['name']),
                   data['signatory']['position']+' '+surname_initials(data['signatory']['name'])),
                  ('Подпись: __________________', 'Подпись: __________________')]
        for col in t.columns: col.width = Mm(87)
        for i, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j, cell in enumerate(row.cells):
                cell.width = Mm(87); cell.text = values[i][j]
                for a in cell.paragraphs:
                    a.paragraph_format.space_after = Pt(3)
                    a.paragraph_format.keep_with_next = i < 3
                    if i == 3: a.paragraph_format.space_before = Pt(18)
                    for r in a.runs: r.font.size = Pt(9)
        floating_picture(t.cell(3,0).paragraphs[0], ROOT/seller['signature_image'], 28, 18, 0)
    def details(title, items):
        a = p(title, 10, True); a.paragraph_format.keep_with_next = True
        t = d.add_table(rows=1, cols=6); t.autofit = False
        widths = [9,72,12,16,27,38]
        for col,w in zip(t.columns,widths): col.width = Mm(w)
        values = [['№','Наименование','Ед.','Кол-во','Цена, ₽','Сумма, ₽']]
        values += [[str(r.get('source_number') or rows.index(r)+1),r['name'],r['unit'],
                    str(r['quantity']).removesuffix('.0'),money(r['unit_price']),money(r['total'])] for r in items]
        for i, vals in enumerate(values):
            row = t.rows[0] if i == 0 else t.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if i == 0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for j, cell in enumerate(row.cells):
                cell.width = Mm(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = str(vals[j])
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'),'E6E8EB' if i == 0 else 'FFFFFF'); cell._tc.get_or_add_tcPr().append(sh)
                for a in cell.paragraphs:
                    a.paragraph_format.space_before = Pt(7); a.paragraph_format.space_after = Pt(7)
                    if j >= 3: a.alignment = 2
                    for r in a.runs: r.font.size = Pt(8); r.bold = i == 0
        for cell in t.rows[-1].cells:
            for a in cell.paragraphs: a.paragraph_format.keep_with_next = True
        p('ИТОГО ПО ГРУППЕ: '+money(sum(Decimal(str(r['total'])) for r in items))+' ₽. '+group_tax(items)+'.',10,True).paragraph_format.keep_with_next = True
    number = act_number(data); day = data['act_date']
    p(('ТЕСТ — ' if test else '')+'АКТ № '+number,20,True)
    title = 'товаров, выполненных работ и оказанных услуг' if len(groups)==2 else ('выполненных работ и оказанных услуг' if rows[0]['kind']=='work' else 'товаров и оборудования')
    p('приёма-передачи '+title,11,True)
    p(day+'   •   г. Москва')
    p('Общая стоимость передачи: '+money(data['total'])+' ₽',13,True)
    from num2words import num2words
    amount = Decimal(str(data['total']))
    p(num2words(int(amount),lang='ru').capitalize()+f' рублей {int(amount*100)%100:02d} копеек. '+data['tax_text']+'.')
    basis = f"Договор № {data['number']} от {data['date']}"
    p('Основание: '+basis+'.')
    p(f"Исполнитель: {seller['legal_name']}, ИНН {seller['inn']}, ОГРНИП {seller['ogrnip']}.")
    p(f"Заказчик: {data['payer_name']}, ИНН {data['payer']['inn']}, КПП {data['payer'].get('kpp') or '—'}. Подписант: {data['signatory']['position']} {data['signatory']['name']}, основание: {data['signatory']['authority']}.")
    p('Объект: '+data['address']+'.')
    if data.get('deadline_condition'): p('Срок исполнения по договору: '+data['deadline_condition'])
    elif data.get('deadline'): p(f"Период выполнения по договору: с {data['date']} по {data['deadline']}.")
    # Keep the approved compact form for a short act; long lists use appendices.
    appendix = len(rows)>3 or sum(len(r['name']) for r in rows)>350
    if not appendix:
        for title, items in groups: details(title,items)
    else:
        for title, items in groups: p(title+': '+money(sum(Decimal(str(r['total'])) for r in items))+' ₽.',10,True)
        p('Состав передачи приведён в приложении № 1, являющемся неотъемлемой частью настоящего акта.')
    p('Поставщик / Исполнитель передал, а Заказчик принял перечисленные товары и результаты выполненных работ и услуг. Товары переданы комплектными, указанные работы и услуги выполнены качественно; относящиеся к этому объёму документы и принадлежности переданы. Претензий по объёму, комплектности и явным недостаткам нет, кроме указанных в особых отметках.')
    p('Настоящий акт подтверждает передачу только перечисленного объёма и не закрывает остальные обязательства по договору. Порядок и сроки оплаты определяются договором. Подписание акта не подтверждает получение оплаты Исполнителем.')
    p('Особые отметки:',9,True)
    for _ in range(2):
        a = p(' '); a.paragraph_format.space_after = Pt(8)
        border = OxmlElement('w:pBdr')
        for name in ('bottom','between'):
            line = OxmlElement('w:'+name)
            for key,value in [('val','single'),('sz','4'),('color','9CAEA8')]: line.set(qn('w:'+key),value)
            border.append(line)
        a._p.get_or_add_pPr().append(border)
    signatures()
    if appendix:
        for title, items in groups:
            d.add_page_break(); p('ПРИЛОЖЕНИЕ № 1',17,True)
            p(f'к акту № {number} от {day}',10,True)
            p(basis+'\nЗаказчик: '+data['payer_name']+'\nОбъект: '+data['address'])
            details(title,items); signatures()
    f = s.footer.paragraphs[0]
    f.text = f"{seller['name']} · ИНН {seller['inn']} · ОГРНИП {seller['ogrnip']} · Стр. "
    for r in f.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('60766E')
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
    h = s.header.paragraphs[0]
    floating_picture(h,ROOT/'templates/closing/kocherzhinsky_ip-smartbild-v1/letterhead-IP.png',210,0,0)
    for el in h._p.xpath('.//wp:anchor/wp:positionH | .//wp:anchor/wp:positionV'): el.set('relativeFrom','page')
    return d
