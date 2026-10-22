#!/usr/bin/env python3
"""Build an installer-facing work statement without deal economics."""
import json
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import xlsx_brand_style as style


POST_SIGNING_STATUSES = {
    "выставление счета", "внесение предоплаты", "закупка", "монтаж",
    "оплата", "отгрузка", "завершена", "рекламация",
}


def normalize_status(value):
    return " ".join(str(value or "").casefold().replace("ё", "е").split())


def is_post_signing_status(status):
    """Unknown stages and Отказ retain rows because their ordering is ambiguous."""
    return normalize_status(status) in POST_SIGNING_STATUSES


def statement_rows(entries, status):
    hide_zero_cost = is_post_signing_status(status)
    rows = []
    for entry in entries or []:
        cost_work = float(entry.get("cost_work") or 0)
        if hide_zero_cost and cost_work <= 0:
            continue
        rows.append({"name": entry.get("name") or "", "unit": entry.get("unit") or "",
                     "qty": entry.get("qty") or 0, "cost_work": cost_work})
    return rows


def build_xlsx(spec):
    rows = statement_rows(spec.get("entries"), spec.get("status"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Ведомость работ"
    ws.sheet_view.showGridLines = False
    headers = ["Наименование", "Ед. изм.", "Кол-во", "Себестоимость работ/ед.", "Итого стоимость работ"]
    subtitle = f"Сделка № {spec.get('deal_id')} · {spec.get('client') or ''} · {spec.get('status') or 'статус не указан'}"
    header_row = style.write_title_banner(ws, "Ведомость работ", subtitle, len(headers))
    style.write_header_row(ws, header_row, headers)
    align_by_col = {1: style.LEFT, 2: style.CENTER, 3: style.CENTER, 4: style.RIGHT, 5: style.RIGHT}
    # freeze_panes creates an otherwise empty cell on the row below the header;
    # write into that row explicitly instead of letting append() leave a visual gap.
    first_data_row = header_row + 1
    for index, row in enumerate(rows):
        row_index = first_data_row + index
        for column, value in enumerate([row["name"], row["unit"], row["qty"], row["cost_work"]], 1):
            ws.cell(row=row_index, column=column, value=value)
        style.style_data_row(ws, row_index, len(headers), align_by_col, stripe=(index % 2 == 1))
        ws.cell(row=row_index, column=4).number_format = style.MONEY_FMT
        ws.cell(row=row_index, column=5, value=f"=C{row_index}*D{row_index}")
        ws.cell(row=row_index, column=5).number_format = style.MONEY_FMT
    last_data_row = first_data_row + len(rows) - 1
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="ИТОГО стоимость работ")
    formula = f"=SUM(E{first_data_row}:E{last_data_row})" if rows else "=0"
    ws.cell(row=total_row, column=5, value=formula).number_format = style.MONEY_FMT
    style.style_total_row(ws, total_row, len(headers))
    ws.row_dimensions[total_row].height = 22
    for column, width in enumerate([58, 11, 11, 23, 23], start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    wb.save(spec["path"])
    return {"ok": True, "path": spec["path"], "row_count": len(rows),
            "filtered": is_post_signing_status(spec.get("status"))}


def main():
    try:
        result = build_xlsx(json.load(sys.stdin))
    except Exception as exc:
        result = {"ok": False, "error": str(exc)}
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
