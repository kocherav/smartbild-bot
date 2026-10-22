"""Read-only, price-free installation supply view. No CRM writes or notifications.

Approved allocations are authoritative; fuzzy suggestions NEVER change coverage.
Preview rollout: CLI only, no installer button/automatic delivery enabled.
"""
import argparse
import json
import re
import textwrap
import html
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import procurement_workflow as procurement


def unit(value):
    value = str(value or '').lower().strip().rstrip('.')
    return {'метр': 'м', 'метры': 'м', 'шт': 'шт', 'штука': 'шт'}.get(value, value)


def conflicts(left, right):
    """Explicit dimensional/electrical contradictions cannot be auto-accepted."""
    def features(text):
        text = str(text).lower().replace('×', 'x').replace('х', 'x').replace(',', '.')
        return (set(re.findall(r'\d+(?:\.\d+)?(?:\s*x\s*\d+(?:\.\d+)?)+', text)),
                set(re.findall(r'\b\d+(?:\.\d+)?\s*(?:в|v|кв|kv)\b', text)))
    a, b = features(left), features(right)
    return any(x and y and x.isdisjoint(y) for x, y in zip(a, b))


def build_view(deal_id, items, registry):
    """Return allowlisted display fields only; no monetary values leave this boundary."""
    rows = [{'id': str(i.get('offer_row_id') or ''), 'name': i.get('name') or '',
             'unit': i.get('unit') or '', 'needed': float(i.get('qty') or 0),
             'confirmed': 0., 'links': [], 'candidates': []} for i in items]
    by_row = {r['id']: r for r in rows if r['id']}
    unmatched, seen, related = [], set(), set()

    def source(raw, invoice, qty, kind):
        return {'name': raw.get('source_name') or raw.get('name') or 'Название не сохранено',
                'unit': raw.get('source_unit') or raw.get('unit') or '',
                'qty': float(qty or 0), 'invoice': str(invoice.get('number') or invoice.get('invoice_number') or ''),
                'kind': kind}

    def suggest(s):
        matches = []
        for r in rows:
            score, _ = procurement.item_similarity(s, r)
            if score > .5:
                warning = 'Характеристики различаются' if conflicts(s['name'], r['name']) else 'Требует подтверждения'
                if not s['unit'] or unit(s['unit']) != unit(r['unit']):
                    warning += '; единицы требуют проверки'
                matches.append((score, r, warning))
        for score, r, warning in sorted(matches, key=lambda x: -x[0])[:3]:
            r['candidates'].append({**s, 'score': round(score * 100), 'warning': warning})
        unmatched.append({**s, 'suggested': bool(matches)})

    for invoice in registry.get('invoices', []):
        for allocation in invoice.get('allocations', []):
            if str(allocation.get('deal_id')) != str(deal_id):
                continue
            related.add(invoice.get('document_key'))
            key = (invoice.get('document_key') or invoice.get('id'), allocation.get('id') or json.dumps(allocation, sort_keys=True))
            if key in seen:
                continue
            seen.add(key)
            for update in allocation.get('updates', []):
                s = source(update, invoice, update.get('qty_used'), 'Подтверждённая привязка')
                row_id = str(update.get('offer_row_id') or '')
                target = by_row.get(row_id)
                # Offer IDs may repeat in a specification. Never guess in that case.
                if not row_id and update.get('offer_id'):
                    matches = [r for r, i in zip(rows, items) if str(i.get('offer_id')) == str(update['offer_id'])]
                    target = matches[0] if len(matches) == 1 else None
                if target is None:
                    suggest(s)
                    continue
                target['confirmed'] += s['qty']
                notes = []
                if not s['unit']:
                    notes.append('Ед. счёта не сохранена; количество взято из утверждённой привязки в ед. сделки')
                elif unit(s['unit']) != unit(target['unit']):
                    notes.append('Единицы различаются: проверить ранее утверждённый пересчёт')
                if conflicts(s['name'], target['name']):
                    notes.append('Характеристики различаются; сохранена ранее утверждённая привязка')
                s['warning'] = '; '.join(notes)
                target['links'].append(s)
            for extra in allocation.get('additional_expenses', []):
                suggest(source(extra, invoice, extra.get('qty') or extra.get('source_qty'), 'Учтено в прочей себестоимости'))
            for extra in allocation.get('additions', []):
                suggest(source(extra, invoice, extra.get('qty'), 'Добавленная позиция: проверить привязку'))

    # Only remaining stock from related invoices; never quantities assigned elsewhere.
    for stock in registry.get('unallocated_stock', []):
        for allocation in stock.get('allocations', []):
            if str(allocation.get('deal_id')) != str(deal_id):
                continue
            # Stock consumption is a distinct confirmed allocation, not the invoice's
            # initial allocation. Do not include its original/full stock quantity.
            key = ('stock', stock.get('id'), allocation.get('id'))
            if key in seen:
                continue
            seen.add(key)
            s = source(stock, stock, allocation.get('qty'), 'Подтверждено из остатков')
            target = by_row.get(str(allocation.get('offer_row_id') or ''))
            if target:
                target['confirmed'] += s['qty']
                s['warning'] = '' if s['unit'] and unit(s['unit']) == unit(target['unit']) else 'Единицы требуют проверки; количество из утверждённой привязки'
                target['links'].append(s)
            else:
                suggest(s)
        if stock.get('document_key') in related and float(stock.get('available_qty') or 0) > 0:
            suggest(source(stock, stock, stock['available_qty'], 'Нераспределённый остаток — НЕ закреплён за объектом'))
    return rows, unmatched


def mobile_sheet(wb, deal_id, rows, unmatched):
    """Compact four-column summary; one record per specification position."""
    ws = wb.create_sheet('Обзор для телефона', 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    for col, width in zip('ABCD', [12, 29, 9, 11]):
        ws.column_dimensions[col].width = width
    def qty(v):
        return f'{v:g}'.replace('.', ',')
    def state(r):
        if any(s.get('warning') for s in r['links']):
            return 0, 'Проверить', 'DDEBF7'
        if r['candidates']:
            return 1, 'Возможная замена', 'FFF2CC'
        if r['confirmed'] <= 0:
            return 2, 'Нет', 'FCE4D6'
        if r['confirmed'] < r['needed']:
            return 3, 'Частично', 'FFF2CC'
        return 4, 'Покрыто', 'C6EFCE'
    states = [state(r)[1] for r in rows]
    for n, text, height in [
        (1, f'МАТЕРИАЛЫ • №{deal_id}', 26),
        (2, f"Покрыто {states.count('Покрыто')} • Частично {states.count('Частично')} • Проверить {sum(s in ('Проверить', 'Возможная замена') for s in states)} • Нет {states.count('Нет')}", 30),
        (3, 'Не подтверждает доставку. * Есть замечание. Название → подробности.', 30),
    ]:
        ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=4)
        c = ws.cell(n, 1, text)
        c.font = Font(name='Calibri', size=14 if n == 1 else 10, bold=n == 1, color='17365D')
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[n].height = height
    ws.append(['Статус', 'Материал по сделке', 'Нужно', 'По счетам'])
    ws.row_dimensions[4].height = 30
    for c in ws[4]:
        c.fill = PatternFill('solid', fgColor='17365D')
        c.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    details = wb['Сопоставления']
    targets = {}
    for n in range(6, details.max_row + 1):
        targets.setdefault(str(details.cell(n, 1).value), n)
    for r in sorted(rows, key=lambda r: state(r)[0]):
        _, status, color = state(r)
        warning = any(s.get('warning') for s in r['links'])
        confirmed = f"{qty(r['confirmed'])} {r['unit']}{'*' if warning else ''}" if r['confirmed'] else '—'
        ws.append([status, r['name'], f"{qty(r['needed'])} {r['unit']}", confirmed])
        n = ws.max_row
        for c in ws[n]:
            c.data_type = 's'
            c.font = Font(name='Calibri', size=11, color='243746')
            c.alignment = Alignment(wrap_text=True, vertical='center')
            c.fill = PatternFill('solid', fgColor='F5F7F9' if n % 2 else 'FFFFFF')
        ws.cell(n, 1).fill = PatternFill('solid', fgColor=color)
        ws.cell(n, 1).font = Font(name='Calibri', size=10, bold=True, color='243746')
        target = targets.get(r['id'])
        if target:
            ws.cell(n, 2).hyperlink = f"#'Сопоставления'!B{target}"
            ws.cell(n, 2).font = Font(name='Calibri', size=11, color='205B87', underline='single')
        lines = max(2, len(textwrap.wrap(r['name'], width=27)))
        ws.row_dimensions[n].height = 8 + lines * 14
    ws.auto_filter.ref = f'A4:D{ws.max_row}'
    ws.freeze_panes = 'B5'
    ws.append([f'Вне привязок: {len(unmatched)} строк → отдельный лист'])
    n = ws.max_row
    ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=4)
    ws.cell(n, 1).hyperlink = "#'Вне подтверждённых привязок'!A1"
    ws.cell(n, 1).font = Font(name='Calibri', size=11, color='205B87', underline='single')
    ws.row_dimensions[n].height = 28
    no_match = [(index, s) for index, s in enumerate(unmatched, 2) if not s['suggested']]
    if no_match:
        ws.append(['Есть в счетах — нет похожей позиции сделки'])
        n = ws.max_row
        ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=4)
        ws.cell(n, 1).font = Font(name='Calibri', size=12, bold=True, color='17365D')
        ws.cell(n, 1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[n].height = 32
        ws.append(['Точные сохранённые названия из счетов. Количество — отнесённая на объект часть либо свободный остаток, не весь счёт.'])
        n = ws.max_row
        ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=4)
        ws.cell(n, 1).font = Font(name='Calibri', size=10, color='526575')
        ws.cell(n, 1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[n].height = 42
        ws.append(['Учёт', 'Название из счёта', 'Кол-во', 'Счёт №'])
        for c in ws[ws.max_row]:
            c.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='17365D')
            c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[ws.max_row].height = 28
        for target, s in no_match:
            stock = s['kind'].startswith('Нераспределённый остаток')
            status = 'Свободный остаток — не объект' if stock else ('Прочая себестоимость' if s['kind'] == 'Учтено в прочей себестоимости' else 'Проверить привязку')
            ws.append([status, s['name'], f"{qty(s['qty'])} {s['unit'] or '(ед. неизвестна)'}", s['invoice']])
            n = ws.max_row
            for c in ws[n]:
                c.data_type = 's'
                c.font = Font(name='Calibri', size=11, color='243746')
                c.alignment = Alignment(wrap_text=True, vertical='center')
                c.fill = PatternFill('solid', fgColor='F5F7F9' if n % 2 else 'FFFFFF')
            ws.cell(n, 1).fill = PatternFill('solid', fgColor='FFF2CC' if stock else 'DDEBF7')
            ws.cell(n, 1).font = Font(name='Calibri', size=10, bold=True, color='243746')
            ws.cell(n, 2).hyperlink = f"#'Вне подтверждённых привязок'!A{target}"
            ws.cell(n, 2).font = Font(name='Calibri', size=11, color='205B87', underline='single')
            lines = max(4, len(textwrap.wrap(s['name'], width=27)), len(textwrap.wrap(s['invoice'], width=10)))
            ws.row_dimensions[n].height = 8 + lines * 14
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:D{ws.max_row}'
    wb.active = 0


def mobile_pdf(sheet, output):
    """Fixed phone-width PDF; no dependence on the iOS XLSX preview canvas."""
    from weasyprint import HTML
    blocks = []
    in_table = False
    merged = {r.min_row for r in sheet.merged_cells.ranges}
    for row in sheet:
        n = row[0].row
        if n in merged:
            if in_table:
                blocks.append('</tbody></table>')
                in_table = False
            tag = 'h1' if n == 1 else 'p'
            blocks.append(f'<{tag}>{html.escape(str(row[0].value or ""))}</{tag}>')
            continue
        if not in_table:
            blocks.append('<table><colgroup><col style="width:19%"><col style="width:47%"><col style="width:16%"><col style="width:18%"></colgroup><thead>')
            blocks.append('<tr>' + ''.join('<th>'+html.escape(str(c.value or ''))+'</th>' for c in row[:4]) + '</tr></thead><tbody>')
            in_table = True
            continue
        cells = []
        for c in row[:4]:
            color = c.fill.fgColor.rgb
            color = '#' + color[-6:] if isinstance(color, str) and len(color) == 8 else '#FFFFFF'
            cells.append(f'<td style="background:{color}">{html.escape(str(c.value or ""))}</td>')
        blocks.append('<tr>'+''.join(cells)+'</tr>')
    if in_table:
        blocks.append('</tbody></table>')
    css = '''@page {size: 390pt 720pt; margin: 12pt;
      @bottom-right {content: counter(page); font-size: 8pt; color:#526575}}
      body {font-family: "DejaVu Sans",sans-serif; font-size:11pt; color:#243746; margin:0}
      h1 {font-size:17pt; color:#17365D} p {font-size:10pt; line-height:1.35}
      table {width:100%; border-collapse:collapse; table-layout:fixed}
      th {background:#17365D; color:white; text-align:left; font-size:10pt}
      td,th {padding:7pt 4pt; overflow-wrap:anywhere; vertical-align:middle}
      td:first-child {font-size:9pt; font-weight:bold} tr {break-inside:avoid}
      thead {display:table-header-group}'''
    HTML(string='<html lang="ru"><meta charset="utf-8"><style>'+css+'</style><body>'+''.join(blocks)+'</body></html>').write_pdf(str(output))


def export(deal_id, output, mobile_only=False, pdf_output=None):
    items = procurement.procurement_items(str(deal_id))
    if not items:
        raise ValueError('Нет локальных позиций сделки; пустой файл не формируется')
    with open(procurement.REGISTRY, encoding='utf-8') as f:
        registry = json.load(f)
    rows, unmatched = build_view(deal_id, items, registry)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Сопоставления'
    ws.append([f'Материалы по счетам — сделка №{deal_id}'])
    ws.append(['Количество по подтверждённым привязкам. Оплата, доставка и передача монтажнику не подтверждаются.'])
    ws.append(['Жёлтый: предположение, в покрытие НЕ включено. Голубой: утверждённая привязка с замечаниями. Зелёный: полное покрытие без замечаний.'])
    ws.append(['Сформировано UTC: ' + datetime.now(timezone.utc).strftime('%d.%m.%Y %H:%M')])
    ws.append(['ID строки', 'Позиция сделки', 'Ед. сделки', 'Нужно', 'Подтверждено по привязкам',
               'Наименование в счёте', 'Кол-во в привязке / остатке', 'Ед. счёта', 'Счёт №', 'Статус / замечание'])
    for r in rows:
        details = r['links'] + r['candidates'] or [{}]
        complete = r['needed'] > 0 and r['confirmed'] >= r['needed'] and not any(s.get('warning') for s in r['links'])
        for index, s in enumerate(details):
            status = s.get('kind', 'Нет подтверждённой закупки')
            if 'score' in s:
                status = f"Возможное соответствие ({s['score']}% сходства). {status}"
            if s.get('warning'):
                status += '. ' + s['warning']
            ws.append([r['id'], r['name'], r['unit'], r['needed'] if index == 0 else None,
                       r['confirmed'] if index == 0 else None, s.get('name', ''), s.get('qty'),
                       s.get('unit') or ('не сохранена' if s else ''), s.get('invoice', ''), status])
            color = 'FFF2CC' if 'score' in s else ('DDEBF7' if s.get('warning') else ('C6EFCE' if complete else None))
            if color:
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill('solid', fgColor=color)
    other = wb.create_sheet('Вне подтверждённых привязок')
    other.append(['Наименование в счёте', 'Количество', 'Ед. счёта', 'Счёт №', 'Статус', 'Найдены кандидаты'])
    for s in unmatched:
        other.append([s['name'], s['qty'], s['unit'] or 'не сохранена', s['invoice'], s['kind'], 'Да, лист Сопоставления' if s['suggested'] else 'Нет'])
    for sheet, header, widths in [(ws, 5, [14, 55, 12, 14, 22, 60, 22, 16, 24, 80]),
                                   (other, 1, [65, 16, 16, 24, 65, 25])]:
        sheet.freeze_panes = f'C{header+1}'
        sheet.auto_filter.ref = f'A{header}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        for n, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(n)].width = width
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # untrusted supplier text must never become an Excel formula
                elif isinstance(cell.value, (float, int)):
                    cell.number_format = '0.###'
        for cell in sheet[header]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
    # Technical table is secondary; keep its top notes readable too.
    for n in range(1, 5):
        ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=10)
        ws.row_dimensions[n].height = 30
    mobile_sheet(wb, deal_id, rows, unmatched)
    if mobile_only:
        # iOS previews can derive the canvas from the widest sheet in a book.
        # A separate one-sheet edition avoids inheriting the 10-column canvas.
        overview = wb.worksheets[0]
        for sheet in list(wb.worksheets)[1:]:
            wb.remove(sheet)
        overview['A3'] = 'Не подтверждает доставку. * Есть замечание; подробности в полной версии.'
        for row in overview:
            for cell in row:
                if cell.hyperlink:
                    cell.hyperlink = None
                    cell.font = Font(name='Calibri', size=11, color='243746')
                if isinstance(cell.value, str) and cell.value.startswith('Вне привязок:'):
                    cell.value = f'Вне привязок: {len(unmatched)} строк. Без похожих позиций — ниже.'
        overview.sheet_view.zoomScale = 100
        overview.sheet_properties.outlinePr.summaryRight = False
        overview.column_dimensions.group('E', 'XFD', hidden=True)
        overview.sheet_view.selection[0].activeCell = 'A1'
        overview.sheet_view.selection[0].sqref = 'A1'
    wb.save(output)
    if pdf_output:
        mobile_pdf(wb.worksheets[0], pdf_output)
    return {'path': str(output), 'positions': len(rows), 'confirmed_positions': sum(r['confirmed'] > 0 for r in rows),
            'candidate_links': sum(len(r['candidates']) for r in rows), 'unmatched_rows': len(unmatched)}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('deal_id', type=int)
    parser.add_argument('output', type=Path)
    parser.add_argument('--mobile-only', action='store_true')
    parser.add_argument('--pdf', type=Path)
    args = parser.parse_args()
    print(json.dumps(export(args.deal_id, args.output, mobile_only=args.mobile_only, pdf_output=args.pdf), ensure_ascii=False))
