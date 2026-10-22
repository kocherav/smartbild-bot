"""Seller selection shared by contract and invoice generation."""
from pathlib import Path
import json
from docx.shared import Inches, Pt
from own_companies import company_by_id

BASE = Path(__file__).resolve().parent

def normalize_party_fonts(document):
    """Use the contract body typography for both parties and appendix signatures."""
    from docx.oxml.ns import qn
    from generate_contract import table_paragraphs
    for index in (0, 2):
        for paragraph in table_paragraphs([document.tables[index]]):
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
                run.font.bold = False
                run.font.italic = False
                fonts = run._element.get_or_add_rPr().rFonts
                for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                    fonts.set(qn('w:' + key), 'Times New Roman')
                for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                    fonts.attrib.pop(qn('w:' + key), None)

def selected_seller(deal_id):
    from procurement_workflow import REGISTRY, load_json, empty_registry, deal_selling_company
    return deal_selling_company(load_json(REGISTRY, empty_registry()), str(deal_id))

def invoice_settings(deal_id):
    company = selected_seller(deal_id)
    if company['id'] == 'cliwell':
        return '1000560', '3'
    if company['id'] == 'kocherzhinsky_ip':
        return '1003357', '4'
    raise ValueError('Для выбранной организации ещё не настроен выпуск документов')

def apply_ip_supplier(document, company):
    from generate_contract import all_paragraphs, set_paragraph_text, keep_row_on_one_page
    if company['id'] == 'cliwell':
        normalize_party_fonts(document)
        return
    if company['id'] != 'kocherzhinsky_ip':
        raise ValueError('Для выбранной организации ещё не настроен договор')
    signature = BASE / company['signature_image']
    if not signature.is_file():
        raise ValueError('Не найден оригинал подписи ИП')
    # Remove template scans before adding the explicitly supplied IP signature.
    for node in document.element.xpath('.//w:drawing | .//w:pict'):
        node.getparent().remove(node)
    for paragraph in all_paragraphs(document):
        if 'совместно в дальнейшем' in paragraph.text:
            delimiter = ', с одной стороны, и '
            if delimiter not in paragraph.text:
                raise ValueError('Не распознана преамбула исходного договора')
            customer = paragraph.text.split(delimiter, 1)[1]
            set_paragraph_text(paragraph, company['legal_name'] + ', ' + company['contract_authority'] +
                               ', именуемый в дальнейшем «Поставщик»' + delimiter + customer)
    bank = company['bank']
    for index in (0, 2):
        table = document.tables[index]
        cell = table.cell(0, 0)
        for child in list(cell._tc):
            if child.tag.endswith('}p') or child.tag.endswith('}tbl'):
                cell._tc.remove(child)
        lines = ['Поставщик:', company['legal_name']]
        if index == 0:
            lines += [company['address'], 'ОГРНИП ' + company['ogrnip'], 'ИНН ' + company['inn'],
                      'р/с ' + company['accounts'][0], bank['name'], 'БИК ' + bank['bik'],
                      'к/с ' + bank['correspondent_account']]
        for line in lines:
            p = cell.add_paragraph(line)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1
        p = cell.add_paragraph()
        # Match the physical size in the supplied invoice (175 px at 158 dpi).
        p.add_run().add_picture(str(signature), width=Inches(175 / 158))
        p.add_run(' /Кочержинский А.В./')
        cell.add_paragraph('«__» ________________ 2026 г.')
        for row in table.rows:
            keep_row_on_one_page(row)
    normalize_party_fonts(document)
    text = '\n'.join(p.text for p in all_paragraphs(document))
    for forbidden in ('КЛИВЭЛ', '7743151323', '1167746383090', 'Кочержинская'):
        if forbidden in text:
            raise ValueError('В документе остались реквизиты прежнего поставщика: ' + forbidden)
