#!/usr/bin/env python3
"""Transactional installation handoff, cumulative reports and OneDrive archive."""
import hashlib
import html
import json
import math
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import urllib.request
import megaplan_transport
from datetime import datetime, timezone
from difflib import SequenceMatcher

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import deal_ledger
import deal_browser
from megaplan_api import MegaplanAPI, MegaplanAPIError
from contract_invoice import v1_auth, v1_call
from registry_transaction import transaction

BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(BASE, "installations.sqlite3")
ONEDRIVE_ROOT = "onedrive:SmartBild/Монтажи"
WORKBOOK_DIR = os.path.join(BASE, "installation_files")

def now(): return datetime.now(timezone.utc).isoformat(timespec="seconds")
def clean(value): return " ".join(str(value or "").split())
def safe(value):
    value=re.sub(r'[<>:"/\\|?*]+',' ',clean(value)).strip(' .')
    return value[:120] or 'Без названия'
def money_value(value):
    return float((value or {}).get('value') or (value or {}).get('valueInMain') or 0) if isinstance(value,dict) else float(value or 0)

def brief_essence(subject,address=''):
    value=clean(subject)
    value=re.sub(r'(?i)^(?:проектирование,?\s*)?(?:поставка|поставку)(?:\s*,?\s*(?:монтаж|установка|пусконаладка|пусконаладочные работы|и\s+монтаж))*\s+', '', value)
    value=re.sub(r'(?i)^и\s+монтаж\s+','',value)
    value=re.sub(r'(?i)^систем(?:ы|а)\s+','',value)
    replacements={'электроснабжения':'электроснабжение','кондиционирования':'кондиционирование','вентиляции':'вентиляция','отопления':'отопление'}
    value=replacements.get(value.casefold(),value)
    tail=''
    match=re.search(r'(?i)(\d+(?:-й|-я|-е)?\s+этаж(?:\s+[A-Za-zА-Яа-яЁё0-9.-]+)*)',address or '')
    if match: tail=re.sub(r'(?i)этаж','эт',clean(match.group(1)))
    return clean(f'{value} {tail}')

def is_project_document(name):
    """Conservative filename classifier: an arbitrary PDF is never a project."""
    value=norm(os.path.basename(str(name or '')))
    extension=os.path.splitext(str(name or ''))[1].casefold()
    negative=('коммерческое предложение','кп ','договор','счет','счёт','упд','накладн',
              'экономик','рентабельност','закуп','монтажная ведомость','отчет','отчёт','акт')
    if any(token in value for token in negative): return False
    if extension in ('.dwg','.dxf','.rvt','.ifc'): return True
    positive=('проектная документация','рабочая документация','рабочий проект','проект ',
              'план трасс','план расположения','однолинейная схема','аксонометр','чертеж','чертёж',
              'раздел ов','раздел эом','раздел эс','раздел вк','раздел ар')
    return extension in ('.pdf','.doc','.docx','.zip') and any(token in value for token in positive)

def connect():
    db=sqlite3.connect(DB); db.row_factory=sqlite3.Row
    db.execute('PRAGMA journal_mode=WAL'); db.execute('PRAGMA foreign_keys=ON')
    db.executescript('''
    CREATE TABLE IF NOT EXISTS installations(
      deal_id TEXT PRIMARY KEY, client TEXT NOT NULL, essence TEXT NOT NULL, address TEXT,
      contact_name TEXT, contact_phone TEXT, installer_name TEXT, installer_ref TEXT,
      status TEXT NOT NULL DEFAULT 'draft', project_required INTEGER,
      folder_name TEXT, folder_remote TEXT, workbook_path TEXT, version INTEGER NOT NULL DEFAULT 1,
      megaplan_comment_id TEXT, megaplan_file_id TEXT,
      created_at TEXT NOT NULL, updated_at TEXT NOT NULL, approved_at TEXT, transferred_at TEXT);
    CREATE TABLE IF NOT EXISTS work_items(
      id INTEGER PRIMARY KEY AUTOINCREMENT, deal_id TEXT NOT NULL REFERENCES installations(deal_id),
      source_key TEXT, name TEXT NOT NULL, unit TEXT, planned_qty REAL NOT NULL,
      rate REAL NOT NULL, sort_order INTEGER NOT NULL, active INTEGER NOT NULL DEFAULT 1,
      UNIQUE(deal_id,source_key));
    CREATE TABLE IF NOT EXISTS reports(
      id INTEGER PRIMARY KEY AUTOINCREMENT, deal_id TEXT NOT NULL REFERENCES installations(deal_id),
      report_no INTEGER NOT NULL, status TEXT NOT NULL DEFAULT 'draft', report_date TEXT NOT NULL,
      raw_text TEXT, installer_chat_id TEXT, revision INTEGER NOT NULL DEFAULT 1,
      submitted_at TEXT, approved_at TEXT, approved_by TEXT, manager_comment TEXT,
      UNIQUE(deal_id,report_no));
    CREATE TABLE IF NOT EXISTS report_items(
      id INTEGER PRIMARY KEY AUTOINCREMENT, report_id INTEGER NOT NULL REFERENCES reports(id),
      work_item_id INTEGER REFERENCES work_items(id), source_text TEXT,
      declared_qty REAL NOT NULL DEFAULT 0, accepted_qty REAL, accepted_amount REAL,
      status TEXT NOT NULL DEFAULT 'draft');
    CREATE TABLE IF NOT EXISTS media(
      id INTEGER PRIMARY KEY AUTOINCREMENT, report_id INTEGER NOT NULL REFERENCES reports(id),
      telegram_file_id TEXT, telegram_unique_id TEXT, local_path TEXT, original_name TEXT,
      mime_type TEXT, size INTEGER, sha256 TEXT, onedrive_remote TEXT, created_at TEXT NOT NULL,
      UNIQUE(report_id,telegram_unique_id));
    CREATE TABLE IF NOT EXISTS payments(
      id INTEGER PRIMARY KEY AUTOINCREMENT, deal_id TEXT NOT NULL REFERENCES installations(deal_id),
      source TEXT NOT NULL, amount REAL NOT NULL, payment_date TEXT NOT NULL,
      external_id TEXT, comment TEXT, attachment_remote TEXT, created_at TEXT NOT NULL,
      UNIQUE(source,external_id));
    CREATE TABLE IF NOT EXISTS audit_log(
      id INTEGER PRIMARY KEY AUTOINCREMENT, deal_id TEXT, entity TEXT, entity_id TEXT,
      action TEXT NOT NULL, actor TEXT, payload TEXT, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS handoff_events(deal_id TEXT PRIMARY KEY, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS handoff_deliveries(deal_id TEXT NOT NULL, chat_id TEXT NOT NULL,
      sent_at TEXT NOT NULL, PRIMARY KEY(deal_id,chat_id));
    CREATE TABLE IF NOT EXISTS publication_jobs(
      deal_id TEXT PRIMARY KEY, revision INTEGER NOT NULL DEFAULT 1,
      status TEXT NOT NULL DEFAULT 'pending', updated_at TEXT NOT NULL, error TEXT);
    CREATE TABLE IF NOT EXISTS report_publications(
      report_id INTEGER PRIMARY KEY REFERENCES reports(id),status TEXT NOT NULL DEFAULT 'pending',
      updated_at TEXT NOT NULL,error TEXT);
    CREATE TABLE IF NOT EXISTS legacy_ledger_jobs(
      source_digest TEXT PRIMARY KEY, payload TEXT NOT NULL,
      status TEXT NOT NULL DEFAULT 'pending', updated_at TEXT NOT NULL,error TEXT);
    CREATE TABLE IF NOT EXISTS project_versions(
      id INTEGER PRIMARY KEY AUTOINCREMENT, deal_id TEXT NOT NULL REFERENCES installations(deal_id),
      version_no INTEGER NOT NULL, status TEXT NOT NULL DEFAULT 'active', created_at TEXT NOT NULL,
      created_by TEXT, UNIQUE(deal_id,version_no));
    CREATE TABLE IF NOT EXISTS project_documents(
      id INTEGER PRIMARY KEY AUTOINCREMENT, version_id INTEGER NOT NULL REFERENCES project_versions(id),
      megaplan_file_id TEXT, name TEXT NOT NULL, megaplan_path TEXT, onedrive_remote TEXT,
      created_at TEXT NOT NULL, UNIQUE(version_id,megaplan_file_id));
    ''')
    columns={x[1] for x in db.execute('PRAGMA table_info(installations)')}
    for name in ('megaplan_comment_id','megaplan_file_id','contact_ref','contact_comment_id','publication_create_state'):
        if name not in columns: db.execute(f'ALTER TABLE installations ADD COLUMN {name} TEXT')
    report_columns={x[1] for x in db.execute('PRAGMA table_info(reports)')}
    work_columns={x[1] for x in db.execute('PRAGMA table_info(work_items)')}
    if 'cost_kind' not in work_columns:
        db.execute("ALTER TABLE work_items ADD COLUMN cost_kind TEXT NOT NULL DEFAULT 'work'")
    publication_columns={x[1] for x in db.execute('PRAGMA table_info(report_publications)')}
    if 'revision' not in publication_columns:
        db.execute('ALTER TABLE report_publications ADD COLUMN revision INTEGER NOT NULL DEFAULT 1')
    for name,definition in (
        ('external_id','TEXT'),('source_path','TEXT'),('imported','INTEGER NOT NULL DEFAULT 0')):
        if name not in report_columns: db.execute(f'ALTER TABLE reports ADD COLUMN {name} {definition}')
    report_item_columns={x[1] for x in db.execute('PRAGMA table_info(report_items)')}
    if 'accepted_amount' not in report_item_columns:
        db.execute('ALTER TABLE report_items ADD COLUMN accepted_amount REAL')
    payment_columns={x[1] for x in db.execute('PRAGMA table_info(payments)')}
    if 'request_id' not in payment_columns:
        db.execute('ALTER TABLE payments ADD COLUMN request_id TEXT')
    for name in ('settlement_amount','coefficient'):
        if name not in payment_columns: db.execute(f'ALTER TABLE payments ADD COLUMN {name} REAL')
    db.execute('CREATE UNIQUE INDEX IF NOT EXISTS payments_request_uq ON payments(source,request_id) WHERE request_id IS NOT NULL')
    db.execute('CREATE UNIQUE INDEX IF NOT EXISTS reports_external_id_uq ON reports(external_id) WHERE external_id IS NOT NULL')
    return db

def audit(db,deal_id,entity,entity_id,action,actor='',payload=None):
    db.execute('INSERT INTO audit_log VALUES(NULL,?,?,?,?,?,?,?)',
               (str(deal_id),entity,str(entity_id or ''),action,str(actor or ''),json.dumps(payload or {},ensure_ascii=False),now()))

def _cached_deal(deal_id):
    """Return the last known deal card without waiting for an unavailable CRM."""
    try:
        with open(deal_browser.CACHE_PATH, encoding='utf-8') as source:
            deals=json.load(source)
    except (OSError, json.JSONDecodeError):
        return None
    deal=next((x for x in deals if str(x.get('id')) == str(deal_id)), None)
    return deal_browser.enrich_cached_deal(deal, deals) if deal else None

def deal_source(deal_id, api=None):
    api=api or MegaplanAPI()
    d=None
    # Preparing a worksheet is a local draft operation. Do not make the manager
    # wait for the HTTP timeout when Megaplan itself is temporarily unreachable.
    if deal_browser.megaplan_reachable():
        try:
            d=api._call('GET',f'/deal/{deal_id}')['data']
        except Exception:
            d=None
    if d is None:
        d=_cached_deal(deal_id)
    if d is None:
        raise ValueError('Нет сохранённых данных сделки; Мегаплан сейчас недоступен')
    contractor=d.get('contractor') or {}
    client=contractor.get('name') or ' '.join(filter(None,[contractor.get('lastName'),contractor.get('firstName'),contractor.get('middleName')]))
    if not client and contractor.get('id'):
        endpoint={'ContractorHuman':'contractorHuman','ContractorCompany':'contractorCompany'}.get(contractor.get('contentType'))
        if endpoint:
            contractor=api._call('GET',f"/{endpoint}/{contractor['id']}")['data']
            client=contractor.get('name') or ' '.join(filter(None,[contractor.get('lastName'),contractor.get('firstName'),contractor.get('middleName')]))
    if not client:raise ValueError('Не удалось определить клиента сделки по ID')
    description=clean(html.unescape(re.sub(r'<[^>]*>',' ',str(d.get('description') or ''))))
    subject=d.get('Category1000068CustomFieldPredmetDogovora') or description or d.get('name') or 'монтажные работы'
    address=d.get('Category1000068CustomFieldAdresObEkta') or ''
    essence=brief_essence(subject,address)
    executors=d.get('Category1000068CustomFieldIspolnitel') or []
    if not isinstance(executors,list): executors=[executors]
    executor=executors[0] if executors else {}
    installer=' '.join(filter(None,[executor.get('lastName'),executor.get('firstName'),executor.get('middleName')])) or executor.get('name') or ''
    contact=d.get('contact') or {}; infos=contact.get('contactInfo') or []
    phone=next((x.get('value') for x in infos if x.get('type') in ('phone','mobile')), '')
    contact_name=contact.get('name') or ' '.join(filter(None,[contact.get('lastName'),contact.get('firstName'),contact.get('middleName')]))
    files=[]
    if deal_browser.megaplan_reachable():
        try:
            files=api._call('GET',f'/deal/{deal_id}/allFiles',params={'limit':100}).get('data') or []
        except Exception:
            files=[]
    projects=[{'id':str(x.get('id')),'name':x.get('name'),'path':x.get('path')} for x in files
              if is_project_document(x.get('name') or '')]
    return {'deal_id':str(deal_id),'client':client,'essence':clean(essence),'address':address,
            'contact_name':contact_name or '','contact_phone':phone or '','installer_name':installer,
            'installer_ref':str(executor.get('id') or ''),'status':(d.get('state') or {}).get('name') or '',
            'project_files':projects}

def ledger_work(deal_id):
    rows=[]
    for i,x in enumerate(deal_ledger.load(str(deal_id)),1):
        rate=float(x.get('cost_work') or 0); qty=float(x.get('qty') or 0)
        if rate<=0 or qty<=0: continue
        rows.append({'source_key':str(x.get('offer_row_id') or x.get('offer_id') or i),'name':clean(x.get('name')),
                     'unit':clean(x.get('unit')) or 'шт.','planned_qty':qty,'rate':rate,'sort_order':i})
    return rows

def folder_name(item):
    return safe(f"Сделка №{item['deal_id']} - {item['client']} - {item['essence']}")

def prepare(deal_id, actor='', api=None):
    src=deal_source(str(deal_id),api); works=ledger_work(str(deal_id))
    if not works: raise ValueError('в сделке не найдено работ с заполненной себестоимостью')
    db=connect(); stamp=now(); fname=folder_name(src)
    db.execute('''INSERT INTO installations(deal_id,client,essence,address,contact_name,contact_phone,installer_name,installer_ref,status,folder_name,folder_remote,created_at,updated_at)
      VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(deal_id) DO UPDATE SET client=excluded.client,essence=excluded.essence,address=excluded.address,
      installer_name=excluded.installer_name,installer_ref=excluded.installer_ref,folder_name=excluded.folder_name,folder_remote=excluded.folder_remote,updated_at=excluded.updated_at''',
      (src['deal_id'],src['client'],src['essence'],src['address'],'','',src['installer_name'],src['installer_ref'],'draft',fname,f'{ONEDRIVE_ROOT}/Активные/{fname}',stamp,stamp))
    for x in works:
        db.execute('''INSERT INTO work_items(deal_id,source_key,name,unit,planned_qty,rate,sort_order) VALUES(?,?,?,?,?,?,?)
          ON CONFLICT(deal_id,source_key) DO UPDATE SET name=excluded.name,unit=excluded.unit,planned_qty=excluded.planned_qty,rate=excluded.rate,sort_order=excluded.sort_order''',
          (src['deal_id'],x['source_key'],x['name'],x['unit'],x['planned_qty'],x['rate'],x['sort_order']))
    audit(db,deal_id,'installation',deal_id,'prepared',actor,{'works':len(works)}); db.commit()
    result=get_installation(deal_id,db); result['project_files']=src['project_files']; result['workbook']=build_workbook(deal_id,db)
    backup_db(); return result

def project_candidates(deal_id,api=None):
    api=api or MegaplanAPI(); files=api._call('GET',f'/deal/{deal_id}/allFiles',params={'limit':100}).get('data') or []
    return [{'id':str(x.get('id') or ''),'name':x.get('name') or 'Файл','path':x.get('path') or ''} for x in files if x.get('id')]

def ensure_installation_record(deal_id,db,api=None):
    row=db.execute('SELECT * FROM installations WHERE deal_id=?',(str(deal_id),)).fetchone()
    if row: return dict(row)
    src=deal_source(str(deal_id),api); stamp=now(); fname=folder_name(src)
    db.execute('''INSERT INTO installations(deal_id,client,essence,address,contact_name,contact_phone,installer_name,installer_ref,status,folder_name,folder_remote,created_at,updated_at)
      VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)''',(src['deal_id'],src['client'],src['essence'],src['address'],src['contact_name'],src['contact_phone'],src['installer_name'],src['installer_ref'],'draft',fname,f'{ONEDRIVE_ROOT}/Активные/{fname}',stamp,stamp))
    return dict(db.execute('SELECT * FROM installations WHERE deal_id=?',(str(deal_id),)).fetchone())

def _download_megaplan(api,item):
    suffix=os.path.splitext(item.get('name') or '')[1] or '.bin'; fd,path=tempfile.mkstemp(prefix='project-',suffix=suffix); os.close(fd)
    url=item['path'] if str(item.get('path') or '').startswith('http') else f"https://{api.domain}{item['path']}"
    request=urllib.request.Request(url,headers={'Authorization':f'Bearer {api.token}'})
    with megaplan_transport.urlopen(request,timeout=60) as response, open(path,'wb') as target: shutil.copyfileobj(response,target)
    return path

def attach_project(deal_id,file_ids=None,local_paths=None,actor='',api=None):
    """Create a new immutable project version from explicit manager selection."""
    api=api or MegaplanAPI(); db=connect(); ins=ensure_installation_record(deal_id,db,api)
    selected=[]; ids={str(x) for x in (file_ids or [])}
    if ids:
        selected=[x for x in project_candidates(deal_id,api) if x['id'] in ids]
        missing=ids-{x['id'] for x in selected}
        if missing: db.close(); raise ValueError('часть выбранных файлов Мегаплана не найдена')
    for local in local_paths or []:
        if not os.path.isfile(local): db.close(); raise ValueError(f'файл не найден: {os.path.basename(local)}')
        uploaded=api.upload_file(local)
        payload={'contentType':'Comment','subject':{'contentType':'Deal','id':str(deal_id)},
                 'content':'Добавлена новая версия проектной документации.',
                 'attaches':[{'contentType':'File','id':str(uploaded['id'])}]}
        api._call('POST',f'/deal/{deal_id}/comments',json_body=payload)
        selected.append({'id':str(uploaded['id']),'name':uploaded.get('name') or os.path.basename(local),
                         'path':uploaded.get('path') or '','local_path':local})
    if not selected: db.close(); raise ValueError('не выбрано ни одного файла проекта')
    version=int(db.execute('SELECT COALESCE(MAX(version_no),0)+1 FROM project_versions WHERE deal_id=?',(str(deal_id),)).fetchone()[0])
    db.execute("UPDATE project_versions SET status='archived' WHERE deal_id=? AND status='active'",(str(deal_id),))
    version_id=db.execute('INSERT INTO project_versions(deal_id,version_no,status,created_at,created_by) VALUES(?,?,?,?,?)',
                          (str(deal_id),version,'active',now(),actor)).lastrowid
    saved=[]
    try:
        for item in selected:
            local=item.get('local_path') or _download_megaplan(api,item)
            remote=f"{ins['folder_remote']}/00 Проектная документация/Версия {version}/{safe(item['name'])}"
            rclone_copy(local,remote)
            db.execute('INSERT INTO project_documents(version_id,megaplan_file_id,name,megaplan_path,onedrive_remote,created_at) VALUES(?,?,?,?,?,?)',
                       (version_id,item['id'],item['name'],item.get('path') or '',remote,now()))
            saved.append({'id':item['id'],'name':item['name'],'remote':remote})
            if not item.get('local_path'):
                try: os.remove(local)
                except OSError: pass
        db.execute('UPDATE installations SET project_required=1,updated_at=? WHERE deal_id=?',(now(),str(deal_id)))
        audit(db,deal_id,'project_version',version_id,'created',actor,{'version':version,'files':[x['name'] for x in saved]}); db.commit()
    except Exception:
        db.rollback(); db.close(); raise
    installation_status=ins.get('status'); db.close(); backup_db(); return {'deal_id':str(deal_id),'version':version,'files':saved,'installation_status':installation_status}

def set_without_project(deal_id,actor=''):
    db=connect(); get_installation(deal_id,db)
    db.execute('UPDATE installations SET project_required=0,updated_at=? WHERE deal_id=?',(now(),str(deal_id)))
    audit(db,deal_id,'installation',deal_id,'without_project_confirmed',actor); db.commit(); db.close(); backup_db()
    return get_installation(deal_id)

def import_legacy_workbook(deal_id,path,report_date,external_id,actor='',api=None):
    """Bootstrap an old installation and one immutable approved report from its source workbook."""
    if not os.path.isfile(path): raise ValueError('исходная ведомость не найдена')
    src=deal_source(str(deal_id),api); book=load_workbook(path,data_only=True); ws=book.active
    header=None
    for row in range(1,min(ws.max_row,80)+1):
        values=[norm(ws.cell(row,col).value) for col in range(1,min(ws.max_column,12)+1)]
        if any('наименование работ' in value for value in values) and any('кол во единиц' in value for value in values):
            header=row; break
    if not header: raise ValueError('в исходной ведомости не найдена таблица работ')
    rows=[]
    for row in range(header+1,ws.max_row+1):
        number=clean(ws.cell(row,1).value); name=clean(ws.cell(row,2).value)
        if 'итого' in number.casefold(): break
        if not name or not re.fullmatch(r'\d+(?:[.,]\d+)?',number): continue
        try:
            planned=float(ws.cell(row,4).value or 0); rate=float(ws.cell(row,5).value or 0)
            accepted_amount=float(ws.cell(row,7).value or 0)
        except (TypeError,ValueError): raise ValueError(f'строка {row}: неверное количество или расценка')
        if planned<0 or rate<0 or accepted_amount<0: raise ValueError(f'строка {row}: отрицательные значения запрещены')
        accepted_qty=accepted_amount/rate if rate else 0
        if accepted_qty>planned+1e-6: raise ValueError(f'строка {row}: принятый объём превышает план')
        rows.append({'source_key':f'legacy:{external_id}:{number}','name':name,'unit':clean(ws.cell(row,3).value) or 'шт.',
                     'planned_qty':planned,'rate':rate,'accepted_qty':accepted_qty,'source_row':row})
    if not rows: raise ValueError('в исходной ведомости не найдено строк работ')
    db=connect()
    duplicate=db.execute('SELECT id FROM reports WHERE external_id=?',(str(external_id),)).fetchone()
    if duplicate:
        db.close(); raise ValueError('этот исторический отчёт уже импортирован')
    stamp=now(); fname=folder_name(src); projects=src.get('project_files') or []
    try:
        db.execute('''INSERT INTO installations(deal_id,client,essence,address,contact_name,contact_phone,installer_name,installer_ref,status,project_required,folder_name,folder_remote,created_at,updated_at)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(deal_id) DO UPDATE SET client=excluded.client,essence=excluded.essence,address=excluded.address,
          contact_name=excluded.contact_name,contact_phone=excluded.contact_phone,installer_name=excluded.installer_name,installer_ref=excluded.installer_ref,
          folder_name=excluded.folder_name,folder_remote=excluded.folder_remote,updated_at=excluded.updated_at''',
          (src['deal_id'],src['client'],src['essence'],src['address'],src['contact_name'],src['contact_phone'],src['installer_name'],src['installer_ref'],'draft',None,fname,f'{ONEDRIVE_ROOT}/Активные/{fname}',stamp,stamp))
        work_ids=[]
        for order,item in enumerate(rows,1):
            cur=db.execute('''INSERT INTO work_items(deal_id,source_key,name,unit,planned_qty,rate,sort_order) VALUES(?,?,?,?,?,?,?)
              ON CONFLICT(deal_id,source_key) DO UPDATE SET name=excluded.name,unit=excluded.unit,planned_qty=excluded.planned_qty,rate=excluded.rate,sort_order=excluded.sort_order,active=1
              RETURNING id''',(src['deal_id'],item['source_key'],item['name'],item['unit'],item['planned_qty'],item['rate'],order))
            work_ids.append(cur.fetchone()[0])
        report_no=db.execute('SELECT COALESCE(MAX(report_no),0)+1 FROM reports WHERE deal_id=?',(src['deal_id'],)).fetchone()[0]
        report_id=db.execute('''INSERT INTO reports(deal_id,report_no,status,report_date,raw_text,submitted_at,approved_at,approved_by,manager_comment,external_id,source_path,imported)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,1)''',(src['deal_id'],report_no,'approved',report_date,
          f'Исторический отчёт, импортированный из {os.path.basename(path)}',stamp,stamp,actor,
          'Импортирован из ранее утверждённого реестра; повторное начисление запрещено.',str(external_id),os.path.abspath(path))).lastrowid
        for item,work_id in zip(rows,work_ids):
            db.execute('''INSERT INTO report_items(report_id,work_item_id,source_text,declared_qty,accepted_qty,status) VALUES(?,?,?,?,?,?)''',
              (report_id,work_id,f"Строка {item['source_row']} исходной ведомости",item['accepted_qty'],item['accepted_qty'],'accepted' if item['accepted_qty'] else 'rejected'))
        total=round(sum(x['accepted_qty']*x['rate'] for x in rows),2)
        audit(db,deal_id,'report',report_id,'legacy_imported',actor,{'external_id':external_id,'source':os.path.basename(path),'accepted_total':total})
        db.commit()
    except Exception:
        db.rollback(); db.close(); raise
    result=get_installation(deal_id,db); result['historical_report_id']=report_id; result['historical_accepted_total']=total
    result['project_files']=projects; result['workbook']=build_workbook(deal_id,db); db.close(); backup_db(); return result

class InstallationNotPrepared(ValueError):
    """Only a confirmed absent local row permits preparation."""

def get_installation(deal_id, db=None):
    own=db is None; db=db or connect(); row=db.execute('SELECT * FROM installations WHERE deal_id=?',(str(deal_id),)).fetchone()
    if not row:
        if own: db.close()
        raise InstallationNotPrepared('монтаж по сделке не подготовлен')
    out=dict(row); out['works']=[dict(x) for x in db.execute("SELECT * FROM work_items WHERE deal_id=? AND active=1 AND cost_kind='work' ORDER BY sort_order,id",(str(deal_id),))]
    out['materials']=[dict(x) for x in db.execute("SELECT * FROM work_items WHERE deal_id=? AND cost_kind='material' ORDER BY sort_order,id",(str(deal_id),))]
    out['reports']=[dict(x) for x in db.execute('SELECT * FROM reports WHERE deal_id=? ORDER BY report_no',(str(deal_id),))]
    out['accepted_total']=sum(float(x['accepted_amount'] if x['accepted_amount'] is not None else float(x['accepted_qty'] or 0)*float(x['rate'])) for x in db.execute('''SELECT ri.accepted_qty,ri.accepted_amount,w.rate FROM report_items ri JOIN reports r ON r.id=ri.report_id JOIN work_items w ON w.id=ri.work_item_id WHERE r.deal_id=? AND r.status IN ('approved','partially_approved')''',(str(deal_id),)))
    out['paid_total']=float(db.execute('SELECT COALESCE(SUM(COALESCE(settlement_amount,amount)),0) FROM payments WHERE deal_id=?',(str(deal_id),)).fetchone()[0])
    out['paid_gross_total']=float(db.execute('SELECT COALESCE(SUM(amount),0) FROM payments WHERE deal_id=?',(str(deal_id),)).fetchone()[0])
    out['debt']=round(out['accepted_total']-out['paid_total'],2)
    out['settlements_complete']=not any(r.get('imported') or str(r.get('external_id') or '').startswith('legacy-accepted:') for r in out['reports'])
    if not out['settlements_complete']:
        out['settlement_scope']='installer_global'
        out['settlement_warning']='По решению руководителя выплаты сборных и исторических отчётов сохраняются в общих взаиморасчётах Сергея без распределения по сделкам. Пообъектный долг не рассчитывается. Общий баланс доступен в разделе «Взаиморасчёты» / «Мои взаиморасчёты».'
    if own: db.close()
    return out

def contact_name(item):
    return clean(' '.join(filter(None,[item.get('lastName'),item.get('firstName'),item.get('middleName')]))) or clean(item.get('name'))

def contact_phone(item):
    return next((clean(x.get('value')) for x in item.get('contactInfo') or [] if x.get('type') in ('phone','mobile') and x.get('value')),'')

def installation_contacts(deal_id,api=None):
    api=api or MegaplanAPI(); deal=api._call('GET',f'/deal/{deal_id}')['data']; company=deal.get('contractor') or {}
    result=[]
    if company.get('contentType')=='ContractorCompany':
        result=api._call('GET',f"/contractorCompany/{company['id']}/contacts",params={'limit':100}).get('data') or []
    current=deal.get('contact') or {}
    if current.get('id') and str(current.get('id')) not in {str(x.get('id')) for x in result}: result.insert(0,current)
    return [{'id':str(x.get('id')),'name':contact_name(x),'phone':contact_phone(x),'position':clean(x.get('position'))} for x in result if x.get('id') and not x.get('isDropped')]

def _publish_installation_contact(db,deal_id,name,phone,ref,api):
    ins=dict(db.execute('SELECT * FROM installations WHERE deal_id=?',(str(deal_id),)).fetchone())
    payload={'contentType':'Comment','subject':{'contentType':'Deal','id':str(deal_id)},
             'content':f"👤 Контакт по монтажу\n{name}\n{phone}\nID контакта Мегаплана: {ref}"}
    comment={}
    if ins.get('contact_comment_id'):
        payload['id']=str(ins['contact_comment_id'])
        try: comment=api._call('POST',f"/comment/{ins['contact_comment_id']}",json_body=payload).get('data') or {}
        except Exception: payload.pop('id',None)
    if not comment: comment=api._call('POST',f'/deal/{deal_id}/comments',json_body=payload).get('data') or {}
    db.execute('UPDATE installations SET contact_comment_id=? WHERE deal_id=?',(str(comment.get('id') or ''),str(deal_id)))

def select_installation_contact(deal_id,contact_ref,actor='',api=None):
    api=api or MegaplanAPI(); contacts=installation_contacts(deal_id,api); item=next((x for x in contacts if x['id']==str(contact_ref)),None)
    if not item: raise ValueError('контакт не относится к клиенту этой сделки')
    if not item['phone']: raise ValueError('у выбранного контакта не указан телефон')
    db=connect(); previous=db.execute('SELECT contact_ref,contact_name,contact_phone,status FROM installations WHERE deal_id=?',(str(deal_id),)).fetchone()
    if not previous: ensure_installation_record(deal_id,db,api); previous=db.execute('SELECT contact_ref,contact_name,contact_phone,status FROM installations WHERE deal_id=?',(str(deal_id),)).fetchone()
    db.execute('UPDATE installations SET contact_ref=?,contact_name=?,contact_phone=?,updated_at=? WHERE deal_id=?',(item['id'],item['name'],item['phone'],now(),str(deal_id)))
    _publish_installation_contact(db,deal_id,item['name'],item['phone'],item['id'],api)
    audit(db,deal_id,'installation_contact',item['id'],'selected',actor,{'before':dict(previous),'after':item}); db.commit(); out=get_installation(deal_id,db); db.close(); backup_db(); return out

def create_installation_contact(deal_id,name,phone,actor='',api=None):
    api=api or MegaplanAPI(); deal=api._call('GET',f'/deal/{deal_id}')['data']; company=deal.get('contractor') or {}
    if company.get('contentType')!='ContractorCompany': raise ValueError('у сделки нет клиента-организации для создания контактного лица')
    parts=clean(name).split();
    if len(parts)<2: raise ValueError('укажите как минимум фамилию и имя')
    payload={'contentType':'ContractorHuman','type':{'contentType':'ContractorType','id':'1000002'},
      'company':{'contentType':'ContractorCompany','id':str(company['id'])},'lastName':parts[0],'firstName':parts[1],
      'middleName':' '.join(parts[2:]),'contactInfo':[{'contentType':'ContactInfo','type':'phone','value':clean(phone),'comment':'Контакт по монтажу'},
      {'contentType':'Address','type':{'contentType':'AddressType','id':'1'},'value':'не указан'}]}
    created=api._call('POST','/contractorHuman',json_body=payload).get('data') or {}
    if not created.get('id'): raise ValueError('Мегаплан не вернул ID нового контакта')
    return select_installation_contact(deal_id,created['id'],actor,api)

def build_workbook(deal_id, db=None):
    own=db is None; db=db or connect(); ins=get_installation(deal_id,db)
    os.makedirs(WORKBOOK_DIR,exist_ok=True); path=os.path.join(WORKBOOK_DIR,f"Монтажная ведомость - сделка {deal_id}.xlsx")
    wb=Workbook(); ws=wb.active; ws.title='Монтажная ведомость'; ws.sheet_view.showGridLines=False
    reports=[x for x in ins['reports'] if x['status'] in ('approved','partially_approved')]
    headers=['ID','Работа','Ед.','План','Расценка','Плановая сумма']+[f"Отчёт №{x['report_no']} - {x['report_date']}" for x in reports]+['Выполнено всего','Осталось','Принято к оплате','Выполнено за этот отчёт, кол-во']
    ws.append([f"Сделка №{deal_id} - {ins['client']} - {ins['essence']}"]); ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers)); ws['A1'].font=Font(bold=True,size=14,color='FFFFFF'); ws['A1'].fill=PatternFill('solid',fgColor='17365D')
    ws.append(headers)
    for c in ws[2]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='2F75B5'); c.alignment=Alignment(wrap_text=True,horizontal='center')
    for w in ins['works']:
        accepted=[]
        accepted_amount=0.0
        for rep in reports:
            q=db.execute('SELECT COALESCE(SUM(accepted_qty),0) FROM report_items WHERE report_id=? AND work_item_id=?',(rep['id'],w['id'])).fetchone()[0]
            accepted.append(float(q or 0))
            accepted_amount+=float(db.execute('SELECT COALESCE(SUM(COALESCE(accepted_amount,accepted_qty*?)),0) FROM report_items WHERE report_id=? AND work_item_id=?',(float(w['rate']),rep['id'],w['id'])).fetchone()[0] or 0)
        total=sum(accepted); row=[w['id'],w['name'],w['unit'],w['planned_qty'],w['rate'],w['planned_qty']*w['rate'],*accepted,total,max(0,float(w['planned_qty'])-total),accepted_amount,None]; ws.append(row)
    first_data_row=3; last_data_row=ws.max_row
    # Telegram/mobile viewers may show stale formula caches.  The file is a
    # controlled export, so persist the verified totals as values; every bot
    # regeneration calculates them anew from the database.
    plan_total=sum(float(w['planned_qty'])*float(w['rate']) for w in ins['works'])
    done_total=sum(float(db.execute('SELECT COALESCE(SUM(accepted_qty),0) FROM report_items ri JOIN reports r ON r.id=ri.report_id WHERE ri.work_item_id=? AND r.status IN (\'approved\',\'partially_approved\')',(w['id'],)).fetchone()[0] or 0) for w in ins['works'])
    accepted_total=float(db.execute("SELECT COALESCE(SUM(COALESCE(ri.accepted_amount,ri.accepted_qty*w.rate)),0) FROM report_items ri JOIN reports r ON r.id=ri.report_id JOIN work_items w ON w.id=ri.work_item_id WHERE r.deal_id=? AND w.cost_kind='work' AND r.status IN ('approved','partially_approved')",(str(deal_id),)).fetchone()[0] or 0)
    report_totals=[]
    for rep in reports:
        report_totals.append(float(db.execute("SELECT COALESCE(SUM(COALESCE(ri.accepted_amount,ri.accepted_qty*w.rate)),0) FROM report_items ri JOIN work_items w ON w.id=ri.work_item_id WHERE ri.report_id=? AND w.cost_kind='work'",(rep['id'],)).fetchone()[0] or 0))
    total_row=['','','','', 'ИТОГО ПО ПЛАНУ', plan_total, *report_totals, done_total, '', accepted_total]
    ws.append(total_row)
    for c in ws[ws.max_row]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='D9EAF7')
    for row in ws.iter_rows(min_row=3):
        for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    remaining_col=8+len(reports)
    complete_fill=PatternFill('solid',fgColor='C6EFCE')
    for row in range(first_data_row,last_data_row+1):
        if float(ws.cell(row,remaining_col).value or 0)<=1e-9:
            for cell in ws[row]: cell.fill=complete_fill
    for col,width in enumerate([8,55,10,12,15,18]+[16]*len(reports)+[16,16,20,28],1): ws.column_dimensions[get_column_letter(col)].width=width
    ws.freeze_panes='A3'; ws.auto_filter.ref=f"A2:{get_column_letter(len(headers))}{max(2,ws.max_row-1)}"
    log=wb.create_sheet('Журнал отчётов'); log.append(['№','Дата','Статус','Отправлен','Принят','Сумма принятого, ₽','Комментарий'])
    for r in ins['reports']:
        accepted_amount=float(db.execute('''SELECT COALESCE(SUM(COALESCE(ri.accepted_amount,ri.accepted_qty*w.rate)),0) FROM report_items ri LEFT JOIN work_items w ON w.id=ri.work_item_id WHERE ri.report_id=?''',(r['id'],)).fetchone()[0] or 0)
        log.append([r['report_no'],r['report_date'],r['status'],r['submitted_at'],r['approved_at'],accepted_amount,r['manager_comment']])
    log.column_dimensions['F'].width=20; log['F1'].font=Font(bold=True)
    if ins['materials']:
        materials=wb.create_sheet('Материалы монтажника')
        materials.append(['ID','Материал','Принято к возмещению, ₽','Источник'])
        for item in ins['materials']:
            amount=db.execute("SELECT COALESCE(SUM(ri.accepted_amount),0) FROM report_items ri JOIN reports r ON r.id=ri.report_id WHERE ri.work_item_id=? AND r.status IN ('approved','partially_approved')",(item['id'],)).fetchone()[0]
            materials.append([item['id'],item['name'],amount,item['source_key']])
        materials.column_dimensions['B'].width=60;materials.column_dimensions['C'].width=26
    extras=db.execute("SELECT w.id,w.name,COALESCE(SUM(ri.accepted_amount),0) amount FROM work_items w JOIN report_items ri ON ri.work_item_id=w.id JOIN reports r ON r.id=ri.report_id WHERE w.deal_id=? AND w.cost_kind='extra_work' AND r.status IN ('approved','partially_approved') GROUP BY w.id",(str(deal_id),)).fetchall()
    if extras:
        sheet=wb.create_sheet('Дополнительные затраты');sheet.append(['ID','Назначение','Принято, ₽'])
        for x in extras:sheet.append([x['id'],x['name'],x['amount']])
        sheet.append(['','ИТОГО',sum(x['amount'] for x in extras)])
        sheet.column_dimensions['B'].width=65;sheet.column_dimensions['C'].width=24
    ws.cell(2,len(headers)).fill=PatternFill('solid',fgColor='548235')
    from openpyxl.comments import Comment
    ws.cell(2,len(headers)).comment=Comment('Заполните только новый объём за этот отчёт в единицах строки. План, прошлое выполнение и расценки не меняйте. Начисление — после утверждения руководителем.','SmartBild')
    wb.save(path); db.execute('UPDATE installations SET workbook_path=?,updated_at=? WHERE deal_id=?',(path,now(),str(deal_id))); db.commit()
    if own: db.close()
    return path

def import_workbook(deal_id,path,actor=''):
    wb=load_workbook(path,data_only=False); ws=wb['Монтажная ведомость']; expected=['ID','Работа','Ед.','План','Расценка']
    if [ws.cell(2,i).value for i in range(1,6)]!=expected: raise ValueError('структура монтажной ведомости изменена')
    db=connect(); db.execute('BEGIN IMMEDIATE'); ids=set(); changes=[]
    try:
        for row in range(3,ws.max_row+1):
            if 'ИТОГО' in clean(ws.cell(row,5).value).upper() or clean(ws.cell(row,6).value).upper()=='ИТОГО': break
            item_id=ws.cell(row,1).value; name=clean(ws.cell(row,2).value); unit=clean(ws.cell(row,3).value); qty=ws.cell(row,4).value; rate=ws.cell(row,5).value
            if not name: continue
            try: qty=float(qty); rate=float(rate)
            except (TypeError,ValueError): raise ValueError(f'строка {row}: количество и расценка должны быть числами')
            if not math.isfinite(qty) or not math.isfinite(rate) or qty<0 or rate<0: raise ValueError(f'строка {row}: нужны конечные неотрицательные числа')
            if item_id:
                if int(item_id) in ids: raise ValueError(f'строка {row}: повторный ID работы')
                old=db.execute('SELECT * FROM work_items WHERE id=? AND deal_id=?',(int(item_id),str(deal_id))).fetchone()
                if not old or old['cost_kind']!='work': raise ValueError(f'строка {row}: неизвестный ID работы')
                ids.add(int(item_id)); new=(name,unit,qty,rate)
                if new!=(old['name'],old['unit'],old['planned_qty'],old['rate']): changes.append({'id':int(item_id),'before':dict(old),'after':{'name':name,'unit':unit,'planned_qty':qty,'rate':rate}})
                db.execute('UPDATE work_items SET name=?,unit=?,planned_qty=?,rate=?,active=1 WHERE id=?',(name,unit,qty,rate,int(item_id)))
            else:
                order=db.execute('SELECT COALESCE(MAX(sort_order),0)+1 FROM work_items WHERE deal_id=?',(str(deal_id),)).fetchone()[0]
                cur=db.execute('INSERT INTO work_items(deal_id,source_key,name,unit,planned_qty,rate,sort_order) VALUES(?,?,?,?,?,?,?)',(str(deal_id),f'manual:{hashlib.sha1((name+str(order)).encode()).hexdigest()[:12]}',name,unit,qty,rate,order)); ids.add(cur.lastrowid); changes.append({'id':cur.lastrowid,'added':{'name':name,'qty':qty,'rate':rate}})
        for old in db.execute("SELECT id,name FROM work_items WHERE deal_id=? AND active=1 AND cost_kind='work'",(str(deal_id),)):
            if old['id'] not in ids: db.execute('UPDATE work_items SET active=0 WHERE id=?',(old['id'],)); changes.append({'id':old['id'],'removed':old['name']})
        enqueue_publication(db,deal_id)
        db.execute('UPDATE installations SET version=version+1,updated_at=? WHERE deal_id=?',(now(),str(deal_id))); audit(db,deal_id,'installation',deal_id,'workbook_imported',actor,{'changes':changes}); db.commit(); out=get_installation(deal_id,db); out['changes']=changes; out['workbook']=build_workbook(deal_id,db); db.close()
        if deal_browser.megaplan_reachable():
            try: refresh_economics(deal_id); out['economics_refreshed']=True
            except Exception as exc: out['economics_refreshed']=False; out['economics_error']=str(exc)
        else:
            out['economics_refreshed']=False
            out['economics_error']='Мегаплан временно недоступен; обновление поставлено в очередь'
        backup_db(); return out
    except Exception:
        try: db.rollback()
        except sqlite3.ProgrammingError: pass
        raise
    finally:
        db.close()

def update_details(deal_id, contact_name=None, contact_phone=None, project_required=None, actor=''):
    db=connect(); fields=[]; values=[]
    for field,value in [('contact_name',contact_name),('contact_phone',contact_phone),('project_required',project_required)]:
        if value is not None: fields.append(f'{field}=?'); values.append(value)
    values.extend([now(),str(deal_id)]); db.execute(f"UPDATE installations SET {','.join(fields)},updated_at=? WHERE deal_id=?",values)
    audit(db,deal_id,'installation',deal_id,'details_updated',actor,{x:y for x,y in [('contact_name',contact_name),('contact_phone',contact_phone),('project_required',project_required)] if y is not None}); db.commit(); out=get_installation(deal_id,db); db.close(); backup_db(); return out

def sync_executor(deal_id,actor='',api=None):
    source=deal_source(str(deal_id),api);db=connect()
    existing=db.execute('SELECT deal_id FROM installations WHERE deal_id=?',(str(deal_id),)).fetchone()
    if not existing:
        db.close();return {'ok':True,'deal_id':str(deal_id),'installation_exists':False,
            'installer_name':source['installer_name'],'installer_ref':source['installer_ref']}
    db.execute('UPDATE installations SET installer_name=?,installer_ref=?,updated_at=? WHERE deal_id=?',
        (source['installer_name'],source['installer_ref'],now(),str(deal_id)))
    audit(db,deal_id,'installation',deal_id,'executor_synced',actor,
        {'installer_name':source['installer_name'],'installer_ref':source['installer_ref']})
    db.commit();result=get_installation(deal_id,db);db.close();backup_db()
    return {'ok':True,'deal_id':str(deal_id),'installation_exists':True,'result':result}

def approve(deal_id,actor=''):
    db=connect(); ins=get_installation(deal_id,db)
    # A site contact improves the hand-off but is not a prerequisite: some
    # installations legitimately start before the customer names one.
    missing=[label for value,label in ((ins['installer_name'],'исполнитель'),(ins['address'],'адрес'),(ins['project_required'],'решение по проектной документации')) if value in (None,'')]
    if missing: raise ValueError('не заполнено: '+', '.join(missing))
    if int(ins['project_required'])==1 and not has_active_project(deal_id,db):
        raise ValueError('проект подтверждён, но в сделке не найдено ни одного проектного файла')
    db.execute("UPDATE installations SET status='approved',approved_at=?,updated_at=? WHERE deal_id=?",(now(),now(),str(deal_id)))
    audit(db,deal_id,'installation',deal_id,'approved',actor); enqueue_publication(db,deal_id)
    db.commit(); out=get_installation(deal_id,db); db.close()
    out.update(local_publication_result(deal_id))
    return out

def transfer(deal_id,actor=''):
    db=connect(); db.execute('BEGIN IMMEDIATE'); ins=get_installation(deal_id,db)
    if ins['status'] in ('transferred','in_progress','review'):
        db.close(); return ins
    if ins['status']!='approved':
        db.close(); raise ValueError('сначала утвердите монтажную ведомость')
    db.execute("UPDATE installations SET status='transferred',transferred_at=?,updated_at=? WHERE deal_id=?",(now(),now(),str(deal_id)))
    audit(db,deal_id,'installation',deal_id,'transferred',actor); enqueue_publication(db,deal_id)
    db.execute('INSERT OR IGNORE INTO handoff_events VALUES(?,?)',(str(deal_id),now()))
    db.commit(); out=get_installation(deal_id,db); db.close()
    out.update(local_publication_result(deal_id))
    return out

def list_mine(installer_name=None):
    """Return objects explicitly handed to an installer.

    ``installations.status`` is the hand-off source of truth.  We used to make
    a second live Megaplan request here and hide an object unless the CRM state
    was literally ``Монтаж``.  That made an already transferred object vanish
    from the installer menu whenever a state was later changed in Megaplan (or
    the API returned a stale/other state).  The object must remain available
    until the manager explicitly completes/returns it in the installation
    workflow.
    """
    # SQLite lower() is ASCII-only in the standard build, so a Cyrillic surname
    # must be matched with Python's Unicode-aware casefold().
    db=connect(); rows=[dict(x) for x in db.execute(
        "SELECT * FROM installations WHERE status IN ('transferred','in_progress','review') ORDER BY transferred_at DESC")]; db.close(); live=[]
    for row in rows:
        if installer_name and installer_name.casefold() not in str(row.get('installer_name') or '').casefold():
            continue
        live.append(row)
    return live


def list_history(installer_name=None):
    """Read-only historical objects imported from approved legacy reports."""
    db=connect(); rows=[dict(x) for x in db.execute(
        "SELECT * FROM installations WHERE status='historical' ORDER BY updated_at DESC")]; db.close()
    return [row for row in rows if not installer_name or installer_name.casefold() in str(row.get('installer_name') or '').casefold()]

def create_report(deal_id,chat_id,text=''):
    db=connect(); ins=get_installation(deal_id,db)
    if ins['status'] not in ('transferred','in_progress','review'): raise ValueError('объект ещё не передан монтажнику')
    existing=db.execute("SELECT * FROM reports WHERE deal_id=? AND status='draft' AND installer_chat_id=?",(str(deal_id),str(chat_id))).fetchone()
    if existing: report_id=existing['id']; db.execute('UPDATE reports SET raw_text=trim(COALESCE(raw_text,\'\') || ? ) WHERE id=?',(('\n'+text) if text else '',report_id))
    else:
        no=db.execute('SELECT COALESCE(MAX(report_no),0)+1 FROM reports WHERE deal_id=?',(str(deal_id),)).fetchone()[0]
        report_id=db.execute('INSERT INTO reports(deal_id,report_no,report_date,raw_text,installer_chat_id) VALUES(?,?,?,?,?)',(str(deal_id),no,datetime.now().strftime('%d.%m.%Y'),text,str(chat_id))).lastrowid
    db.execute("UPDATE installations SET status='in_progress',updated_at=? WHERE deal_id=?",(now(),str(deal_id))); enqueue_publication(db,deal_id); db.commit(); row=dict(db.execute('SELECT * FROM reports WHERE id=?',(report_id,)).fetchone()); db.close(); return row

def append_report_text(report_id,text):
    db=connect()
    changed=db.execute("UPDATE reports SET raw_text=trim(COALESCE(raw_text,'') || ?) WHERE id=? AND status IN ('draft','returned')",('\n'+clean(text),int(report_id)))
    if not changed.rowcount:
        db.close(); raise ValueError('Текст отправленного или принятого отчёта нельзя изменить; создайте новый отчёт')
    db.commit(); row=dict(db.execute('SELECT * FROM reports WHERE id=?',(int(report_id),)).fetchone()); db.close(); return row

def add_media(report_id,path,file_id='',unique_id='',name='',mime=''):
    from report_media_store import retain
    retained,digest,size=retain(path,os.path.join(os.path.dirname(DB),'report_sources'))
    db=connect()
    try:
        db.execute('BEGIN IMMEDIATE')
        report=db.execute('SELECT status FROM reports WHERE id=?',(int(report_id),)).fetchone()
        if not report or report['status'] not in ('draft','returned'):raise ValueError('В отправленный или принятый отчёт нельзя добавлять материалы')
        db.execute('INSERT OR IGNORE INTO media(report_id,telegram_file_id,telegram_unique_id,local_path,original_name,mime_type,size,sha256,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(int(report_id),file_id,unique_id or digest,retained,name or os.path.basename(path),mime,size,digest,now())); db.commit()
        return {'size':size,'sha256':digest}
    finally:db.close()

def norm(value): return re.sub(r'[^0-9a-zа-я]+',' ',str(value or '').casefold().replace('ё','е')).strip()
def get_report(report_id):
    db=connect()
    try:
        report=db.execute('SELECT * FROM reports WHERE id=?',(int(report_id),)).fetchone()
        if not report: raise ValueError('отчёт не найден')
        items=[dict(x) for x in db.execute('''SELECT ri.*,w.name,w.unit,w.rate,w.planned_qty FROM report_items ri
            LEFT JOIN work_items w ON w.id=ri.work_item_id WHERE ri.report_id=?''',(int(report_id),))]
        return {'report':dict(report),'items':items}
    finally: db.close()

def pending_reports():
    db=connect()
    try:return [dict(row) for row in db.execute("SELECT id,deal_id,report_no,report_date,status FROM reports WHERE status='submitted' ORDER BY submitted_at,id")]
    finally:db.close()


def parse_report(report_id):
    db=connect(); db.execute('BEGIN IMMEDIATE'); report=db.execute('SELECT * FROM reports WHERE id=?',(int(report_id),)).fetchone()
    if not report or report['status'] in ('approved','partially_approved'):
        db.close(); raise ValueError('Принятый отчёт нельзя разбирать повторно; начисления сохранены')
    works=[dict(x) for x in db.execute('SELECT * FROM work_items WHERE deal_id=? AND active=1',(report['deal_id'],))]
    if not works:
        db.close(); raise ValueError('В ведомости нет действующих работ для сопоставления')
    db.execute('DELETE FROM report_items WHERE report_id=?',(int(report_id),)); unmatched=[]
    raw=str(report['raw_text'] or '')
    word_numbers={'один':'1','одна':'1','одно':'1','два':'2','две':'2','три':'3','четыре':'4','пять':'5','шесть':'6','семь':'7','восемь':'8','девять':'9','десять':'10'}
    for word,number in word_numbers.items(): raw=re.sub(rf'(?i)\b{word}\b',number,raw)
    chunks=re.split(r'[\n;,]+|(?<=[.!?])\s+',raw)
    for line in [clean(x).strip(' .') for x in chunks if clean(x).strip(' .')]:
        nums=re.findall(r'(?<!\d)(\d+(?:[.,]\d+)?)',line)
        if not nums: unmatched.append(line); continue
        qty=float(nums[-1].replace(',','.')); target=norm(re.sub(r'\d+(?:[.,]\d+)?',' ',line)); scored=[]
        for w in works:
            wn=norm(w['name']); score=SequenceMatcher(None,target,wn).ratio()
            tokens=set(target.split()); wt=set(wn.split()); score=max(score,len(tokens&wt)/max(1,len(tokens|wt)))
            scored.append((score,w))
        score,w=max(scored,key=lambda x:x[0])
        if score<0.28: unmatched.append(line); db.execute('INSERT INTO report_items(report_id,source_text,declared_qty,status) VALUES(?,?,?,?)',(int(report_id),line,qty,'unmatched'))
        else: db.execute('INSERT INTO report_items(report_id,work_item_id,source_text,declared_qty,status) VALUES(?,?,?,?,?)',(int(report_id),w['id'],line,qty,'matched'))
    db.commit(); items=[dict(x) for x in db.execute('''SELECT ri.*,w.name,w.unit,w.planned_qty,w.rate FROM report_items ri LEFT JOIN work_items w ON w.id=ri.work_item_id WHERE report_id=?''',(int(report_id),))]; db.close(); return {'report':dict(report),'items':items,'unmatched':unmatched}

def submit_report(report_id):
    current=get_report(report_id)
    if current['report']['status']=='submitted':return {**current,'publication_pending':True}
    if current['report']['status'] not in ('draft','returned'):raise ValueError('Этот отчёт уже принят или закрыт')
    parsed=parse_report(report_id); db=connect()
    try:
        db.execute('BEGIN IMMEDIATE')
        changed=db.execute("UPDATE reports SET status='submitted',submitted_at=? WHERE id=? AND status IN ('draft','returned')",(now(),int(report_id)))
        if not changed.rowcount:raise ValueError('Статус отчёта изменился; повторная отправка остановлена')
        db.execute("UPDATE installations SET status='review',updated_at=? WHERE deal_id=?",(now(),parsed['report']['deal_id']))
        enqueue_report_publication(db,report_id)
        db.commit()
    finally:db.close()
    return {**get_report(report_id),'unmatched':parsed.get('unmatched') or [],'publication_pending':True}

def approve_report(report_id,actor='',partial=None,comment=''):
    db=connect()
    try:
        db.execute('BEGIN IMMEDIATE')
        report=db.execute('SELECT * FROM reports WHERE id=?',(int(report_id),)).fetchone()
        if not report: raise ValueError('отчёт не найден')
        if report['status'] in ('approved','partially_approved'):
            raise ValueError('отчёт уже принят; повторное начисление или изменение принятого запрещено')
        rows=db.execute('SELECT * FROM report_items WHERE report_id=?',(int(report_id),)).fetchall()
        if not rows: raise ValueError('в отчёте нет разобранных работ')
        if partial and set(partial)-{str(x['id']) for x in rows}:
            raise ValueError('В корректировке указан ID строки, отсутствующий в этом отчёте')
        accepted_by_work={}
        for row in rows:
            qty=float((partial or {}).get(str(row['id']),row['declared_qty']) if partial is not None else row['declared_qty']) if row['work_item_id'] else 0
            if not math.isfinite(qty) or qty < 0: raise ValueError('принятый объём должен быть конечным неотрицательным числом')
            amount=0
            if row['work_item_id']:
                work=db.execute('SELECT * FROM work_items WHERE id=? AND deal_id=?',(row['work_item_id'],report['deal_id'])).fetchone()
                if not work or not work['active']: raise ValueError('работа не входит в действующую ведомость объекта')
                plan=float(work['planned_qty']); rate=float(work['rate'])
                if not math.isfinite(rate) or rate < 0: raise ValueError('неверная расценка работы')
                prior=float(db.execute('''SELECT COALESCE(SUM(ri.accepted_qty),0) FROM report_items ri JOIN reports r ON r.id=ri.report_id WHERE ri.work_item_id=? AND r.status IN ('approved','partially_approved') AND r.id<>?''',(row['work_item_id'],int(report_id))).fetchone()[0])
                cumulative=accepted_by_work.get(row['work_item_id'],0)+qty
                if cumulative>plan-prior+1e-9:
                    raise ValueError(f"объём превышает остаток по работе ID {row['work_item_id']}: доступно {max(0,plan-prior):g}")
                accepted_by_work[row['work_item_id']]=cumulative
                amount=round(qty*rate,2)
            db.execute('UPDATE report_items SET accepted_qty=?,accepted_amount=?,status=? WHERE id=?',
                       (qty,amount,'accepted' if qty else 'rejected',row['id']))
        status='partially_approved' if partial is not None else 'approved'
        db.execute('UPDATE reports SET status=?,approved_at=?,approved_by=?,manager_comment=? WHERE id=?',(status,now(),actor,comment,int(report_id)))
        db.execute("UPDATE installations SET status='in_progress',updated_at=? WHERE deal_id=?",(now(),report['deal_id']))
        audit(db,report['deal_id'],'report',report_id,'approved',actor,{'partial':partial is not None})
        enqueue_publication(db,report['deal_id'])
        enqueue_report_publication(db,report_id)
        db.commit(); deal_id=report['deal_id']
    except Exception:
        db.rollback(); raise
    finally:
        db.close()
    # Financial acceptance is already committed. Publication failures must not
    # turn it into a false "not accepted" response or permit re-acceptance.
    return local_publication_result(deal_id)


def local_publication_result(deal_id):
    failures={}
    try: build_workbook(deal_id)
    except Exception as exc: failures['workbook']=type(exc).__name__
    result=get_installation(deal_id)
    result.update(economics_refreshed=False, economics_error='', publication_pending=True, publication_errors=failures)
    return result


def enqueue_publication(db,deal_id):
    # Called in the same transaction as the financial fact; a crash cannot
    # commit acceptance while forgetting that its views need rebuilding.
    db.execute('''INSERT INTO publication_jobs(deal_id,updated_at) VALUES(?,?)
        ON CONFLICT(deal_id) DO UPDATE SET revision=revision+1,status='pending',error=NULL,updated_at=excluded.updated_at''',
        (str(deal_id),now()))

def enqueue_report_publication(db,report_id):
    db.execute("INSERT INTO report_publications(report_id,updated_at) VALUES(?,?) ON CONFLICT(report_id) DO UPDATE SET revision=revision+1,status='pending',error=NULL,updated_at=excluded.updated_at",(int(report_id),now()))


def publish_pending(deal_id):
    from registry_transaction import locked
    # These are derived views only. No invoice, accrual or payment is recreated.
    with locked(DB+'.publication'):
        db=connect(); job=db.execute('SELECT * FROM publication_jobs WHERE deal_id=?',(str(deal_id),)).fetchone(); db.close()
        if not job or job['status']=='done': return {}
        failures={}
        reachable=deal_browser.megaplan_reachable()
        actions=[('workbook',lambda:build_workbook(deal_id)),('onedrive',lambda:sync_installation(deal_id))]
        if reachable:
            actions.extend([('megaplan',lambda:publish_workbook(deal_id)),
                            ('economics',lambda:refresh_economics(deal_id))])
        else:
            # Keep the job pending, but never occupy a worker for minutes on
            # predictable CRM connection timeouts.
            failures={'megaplan':'offline','economics':'offline'}
        actions.append(('backup',backup_db))
        for step,action in actions:
            try: action()
            except Exception as exc: failures[step]=type(exc).__name__
        db=connect()
        db.execute('UPDATE publication_jobs SET status=?,error=?,updated_at=? WHERE deal_id=? AND revision=?',
                   ('pending' if failures else 'done',json.dumps(failures),now(),str(deal_id),job['revision']))
        if failures: audit(db,deal_id,'publication',deal_id,'publication_failed','system',failures)
        db.commit(); db.close()
        return failures


def retry_publications(limit=1):
    from legacy_report_commit import sync_ledger
    sync_ledger(limit=limit)
    retry_report_publications(limit)
    db=connect(); rows=db.execute("SELECT deal_id FROM publication_jobs WHERE status='pending' AND (error IS NULL OR (julianday('now')-julianday(updated_at))*86400>=300) ORDER BY updated_at LIMIT ?",(max(1,min(int(limit),5)),)).fetchall(); db.close()
    return {row['deal_id']:publish_pending(row['deal_id']) for row in rows}

def retry_report_publications(limit=1):
    from registry_transaction import locked
    with locked(DB+'.report-publications'):
        db=connect(); rows=db.execute("SELECT report_id,revision FROM report_publications WHERE status='pending' AND (error IS NULL OR (julianday('now')-julianday(updated_at))*86400>=300) ORDER BY updated_at LIMIT ?",(max(1,min(int(limit),5)),)).fetchall(); db.close()
        result={}
        for row in rows:
            error=None
            try:sync_report(row['report_id']); backup_db()
            except Exception as exc:error=type(exc).__name__
            db=connect(); db.execute('UPDATE report_publications SET status=?,error=?,updated_at=? WHERE report_id=? AND revision=?',('pending' if error else 'done',error,now(),row['report_id'],row['revision'])); db.commit(); db.close()
            result[str(row['report_id'])]=error
        return result

def return_report(report_id,comment,actor=''):
    if not clean(comment): raise ValueError('укажите причину возврата')
    db=connect()
    try:
        db.execute('BEGIN IMMEDIATE')
        report=db.execute('SELECT * FROM reports WHERE id=?',(int(report_id),)).fetchone()
        if not report: raise ValueError('отчёт не найден')
        if report['status'] in ('approved','partially_approved'):
            raise ValueError('Принятый отчёт нельзя вернуть без отдельной корректировки начислений')
        db.execute("UPDATE reports SET status='returned',manager_comment=? WHERE id=?",(clean(comment),int(report_id)))
        db.execute("UPDATE installations SET status='in_progress',updated_at=? WHERE deal_id=?",(now(),report['deal_id']))
        audit(db,report['deal_id'],'report',report_id,'returned',actor,{'comment':clean(comment)})
        enqueue_report_publication(db,report_id)
        db.commit()
        return {**dict(report),'status':'returned','manager_comment':clean(comment),'publication_pending':True}
    except Exception:
        db.rollback(); raise
    finally:db.close()

def rclone_copy(local,remote):
    result=subprocess.run(['rclone','copyto',local,remote,'--retries','3','--low-level-retries','5'],capture_output=True,text=True,timeout=90)
    if result.returncode: raise ValueError('OneDrive: '+(result.stderr or 'ошибка загрузки')[-500:])

def sync_installation(deal_id):
    ins=get_installation(deal_id); path=build_workbook(deal_id); remote=f"{ins['folder_remote']}/01 Монтажная ведомость/{os.path.basename(path)}"; rclone_copy(path,remote); return remote
def has_active_project(deal_id,db=None):
    own=db is None; db=db or connect(); value=db.execute('''SELECT COUNT(*) FROM project_documents d JOIN project_versions v ON v.id=d.version_id WHERE v.deal_id=? AND v.status='active' ''',(str(deal_id),)).fetchone()[0]
    if own: db.close()
    return bool(value)

def project_files(deal_id):
    db=connect(); version=db.execute("SELECT * FROM project_versions WHERE deal_id=? AND status='active' ORDER BY version_no DESC LIMIT 1",(str(deal_id),)).fetchone()
    if not version: db.close(); return []
    docs=[dict(x) for x in db.execute('SELECT * FROM project_documents WHERE version_id=? ORDER BY id',(version['id'],))]; db.close(); result=[]
    for item in docs:
        suffix=os.path.splitext(item['name'])[1] or '.bin'; fd,path=tempfile.mkstemp(prefix=f'project-{deal_id}-v{version["version_no"]}-',suffix=suffix); os.close(fd)
        copy=subprocess.run(['rclone','copyto',item['onedrive_remote'],path],capture_output=True,text=True,timeout=600)
        if copy.returncode: raise ValueError('не удалось получить файл проекта из OneDrive')
        result.append({'path':path,'name':item['name'],'id':item['megaplan_file_id'],'version':version['version_no']})
    return result
def sync_project_files(deal_id):
    # Files are copied to OneDrive atomically when a version is created.
    return []
def publish_workbook(deal_id):
    db=connect()
    try:
        ins=get_installation(deal_id,db)
        if ins.get('publication_create_state')=='uncertain':
            raise ValueError('Требуется сверка первой публикации ведомости: результат предыдущего создания неизвестен')
        api=MegaplanAPI(); path=build_workbook(deal_id,db); uploaded=api.upload_file(path)
        payload={'contentType':'Comment','subject':{'contentType':'Deal','id':str(deal_id)},'content':'Актуальная монтажная ведомость. Файл обновляется после каждого принятого отчёта монтажника.','attaches':[{'contentType':'File','id':str(uploaded['id'])}]}
        if ins.get('megaplan_comment_id'):
            payload['id']=str(ins['megaplan_comment_id'])
            # An update timeout is not evidence that the existing comment is
            # absent. Retrying its ID is safe; creating another comment is not.
            try:
                comment=api._call('POST',f"/comment/{ins['megaplan_comment_id']}",json_body=payload).get('data') or {}
            except MegaplanAPIError as exc:
                if exc.status!=403:raise
                previous=api._call('GET',f"/comment/{ins['megaplan_comment_id']}").get('data') or {}
                if previous.get('possibleActions')!=['act_read'] or str((previous.get('subject') or {}).get('id'))!=str(deal_id):raise
                # CRM disallows editing this old comment. Publish a version,
                # retaining the immutable original rather than retrying forever.
                audit(db,deal_id,'publication',ins['megaplan_comment_id'],'readonly_comment_superseded','system',
                    {'previous_comment_id':ins['megaplan_comment_id'],'previous_file_id':ins.get('megaplan_file_id'),'pending_file_id':str(uploaded['id'])})
                db.execute("UPDATE installations SET publication_create_state='uncertain',megaplan_comment_id=NULL,megaplan_file_id=? WHERE deal_id=?",(str(uploaded['id']),str(deal_id)));db.commit()
                payload.pop('id',None)
                payload['content']+=' Новая версия; предыдущая публикация №'+str(ins['megaplan_comment_id'])+' сохранена в истории (Мегаплан запрещает её редактирование).'
                comment=api._call('POST',f'/deal/{deal_id}/comments',json_body=payload).get('data') or {}
                if not comment.get('id'):raise ValueError('Мегаплан не подтвердил новую версию ведомости')
        else:
            # Persist before the external write: loss of its response must never
            # cause another creation on the next background publication retry.
            db.execute("UPDATE installations SET publication_create_state='uncertain',megaplan_file_id=? WHERE deal_id=?",(str(uploaded['id']),str(deal_id))); db.commit()
            comment=api._call('POST',f'/deal/{deal_id}/comments',json_body=payload).get('data') or {}
        comment_id=str(comment.get('id') or ins.get('megaplan_comment_id') or '')
        if not comment_id: raise ValueError('Мегаплан не подтвердил ID публикации ведомости')
        db.execute("UPDATE installations SET megaplan_comment_id=?,megaplan_file_id=?,updated_at=?,publication_create_state='confirmed' WHERE deal_id=?",(comment_id,str(uploaded['id']),now(),str(deal_id))); db.commit()
        return {'comment_id':comment_id,'file_id':uploaded['id']}
    finally:
        db.close()

def refresh_economics(deal_id):
    """Rebuild the one managed economics workbook after any installation cost event.

    Megaplan receives the current working copy.  OneDrive receives an immutable
    dated version; a temporary archive outage is reported but never rolls back an
    approved report or a recorded payment.
    """
    script=os.path.join(BASE,'build_profitability_report.py')
    result=subprocess.run([sys.executable,script,str(deal_id)],cwd=BASE,capture_output=True,text=True,timeout=180)
    lines=(result.stdout or '').strip().splitlines(); parsed=json.loads(lines[-1]) if lines else {}
    if result.returncode or not parsed.get('ok'): raise ValueError(parsed.get('error') or (result.stderr or 'ошибка пересчёта экономики')[-500:])
    from deal_browser import replace_economics_attachment
    stored=replace_economics_attachment(MegaplanAPI(),str(deal_id),parsed['path'])
    parsed['stored']=stored
    try:
        from financial_archive import archive_document
        parsed['archive']=archive_document(str(deal_id),parsed['path'],'economics')
    except Exception as exc:
        parsed['archive_warning']=str(exc)
    return parsed
@transaction(lambda:DB+'.payments')
def record_payment(deal_id,source,amount,payment_date,external_id,comment='',allow_overpay=False):
    amount=float(amount)
    if not math.isfinite(amount) or amount<=0: raise ValueError('сумма выплаты должна быть конечной и больше нуля')
    if source not in ('cash','bank') or not str(external_id or '').strip(): raise ValueError('Не указан источник или устойчивый ID выплаты')
    datetime.strptime(payment_date,'%d.%m.%Y')
    request_id=str(external_id)
    db=connect(); prior=db.execute('SELECT * FROM payments WHERE source=? AND (request_id=? OR external_id=?)',(source,request_id,request_id)).fetchone(); db.close()
    if prior:
        if str(prior['deal_id'])!=str(deal_id) or abs(float(prior['amount'])-amount)>.001 or prior['payment_date']!=payment_date:
            raise ValueError('Этот ID выплаты уже связан с другими данными')
        out=get_installation(deal_id); out['operation_id']=str(prior['external_id']).removeprefix('megaplan:'); out['duplicate']=True; return out
    ins=get_installation(deal_id)
    # Q-PAY-01: owner confirmed 94% for NEW bank payments. NULL values in
    # historical rows retain their original credited amount; no backfill.
    from decimal import Decimal, ROUND_HALF_UP
    coefficient=0.94 if source=='bank' else 1.0
    settlement_amount=float((Decimal(str(amount))*Decimal(str(coefficient))).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))
    if settlement_amount>float(ins['debt'])+0.01 and not allow_overpay:
        raise ValueError(f"выплата превышает задолженность {float(ins['debt']):.2f} руб.")
    operation_id=''
    if source=='cash':
        api=MegaplanAPI(); auth=v1_auth(api)
        form={'Model[OperationType]':'35','Model[Kind]':'expense','Model[Sum][Value]':f'{amount:.2f}',
              'Model[Sum][Currency]':'1','Model[RealAccount]':'19','Model[Contractor]':str(ins.get('installer_ref') or ''),
              'Model[Date]':datetime.strptime(payment_date,'%d.%m.%Y').strftime('%Y-%m-%d 12:00:00'),
              'Model[DocumentNumber]':str(external_id),'Model[Description]':f"Выплата монтажнику по сделке №{deal_id}. {comment}"[:255]}
        response=v1_call(api,'/BumsFinApiV01/Operation/save.api',form,auth); item=(response.get('data') or {}).get('item') or (response.get('data') or {}).get('Operation') or response.get('data') or {}; operation_id=str(item.get('Id') or item.get('id') or '')
        if not operation_id: raise ValueError('Мегаплан не вернул ID кассовой операции')
        external_id=f'megaplan:{operation_id}'
    db=connect()
    try:
        db.execute('INSERT INTO payments(deal_id,source,amount,payment_date,external_id,comment,created_at,request_id,settlement_amount,coefficient) VALUES(?,?,?,?,?,?,?,?,?,?)',(str(deal_id),source,amount,payment_date,str(external_id),comment,now(),request_id,settlement_amount,coefficient))
    except sqlite3.IntegrityError: db.close(); raise ValueError('эта выплата уже учтена')
    enqueue_publication(db,deal_id)
    audit(db,deal_id,'payment',external_id,'recorded',source,{'amount':amount,'settlement_amount':settlement_amount,'coefficient':coefficient,'date':payment_date}); db.commit(); db.close()
    economics_error=''
    try:
        refresh_economics(deal_id)
    except Exception as exc:
        # A confirmed payment is a financial fact and must survive a report/archive
        # failure.  Preserve it and leave a precise audit trail for retry.
        economics_error=str(exc); db=connect(); audit(db,deal_id,'economics',deal_id,'refresh_failed','system',{'event':'installer_payment','error':economics_error}); db.commit(); db.close()
    out=get_installation(deal_id); out['operation_id']=operation_id; out['economics_refreshed']=not economics_error; out['economics_error']=economics_error
    out.update(payment_gross_amount=amount,payment_settlement_amount=settlement_amount,payment_coefficient=coefficient)
    return out
def sync_report(report_id):
    db=connect(); report=db.execute('SELECT * FROM reports WHERE id=?',(int(report_id),)).fetchone(); ins=dict(db.execute('SELECT * FROM installations WHERE deal_id=?',(report['deal_id'],)).fetchone()); media=[dict(x) for x in db.execute('SELECT * FROM media WHERE report_id=?',(int(report_id),))]; db.close(); base=f"{ins['folder_remote']}/02 Отчёты монтажника/Отчёт №{report['report_no']:03d} - {report['report_date']}"
    for item in media:
        kind='Видео' if str(item['mime_type']).startswith('video') else 'Фото' if str(item['mime_type']).startswith('image') else 'Документы'; remote=f"{base}/{kind}/{safe(item['original_name'])}"; rclone_copy(item['local_path'],remote); db=connect(); db.execute('UPDATE media SET onedrive_remote=? WHERE id=?',(remote,item['id'])); db.commit(); db.close()
    fd,path=tempfile.mkstemp(suffix='.txt',prefix='installation-report-'); os.close(fd); open(path,'w',encoding='utf-8').write(report['raw_text'] or ''); rclone_copy(path,f"{base}/Отчёт монтажника.txt"); os.remove(path); return base
def backup_db():
    if not os.path.exists(DB): return
    stamp=datetime.now().strftime('%Y-%m-%d'); target=f'{ONEDRIVE_ROOT}/Служебные реестры/installations-{stamp}.sqlite3'
    tmp=DB+'.backup'; src=sqlite3.connect(DB); dst=sqlite3.connect(tmp); src.backup(dst); dst.close(); src.close()
    try:rclone_copy(tmp,target)
    finally:
        try:os.remove(tmp)
        except OSError:pass

def main():
    req=json.load(sys.stdin); action=req.get('action')
    import legacy_report_commit
    funcs={'prepare':lambda:prepare(req['deal_id'],req.get('actor','')),'get':lambda:get_installation(req['deal_id']),
      'preview_legacy_report':lambda:legacy_report_commit.preview(req['session']),
      'accept_legacy_report':lambda:legacy_report_commit.accept(req['session'],req['source_digest'],req['expected_fingerprint'],req.get('actor','')),
      'sync_legacy_ledger':lambda:legacy_report_commit.sync_ledger(req.get('source_digest')),
      'retry_publications':lambda:retry_publications(req.get('limit',1)),
      'import_workbook':lambda:import_workbook(req['deal_id'],req['path'],req.get('actor','')),
      'sync_executor':lambda:sync_executor(req['deal_id'],req.get('actor','')),
      'import_legacy_workbook':lambda:import_legacy_workbook(req['deal_id'],req['path'],req['report_date'],req['external_id'],req.get('actor','')),
      'update_details':lambda:update_details(req['deal_id'],req.get('contact_name'),req.get('contact_phone'),req.get('project_required'),req.get('actor','')),
      'contact_candidates':lambda:installation_contacts(req['deal_id']),
      'select_contact':lambda:select_installation_contact(req['deal_id'],req['contact_ref'],req.get('actor','')),
      'create_contact':lambda:create_installation_contact(req['deal_id'],req['name'],req['phone'],req.get('actor','')),
      'approve':lambda:approve(req['deal_id'],req.get('actor','')),'transfer':lambda:transfer(req['deal_id'],req.get('actor','')),
      'list_mine':lambda:list_mine(req.get('installer_name')),'list_history':lambda:list_history(req.get('installer_name')),'create_report':lambda:create_report(req['deal_id'],req['chat_id'],req.get('text','')),
      'append_report_text':lambda:append_report_text(req['report_id'],req['text']),'add_media':lambda:add_media(req['report_id'],req['path'],req.get('file_id',''),req.get('unique_id',''),req.get('name',''),req.get('mime','')),
      'get_report':lambda:get_report(req['report_id']),
      'pending_reports':pending_reports,
      'parse_report':lambda:parse_report(req['report_id']),'submit_report':lambda:submit_report(req['report_id']),
      'approve_report':lambda:approve_report(req['report_id'],req.get('actor',''),req.get('partial'),req.get('comment','')),
      'return_report':lambda:return_report(req['report_id'],req['comment'],req.get('actor','')),
      'project_candidates':lambda:project_candidates(req['deal_id']),
      'attach_project':lambda:attach_project(req['deal_id'],req.get('file_ids'),req.get('local_paths'),req.get('actor','')),
      'without_project':lambda:set_without_project(req['deal_id'],req.get('actor','')),
      'project_files':lambda:project_files(req['deal_id']),
      'record_payment':lambda:record_payment(req['deal_id'],req['source'],req['amount'],req['payment_date'],req['external_id'],req.get('comment',''),req.get('allow_overpay',False)),
      'backup':lambda:(backup_db() or {'done':True})}
    if action not in funcs: raise ValueError('неизвестное действие')
    result=funcs[action](); print(json.dumps({'ok':True,'result':result},ensure_ascii=False))
if __name__=='__main__':
    try:main()
    except InstallationNotPrepared as exc:
        print(json.dumps({'ok':False,'error':str(exc),'error_code':'installation_not_prepared'},ensure_ascii=False)); raise SystemExit(1)
    except Exception as exc: print(json.dumps({'ok':False,'error':str(exc)},ensure_ascii=False)); raise SystemExit(1)
